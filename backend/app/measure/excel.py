# ruff: noqa: E501  (cell layouts read better unwrapped)
"""The measurement sheet as Excel, with live formulas (Qty = No × L × H, Total = ROUND(Qty × 10.7639, 3)
for sqft, the abstract per item by SUMIF), and reading old Excel sheets and proformas back.

`evaluate` works out the formulas this file writes (and the usual ones in old sheets: + - * / ^,
ROUND, SQRT, SUM, SUMIF, cell references): the tests check the workbook's totals against the app's,
and an old file saved without its calculated values can still be read."""

import io
import re
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.measure import geom
from app.measure.models import MeasureLine

HEAD = ["Sr / WO item", "Particulars", "No", "Length (m)", "Width (m)", "Height / Depth (m)", "Qty", "Total qty", "Unit"]
FIRST = 7  # the first line's row
DASH = "–"
KEY = 14  # the hidden column with each line's item key (N), clear of the abstract's columns


def _f(v) -> float | None:
    return None if v is None else float(Decimal(v))


def _qty_formula(r: int, ln: MeasureLine) -> str | float:
    sign = "-" if ln.kind == "deduction" else ""
    if ln.part == "walls" and ln.height_m is not None and ln.width_m is None:
        return f"={sign}C{r}*D{r}*F{r}"
    if ln.part == "tops":
        return f"={sign}C{r}*D{r}*E{r}"
    if ln.part == "count":
        return f"=C{r}"
    if ln.part == "walls" and ln.length_m is not None:
        return f"={sign}C{r}*D{r}"
    return _f(ln.qty)  # an outline's own area (or a slope face, written below with its formula)


def _total_formula(r: int, unit: str | None) -> str:
    u = (unit or "").strip().lower()
    if u == "sqft":
        return f"=ROUND(G{r}*$H$3,3)"
    if u == "rft":
        return f"=ROUND(G{r}/0.3048,3)"
    return f"=ROUND(G{r},3)"


def export(sheet, site, lines: list[MeasureLine], cumulative: list[dict], areas: dict, slope: dict, company: str, sqft: Decimal = geom.SQFT_PER_SQM) -> bytes:
    """The billing team's layout: the area's header line, then its lines; deductions negative with
    their reason; the abstract per work-order item (previous, this bill, up to date)."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet.code[:31]
    bold = Font(bold=True)
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"], ws["A1"].font = company, Font(bold=True, size=13)
    ws["A2"], ws["A2"].font = "MEASUREMENT SHEET", bold
    ws["A3"] = f"Site: {site.code} {site.name}"
    ws["D3"] = f"Sheet {sheet.code} · measured on {sheet.measured_on:%d-%m-%Y}"
    ws["G3"], ws["H3"], ws["I3"] = "1 sqm =", float(sqft), "sqft"
    ws["G3"].font = ws["H3"].font = bold
    ws["A4"] = "Quantities from the drawing at full precision; each line's total rounded to 3 decimals (half up, never truncated)."
    for i, h in enumerate(HEAD, start=1):
        c = ws.cell(row=6, column=i, value=h)
        c.font, c.border = bold, box
        c.fill = PatternFill("solid", fgColor="DDE7E1")
    ws.cell(row=6, column=KEY, value="item key").font = Font(color="999999")
    r = FIRST
    block = FIRST  # where the item's lines start (a study converting per item totals them once)
    for ln in lines:
        if ln.kind == "area":
            if ln.wo_line_id and not ln.area_id:
                block = r + 1
                ws.cell(row=r, column=1, value=ln.item_no).font = bold
            ws.cell(row=r, column=2, value=ln.particulars).font = bold
            r += 1
            continue
        if ln.kind == "total":
            ws.cell(row=r, column=2, value=ln.particulars).font = bold
            ws.cell(row=r, column=7, value=f"=SUM(G{block}:G{r - 1})").font = bold
            ws.cell(row=r, column=8, value=_total_formula(r, ln.unit)).font = bold
            ws.cell(row=r, column=9, value=ln.unit)
            ws.cell(row=r, column=KEY, value=f"wo:{ln.wo_line_id}").font = Font(color="999999")
            ws.cell(row=r, column=7).number_format = "0.000000"
            ws.cell(row=r, column=8).number_format = "0.000"
            r += 1
            continue
        ws.cell(row=r, column=1, value=ln.item_no)
        text = ln.particulars + (f" — {ln.reason}" if ln.kind == "deduction" and ln.reason else "")
        ws.cell(row=r, column=2, value=text).alignment = Alignment(wrap_text=True, vertical="top")
        if ln.kind == "not_done":
            ws.cell(row=r, column=7, value=DASH)
            ws.cell(row=r, column=8, value=DASH)
        else:
            ws.cell(row=r, column=3, value=_f(ln.no))
            ws.cell(row=r, column=4, value=_f(ln.length_m))
            ws.cell(row=r, column=5, value=_f(ln.width_m))
            ws.cell(row=r, column=6, value=_f(ln.height_m))
            q = _qty_formula(r, ln)
            if ln.part == "slope" and ln.area_id in slope:
                plan, rise, run, length = slope[ln.area_id]
                q = f"={plan!r}*{length!r}/{run!r}" if length else f"={plan!r}*SQRT({rise!r}^2+{run!r}^2)/{run!r}"
            ws.cell(row=r, column=7, value=q)
            if ln.total is not None:
                ws.cell(row=r, column=8, value=_total_formula(r, ln.unit))
        ws.cell(row=r, column=9, value=ln.unit)
        ws.cell(row=r, column=KEY, value=f"wo:{ln.wo_line_id}" if ln.wo_line_id and ln.total is not None else None).font = Font(color="999999")
        for col in range(3, 9):
            ws.cell(row=r, column=col).number_format = "0.000" if col == 8 else "0.0000"
        r += 1
    last = r - 1
    r += 1
    ws.cell(row=r, column=1, value="Abstract").font = bold
    r += 1
    # the bill's figures: the drawing's (live, from the lines), the client's approved where they
    # approved it (a correction included); "This bill" is the approved one, as on the RA bill
    heads = ["WO item", "Description", "Unit", "Previous", "On the drawing", "Approved by the client", "This bill", "Up to date", "WO qty", "Rate", "This bill ₹"]
    for i, h in enumerate(heads, start=1):
        c = ws.cell(row=r, column=i, value=h)
        c.font, c.border = bold, box
    r += 1
    abstract_first = r
    key = get_column_letter(KEY)
    for row in cumulative:
        ws.cell(row=r, column=1, value=row["item_no"])
        ws.cell(row=r, column=2, value=(row["description"] or "")[:200]).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=3, value=row["unit"])
        ws.cell(row=r, column=4, value=_f(row["previous"]))
        ws.cell(row=r, column=5, value=f'=SUMIF(${key}${FIRST}:${key}${max(last, FIRST)},"wo:{row["wo_line_id"]}",$H${FIRST}:$H${max(last, FIRST)})')
        if row.get("approved") is not None:
            ws.cell(row=r, column=6, value=_f(row["approved"]))
            ws.cell(row=r, column=7, value=f"=F{r}")
        else:
            ws.cell(row=r, column=7, value=f"=E{r}")
        ws.cell(row=r, column=8, value=f"=D{r}+G{r}")
        ws.cell(row=r, column=9, value=_f(row["wo_qty"]))
        ws.cell(row=r, column=10, value=_f(row["rate"]))
        ws.cell(row=r, column=11, value=f"=ROUND(G{r}*J{r},2)")
        for col in (4, 5, 6, 7, 8, 9):
            ws.cell(row=r, column=col).number_format = "0.000"
        ws.cell(row=r, column=11).number_format = "#,##0.00"
        r += 1
    ws.cell(row=r, column=10, value="Basic").font = bold
    ws.cell(row=r, column=11, value=f"=SUM(K{abstract_first}:K{r - 1})").font = bold
    ws.cell(row=r, column=11).number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHIJK", (12, 60, 7, 11, 12, 13, 12, 12, 10, 9, 12), strict=True):
        ws.column_dimensions[col].width = w
    ws.column_dimensions[key].hidden = True
    ws.freeze_panes = "C7"
    ws.print_title_rows = "6:6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- formulas ------------------------------------------------------------------------------------

_TOKEN = re.compile(
    r'\s*(?:(\d+(?:\.\d*)?(?:[eE][+-]?\d+)?)|("(?:[^"]|"")*")'
    r"|((?:'(?:[^']|'')+'|[A-Za-z_][\w.]*)!\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)"  # another sheet's cell or range
    r'|(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)|([A-Z]+(?=\())|(\*\*|[-+*/^(),:<>=&]))'
)


def _col(s: str) -> int:
    n = 0
    for ch in s:
        n = n * 26 + ord(ch) - 64
    return n


def _ref(s: str) -> tuple[int, int]:
    m = re.fullmatch(r"\$?([A-Z]{1,3})\$?(\d+)", s)
    return int(m.group(2)), _col(m.group(1))


class _Eval:
    def __init__(self, ws, values=None):
        """ws: the sheet with formulas; values: the same workbook's calculated values, used first
        where the file has them (a reference to another sheet reads that sheet)."""
        self.ws, self.values, self.cache, self.busy = ws, values, {}, set()

    def cell(self, row: int, col: int, sheet=None):
        ws = sheet if sheet is not None else self.ws
        k = (ws.title, row, col)
        if k in self.cache:
            return self.cache[k]
        if k in self.busy:
            raise ValueError("circular reference")
        if self.values is not None and ws.title in self.values.sheetnames:
            cached = self.values[ws.title].cell(row=row, column=col).value
            if cached is not None:
                self.cache[k] = cached
                return cached
        v = ws.cell(row=row, column=col).value
        if isinstance(v, str) and v.startswith("="):
            self.busy.add(k)
            here, self.ws = self.ws, ws
            try:
                v = self.formula(v[1:])
            finally:
                self.ws = here
                self.busy.discard(k)
        self.cache[k] = v
        return v

    def _sheet_ref(self, t: str):
        name, ref = t.rsplit("!", 1)
        name = name[1:-1].replace("''", "'") if name.startswith("'") else name
        wb = self.ws.parent
        if name not in wb.sheetnames:
            raise ValueError(f"no sheet {name}")
        return wb[name], ref

    def num(self, v) -> float:
        if v is None or v == "":
            return 0.0
        if isinstance(v, int | float):
            return float(v)
        try:
            return float(str(v).replace(",", ""))
        except ValueError:
            return 0.0

    def formula(self, text: str):
        saved = getattr(self, "toks", None), getattr(self, "i", 0)  # a cell it refers to has its own
        self.toks = [next(t for t in m.groups() if t is not None) for m in _TOKEN.finditer(text) if m.group(0).strip()]
        self.i = 0
        try:
            return self.expr()
        finally:
            self.toks, self.i = saved

    def peek(self):
        return self.toks[self.i] if self.i < len(self.toks) else None

    def take(self):
        t = self.toks[self.i]
        self.i += 1
        return t

    def expr(self):
        v = self.term()
        while self.peek() in ("+", "-"):
            op = self.take()
            w = self.term()
            v = self.num(v) + self.num(w) if op == "+" else self.num(v) - self.num(w)
        return v

    def term(self):
        v = self.power()
        while self.peek() in ("*", "/"):
            op = self.take()
            w = self.power()
            v = self.num(v) * self.num(w) if op == "*" else self.num(v) / self.num(w)
        return v

    def power(self):
        v = self.unary()
        while self.peek() == "^":
            self.take()
            v = self.num(v) ** self.num(self.unary())
        return v

    def unary(self):
        if self.peek() == "-":
            self.take()
            return -self.num(self.unary())
        if self.peek() == "+":
            self.take()
            return self.num(self.unary())
        return self.atom()

    def atom(self):
        t = self.take()
        if t == "(":
            v = self.expr()
            self.take()
            return v
        if t[0].isdigit():
            return float(t)
        if t.startswith('"'):
            return t[1:-1].replace('""', '"')
        if self.peek() == "(" and re.fullmatch(r"[A-Z]+", t):
            self.take()
            args = []
            while self.peek() != ")":
                args.append(self.arg())
                if self.peek() == ",":
                    self.take()
            self.take()
            return self.call(t, args)
        if "!" in t:
            sheet, ref = self._sheet_ref(t)
            return self.range(ref, sheet) if ":" in ref else self.cell(*_ref(ref), sheet=sheet)
        if ":" in t:
            return self.range(t)
        return self.cell(*_ref(t))

    def arg(self):
        t = self.peek()
        if t and ":" in t and "!" not in t and re.fullmatch(r"\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+", t):
            self.take()
            return self.range(t)
        if t and "!" in t and ":" in t.rsplit("!", 1)[1]:
            self.take()
            sheet, ref = self._sheet_ref(t)
            return self.range(ref, sheet)
        return self.expr()

    def range(self, t: str, sheet=None) -> list:
        a, b = t.split(":")
        (r1, c1), (r2, c2) = _ref(a), _ref(b)
        return [self.cell(r, c, sheet) for r in range(min(r1, r2), max(r1, r2) + 1) for c in range(min(c1, c2), max(c1, c2) + 1)]

    def call(self, name: str, args: list):
        if name == "ROUND":
            x, n = self.num(args[0]), int(self.num(args[1]))
            q = Decimal(repr(x)).quantize(Decimal(1).scaleb(-n), rounding="ROUND_HALF_UP")
            return float(q)
        if name == "SQRT":
            return self.num(args[0]) ** 0.5
        if name == "SUM":
            return sum(self.num(v) for a in args for v in (a if isinstance(a, list) else [a]) if not isinstance(v, str) or v.strip())
        if name == "SUMIF":
            keys, crit, vals = args[0], args[1], args[2] if len(args) > 2 else args[0]
            return sum(self.num(v) for k, v in zip(keys, vals, strict=False) if k == crit)
        if name == "ABS":
            return abs(self.num(args[0]))
        raise ValueError(f"{name} is not read")


def evaluate(ws, row: int, col: int):
    """A cell's value, its formula worked out (the cells it refers to as well)."""
    return _Eval(ws).cell(row, col)


# --- reading old sheets --------------------------------------------------------------------------

_HEADS = {
    "item": ("sr / wo item", "wo item", "item no", "item", "sr no", "sr. no", "sr", "s no", "s.no", "sl no", "sr no./ wo item no.", "sr no./wo item no."),
    "particulars": ("particulars", "description", "details"),
    "no": ("no", "nos", "no."),
    "length": ("length", "length (m)", "l"),
    "width": ("width", "width (m)", "breadth", "b", "w"),
    "height": ("height / depth (m)", "height/depth", "height", "depth", "h"),
    "qty": ("qty", "quantity"),
    "total": ("total qty", "total quantity", "total"),
    "unit": ("unit", "uom", "u.o.m."),
    "rate": ("rate",),
    "amount": ("amount", "amount (rs)", "amount ₹", "this bill ₹"),
    "previous": ("previous", "previous qty"),
    "this": ("this bill", "this bill qty", "present"),
    "upto": ("up to date", "upto date", "cumulative", "total up to date"),
}
UNIT_WORDS = {"sq ft": "sqft", "sqft": "sqft", "sq. ft": "sqft", "sq mt": "sqm", "sq m": "sqm", "sqm": "sqm", "sq. mt": "sqm", "nos": "nos", "no": "nos", "rmt": "rmt", "r mt": "rmt"}


def _norm(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().lower()


def _head_key(h: str) -> str | None:
    for key, names in _HEADS.items():
        if h in names:
            return key
    if h.startswith(("sr no", "sr. no")) or "wo item" in h:
        return "item"  # "Sr No./ WO item No." and the like
    return None


def _header(ws, need: tuple[str, ...]) -> tuple[int, dict[str, int]] | None:
    for r in range(1, min(ws.max_row, 40) + 1):
        cols: dict[str, int] = {}
        for c in range(1, min(ws.max_column, 30) + 1):
            key = _head_key(_norm(ws.cell(row=r, column=c).value))
            if key and key not in cols:
                cols[key] = c
        if all(k in cols for k in need):
            return r, cols
    return None


def _value(ev: _Eval, ws, r: int, c: int | None):
    """A cell's value: the file's calculated one, else its formula worked out (another sheet's
    cells followed)."""
    if c is None:
        return None
    v = ws.cell(row=r, column=c).value
    if v is None or (isinstance(v, str) and v.startswith("=")):
        try:
            v = ev.cell(r, c)
        except Exception:  # noqa: BLE001
            return None
    return v


def _number(v) -> Decimal | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, int | float):
        return Decimal(repr(float(v))) if isinstance(v, float) else Decimal(v)
    s = str(v).replace(",", "").replace("₹", "").strip()
    try:
        return Decimal(s) if s and s not in ("-", DASH) else None
    except ArithmeticError:
        return None


def _item_no(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def _unit(v) -> str:
    s = _norm(v).rstrip(".")
    return UNIT_WORDS.get(s, s)


def _workbooks(data: bytes):
    """The calculated values where the file has them; the formulas to work out where not."""
    return load_workbook(io.BytesIO(data), data_only=True), load_workbook(io.BytesIO(data))


def read_sheet(data: bytes) -> dict:
    """An old measurement sheet: per work-order item its quantity: the item's own total row ("Total
    Quantity =", in the unit printed under it) where the sheet has one, else its lines added up.
    The area headers, the abstract and the sheet's own total rows are never added again. A line
    with no item number is counted under the item above it, and listed to check unless the item's
    own total row confirms it. Every row with a number that could not be placed is listed."""
    try:
        values, formulas = _workbooks(data)
    except Exception as exc:  # noqa: BLE001
        raise ValueError("This is not an Excel workbook (.xlsx) that can be read") from exc
    items: dict[str, dict] = {}
    skipped: list[str] = []
    assumed: dict[str, list[str]] = {}  # per item: its lines with no item number
    for name in values.sheetnames:
        ws, wf = values[name], formulas[name]
        found = _header(ws, ("particulars", "qty"))
        priced = _header(ws, ("rate", "amount"))
        if priced and (not found or priced[0] <= found[0]):
            continue  # a proforma's page (its priced header first): read by read_proforma
        if not found:
            if ws.max_row > 2:
                skipped.append(f"Sheet '{name}': no header row with Particulars and Qty")
            continue
        hr, cols = found
        ev = _Eval(wf, values)
        current = None
        r = hr
        while r < ws.max_row:
            r += 1
            part = str(ws.cell(row=r, column=cols["particulars"]).value or "").strip()
            first = _norm(ws.cell(row=r, column=1).value)
            if first.startswith(("abstract", "summary")) or _norm(part).startswith(("abstract", "summary")):
                break  # the abstract repeats the totals: not added again
            item = _item_no(ws.cell(row=r, column=cols["item"]).value) if "item" in cols else None
            qty = _number(_value(ev, ws, r, cols.get("qty")))
            total = _number(_value(ev, ws, r, cols.get("total")))
            if _norm(part).startswith(("total", "sub total", "subtotal", "grand total", "net total")):
                # the item's own total: its quantity in the unit printed on the row below (Sq Ft)
                if current and (total is not None or qty is not None):
                    row = items.setdefault(current, _new(current, part))
                    unit_below = _unit(ws.cell(row=r + 1, column=cols["total"]).value) if "total" in cols and r + 1 <= ws.max_row else ""
                    row["total_row"] = {"qty": qty, "total": total, "unit": unit_below or row["unit"], "row": r}
                continue
            value = total if total is not None else qty
            if value is None:
                if item:  # an item's heading; an area's heading or "work not done" inside it go on
                    current = item
                    items.setdefault(item, _new(item, part))
                continue
            if not item and not current:
                skipped.append(f"Sheet '{name}' row {r}: {value} with no work-order item ({part[:60] or 'no particulars'})")
                continue
            if not item:
                assumed.setdefault(current, []).append(f"Sheet '{name}' row {r}: no item number, counted under item {current} (the line above): {part[:50]}")
            item = item or current
            current = item
            unit = _unit(ws.cell(row=r, column=cols["unit"]).value) if "unit" in cols else ""
            row = items.setdefault(item, _new(item, part))
            row["lines_qty"] += qty if qty is not None else value
            row["lines_total"] += value
            row["rows"] += 1
            row["unit"] = row["unit"] or unit
    notes: list[str] = []
    for key, row in list(items.items()):
        t = row.pop("total_row", None)
        if t:
            row["qty"] = t["total"] if t["total"] is not None else t["qty"]
            row["unit"] = (t["unit"] or row["unit"]) if t["total"] is not None else row["unit"]
            row["sqm"] = t["qty"] if t["total"] is not None else None
            agree = (t["qty"] is not None and abs(t["qty"] - row["lines_qty"]) <= Decimal("0.0005")) or (
                t["qty"] is None and t["total"] is not None and abs(t["total"] - row["lines_total"]) <= Decimal("0.0005"))  # fmt: skip
            if agree:
                assumed.pop(key, None)  # the item's own total row confirms the lines counted under it
            else:
                notes.append(f"Item {key}: its lines add to {row['lines_qty']} but its total row says {t['qty'] if t['qty'] is not None else t['total']}")
        else:
            row["qty"] = row["lines_total"]
        if not row["rows"] and not t:
            items.pop(key)  # an item heading with nothing measured ("work not done")
    return {"items": items, "skipped": skipped, "assumed": [a for rows in assumed.values() for a in rows], "notes": notes}


def _new(item: str, part: str) -> dict:
    return {"item_no": item, "description": part, "unit": "", "qty": Decimal(0), "lines_qty": Decimal(0), "lines_total": Decimal(0), "rows": 0}


def read_proforma(data: bytes) -> dict:
    """An old proforma: per item its quantity (up to date where the bill shows it, else this
    bill's), rate and amount, the item number taken from the row that starts the item; the basic
    and the grand total as printed."""
    try:
        values, formulas = _workbooks(data)
    except Exception as exc:  # noqa: BLE001
        raise ValueError("This is not an Excel workbook (.xlsx) that can be read") from exc
    items: dict[str, dict] = {}
    skipped: list[str] = []
    basic = grand = None
    for name in values.sheetnames:
        ws, wf = values[name], formulas[name]
        found = _header(ws, ("rate", "amount"))
        if not found:
            continue
        hr, cols = found
        ev = _Eval(wf, values)
        qcol = cols.get("upto") or cols.get("qty") or cols.get("this")
        current = None
        for r in range(hr + 1, ws.max_row + 1):
            raw = [ws.cell(row=r, column=c).value for c in range(1, min(ws.max_column, 12) + 1)]
            label = " ".join(_norm(v) for v in raw if isinstance(v, str))
            amount = _number(_value(ev, ws, r, cols["amount"]))
            if "grand total" in label or "total amount" in label or "invoice total" in label:
                grand = amount if amount is not None else grand
                continue
            if any(w in label for w in ("basic", "taxable", "sub total", "subtotal")) or re.search(r"(^| )total( |$)", label):
                basic = amount if amount is not None and basic is None else basic
                continue
            if any(w in label for w in ("cgst", "sgst", "igst", "round off", "round-off", "gst")):
                continue
            item = _item_no(ws.cell(row=r, column=cols["item"]).value) if "item" in cols else None
            if item:
                current = item
            qty = _number(_value(ev, ws, r, qcol)) if qcol else None
            rate = _number(_value(ev, ws, r, cols["rate"]))
            if qty is not None and rate is not None and (item or current):
                key = item or current
                desc = str(ws.cell(row=r, column=cols["particulars"]).value or "").strip() if "particulars" in cols else ""
                unit = _unit(_value(ev, ws, r, cols.get("unit"))) if "unit" in cols else ""
                items[key] = {"item_no": key, "description": desc, "unit": unit, "qty": qty, "rate": rate, "amount": amount}
                current = None  # the item's quantity is read: the next starts with its own number
            elif amount is not None or qty is not None:
                skipped.append(f"Sheet '{name}' row {r}: " + " ".join(str(v) for v in raw if v is not None)[:80] + " (no item number, or no quantity and rate)")
    return {"items": items, "basic": basic, "grand": grand, "skipped": skipped}


def match(items: dict[str, dict], wo_lines) -> tuple[dict[int, dict], list[str]]:
    """The read items to work-order lines: by item number, else by the start of the description."""
    by_no = {_norm(w.item_no): w for w in wo_lines if w.item_no}
    out: dict[int, dict] = {}
    skipped = []
    for key, it in items.items():
        w = by_no.get(_norm(key)) or by_no.get(_norm(key).rstrip(".0") if _norm(key).endswith(".0") else "")
        if w is None and it.get("description"):
            d = _norm(it["description"])[:25]
            w = next((x for x in wo_lines if d and _norm(x.description).startswith(d)), None)
        if w is None:
            skipped.append(f"Item {key} ({(it.get('description') or '')[:50]}): not on the work order")
            continue
        if w.id in out:
            out[w.id]["qty"] += it["qty"]
        else:
            out[w.id] = dict(it)
    return out, skipped
