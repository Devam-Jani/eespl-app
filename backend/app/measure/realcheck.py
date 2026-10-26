# ruff: noqa: E501
"""The billing team's real example, measured in the app and compared (the training company only).

Everything specific to the example is in the folder with its files (gitignored data/), never in
the code: the drawing (.dwg), the old Excel sheet and proforma (.xlsx), and check.json, which says
what a person clicks: where, which outline to choose there (by layer), which line to pick. This
module replays those clicks through the same calls the Measure and bill page makes:

1. a site with the order's items (rates as billed), the old RA bill imported as the previous;
2. the drawing read (its unit set by hand when it has none), the areas clicked: the plan's own
   outlines (the app's way) and the copies the billing team drew beside the plans (their way);
3. heights from the sections (two levels, numbers added up, or a picked line's true length), the
   slope from its picked line, the wall between the team's two copies from the line drawn for it;
4. two study sheets (never sent or billed): the billing team's way, which must reproduce their
   quantities, and the app's default way; the comparison item by item, in sqm, sqft and rupees.

    docker compose exec api python -m app.cli real-billing-check /tmp/real --out /tmp/real/pool-a-comparison.xlsx
"""

import io
import json
import shutil
from datetime import date
from decimal import Decimal
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.measure import geom, importer
from app.measure import service as ms
from app.measure.models import MeasureArea, MeasureEdge, MeasureItem, MeasureSheet

D = Decimal


class Missing(SystemExit):
    pass


def files(folder: Path) -> tuple[Path, Path, dict]:
    """The example's files, found by the globs in check.json (their names have double spaces and
    brackets: never typed)."""
    cfg_path = folder / "check.json"
    if not cfg_path.exists():
        raise Missing(f"The real example is not in {folder}: check.json is missing")
    cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
    base = folder / cfg.get("folder", ".")
    xl = sorted(p for p in base.glob(cfg.get("sheet_glob", "*.xlsx")) if not p.name.startswith("~$"))
    dw = sorted(base.glob(cfg.get("drawing_glob", "*.dwg")))
    if not base.is_dir() or not xl or not dw:
        raise Missing(f"The real example is not in {base}: the old sheet (.xlsx) or the drawing (.dwg) is missing")
    return xl[0], dw[0], cfg


def _user(db: Session):
    from app.models import User
    from app.team.demo import email_for

    u = db.scalar(select(User).where(User.email == email_for("billing")))
    if u is None:
        u = db.scalar(select(User).where(User.is_active.is_(True)).order_by(User.created_at).limit(1))
    return u


def _site(db: Session, cfg: dict, user) -> tuple:
    from app.closeout.models import ClientWorkOrder, ClientWorkOrderLine
    from app.finance.models import ClientContract, ContractLine
    from app.masters.models import Client
    from app.measure.demo import reset
    from app.sites.models import Site

    cl = db.scalar(select(Client).where(Client.name == cfg["client"]))
    if cl is None:
        cl = Client(name=cfg["client"], type="other", state="Gujarat", notes="Training: the billing team's real example, measured in the app")
        db.add(cl)
        db.flush()
    site = db.scalar(select(Site).where(Site.code == cfg["site_code"]))
    if site is None:
        site = Site(code=cfg["site_code"], name=cfg["site"], state="Gujarat", status="active", client_id=cl.id)
        db.add(site)
        db.flush()
    else:
        reset(db, site.id)
    wo = ClientWorkOrder(code=f"CWO-{cfg['site_code']}", client_ref=cfg.get("order"), site_id=site.id, client_id=cl.id, lines_from="typed", status="accepted")
    for i, it in enumerate(cfg["items"]):
        wo.lines.append(ClientWorkOrderLine(sort_order=i, item_no=it["no"], description=it["description"], unit=it["unit"], qty=D(it["qty"]), rate=D(it["rate"]), layers=1, per_layer=False))
    db.add(wo)
    db.flush()
    c = ClientContract(site_id=site.id, client_id=cl.id, contract_value=sum((D(it["qty"]) * D(it["rate"]) for it in cfg["items"]), D(0)), retention_percent=D(0), place_of_supply="Gujarat", created_by=user.id)
    for w in wo.lines:
        c.lines.append(ContractLine(sort_order=w.sort_order, item_no=w.item_no, description=w.description, unit=w.unit, qty=w.qty, rate=w.rate, wo_line_id=w.id))
    db.add(c)
    db.flush()
    return site, wo


def _drawing(db: Session, site, dwg: Path, cfg: dict, user):
    from app.config import settings
    from app.sites.models import Drawing, DrawingRevision
    from app.takeoff import pipeline
    from app.takeoff.models import CadFile

    rel = f"sites/{site.id}/drawings/{site.code}-r1{dwg.suffix.lower()}"
    dst = Path(settings.media_dir) / rel
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(dwg, dst)
    d = Drawing(site_id=site.id, title=f"{cfg['site']}: structural drawing (the billing team's)", discipline="structural", created_by=user.id)
    db.add(d)
    db.flush()
    r = DrawingRevision(drawing_id=d.id, rev_no=1, stored_path=rel, filename=dwg.name, size_bytes=dst.stat().st_size, created_by=user.id)
    db.add(r)
    db.flush()
    d.current_revision_id = r.id
    # the drawing has no unit ($INSUNITS 0): set by hand, as the takeoff page asks
    f = CadFile(revision_id=r.id, drawing_id=d.id, site_id=site.id, status="queued", unit_m=cfg["unit_m"], units_set_by_hand=True, created_by=user.id)
    db.add(f)
    db.commit()
    f = pipeline.process(db, f)
    if f.status != "ready":
        raise SystemExit(f"The drawing was not read: {f.error}")
    return f


def _pick(db: Session, site, f, spec: dict, user) -> MeasureArea:
    """A click inside the outline, then the outline chosen among those around it (by layer)."""
    taken = {a.handle for a in ms.areas_of(db, site.id) if a.handle}
    around = ms.candidates(f, spec["click"][0], spec["click"][1], taken)
    chosen = [c for c in around if c["layer"] == spec["layer"]]
    if not chosen:
        raise SystemExit(f"No outline on layer {spec['layer']} around the click for {spec['code']}: {[(c['layer'], round(c['area_sqm'], 4)) for c in around]}")
    o = ms.outline_of(f, chosen[0]["handle"])
    return ms.make_area(db, site.id, o["outline"], code=spec["code"], name=spec["name"], cad_file=f, handle=o["handle"], layer=o["layer"], user_id=user.id,
                        gap_m=o.get("gap_m"), drawn_m=o.get("drawn_m"))  # fmt: skip


def _line(f, near: list[float]) -> dict:
    from shapely.geometry import LineString, Point

    pt = Point(*near)
    best = None
    for ln in ms.drawing(f)["lines"]:
        path = geom.flatten_open(ln["points"])
        if len(path) < 2:
            continue
        dist = LineString(path).distance(pt)
        if best is None or dist < best[0]:
            best = (dist, ln["handle"])
    if best is None or best[0] > 0.05:
        raise SystemExit(f"No line within 50 mm of {near}")
    return ms.line_of(f, best[1])


def run(db: Session, folder: Path) -> dict:
    xl, dwg, cfg = files(folder)
    user = _user(db)
    site, wo = _site(db, cfg, user)
    wl = {w.item_no: w for w in wo.lines}
    # 1. the old RA bill: its quantities are the previous
    data = xl.read_bytes()
    imp = importer.run(db, site, sheet_file=(xl.name, data), proforma_file=(xl.name, data), measured_on=date.fromisoformat(cfg["measured_on"]), user=user, ra_no=cfg.get("ra_no"))
    db.commit()
    # 2. the drawing and the areas, clicked
    f = _drawing(db, site, dwg, cfg, user)
    areas: dict[str, MeasureArea] = {}
    for spec in cfg["areas"]:
        areas[spec["code"]] = _pick(db, site, f, spec, user)
    db.commit()
    texts = {t[3].strip(): t for t in ms.texts_m(f)}
    for code, h in cfg["heights"].items():
        a = areas[code]
        if h.get("line_near"):
            ln = _line(f, h["line_near"])
            a.height_mm = geom.d(ln["length_m"] * 1000)
            a.height_source = f"{h['section']}: the line on layer {ln['layer']}, {D(a.height_mm).quantize(D('0.01'))} mm long (an arc measured as an arc)"
        elif h.get("sum_of"):
            a.height_mm, a.height_source = geom.height_sum(h["section"], h["sum_of"])
        else:
            for want in (h["top"], h["bottom"]):
                if want not in texts:
                    raise SystemExit(f"The level '{want}' is not written on the drawing")
            a.height_mm, a.height_source = geom.height_from(h["section"], h["top"], h["bottom"])
    for code, s in cfg["slopes"].items():
        a = areas[code]
        ln = _line(f, s["line_near"])
        a.kind, a.slope_mm, a.run_mm, a.rise_mm = "slope", geom.d(ln["length_m"] * 1000), geom.d(ln["run_m"] * 1000), geom.d(ln["rise_m"] * 1000)
        a.slope_source = f"{s['section']}: the slope line on layer {ln['layer']}, {D(a.slope_mm).quantize(D('0.01'))} mm long over a {D(a.run_mm).quantize(D('0.01'))} mm run"
    for de in cfg.get("drawn_edges", []):
        x, y = (areas[c] for c in de["between"])
        ln = _line(f, de["line_near"])
        path = geom.flatten_open(ln["points"])
        a_id, b_id = sorted((x.id, y.id))
        db.add(MeasureEdge(site_id=site.id, area_a_id=a_id, area_b_id=b_id, seq=100, length_m=geom.d(ln["length_m"]), manual=True,
                           line=[round(v, 6) for p in path[:: max(1, len(path) // 96)] for v in p], source=f"the line on layer {ln['layer']} drawn between them"))  # fmt: skip
    db.flush()
    ms.refresh(db, site.id)
    drawn = {(e.area_a_id, e.area_b_id) for e in ms.edges_of(db, site.id) if e.manual}
    for e in ms.edges_of(db, site.id):
        # asked once: the slope runs into the pool, no wall (the billing team's answer); where the
        # wall is drawn as its own line, the outlines running close elsewhere are not a shared edge
        e.answer = "none" if not e.manual and (e.area_a_id, e.area_b_id) in drawn else cfg["answer"]
    # 3. how each item is measured
    sleeves = {str(areas[c].id): str(n) for c, n in cfg["sleeves_counted_on_site"].items()}
    for it in cfg["items"]:
        w = wl[it["no"]]
        db.add(MeasureItem(site_id=site.id, wo_line_id=w.id, parts=it["parts"], layers=1, manual=sleeves if "count" in it["parts"] else {},
                           skip={str(areas[c].id): ["tops"] for c in cfg["team"]["no_wall_tops"]} if "tops" in it["parts"] else {}))  # fmt: skip
    db.commit()
    # 4. both ways, as study sheets
    team_ids = [areas[s["code"]].id for s in cfg["areas"] if s["way"] == "team"]
    app_ids = [areas[s["code"]].id for s in cfg["areas"] if s["way"] == "app"]
    bb = next(it for it in cfg["items"] if it["no"] == "3")
    team = ms.new_sheet(db, site, date.fromisoformat(cfg["measured_on"]), "The billing team's way (reproduces RA 10)", user.id)
    team.status = "study"
    team.options = {"areas": team_ids, "sizes": "truncate3", "sqft_per_sqm": "10.76", "convert": "item", "periphery": "drawn",
                    "items_as": {str(wl[bb["no"]].id): {"layers_as_no": True, "layers_by_area": {str(areas[c].id): n for c, n in cfg["team"]["layers"].items()},
                                                        "skip": {str(areas[c].id): ["floor"] for c in cfg["team"].get("no_brick_bat", [])}}}}  # fmt: skip
    app = ms.new_sheet(db, site, date.fromisoformat(cfg["measured_on"]), "The app's default way", user.id)
    app.status, app.options = "study", {"areas": app_ids}
    db.commit()
    result = compare(db, cfg, site, areas, team, app, wl, imp)
    db.commit()
    return result


def _per_item(built: ms.Built) -> dict[int, dict]:
    out: dict[int, dict] = {}
    for ln in built.lines:
        w = ln.get("wo_line_id")
        if not w or ln["kind"] == "area":
            continue
        row = out.setdefault(w, {"sqm": D(0), "total": D(0)})
        if ln["kind"] == "total":
            row["total"] = D(ln["total"])
        else:
            row["sqm"] += D(ln["qty"] or 0)
            if ln.get("total") is not None:
                row["total"] += D(ln["total"])
    return out


def _per_part(built: ms.Built, codes: dict[int, str]) -> dict[tuple[str, str, str], D]:
    out: dict[tuple[str, str, str], D] = {}
    for ln in built.lines:
        if ln["kind"] in ("line", "deduction") and ln.get("area_id"):
            k = (ln.get("item_no") or "", codes[ln["area_id"]].removeprefix("T"), ln.get("part") or "")
            out[k] = out.get(k, D(0)) + D(ln["qty"])
    return out


def compare(db: Session, cfg: dict, site, areas: dict, team: MeasureSheet, app: MeasureSheet, wl: dict, imp: dict) -> dict:
    tb, ab = ms.build(db, team), ms.build(db, app)
    t_items, a_items = _per_item(tb), _per_item(ab)
    codes = {a.id: c for c, a in areas.items()}
    theirs = cfg["their_sheet"]
    rows = []
    for no, w in wl.items():
        t, a = t_items.get(w.id, {"sqm": D(0), "total": D(0)}), a_items.get(w.id, {"sqm": D(0), "total": D(0)})
        their = theirs.get(no, {})
        their_total = D(their.get("sqft") or their.get("nos") or 0)
        rate = D(w.rate)
        rows.append({"item": no, "unit": w.unit, "rate": rate, "their_sqm": D(their["sqm"]) if "sqm" in their else None, "their_total": their_total,
                     "team_sqm": t["sqm"] if w.unit == "sqft" else None, "team_total": t["total"], "reproduced": t["total"] == their_total,
                     "app_sqm": a["sqm"] if w.unit == "sqft" else None, "app_total": a["total"],
                     "diff_sqm": (a["sqm"] - t["sqm"]) if w.unit == "sqft" else None, "diff_total": a["total"] - their_total,
                     "diff_rupees": ((a["total"] - their_total) * rate).quantize(D("0.01"))})  # fmt: skip
    parts_t, parts_a = _per_part(tb, codes), _per_part(ab, codes)
    by_part = []
    for k in sorted(set(parts_t) | set(parts_a)):
        w = wl.get(k[0])
        t, a = parts_t.get(k, D(0)), parts_a.get(k, D(0))
        rate = D(w.rate) if w else D(0)
        sq = (a * geom.SQFT_PER_SQM - t * D("10.76")) if w and w.unit == "sqft" else (a - t)
        by_part.append({"item": k[0], "area": k[1], "part": k[2], "team": t, "app": a, "diff_sqm": a - t if w and w.unit == "sqft" else None,
                        "diff_unit": sq, "diff_rupees": (sq * rate).quantize(D("0.01"))})  # fmt: skip
    drawing = []
    for code in ("P", "Q", "R"):
        p, tq = areas[code], areas["T" + code]
        drawing.append({"area": code, "plan_area": D(p.area_sqm), "plan_net": D(p.net_area_sqm), "plan_periphery": D(p.net_periphery_m),
                        "team_area": D(tq.area_sqm), "team_periphery": D(tq.periphery_m), "team_drawn": D(tq.drawn_length_m) if tq.drawn_length_m is not None else None,
                        "gap_mm": D(tq.closed_gap_mm) if tq.closed_gap_mm else None, "height_mm": D(p.height_mm) if p.height_mm is not None else None,
                        "slope_mm": D(p.slope_mm) if p.slope_mm else (D(tq.height_mm) if code == "Q" and tq.height_mm else None)})  # fmt: skip
    edges = [{"between": f"{codes[e.area_a_id]}|{codes[e.area_b_id]}", "length_m": D(e.length_m), "how": e.source or "found along the outlines"} for e in ms.edges_of(db, site.id)]
    return {"site_id": site.id, "team_sheet": team.id, "app_sheet": app.id, "ra": imp["ra_code"], "imported": imp, "items": rows, "parts": by_part, "drawing": drawing,
            "edges": edges, "team_total_rupees": sum((r["team_total"] * r["rate"] for r in rows), D(0)), "app_total_rupees": sum((r["app_total"] * r["rate"] for r in rows), D(0))}  # fmt: skip


def workbook(result: dict) -> bytes:
    """The comparison, for the billing team (their files are never changed)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    bold = Font(bold=True)

    def sheet(title, heads, rows):
        ws = wb.create_sheet(title)
        ws.append(heads)
        for c in ws[1]:
            c.font = bold
        for r in rows:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in r])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(10, min(48, max(len(str(c.value or "")) for c in col) + 2))
        return ws

    wb.remove(wb.active)
    sheet("Items", ["Item", "Unit", "Rate ₹", "Their sheet (sqm)", "Their sheet (unit)", "Billing team's way in the app (sqm)", "Billing team's way (unit)", "Reproduced",
                    "App's default way (sqm)", "App's default way (unit)", "Difference (sqm)", "Difference (unit)", "Difference ₹"],
          [[r["item"], r["unit"], r["rate"], r["their_sqm"], r["their_total"], r["team_sqm"], r["team_total"], "yes" if r["reproduced"] else "NO", r["app_sqm"], r["app_total"],
            r["diff_sqm"], r["diff_total"], r["diff_rupees"]] for r in result["items"]])  # fmt: skip
    sheet("By area and part", ["Item", "Area", "Part", "Billing team's way (sqm or nos)", "App's default way (sqm or nos)", "Difference (sqm)", "Difference (unit)", "Difference ₹"],
          [[p["item"], p["area"], p["part"], p["team"], p["app"], p["diff_sqm"], p["diff_unit"], p["diff_rupees"]] for p in result["parts"]])  # fmt: skip
    sheet("Drawing", ["Area", "Plan outline (sqm)", "Plan, after the trim (sqm)", "Plan periphery (m)", "Billing team's outline (sqm)", "Its periphery, closed (m)",
                      "Its line as drawn (m)", "Gap left open (mm)", "Height (mm)", "Slope line (mm)"],
          [[d["area"], d["plan_area"], d["plan_net"], d["plan_periphery"], d["team_area"], d["team_periphery"], d["team_drawn"], d["gap_mm"], d["height_mm"], d["slope_mm"]] for d in result["drawing"]])  # fmt: skip
    sheet("Shared edges", ["Between", "Length (m)", "How it was found"], [[e["between"], e["length_m"], e["how"]] for e in result["edges"]])
    imp = result["imported"]
    sheet("Imported RA", ["Item", "Read (unit)", "Kept (3 decimals)", "Read (sqm)"], [[r["item_no"], r["read"], r["qty"], r.get("sqm")] for r in imp["items"]] +
          [[], ["Basic as printed", imp["totals"]["read_basic"]], ["Grand total as printed", imp["totals"]["read_grand"]], ["Lines to the paisa", imp["totals"]["lines"]]] +
          [[], ["Not read"]] + [[s] for s in imp["skipped"]] + [[], ["To check"]] + [[s] for s in imp["checks"]])  # fmt: skip
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
