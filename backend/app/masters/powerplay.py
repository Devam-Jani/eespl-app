"""Import Excel exports from EESPL's old Powerplay account: material list, vendor list and team
members.

Powerplay's exports are not documented, so columns are found by header name: each field
accepts the names seen in Powerplay's screens (Material Name, Additional Specification, Vendor,
Vendor Category, Member, Role ...) and common alternatives. The header row may sit below a
title block; the first 15 rows are searched for it.

All three imports are idempotent: rows match existing records by name (case-insensitive,
spaces collapsed; team members by email first), duplicates inside a file are merged, and an
existing record only gets fields it does not have yet, so edits made in the app are kept.
Nobody is ever emailed: imported team members are inactive users without a password.
"""

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pydantic import EmailStr as _EmailStr
from pydantic import TypeAdapter, ValidationError
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app import audit
from app.masters.models import (
    Category,
    Product,
    Vendor,
    VendorBankAccount,
    VendorContact,
)
from app.masters.schemas import (
    validate_account_number,
    validate_gstin,
    validate_ifsc,
    validate_pan,
)
from app.masters.units import load_aliases, normalise_unit
from app.models import User

HEADER_SEARCH_ROWS = 15

MATERIAL_COLUMNS = {
    "name": ["material name", "material", "item name", "name", "material / item"],
    "category": ["category", "material category"],
    "brand": ["brand", "make", "brand name"],
    "unit": ["unit", "uom", "units", "unit of measurement", "unit of measure"],
    "specs": ["additional specification", "additional specifications", "specification",
              "specifications", "specs", "description"],
}  # fmt: skip
VENDOR_COLUMNS = {
    "name": ["vendor name", "vendor", "name", "company name", "firm name"],
    "category": ["vendor category", "category", "type", "vendor type"],
    "gstin": ["gstin", "gst number", "gst no", "gstin/uin", "gst"],
    "pan": ["pan", "pan number", "pan no"],
    "contact": ["contact person", "contact name", "poc", "contact"],
    "phone": ["phone", "mobile", "phone number", "mobile number", "contact number"],
    "email": ["email", "email id", "email address"],
    "address": ["address", "billing address"],
    "city": ["city"],
    "state": ["state"],
    "account_name": ["account holder name", "account name", "beneficiary name"],
    "account_number": ["account number", "bank account number", "account no", "a/c no"],
    "ifsc": ["ifsc", "ifsc code"],
    "bank": ["bank name", "bank"],
    "branch": ["branch", "branch name"],
}  # fmt: skip
TEAM_COLUMNS = {
    "name": ["member", "member name", "name", "full name"],
    "phone": ["phone", "mobile", "phone number", "mobile number", "contact number"],
    "email": ["email", "email id", "email address"],
    "role": ["role", "designation", "job title"],
}

POWERPLAY_STAFF_DOMAIN = "getpowerplay.in"
_email = TypeAdapter(_EmailStr)


class PowerplayFormatError(ValueError):
    pass


@dataclass
class ImportResult:
    kind: str
    sheet: str = ""
    header: list[str] = field(default_factory=list)
    mapping: dict[str, str] = field(default_factory=dict)  # our field -> their column
    rows_in_file: int = 0
    created: int = 0
    already_present: int = 0
    duplicates_merged: int = 0
    skipped: list[tuple[int, str]] = field(default_factory=list)
    warnings: list[tuple[int, str]] = field(default_factory=list)
    categories_created: list[str] = field(default_factory=list)


def name_key(text: str) -> str:
    return " ".join(text.split()).casefold()


def _header_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9/ ]", "", " ".join(str(value or "").lower().split())).strip()


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # phone numbers read as 9033371007.0
    text = " ".join(str(value).split())
    return text or None


def read_sheet(path: str | Path, columns: dict[str, list[str]], result: ImportResult):
    """Find the header row and yield (row number, {field: value}) for every data row."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            for header_index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
                keys = [_header_key(v) for v in row]
                mapping: dict[str, int] = {}
                for name, aliases in columns.items():
                    for alias in aliases:
                        if alias in keys and keys.index(alias) not in mapping.values():
                            mapping[name] = keys.index(alias)
                            break
                if "name" not in mapping:
                    continue
                result.sheet = ws.title
                result.header = [str(v) for v in row if v is not None]
                result.mapping = {f: str(row[i]) for f, i in mapping.items()}
                data = []
                for number, values in enumerate(rows[header_index + 1 :], start=header_index + 2):
                    if not values or all(v is None for v in values):
                        continue
                    data.append(
                        (
                            number,
                            {f: values[i] if i < len(values) else None for f, i in mapping.items()},
                        )  # fmt: skip
                    )
                return data
    finally:
        wb.close()
    raise PowerplayFormatError(
        f"No header row with a name column ({', '.join(columns['name'])}) in the first "
        f"{HEADER_SEARCH_ROWS} rows of any sheet"
    )


# --- materials -> products -------------------------------------------------------------------

_BRAND_IN_SPECS = re.compile(r"\bbrand\s*[:\-]\s*([^,;|\n]+)", re.IGNORECASE)
_CATEGORY_IN_SPECS = re.compile(r"\bcategory\s*[:\-]\s*([^,;|\n]+)", re.IGNORECASE)


# Powerplay spellings of our seeded material categories
CATEGORY_ALIASES = {"others": "other", "misc": "other", "miscellaneous": "other"}


def _material_category(db: Session, name: str, result: ImportResult, cache: dict) -> int:
    key = CATEGORY_ALIASES.get(name_key(name), name_key(name))
    if key not in cache:
        category = db.scalar(
            select(Category).where(Category.kind == "material", func.lower(Category.name) == key)
        )
        if category is None:
            category = Category(kind="material", name=name.strip()[:100])
            db.add(category)
            db.flush()
            result.categories_created.append(category.name)
        cache[key] = category.id
    return cache[key]


def _next_codes(db: Session):
    used = set(db.scalars(select(Product.code).where(Product.code.like("PP-%"))))
    n = 0
    while True:
        n += 1
        code = f"PP-{n:04d}"
        if code not in used:
            used.add(code)
            yield code


def import_materials(db: Session, path: str | Path) -> ImportResult:
    """Materials become products (no price). Brand from a Brand column or "Brand: X" in the
    specification; category from the Category column or "Category: X" in the specification
    (new material categories are created as needed)."""
    result = ImportResult("materials")
    rows = read_sheet(path, MATERIAL_COLUMNS, result)
    aliases = load_aliases(db)
    existing = {name_key(p.name): p for p in db.scalars(select(Product))}
    seen: dict[str, int] = {}
    categories: dict[str, int] = {}
    codes = _next_codes(db)
    for number, r in rows:
        name = _text(r.get("name"))
        if not name:
            result.skipped.append((number, "no material name"))
            continue
        result.rows_in_file += 1
        key = name_key(name)
        if key in seen:
            result.duplicates_merged += 1
            continue
        seen[key] = number
        specs = _text(r.get("specs")) or ""
        brand = _text(r.get("brand"))
        if not brand and (m := _BRAND_IN_SPECS.search(specs)):
            brand = m.group(1).strip()
        category = _text(r.get("category"))
        if not category and (m := _CATEGORY_IN_SPECS.search(specs)):
            category = m.group(1).strip()
        raw_unit = _text(r.get("unit"))
        unit = normalise_unit(raw_unit, aliases) if raw_unit else None

        product = existing.get(key)
        if product is not None:
            result.already_present += 1
            if not product.brand and brand:
                product.brand = brand[:100]
            if product.category_id is None and category:
                product.category_id = _material_category(db, category, result, categories)
            continue
        if unit is None:
            result.skipped.append((number, f"unknown unit {raw_unit!r}" if raw_unit else "no unit"))
            continue
        product = Product(
            code=next(codes),
            name=name[:200],
            brand=brand[:100] if brand else None,
            unit=unit,
            category_id=_material_category(db, category, result, categories) if category else None,
        )
        db.add(product)
        existing[key] = product
        result.created += 1
    db.flush()
    _audit(db, result, path)
    db.commit()
    return result


# --- vendors -----------------------------------------------------------------------------------


def vendor_type(category: str | None) -> str:
    c = (category or "").lower()
    if "material" in c or "supplier" in c:
        return "material_supplier"
    if "labour" in c or "labor" in c or "manpower" in c:
        return "labour_contractor"
    if "sub" in c and "contract" in c:
        return "subcontractor"
    if "transport" in c:
        return "transporter"
    return "material_supplier" if not c else "other"


def _valid(validator, value: str | None, number: int, label: str, result: ImportResult):
    if not value:
        return None
    try:
        return validator(value)
    except ValueError as exc:
        result.warnings.append((number, f"{label} {value!r} ignored: {exc}"))
        return None


def import_vendors(db: Session, path: str | Path) -> ImportResult:
    result = ImportResult("vendors")
    rows = read_sheet(path, VENDOR_COLUMNS, result)
    existing = {name_key(v.name): v for v in db.scalars(select(Vendor))}
    seen: set[str] = set()
    for number, r in rows:
        name = _text(r.get("name"))
        if not name:
            result.skipped.append((number, "no vendor name"))
            continue
        result.rows_in_file += 1
        key = name_key(name)
        if key in seen:
            result.duplicates_merged += 1
            continue
        seen.add(key)

        gstin = _valid(validate_gstin, _text(r.get("gstin")), number, "GSTIN", result)
        pan = _valid(validate_pan, _text(r.get("pan")), number, "PAN", result)
        if gstin and pan and gstin[2:12] != pan:
            result.warnings.append((number, "PAN does not match the GSTIN; PAN ignored"))
            pan = None
        email = _text(r.get("email"))
        if email:
            try:
                email = _email.validate_python(email)
            except ValidationError:
                result.warnings.append((number, f"email {email!r} ignored: not a valid address"))
                email = None
        contact_name = _text(r.get("contact"))
        phone = _text(r.get("phone"))

        vendor = existing.get(key)
        if vendor is None:
            vendor = Vendor(name=name[:200], type=vendor_type(_text(r.get("category"))))
            db.add(vendor)
            existing[key] = vendor
            result.created += 1
        else:
            result.already_present += 1
        fill = {
            "gstin": gstin,
            "pan": pan,
            "address": _text(r.get("address")),
            "city": _text(r.get("city")),
            "state": _text(r.get("state")),
        }
        for field_name, value in fill.items():
            if value and not getattr(vendor, field_name):
                setattr(vendor, field_name, value)
        if (contact_name or phone or email) and not vendor.contacts:
            vendor.contacts.append(
                VendorContact(
                    name=(contact_name or name)[:200], phone=phone, email=email, is_primary=True
                )  # fmt: skip
            )
        number_raw = _text(r.get("account_number"))
        ifsc_raw = _text(r.get("ifsc"))
        if (number_raw or ifsc_raw) and not vendor.bank_accounts:
            account = _valid(validate_account_number, number_raw, number, "account number", result)
            ifsc = _valid(validate_ifsc, ifsc_raw, number, "IFSC", result)
            if account and ifsc:
                vendor.bank_accounts.append(
                    VendorBankAccount(
                        account_name=(_text(r.get("account_name")) or name)[:200],
                        account_number=account,
                        ifsc=ifsc,
                        bank=_text(r.get("bank")),
                        branch=_text(r.get("branch")),
                        is_primary=True,
                    )
                )
            else:
                result.warnings.append((number, "bank details incomplete or invalid; not added"))
    db.flush()
    _audit(db, result, path)
    db.commit()
    return result


# --- team members -> users -------------------------------------------------------------------


def import_team(db: Session, path: str | Path) -> ImportResult:
    """Team members become users with is_active = false and no password; an admin adds the
    email (if missing), sets a password and roles, and activates them. No email is sent."""
    result = ImportResult("team")
    rows = read_sheet(path, TEAM_COLUMNS, result)
    users = list(db.scalars(select(User)))
    by_email = {u.email: u for u in users if u.email}
    by_name = {name_key(u.full_name): u for u in users}
    seen: set[str] = set()
    for number, r in rows:
        name = _text(r.get("name"))
        if not name:
            result.skipped.append((number, "no member name"))
            continue
        email = _text(r.get("email"))
        if email:
            try:
                email = str(_email.validate_python(email)).lower()
            except ValidationError:
                result.warnings.append((number, f"email {email!r} ignored: not a valid address"))
                email = None
        if email and email.endswith("@" + POWERPLAY_STAFF_DOMAIN):
            result.skipped.append((number, "Powerplay's own support account"))
            continue
        result.rows_in_file += 1
        key = email or name_key(name)
        if key in seen:
            result.duplicates_merged += 1
            continue
        seen.add(key)

        user = (by_email.get(email) if email else None) or by_name.get(name_key(name))
        if user is not None:
            result.already_present += 1
            if not user.phone and _text(r.get("phone")):
                user.phone = _text(r.get("phone"))
            if not user.job_title and _text(r.get("role")):
                user.job_title = _text(r.get("role"))[:100]
            if not user.email and email and email not in by_email:
                user.email = email
                by_email[email] = user
            continue
        user = User(
            full_name=name[:200],
            email=email,
            phone=_text(r.get("phone")),
            job_title=(_text(r.get("role")) or "")[:100] or None,
            password_hash=None,
            is_active=False,
        )
        db.add(user)
        if email:
            by_email[email] = user
        by_name[name_key(name)] = user
        result.created += 1
    db.flush()
    _audit(db, result, path)
    db.commit()
    return result


def _audit(db: Session, result: ImportResult, path: str | Path) -> None:
    audit.record(
        db,
        f"powerplay.import_{result.kind}",
        result.kind,
        None,
        after={
            "file": Path(path).name,
            "rows": result.rows_in_file,
            "created": result.created,
            "already_present": result.already_present,
            "duplicates_merged": result.duplicates_merged,
            "skipped": len(result.skipped),
        },
    )
