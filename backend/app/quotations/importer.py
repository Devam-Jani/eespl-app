# ruff: noqa: E501  (HTML, patterns and offer wording read better unwrapped)
"""Import a hand-made techno-commercial offer (.docx) into the quotation libraries, once:

    python -m app.cli import-offer data/samples/offer-ladani.docx

Reads the document in order: the cover letter (-> a letter template with {placeholders}, the
letterhead with its logo and signatory), the technical specification items (-> spec blocks with
stage sections and numbered steps, images with captions), the budgetary offer tables (-> offer
lines with UoM and rate, Opt.1 / Opt.2 rows, "if required", "client's scope"), the general terms
(-> T&C clauses in their groups and a T&C template for the letterhead) and "Our esteemed
clients" (-> references). Every row is marked "imported, check". Known typos are fixed and
listed.
"""

import re
import uuid
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import settings as app_settings
from app.masters.models import TcClause, TcTemplate, TcTemplateClause
from app.quotations import library as lib
from app.quotations.markup import normalize, plain
from app.quotations.models import (
    STAGES,
    Letterhead,
    LetterTemplate,
    OfferItem,
    OfferItemLine,
    OfferItemSpec,
    OfferLine,
    Reference,
    SpecBlock,
)
from app.survey.models import AreaType

W_P, W_TBL = qn("w:p"), qn("w:tbl")
ACRONYMS = {
    "HDPE",
    "PU",
    "RCC",
    "PCC",
    "BRONCO",
    "GSM",
    "PVC",
    "EBA",
    "PMC",
    "OH",
    "UG",
    "CTE",
    "HB",
    "NSPG",
}
UOM_MAP = {
    "SQMT": "sqm",
    "SQM": "sqm",
    "SQFT": "sqft",
    "RFT": "rft",
    "RMT": "rmt",
    "NOS": "nos",
    "NO": "nos",
    "KG": "kg",
    "LTR": "ltr",
    "LS": "ls",
}
# (wrong, right, applies to: None = everywhere, else item numbers)
TYPOS: list[tuple[str, str, tuple[int, ...] | None]] = [
    ("METHEDOLOGY", "METHODOLOGY", None),
    ("HORIZINTAL", "HORIZONTAL", None),
    (" on on ", " on ", None),
    ("Al low", "Allow", None),
    ("CEMSHEILD", "CEMSHIELD", None),
    ("loos particles", "loose particles", None),
    ("Tie rode hole", "Tie rod hole", None),
    ("filing the same", "filling the same", None),
]
ROOF_FIX_WORDS = ("tank", "wall", "pool")  # "over the Roof" there means "over the surface"
AREA_KEYWORDS = [
    ("raft", "Raft"),
    ("retaining", "Basement retaining wall"),
    ("garden", "Podium / planter"),
    ("planter", "Podium / planter"),
    ("pool", "Swimming pool"),
    ("terrace", "Terrace"),
    ("sunken", "Sunken slab"),
]
M = r"(?:\*\*|==)*"  # markup markers that may sit around or inside the brackets
CLIENT_SCOPE = re.compile(
    r"[ \t]*"
    + M
    + r"[ \t]*\([ \t]*"
    + M
    + r"[ \t]*Client[’'`]?s?[ \t]+scope[ \t]*"
    + M
    + r"[ \t]*\)[ \t]*"
    + M
    + r"(?=[ \t]*\.?[ \t]*$)",
    re.I | re.M,
)
S = r"[ \t]*"  # spaces only: a newline is a paragraph break and stays
IF_REQUIRED = re.compile(rf"{S}{M}{S}\({S}{M}{S}If required{S}\)?{S}(:?){S}{M}{S}\)?", re.I)
# the T&C groups this offer uses, as the T&C library's category codes
TC_CODES = {
    "Client's obligations": "client_scope",
    "Force majeure": "force_majeure",
    "Taxes": "taxes",
    "Validity": "validity",
    "Applicator": "applicator",
    "Terms of payment": "payment",
}


@dataclass
class Para:
    text: str  # markup
    plain: str
    numbered: bool
    num_id: str | None
    images: list[tuple[bytes, str]]
    all_bold: bool


@dataclass
class Result:
    counts: dict[str, int] = field(default_factory=dict)
    fixes: dict[str, int] = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)


# --- reading the document ------------------------------------------------------------------------


def _markup(p: Paragraph) -> tuple[str, bool]:
    """Runs -> markup (adjacent runs with the same look merged; blank bold runs dropped)."""
    parts: list[tuple[str, bool, bool]] = []
    for r in p.runs:
        t = r.text
        if not t:
            continue
        bold = bool(r.bold)
        hl = r.font.highlight_color is not None
        if not t.strip():
            bold = hl = False  # a space never carries bold / highlight markers
        if parts and parts[-1][1] == bold and parts[-1][2] == hl:
            parts[-1] = (parts[-1][0] + t, bold, hl)
        else:
            parts.append((t, bold, hl))
    out = []
    for t, bold, hl in parts:
        lead = t[: len(t) - len(t.lstrip())]
        trail = t[len(t.rstrip()) :]
        core = t.strip()
        if hl:
            core = f"=={core}=="
        if bold:
            core = f"**{core}**"
        out.append(lead + core + trail)
    text = re.sub(r"[ \t]+", " ", "".join(out)).strip()
    letters = [(t, b) for t, b, _ in parts if t.strip()]
    return text, bool(letters) and all(b for _, b in letters)


def _images(p: Paragraph, part) -> list[tuple[bytes, str]]:
    found = []
    for blip in p._p.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if rid and rid in part.related_parts:
            img = part.related_parts[rid]
            ext = Path(img.partname).suffix.lstrip(".") or "png"
            found.append((img.blob, ext))
    return found


def read(path: Path):
    d = Document(str(path))
    items: list[Any] = []
    for el in d.element.body:
        if el.tag == W_P:
            p = Paragraph(el, d)
            text, all_bold = _markup(p)
            num = el.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
            num_id = None
            if num is not None and num.find(qn("w:numId")) is not None:
                num_id = num.find(qn("w:numId")).get(qn("w:val"))
            items.append(
                Para(text, plain(text), num is not None, num_id, _images(p, d.part), all_bold)
            )
        elif el.tag == W_TBL:
            rows = []
            for tr in el.findall(qn("w:tr")):
                cells = []
                for tc in tr.findall(qn("w:tc")):
                    paras = [_markup(Paragraph(pe, d))[0] for pe in tc.findall(qn("w:p"))]
                    cells.append([x for x in paras if plain(x)])
                rows.append(cells)
            items.append(rows)
    header_images = []
    for sec in d.sections:
        for p in sec.header.paragraphs:
            header_images += _images(p, sec.header.part)
    return items, header_images


# --- helpers -------------------------------------------------------------------------------------


def sentence(text: str) -> str:
    """ "RAFT TREATMENT WATERPROOFING USING HDPE" -> "Raft treatment waterproofing using HDPE"."""
    words = []
    for w in text.split():
        core = re.sub(r"[^A-Za-z]", "", w)
        words.append(w if core.upper() in ACRONYMS or any(c.isdigit() for c in w) else w.lower())
    s = " ".join(words)
    return s[:1].upper() + s[1:]


class Fixer:
    def __init__(self, result: Result):
        self.result = result

    def __call__(self, text: str, item_no: int | None = None, item_name: str = "") -> str:
        for wrong, right, only in TYPOS:
            if only is not None and item_no not in only:
                continue
            n = text.count(wrong)
            if n:
                text = text.replace(wrong, right)
                self.result.fixes[f"{wrong.strip()} -> {right.strip()}"] = (
                    self.result.fixes.get(f"{wrong.strip()} -> {right.strip()}", 0) + n
                )
        if any(w in item_name.lower() for w in ROOF_FIX_WORDS) and "over the Roof" in text:
            n = text.count("over the Roof")
            text = text.replace("over the Roof", "over the surface")
            key = "over the Roof -> over the surface (tanks, walls, pools)"
            self.result.fixes[key] = self.result.fixes.get(key, 0) + n
        return text


def _save_image(blob: bytes, ext: str) -> str:
    rel = f"quotations/library/{uuid.uuid4().hex}.{ext.lower()}"
    p = Path(app_settings.media_dir) / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(blob)
    return rel


def _stage(heading: str) -> str:
    h = re.sub(r"[:\s]+$", "", heading).strip()
    for st in STAGES:
        if h.lower().startswith(st.lower()):
            return st if h.lower() == st.lower() else sentence(h)
    return sentence(h)


def _decimal(text: str) -> Decimal | None:
    try:
        return Decimal(re.sub(r"[^\d.]", "", text)) if re.search(r"\d", text) else None
    except InvalidOperation:
        return None


def _flags(text: str) -> tuple[str, bool, bool]:
    """Take "(Client's scope)" (at the end of a step) and "(If required)" out of the text and
    return them as flags; the renderer prints them highlighted again. Markers around or inside
    the brackets go with them; what is left is re-balanced."""
    client = bool(CLIENT_SCOPE.search(text))
    text = CLIENT_SCOPE.sub("", text)
    req = bool(re.search(r"\(\s*(?:\*\*|==)*\s*If required", text, re.I))
    if req:
        text = IF_REQUIRED.sub(r"\1", text)
    return normalize(text).strip(), client, req


SPEC_RE = re.compile(
    r"^(?:OPTION\s*:?\s*(\d+)\s*[-–]?\s*)?ITEM NO[:.]?\s*(\d+)\s*[-–]?\s*(TECHNICAL SPECIFICATION(?:\s*&\s*METH\w+)?)\s+FOR\s+(.*?)[:\s]*$",
    re.I,
)
BUDGET_RE = re.compile(r"^ITEM NO[:.]?\s*(\d+)\s*[-–]?\s*BUDGETARY OFFER FOR\s+(.*?)[:\s]*$", re.I)
OPTION_HEAD = re.compile(r"^OPTION\s*:?\s*(\d+)\s*[-–]?\s*(.*)$", re.I)


def _heading_like(p: Para) -> bool:
    t = p.plain
    if not t or p.numbered or len(t) > 90 or t.lower().startswith("image"):
        return False
    letters = re.sub(r"[^A-Za-z]", "", t)
    return p.all_bold or (letters.isupper() and len(letters) > 3)


# --- the import ----------------------------------------------------------------------------------


def import_offer(db: Session, path: Path, user_id=None, again: bool = False) -> Result:
    res = Result()
    fix = Fixer(res)
    items, header_images = read(path)
    i = 0

    def next_para():
        nonlocal i
        while (
            i < len(items)
            and isinstance(items[i], Para)
            and not items[i].plain
            and not items[i].images
        ):
            i += 1

    # ---- the cover letter
    letter: list[Para] = []
    while i < len(items) and not (isinstance(items[i], Para) and SPEC_RE.match(items[i].plain)):
        if isinstance(items[i], Para) and items[i].plain:
            letter.append(items[i])
        i += 1
    get = {
        "firm": next((p for p in letter if p.plain.lower().startswith("firm name")), None),
        "attn": next((p for p in letter if p.plain.lower().startswith("kind attn")), None),
        "subject": next((p for p in letter if p.plain.upper().startswith("SUB")), None),
        "for": next((p for p in letter if p.plain.startswith("For ")), None),
    }
    firm = re.sub(r"^firm name\s*:\s*", "", get["firm"].plain, flags=re.I) if get["firm"] else ""
    sig_firm = get["for"].plain if get["for"] else "For EESPL"
    company = sig_firm[4:].strip()
    head_name = company.split()[0].title() if company else "Imported"
    if db.scalar(select(Letterhead).where(Letterhead.name == head_name)) and not again:
        raise ValueError(
            f"Already imported: the letterhead {head_name} exists (use --again to import another copy)"
        )
    if again:
        head_name = f"{head_name} {uuid.uuid4().hex[:4]}"
    idx_for = letter.index(get["for"]) if get["for"] else len(letter)
    sign = letter[idx_for + 1].plain if idx_for + 1 < len(letter) else ""
    sig_name, _, sig_desig = (x.strip() for x in sign.partition("|"))
    idx_sub = letter.index(get["subject"]) if get["subject"] else 0
    body_paras = [p.text for p in letter[idx_sub + 1 : idx_for]]
    brand_m = re.search(r"using\s+\*\*([A-Z][A-Z ]+?)\*\*", " ".join(body_paras), re.I)
    brand = brand_m.group(1).strip() if brand_m else None
    body = "\n".join(body_paras)
    if brand:
        body = body.replace(brand, "{brand}")
    project_words = [
        w
        for w in re.findall(r"\*\*([^*]+)\*\*", body)
        if firm and w.strip().lower() == firm.lower()
    ]
    for w in project_words:
        body = body.replace(f"**{w}**", "**{project}**")
    subject = (
        get["subject"].text if get["subject"] else "SUB: OFFER FOR {areas_list} WATERPROOFING WORK."
    )
    subject = re.sub(r"==[^=]+==", "{areas_list}", subject, count=1).replace("**", "")
    enc_start = next((k for k, p in enumerate(letter) if p.plain.lower().startswith("encl")), None)
    enclosures = (
        "\n".join(p.plain.rstrip(".") for p in letter[enc_start + 1 :])
        if enc_start is not None
        else ""
    )
    lt = LetterTemplate(
        name=f"{head_name} offer letter (imported)",
        opening="Date: **{date}**\nFirm Name: **{client_firm}**\n**{client_city}**\nKind Attn.: **{attention}**\n\nDear Sir,",
        subject=subject,
        body=body,
        enclosures=enclosures,
        needs_check=True,
        created_by=user_id,
    )
    lib.check_letter(lt.subject, lt.body, lt.opening)
    db.add(lt)
    db.flush()
    lib.save_version(db, lt, "import", user_id, note=f"imported from {path.name}")
    logo = _save_image(*header_images[0]) if header_images else None
    head = Letterhead(
        name=head_name,
        company_name=company or head_name,
        logo_path=logo,
        signatory_firm=sig_firm,
        signatory_name=sig_name or None,
        signatory_designation=sig_desig or None,
        brand=brand,
        letter_template_id=lt.id,
        needs_check=True,
        created_by=user_id,
    )
    db.add(head)
    db.flush()
    res.counts["letterheads"] = 1
    res.counts["letter_templates"] = 1

    # ---- technical specifications
    blocks: dict[int, list[SpecBlock]] = {}
    names: dict[int, str] = {}
    current: SpecBlock | None = None
    item_no = 0
    section: dict | None = None
    pending_images: list[str] = []
    while i < len(items):
        it = items[i]
        if isinstance(it, Para) and BUDGET_RE.match(it.plain):
            break
        i += 1
        if not isinstance(it, Para):
            continue
        if it.images and current is not None:
            pending_images += [_save_image(b, e) for b, e in it.images]
            continue
        if not it.plain or re.fullmatch(r"OR", it.plain.strip(), re.I):
            continue
        m = SPEC_RE.match(it.plain)
        if m:
            opt, n, prefix, title = m.group(1), int(m.group(2)), m.group(3), m.group(4)
            item_no = n
            title_clean = re.sub(r"\s+", " ", title.replace("–", "-")).strip()
            if n not in names:
                names[n] = sentence(
                    re.split(r"\s+WATERPROOFING", title_clean, flags=re.I)[0].strip(" -")
                )
            current = SpecBlock(
                title=fix(sentence(title_clean), n, names[n]),
                heading_prefix=fix(re.sub(r"\s+", " ", prefix.upper()), n) + " FOR",
                option_label=opt,
                sections=[],
                images=[],
                needs_check=True,
                created_by=user_id,
            )
            blocks.setdefault(n, []).append(current)
            section = None
            continue
        if current is None:
            continue
        if pending_images and it.plain.lower().startswith("image"):
            current.images = current.images + [
                {"path": p, "caption": it.plain} for p in pending_images
            ]
            pending_images = []
            continue
        if _heading_like(it):
            hm = OPTION_HEAD.match(it.plain)
            option, heading = (hm.group(1), hm.group(2)) if hm else (None, it.plain)
            heading = fix(heading, item_no, names.get(item_no, ""))
            section = {"heading": _stage(heading), "option": option, "steps": []}
            current.sections = current.sections + [section]
            continue
        text, client, req = _flags(fix(it.text, item_no, names.get(item_no, "")))
        if section is None:
            section = {"heading": "", "option": None, "steps": []}
            current.sections = current.sections + [section]
        if it.numbered or not section["steps"]:
            section["steps"].append({"text": text, "client_scope": client, "if_required": req})
        else:  # a paragraph that continues the step above
            last = section["steps"][-1]
            last["text"] += "\n" + text
            last["client_scope"] = last["client_scope"] or client
            last["if_required"] = last["if_required"] or req
        current.sections = list(current.sections)  # mark the JSON as changed
    if pending_images and current is not None:
        current.images = current.images + [{"path": p, "caption": ""} for p in pending_images]

    # ---- budgetary offer
    budget: dict[int, tuple[str, list[dict]]] = {}
    while i < len(items):
        it = items[i]
        if isinstance(it, Para) and "TERMS AND CONDITIONS" in it.plain.upper():
            break
        i += 1
        if not isinstance(it, Para):
            continue
        m = BUDGET_RE.match(it.plain)
        if not m:
            continue
        n, title = int(m.group(1)), m.group(2)
        while i < len(items) and not isinstance(items[i], list):
            i += 1
        table = items[i] if i < len(items) else []
        i += 1
        budget[n] = (
            sentence(re.sub(r"\s+", " ", title).strip(" -")),
            _lines(table[1:], fix, n, names.get(n, ""), res),
        )

    # ---- terms and conditions
    clauses: list[tuple[str, str]] = []
    category = "Client's obligations"
    list_id: str | None = None
    while i < len(items):
        it = items[i]
        if isinstance(it, Para) and "esteemed clients" in it.plain.lower():
            break
        i += 1
        if not isinstance(it, Para) or not it.plain:
            continue
        up = it.plain.upper()
        if "GENERAL TERMS AND CONDITIONS" in up:
            continue
        if not it.numbered and "OBLIGATION" in up and len(up) < 40:
            category = "Client's obligations"
            continue
        if it.numbered and (list_id is None or it.num_id == list_id):
            list_id = list_id or it.num_id
            text = fix(it.text)
            cat = category
            if up.startswith("FORCE MAJEURE"):
                cat = "Force majeure"
            elif up.startswith("TAXES"):
                cat = "Taxes"
            elif up.startswith("VALIDITY"):
                cat = "Validity"
            elif "AUTHORIZED APPLICATOR" in up and "WORK ORDER" in up:
                cat = "Applicator"
                text = re.sub(
                    r"ETHIOS ENVIRO SOLUTIONS PVT\.? LTD\.?", "{applicator}", text, flags=re.I
                )
            elif up.startswith("TERMS OF PAYMENT"):
                cat = "Terms of payment"
            clauses.append((cat, text))
            continue
        if clauses:  # the text under TAXES: / VALIDITY:, or a payment sub-point
            cat, text = clauses[-1]
            line = ("- " if it.numbered else "") + fix(it.text)
            clauses[-1] = (cat, text + "\n" + line)

    # ---- references
    state, title = None, None
    refs: list[Reference] = []
    while i < len(items):
        it = items[i]
        i += 1
        if isinstance(it, Para) and "esteemed clients" in it.plain.lower():
            title = it.plain
            sm = re.search(r"\bin\s+([A-Za-z][A-Za-z ]+)$", it.plain)
            state = sm.group(1).strip() if sm else None
            continue
        if isinstance(it, list) and title:
            for order, row in enumerate(it[1:], 1):
                cells = [" ".join(plain(x) for x in c) for c in row]
                if len(cells) < 5 or not any(cells[1:]):
                    continue
                area = cells[4]
                value = _decimal(area.replace(",", ""))
                unit = re.sub(r"[\d,.\s]", "", area).upper() or None
                refs.append(
                    Reference(
                        client_name=cells[1],
                        project=cells[2],
                        application=" / ".join(plain(x) for x in row[3]),
                        area_value=value,
                        area_unit=unit,
                        state=state,
                        source="imported",
                        sort_order=order,
                        needs_check=True,
                        created_by=user_id,
                    )
                )
            title = None

    # ---- write the libraries
    types = {t.name: t.id for t in db.scalars(select(AreaType))}

    def area_type(name: str) -> int | None:
        for k, t in AREA_KEYWORDS:
            if k in name.lower():
                return types.get(t)
        return None

    n_specs = n_lines = n_options = 0
    for n in sorted(set(blocks) | set(budget)):
        name = names.get(n) or (budget[n][0] if n in budget else f"Item {n}")
        at = area_type(name)
        item = OfferItem(
            name=name,
            area_type_id=at,
            budget_title=budget[n][0] if n in budget else name,
            sort_order=n * 10,
            needs_check=True,
            created_by=user_id,
        )
        options = set()
        for k, b in enumerate(blocks.get(n, [])):
            b.area_type_id = at
            db.add(b)
            db.flush()
            lib.save_version(db, b, "import", user_id, note=f"imported from {path.name}")
            item.specs.append(OfferItemSpec(spec_block_id=b.id, sort_order=k))
            n_specs += 1
            options |= {b.option_label} if b.option_label else set()
            options |= {s["option"] for s in b.sections if s.get("option")}
        for k, ln in enumerate(budget.get(n, ("", []))[1]):
            ol = OfferLine(
                description=ln["description"],
                uom=ln["uom"],
                rate_source="fixed",
                default_rate=ln["rate"],
                if_required=ln["if_required"],
                client_scope=ln["client_scope"],
                needs_check=True,
                created_by=user_id,
            )
            db.add(ol)
            db.flush()
            lib.save_version(db, ol, "import", user_id, note=f"imported from {path.name}")
            item.lines.append(OfferItemLine(offer_line_id=ol.id, option=ln["option"], sort_order=k))
            n_lines += 1
            options |= {ln["option"]} if ln["option"] else set()
        db.add(item)
        db.flush()
        lib.save_version(db, item, "import", user_id, note=f"imported from {path.name}")
        n_options += len(options)
        if at is None:
            res.notes.append(f"Item {n} “{name}”: no area type matched, pick one in the library")
    res.counts.update(
        items=len(set(blocks) | set(budget)),
        spec_blocks=n_specs,
        options=n_options,
        offer_lines=n_lines,
    )

    # T&C: reuse existing clauses with the same text, else add them (marked for review)
    tpl = TcTemplate(name=f"{head_name} techno-commercial offer (imported)", created_by=user_id)
    db.add(tpl)
    db.flush()
    for order, (cat, text) in enumerate(clauses):
        c = db.scalar(select(TcClause).where(TcClause.text == text))
        if c is None:
            c = TcClause(
                text=text,
                category=TC_CODES.get(cat, "general"),
                needs_review=True,
                created_by=user_id,
            )
            db.add(c)
            db.flush()
        db.add(TcTemplateClause(template_id=tpl.id, clause_id=c.id, sort_order=order))
    head.tc_template_id = tpl.id
    head.references_state = state
    lib.save_version(db, head, "import", user_id, note=f"imported from {path.name}")
    res.counts["tc_clauses"] = len(clauses)
    for r in refs:
        db.add(r)
        db.flush()
        lib.save_version(db, r, "import", user_id, note=f"imported from {path.name}")
    res.counts["references"] = len(refs)
    return res


def _lines(rows: list[list[list[str]]], fix: Fixer, n: int, name: str, res: Result) -> list[dict]:
    """Budget table rows -> offer lines. A row with no Sr, UoM and rate continues the row above;
    a row whose UoM cell holds two UoMs for "Option 1 ... OR ... Option 2" is split into Opt.1 /
    Opt.2 rows (the duplicated "SQ FTSQ FT")."""
    out: list[dict] = []
    for cells in rows:
        if len(cells) < 3:
            continue
        sr = plain(" ".join(cells[0]))
        desc_paras = [fix(x, n, name) for x in cells[1]]
        uom_cell = [plain(x) for x in cells[2]]
        rate_cell = [plain(x) for x in cells[3]] if len(cells) > 3 else []
        if not sr and not uom_cell and not rate_cell and out:
            out[-1]["description"] += "\n" + "\n".join(desc_paras)
            continue
        client = any("client" in u.lower() for u in uom_cell + rate_cell)
        opt_m = re.match(r"Opt\.?\s*(\d+)", sr, re.I)
        option = opt_m.group(1) if opt_m else None
        if len(uom_cell) > 1 and len(set(uom_cell)) == 1:
            key = f"duplicated UoM {' / '.join(uom_cell)} -> one UoM, split into Opt. rows"
            res.fixes[key] = res.fixes.get(key, 0) + 1
            starts = [
                k for k, p in enumerate(desc_paras) if re.match(r"Option\s*:?\s*\d", plain(p), re.I)
            ]
            if len(starts) >= 2 and len(rate_cell) >= 2:
                common = desc_paras[: starts[0]]
                for j, s in enumerate(starts[:2]):
                    end = starts[j + 1] if j + 1 < len(starts) else len(desc_paras)
                    part = [p for p in desc_paras[s:end] if plain(p).upper() != "OR"]
                    part[0] = re.sub(
                        r"^(\*\*)?(==)?\s*Option\s*:?\s*\d+\s*[-–]?\s*(==)?(\*\*)?\s*",
                        "",
                        part[0],
                        flags=re.I,
                    )
                    text, c2, req = _flags("\n".join(common + part))
                    out.append(
                        {
                            "description": text,
                            "uom": _uom(uom_cell[0]),
                            "rate": _decimal(rate_cell[j]),
                            "option": str(j + 1),
                            "if_required": req,
                            "client_scope": c2,
                        }
                    )
                continue
            uom_cell = uom_cell[:1]
        text, c2, req = _flags("\n".join(desc_paras))
        out.append(
            {
                "description": text,
                "uom": _uom(uom_cell[0]) if uom_cell and not client else "sqft",
                "rate": None if client else _decimal(rate_cell[0] if rate_cell else ""),
                "option": option,
                "if_required": req,
                "client_scope": client or c2,
            }
        )
    return out


def _uom(text: str) -> str:
    key = re.sub(r"[^A-Z]", "", text.upper())
    return UOM_MAP.get(key, "sqft")
