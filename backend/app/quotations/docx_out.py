"""The quotation as an editable Word file (python-docx), from a template the team can restyle.

The template (app/quotations/templates/offer-template.docx, or the one uploaded in Settings ›
Quotation library) carries the named styles below; restyle them in Word (fonts, sizes, colours,
spacing) and upload it again. The page rules: item headings and the "OR" keep with the next
paragraph, table header rows repeat on every page, rows never split across pages.
"""

import io
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.config import settings as app_settings
from app.quotations.markup import runs
from app.quotations.render import Doc

TEMPLATE = Path(__file__).parent / "templates" / "offer-template.docx"
STYLES = {
    # name: (size pt, bold, keep with next, space before pt, alignment)
    "Offer Body": (10, False, False, 0, None),
    "Offer Subject": (10, True, False, 8, None),
    "Offer Item Heading": (10.5, True, True, 12, None),
    "Offer Stage": (10, True, True, 6, None),
    "Offer Step": (10, False, False, 0, None),
    "Offer OR": (10, True, True, 6, WD_ALIGN_PARAGRAPH.CENTER),
    "Offer Caption": (9.5, True, False, 2, WD_ALIGN_PARAGRAPH.CENTER),
    "Offer Table Text": (9.5, False, False, 0, None),
    "Offer Clause": (10, False, False, 0, None),
}


def make_template(path: Path = TEMPLATE) -> Path:
    """Write the default template with the named styles (python -m app.cli make-offer-template)."""
    d = Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(3.2), Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.8)
    normal = d.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    for name, (size, bold, keep, before, align) in STYLES.items():
        st = d.styles.add_style(name, 1)  # paragraph style
        st.base_style = normal
        st.font.size = Pt(size)
        st.font.bold = bold
        pf = st.paragraph_format
        pf.keep_with_next = keep
        pf.space_before = Pt(before)
        pf.space_after = Pt(3)
        if align is not None:
            pf.alignment = align
        if name in ("Offer Step", "Offer Clause"):
            pf.left_indent = Cm(0.9)
            pf.first_line_indent = Cm(-0.6)
            pf.widow_control = True
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))
    return path


def _template(custom: str | None) -> Document:
    """The team's uploaded template (a media path) when there is one, else the built-in one."""
    if custom:
        p = Path(app_settings.media_dir) / custom
        if p.exists():
            return Document(str(p))
    if not TEMPLATE.exists():
        make_template()
    return Document(str(TEMPLATE))


def _style(d, name: str) -> str:
    try:
        d.styles[name]
        return name
    except KeyError:
        return "Normal"


def _runs(p, text: str, bold: bool = False) -> None:
    for r in runs(text, bold):
        run = p.add_run(r.text)
        run.bold = r.bold or None
        if r.highlight:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _para(d_or_cell, text: str, style: str, bold: bool = False, keep: bool | None = None):
    p = d_or_cell.add_paragraph(style=style)
    _runs(p, text, bold)
    if keep is not None:
        p.paragraph_format.keep_with_next = keep
    return p


def _field(p, code: str) -> None:
    """A Word field (PAGE, NUMPAGES) in a paragraph."""
    for kind, text in (
        ("begin", None),
        (None, code),
        ("separate", None),
        (None, "1"),
        ("end", None),
    ):
        r = p.add_run()
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == "1":
            el = OxmlElement("w:t")
            el.text = "1"
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {text} "
        r._r.append(el)


def _row_rules(row, header: bool = False) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))
    if header:
        tr_pr.append(OxmlElement("w:tblHeader"))


def _borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "666666")
        borders.append(el)
    tbl_pr.append(borders)


def _shade(cell, fill: str = "EEF3F1") -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _cell(cell, text: str, style: str, bold: bool = False, align=None) -> None:
    cell.text = ""
    paras = text.split("\n") or [""]
    first = cell.paragraphs[0]
    first.style = style
    _runs(first, paras[0], bold)
    if align is not None:
        first.alignment = align
    for line in paras[1:]:
        p = cell.add_paragraph(style=style)
        _runs(p, line, bold)
        if align is not None:
            p.alignment = align


def _band_mm(path) -> float:
    from PIL import Image  # noqa: PLC0415

    with Image.open(path) as im:
        w, h = im.size
    return 210 * h / w if w else 0


def _full_width(p, sec) -> None:
    """A header / footer paragraph that runs edge to edge (the band images)."""
    p.paragraph_format.left_indent = -sec.left_margin
    p.paragraph_format.right_indent = -sec.right_margin
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)


def _watermark(header, path) -> None:
    """The faint logo behind the text on every page (a VML picture, as Word makes it)."""
    from docx.oxml import parse_xml  # noqa: PLC0415
    from PIL import Image  # noqa: PLC0415

    rid, _img = header.part.get_or_add_image(str(path))
    with Image.open(path) as im:
        w, h = im.size
    width = 425.0
    height = width * h / w if w else width
    xml = (
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<w:r><w:pict><v:shape id="EESPLWatermark" o:spid="_x0000_s2049" type="#_x0000_t75" '
        f'style="position:absolute;margin-left:0;margin-top:0;width:{width:.0f}pt;height:{height:.0f}pt;'
        "z-index:-251656192;mso-position-horizontal:center;mso-position-horizontal-relative:margin;"
        'mso-position-vertical:center;mso-position-vertical-relative:margin" o:allowincell="f">'
        f'<v:imagedata r:id="{rid}" o:title="watermark"/></v:shape></w:pict></w:r></w:p>'
    )
    header._element.append(parse_xml(xml))


def _image_header_footer(d, doc: Doc) -> None:
    """An image letterhead: the header band at the top edge, the footer band at the bottom edge,
    the watermark behind the text. On pre-printed paper only the space is kept."""
    sec = d.sections[0]
    top, bottom = (
        _band_mm(doc.header_image) if doc.header_image else 25,
        _band_mm(doc.footer_image) if doc.footer_image else 15,
    )
    sec.header_distance = sec.footer_distance = Cm(0)
    sec.top_margin = Cm((top + 6) / 10)
    sec.bottom_margin = Cm((bottom + 12) / 10)
    header, footer = sec.header, sec.footer
    header.is_linked_to_previous = footer.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.text = ""
    _full_width(hp, sec)
    if not doc.preprinted:
        if doc.header_image:
            hp.add_run().add_picture(str(doc.header_image), width=Cm(21.0))
        if doc.watermark:
            _watermark(header, doc.watermark)
    fp = footer.paragraphs[0]
    fp.text = ""
    for text in (
        f"{doc.footer}    ",
        (doc.footer_text.replace("\n", " ") + "    ") if doc.footer_text else "",
    ):
        if text:
            fp.add_run(text).font.size = Pt(8)
    fp.add_run("Page ").font.size = Pt(8)
    _field(fp, "PAGE")
    fp.add_run(" of ").font.size = Pt(8)
    _field(fp, "NUMPAGES")
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if doc.footer_image and not doc.preprinted:
        bp = footer.add_paragraph()
        _full_width(bp, sec)
        bp.add_run().add_picture(str(doc.footer_image), width=Cm(21.0))


def _header_footer(d, doc: Doc) -> None:
    if doc.header_image or doc.footer_image:
        _image_header_footer(d, doc)
        return
    sec = d.sections[0]
    header = sec.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.text = ""
    if doc.preprinted:  # letterhead stationery: the header space stays blank
        pass
    elif doc.logo:
        hp.add_run().add_picture(str(doc.logo), height=Cm(1.8))
        hp.add_run("   ")
    if not doc.preprinted:
        name = hp.add_run(doc.company_name)
        name.bold = True
        name.font.size = Pt(12)
        name.font.color.rgb = RGBColor.from_string(doc.primary.lstrip("#").upper()[:6] or "0F6E5A")
    if doc.header_text and not doc.preprinted:
        hp2 = header.add_paragraph()
        r = hp2.add_run(doc.header_text)
        r.font.size = Pt(8)
    footer = sec.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.text = ""
    r = fp.add_run(f"{doc.footer}    ")
    r.font.size = Pt(8)
    if doc.footer_text:
        r2 = fp.add_run(doc.footer_text.replace("\n", " ") + "    ")
        r2.font.size = Pt(8)
    fp.add_run("Page ").font.size = Pt(8)
    _field(fp, "PAGE")
    fp.add_run(" of ").font.size = Pt(8)
    _field(fp, "NUMPAGES")
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _clear_body(d) -> None:
    body = d.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)


def docx(doc: Doc, template: str | None = None) -> bytes:
    d = _template(template)
    _clear_body(d)
    s = {k: _style(d, k) for k in STYLES}
    _header_footer(d, doc)
    # the cover letter
    for line in doc.opening:
        _para(d, line, s["Offer Body"])
    _para(d, doc.subject, s["Offer Subject"])
    for line in doc.body:
        _para(d, line, s["Offer Body"])
    _para(d, "", s["Offer Body"])
    _para(d, doc.signatory_firm, s["Offer Body"], bold=True, keep=True)
    _para(d, doc.signatory_line, s["Offer Body"], bold=True)
    if doc.enclosures:
        _para(d, "Encl.:", s["Offer Body"], keep=True)
        for i, e in enumerate(doc.enclosures, 1):
            _para(d, f"{i}.\t{e}", s["Offer Step"])
    # technical specifications
    first = True
    for _n, specs in doc.specs:
        for spec in specs:
            if spec.or_before:
                _para(d, "==OR==", s["Offer OR"], keep=True)
            h = _para(d, spec.heading, s["Offer Item Heading"], bold=True, keep=True)
            if first:
                h.paragraph_format.page_break_before = True
                first = False
            if spec.subtitle:
                _para(d, spec.subtitle, s["Offer Stage"], bold=True, keep=True)
            for sec in spec.sections:
                if sec.or_before:
                    _para(d, "==OR==", s["Offer OR"], keep=True)
                if sec.heading:
                    _para(d, sec.heading + ":", s["Offer Stage"], bold=True, keep=True)
                for st in sec.steps:
                    paras = st.text.split("\n")
                    _para(d, f"{st.number}.\t{paras[0]}", s["Offer Step"])
                    for extra in paras[1:]:
                        _para(d, f"\t{extra}", s["Offer Step"])
            for im in spec.images:
                p = d.add_paragraph(style=s["Offer Caption"])
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(im["file"]), width=Cm(11))
                _para(d, im.get("caption") or "", s["Offer Caption"])
    # budgetary offer
    amounts = doc.show_amounts
    first = True
    for b in doc.budgets:
        h = _para(d, b.heading, s["Offer Item Heading"], bold=True, keep=True)
        if first:
            h.paragraph_format.page_break_before = True
            first = False
        heads = ["Sr No.", "Item Description", "UoM"] + (
            ["Qty", "Rate", "Amount"] if amounts else ["Rate"]
        )
        t = d.add_table(rows=1, cols=len(heads))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t)
        widths = [Cm(1.4), Cm(11.2 if not amounts else 8.6), Cm(1.8)] + (
            [Cm(1.6), Cm(1.8), Cm(2.2)] if amounts else [Cm(2.0)]
        )
        for i, text in enumerate(heads):
            c = t.rows[0].cells[i]
            _cell(
                c,
                text,
                s["Offer Table Text"],
                bold=True,
                align=WD_ALIGN_PARAGRAPH.RIGHT if i >= 3 else None,
            )
            _shade(c)
        _row_rules(t.rows[0], header=True)
        # keep the heading with the first row: the header row keeps with the next row too
        for p in t.rows[0].cells[1].paragraphs:
            p.paragraph_format.keep_with_next = True
        for r in b.rows:
            row = t.add_row()
            _row_rules(row)
            cells = row.cells
            _cell(cells[0], r.sr, s["Offer Table Text"])
            _cell(cells[1], r.description, s["Offer Table Text"])
            if r.client_scope:
                merged = cells[2].merge(cells[-1])
                _cell(
                    merged,
                    r.rate,
                    s["Offer Table Text"],
                    bold=True,
                    align=WD_ALIGN_PARAGRAPH.CENTER,
                )
                continue
            _cell(cells[2], r.uom, s["Offer Table Text"])
            if amounts:
                _cell(cells[3], r.qty, s["Offer Table Text"], align=WD_ALIGN_PARAGRAPH.RIGHT)
                _cell(cells[4], r.rate, s["Offer Table Text"], align=WD_ALIGN_PARAGRAPH.RIGHT)
                _cell(cells[5], r.amount, s["Offer Table Text"], align=WD_ALIGN_PARAGRAPH.RIGHT)
            else:
                _cell(cells[3], r.rate, s["Offer Table Text"], align=WD_ALIGN_PARAGRAPH.RIGHT)
        if amounts and b.subtotal:
            row = t.add_row()
            _row_rules(row)
            label = row.cells[0].merge(row.cells[4])
            _cell(label, "Item total", s["Offer Table Text"], bold=True)
            _cell(
                row.cells[5],
                b.subtotal,
                s["Offer Table Text"],
                bold=True,
                align=WD_ALIGN_PARAGRAPH.RIGHT,
            )
        for row in t.rows:
            for i, w in enumerate(widths):
                if i < len(row.cells):
                    row.cells[i].width = w
        _para(d, "", s["Offer Body"])
    if amounts and doc.total:
        p = _para(d, f"Total (excluding GST): ₹ {doc.total}", s["Offer Subject"], bold=True)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2 = _para(d, doc.total_note, s["Offer Body"])
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    # terms and conditions
    if doc.terms:
        _para(d, "GENERAL TERMS AND CONDITIONS", s["Offer Item Heading"], bold=True, keep=True)
        for g in doc.terms:
            if g.heading:
                _para(d, g.heading, s["Offer Stage"], bold=True, keep=True)
            for n, text in g.clauses:
                letter = 0
                for i, line in enumerate(text.split("\n")):
                    if line.startswith("- "):
                        letter += 1
                        p = _para(d, f"{chr(96 + letter)})\t{line[2:]}", s["Offer Clause"])
                        p.paragraph_format.left_indent = Cm(1.6)
                    else:
                        _para(d, f"{n}.\t{line}" if i == 0 else f"\t{line}", s["Offer Clause"])
    # references
    if doc.references:
        _para(d, doc.references_title, s["Offer Item Heading"], bold=True, keep=True)
        heads = ["Sr. No.", "Client Name", "Project", "Application Method & Area", "Area Covered"]
        t = d.add_table(rows=1, cols=5)
        _borders(t)
        for i, text in enumerate(heads):
            _cell(t.rows[0].cells[i], text, s["Offer Table Text"], bold=True)
            _shade(t.rows[0].cells[i])
        _row_rules(t.rows[0], header=True)
        for i, r in enumerate(doc.references, 1):
            row = t.add_row()
            _row_rules(row)
            for j, text in enumerate(
                [
                    str(i),
                    r.get("client_name") or "",
                    r.get("project") or "",
                    r.get("application") or "",
                    r["area_text"],
                ]
            ):
                _cell(
                    row.cells[j],
                    text,
                    s["Offer Table Text"],
                    align=WD_ALIGN_PARAGRAPH.RIGHT if j == 4 else None,
                )
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
