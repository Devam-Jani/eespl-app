# ruff: noqa: E501  (HTML and wording read better unwrapped)
"""The methodology as Word and PDF on the EESPL letterhead (the Prompt 17 quotation builder makes
the letterhead and the cover letter; the items follow in our offer format):

per item "ITEM NO: n - TECHNICAL SPECIFICATION & METHODOLOGY FOR <element> USING <brand> SYSTEM",
the condition and the reason, the stages with numbered steps (Option 1 / OR / Option 2), the
material specification table (property, standard, required, ours), the tests and hold points,
the checklist (Yes / No per stage), the guarantee and the client's scope.

The client's copy has nothing internal (no compliance colours, gaps, sources or pre-bid
questions); the internal copy adds the compliance summary and the pre-bid query sheet."""

import io
import re
from html import escape

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.execution.common import media
from app.masters.models import CompanyProfile
from app.methodology import vocab
from app.methodology.build import _ref, required_types
from app.methodology.models import (
    ClientStandard,
    Methodology,
    MethodologyFile,
    MethodologyItem,
    MethodSystem,
    PreBidQuery,
    StandardRule,
)
from app.quotations import docx_out, html_out
from app.quotations.models import Letterhead
from app.quotations.render import Doc, _media
from app.tenders.models import Tender
from app.timefmt import ist_today

STATUS_LABEL = {"green": "Complies", "red": "Gap", "confirm": "To confirm"}


def base_doc(db: Session, m: Methodology) -> Doc:
    head = (
        db.get(Letterhead, m.letterhead_id)
        if m.letterhead_id
        else db.scalars(
            select(Letterhead).where(Letterhead.is_active).order_by(Letterhead.id)
        ).first()
    )
    profile = db.get(CompanyProfile, 1)
    applicator = (
        (profile.legal_name or profile.trade_name) if profile else None
    ) or "Ethios Enviro Solutions Pvt. Ltd."
    tender = db.get(Tender, m.tender_id) if m.tender_id else None
    std = db.get(ClientStandard, m.standard_id) if m.standard_id else None
    client = (
        (tender.client.name if tender and tender.client else None)
        or (std.client_name if std else None)
        or "the client"
    )
    project = tender.name if tender else m.title
    return Doc(
        code=m.code,
        revision=m.revision,
        footer=f"{m.code} R{m.revision} v{m.version}",
        company_name=head.company_name if head else applicator,
        logo=_media(head.logo_path) if head else _media(profile.logo_path if profile else None),
        header_image=_media(head.header_image_path) if head else None,
        footer_image=_media(head.footer_image_path) if head else None,
        watermark=_media(head.watermark_path) if head else None,
        preprinted=bool(head and head.preprinted),
        header_text=(head.header_text if head else "") or "",
        footer_text=(head.footer_text if head else "") or "",
        primary=head.primary_color if head else "#0F6E5A",
        accent=head.accent_color if head else "#E69F00",
        opening=[f"Date: {ist_today():%d.%m.%Y}", f"To: {client}", f"Project: {project}"],
        subject=f"SUB: TECHNICAL SPECIFICATION & METHODOLOGY FOR WATERPROOFING WORKS — {project.upper()}",
        body=[
            "Dear Sir,",
            "",
            "Further to the tender, we submit herewith our technical specification and application methodology for the waterproofing works, item by item, "
            f"with the material specifications, the tests and hold points, the checklists and the guarantee{', aligned to ' + std.name if std else ''}.",
            "Please revert for any clarification required.",
            "",
            "Thanking you.",
        ],  # fmt: skip
        signatory_firm=(head.signatory_firm if head else f"For {applicator.upper()}"),
        signatory_line="",
        enclosures=["Technical specification & methodology (item wise)"],
    )


# --- what each item shows ---------------------------------------------------------------------------


def item_view(db: Session, m: Methodology, it: MethodologyItem, n: int, internal: bool) -> dict:
    std = db.get(ClientStandard, m.standard_id) if m.standard_id else None
    rule = db.get(StandardRule, it.rule_id) if it.rule_id else None
    system = db.get(MethodSystem, it.system_id) if it.system_id else None
    brand = (system.brand if system else m.brand) or ""
    element = (it.name or vocab.ELEMENT_LABEL.get(it.element or "", "")).upper()
    heading = f"ITEM NO: {n} - TECHNICAL SPECIFICATION & METHODOLOGY FOR {element} USING {brand.upper() + ' ' if brand else ''}SYSTEM"
    conds = []
    for k, v in (it.conditions or {}).items():
        if k.startswith("_") or not isinstance(v, dict):
            continue
        val = v.get("value")
        sized = [p["area_sqm"] for p in it.places or [] if p.get("area_sqm")]
        if k == "area_sqm" and len(sized) > 1:  # the toilets: the range the drawings give
            conds.append(
                (
                    f"Areas from the drawings ({len(sized)} labelled)",
                    f"{min(sized):g}–{max(sized):g} sqm",
                    v.get("source") if internal else None,
                )
            )
            continue
        label = {"area_sqm": "Area (sqm)", "sunk_depth_mm": "Sunk depth (mm)", "tank_area_sqm": "Tank plan area (sqm)", "tank_depth_m": "Tank depth (m)",
                 "exposure": "Open to sky / covered", "toilet_kind": "Sunken / underslung", "water_table": "Water table", "confined": "Confined", "face": "Face treated"}.get(k, k)  # fmt: skip
        conds.append(
            (
                label,
                f"{val:g}" if isinstance(val, float) else str(val),
                v.get("source") if internal else None,
            )
        )
    places = it.places or []
    reason = []
    if rule:
        req = " + ".join(
            f"{vocab.SYSTEM_LABEL.get(t, t)}{' on ' + w if w else ''}"
            for w, t in required_types(rule)
        )
        reason.append(
            f"{std.client_name or std.name if std else 'The client'}'s standard ({_ref(rule)}) asks for {req} for {rule.element_label.lower()}{' — ' + rule.condition_label if rule.condition_label else ''}."
        )
    if system:
        reason.append(f"We propose {system.name}.")
    if internal and it.reason:
        reason.append(f"Why: {it.reason}")
    # the spec table: the client's minimum and ours
    rows = []
    comp = {
        r.get("property"): r for r in (it.compliance or {}).get("rows", []) if r.get("property")
    }
    for r in comp.values():
        ours = (r.get("ours") or "").split(" — ")[0]
        rows.append({"property": r["topic"], "standard": r.get("standard") or "", "required": r["required"].split(" (")[0], "ours": ours if ours != "to confirm" or internal else "—",
                     "status": r["status"]})  # fmt: skip
    tests = []
    if rule:
        for t in rule.tests:
            if t["kind"] == "sampling":
                tests.append(("Material sampling", t["text"][:220]))
            elif t.get("hours"):
                tests.append(
                    (
                        {"pre_pond": "Pre-pond test", "pond": "Pond test", "fill": "Fill test"}[
                            t["kind"]
                        ],
                        f"{t['hours']} hours"
                        + (f", {t['depth_mm']} mm of water" if t.get("depth_mm") else ""),
                    )
                )
    elif system:
        tests += [
            (t["name"], f"{t['hours']} hours" if t.get("hours") else "as per the client's standard")
            for t in system.tests
        ]
    holds = [
        h["text"] for h in (rule.hold_points if rule else (system.hold_points if system else []))
    ]
    checklist = rule.checklist if rule and rule.checklist else (system.checklist if system else [])
    guarantee = (
        (system.guarantee_years if system and system.guarantee_years else None)
        or (rule.guarantee_years if rule else None)
        or (std.guarantee_years if std else None)
    )
    client_scope = [x["text"] for s in it.stages for x in s["steps"] if x.get("scope") == "client"]
    options = [db.get(MethodSystem, i) for i in it.option_ids or []]
    return {"heading": heading, "conds": conds, "places": places, "reason": reason, "stages": it.stages, "options": [o for o in options if o],
            "specs": rows, "tests": tests, "holds": holds, "checklist": checklist, "guarantee": guarantee, "client_scope": client_scope,
            "compliance": it.compliance if internal else None, "item": it}  # fmt: skip


def places_text(places: list[dict], internal: bool) -> str:
    """The client's copy: how many, the size range and how many of each size ("24 toilets,
    2.89 to 3.26 sqm: 1365 × 2135 (8), 1455 × 2135 (7) ..."); the internal copy: every label."""
    if internal:
        return ", ".join(sorted({p["label"] for p in places if p.get("label")}))
    sized = [p for p in places if p.get("area_sqm")]
    if not sized:
        return f"{len(places)} place(s)"
    counts: dict[str, int] = {}
    for p in sized:
        m = re.search(r"(\d{3,5})\s*[xX×*]\s*(\d{3,5})", p.get("label") or "")
        key = (
            " × ".join(str(n) for n in sorted(int(g) for g in m.groups()))
            if m
            else p.get("label") or "?"
        )
        counts[key] = counts.get(key, 0) + 1
    name = (sized[0].get("name") or "place").split()[0].lower()
    areas = [float(p["area_sqm"]) for p in sized]
    sizes = ", ".join(
        f"{k} ({n})" for k, n in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))
    )
    return f"{len(sized)} {name}{'s' if len(sized) > 1 and not name.endswith('s') else ''}, {min(areas):.2f} to {max(areas):.2f} sqm: {sizes}"


def _numbered(stages: list[dict]) -> list[tuple[str, list[tuple[int, str]]]]:
    out, n = [], 0
    for s in stages:
        steps = []
        for x in s["steps"]:
            n += 1
            t = x["text"]
            if x.get("scope") == "client" and "scope" not in t.lower():
                t += " (Client's scope)"
            steps.append((n, t))
        out.append((s["heading"], steps))
    return out


# --- Word ------------------------------------------------------------------------------------------------


def docx(db: Session, m: Methodology, internal: bool = False) -> bytes:
    doc = base_doc(db, m)
    from app.quotations import service as qsvc

    data = docx_out.docx(doc, qsvc.settings(db).word_template_path)
    d = Document(io.BytesIO(data))
    s = {k: docx_out._style(d, k) for k in docx_out.STYLES}
    P = docx_out._para
    first = True
    for n, it in enumerate(m.items, 1):
        v = item_view(db, m, it, n, internal)
        h = P(d, v["heading"], s["Offer Item Heading"], bold=True, keep=True)
        if first:
            h.paragraph_format.page_break_before = True
            first = False
        P(d, "CONDITION AND REASON:", s["Offer Stage"], bold=True, keep=True)
        for label, val, src in v["conds"]:
            P(d, f"•\t{label}: {val}" + (f"  [{src}]" if src else ""), s["Offer Step"])
        if v["places"]:
            labels = places_text(v["places"], internal)
            P(d, f"•\tFrom the drawings: {labels}", s["Offer Step"])
        for r in v["reason"]:
            P(d, r, s["Offer Body"])
        _steps_docx(d, s, v["stages"])
        for o in v["options"]:
            P(d, "==OR==", s["Offer OR"], keep=True)
            P(d, f"OPTION 2 - {o.name.upper()}", s["Offer Stage"], bold=True, keep=True)
            _steps_docx(d, s, o.stages)
        if v["specs"]:
            P(d, "MATERIAL SPECIFICATION:", s["Offer Stage"], bold=True, keep=True)
            heads = ["Property", "Standard", "Required", "Ours"] + (["Status"] if internal else [])
            _table(
                d,
                s,
                heads,
                [
                    [r["property"], r["standard"], r["required"], r["ours"]]
                    + ([STATUS_LABEL[r["status"]]] if internal else [])
                    for r in v["specs"]
                ],
            )
        if v["tests"] or v["holds"]:
            P(d, "TESTS AND HOLD POINTS:", s["Offer Stage"], bold=True, keep=True)
            for name, what in v["tests"]:
                P(d, f"•\t{name}: {what}", s["Offer Step"])
            for h_ in v["holds"]:
                P(d, f"•\tHold point: {h_}", s["Offer Step"])
        if v["checklist"]:
            P(d, "CHECKLIST:", s["Offer Stage"], bold=True, keep=True)
            rows = []
            for c in v["checklist"]:
                rows.append([c["stage"].upper(), "", ""])
                rows += [[q, "☐", "☐"] for q in c["items"]]
            _table(
                d,
                s,
                ["Description", "Yes", "No"],
                rows,
                bold_rows={i for i, r in enumerate(rows) if r[1] == ""},
            )
        P(
            d,
            f"GUARANTEE: {v['guarantee']} years"
            if v["guarantee"]
            else "GUARANTEE: as per the tender terms",
            s["Offer Stage"],
            bold=True,
        )
        if v["client_scope"]:
            P(d, "CLIENT'S SCOPE:", s["Offer Stage"], bold=True, keep=True)
            for t in v["client_scope"]:
                P(d, f"•\t{t}", s["Offer Step"])
        if internal and v["compliance"]:
            P(
                d,
                f"COMPLIANCE (INTERNAL): {v['compliance'].get('green', 0)} complies · {v['compliance'].get('red', 0)} gaps · {v['compliance'].get('confirm', 0)} to confirm",
                s["Offer Stage"],
                bold=True,
                keep=True,
            )
            _table(
                d,
                s,
                ["Requirement", "Client", "Ours", "Status", "Gap"],
                [
                    [
                        r["topic"],
                        r["required"],
                        r.get("ours") or "",
                        STATUS_LABEL[r["status"]],
                        r.get("gap") or "",
                    ]
                    for r in v["compliance"]["rows"]
                ],
            )
    if internal:
        _prebid_docx(db, d, s, m)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def _steps_docx(d, s, stages: list[dict]) -> None:
    for heading, steps in _numbered(stages):
        docx_out._para(d, heading + ":", s["Offer Stage"], bold=True, keep=True)
        for num, t in steps:
            docx_out._para(d, f"{num}.\t{t}", s["Offer Step"])


def _table(
    d, s, heads: list[str], rows: list[list[str]], bold_rows: set[int] | None = None
) -> None:
    t = d.add_table(rows=1, cols=len(heads))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    docx_out._borders(t)
    for i, h in enumerate(heads):
        docx_out._cell(t.rows[0].cells[i], h, s["Offer Table Text"], bold=True)
        docx_out._shade(t.rows[0].cells[i])
    docx_out._row_rules(t.rows[0], header=True)
    for k, r in enumerate(rows):
        row = t.add_row()
        docx_out._row_rules(row)
        for i, c in enumerate(r):
            docx_out._cell(
                row.cells[i],
                str(c or ""),
                s["Offer Table Text"],
                bold=bool(bold_rows and k in bold_rows),
            )
    widths = {
        3: [Cm(12.4), Cm(1.6), Cm(1.6)],
        4: [Cm(4.2), Cm(3.0), Cm(4.2), Cm(4.2)],
        5: [Cm(3.6), Cm(3.4), Cm(3.6), Cm(2.0), Cm(3.0)],
    }.get(len(heads))
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = w
    docx_out._para(d, "", s["Offer Body"])


def _prebid_docx(db: Session, d, s, m: Methodology) -> None:
    qs = list(
        db.scalars(
            select(PreBidQuery)
            .where(PreBidQuery.methodology_id == m.id)
            .order_by(PreBidQuery.kind, PreBidQuery.id)
        )
    )
    h = docx_out._para(
        d, "PRE-BID QUERIES (INTERNAL)", s["Offer Item Heading"], bold=True, keep=True
    )
    h.paragraph_format.page_break_before = True
    _table(
        d,
        s,
        ["Sr", "Query", "Reference", "Kind", "Status"],
        [[str(i), q.text, q.reference or "", q.kind, q.status] for i, q in enumerate(qs, 1)],
    )


# --- PDF -------------------------------------------------------------------------------------------------

EXTRA_CSS = """
.cond { margin: 2px 0 2px 18px; } .why { margin: 4px 0; }
table.spec td, table.spec th { font-size: 9pt; } td.st-green { color: #1b7a3d; font-weight: 600; } td.st-red { color: #b3261e; font-weight: 600; } td.st-confirm { color: #8a6100; font-weight: 600; }
tr.cl-stage td { font-weight: 700; background: #f2f5f4; } td.box { text-align: center; width: 1.2cm; }
p.or { text-align: center; font-weight: 700; }
"""


def html(db: Session, m: Methodology, internal: bool = False, for_pdf: bool = True) -> str:
    doc = base_doc(db, m)
    page = html_out.html(doc, for_pdf=for_pdf)
    parts = []
    for n, it in enumerate(m.items, 1):
        v = item_view(db, m, it, n, internal)
        parts.append('<div class="page-break"></div>' if n == 1 else "")
        parts.append(f'<h2 class="item">{escape(v["heading"])}</h2>')
        parts.append('<h3 class="stage">CONDITION AND REASON:</h3>')
        parts += [
            f'<p class="cond">• {escape(lab)}: <b>{escape(val)}</b>'
            + (f" <i>[{escape(src)}]</i>" if src else "")
            + "</p>"
            for lab, val, src in v["conds"]
        ]
        if v["places"]:
            labels = places_text(v["places"], internal)
            parts.append(f'<p class="cond">• From the drawings: {escape(labels)}</p>')
        parts += [f'<p class="why">{escape(r)}</p>' for r in v["reason"]]
        parts.append(_steps_html(v["stages"]))
        for o in v["options"]:
            parts.append('<p class="or"><mark>OR</mark></p>')
            parts.append(f'<h3 class="stage">OPTION 2 - {escape(o.name.upper())}</h3>')
            parts.append(_steps_html(o.stages))
        if v["specs"]:
            parts.append(
                '<h3 class="stage">MATERIAL SPECIFICATION:</h3><table class="spec"><thead><tr><th>Property</th><th>Standard</th><th>Required</th><th>Ours</th>'
                + ("<th>Status</th>" if internal else "")
                + "</tr></thead><tbody>"
            )
            for r in v["specs"]:
                parts.append(
                    f"<tr><td>{escape(r['property'])}</td><td>{escape(r['standard'])}</td><td>{escape(r['required'])}</td><td>{escape(r['ours'])}</td>"
                    + (
                        f'<td class="st-{r["status"]}">{STATUS_LABEL[r["status"]]}</td>'
                        if internal
                        else ""
                    )
                    + "</tr>"
                )
            parts.append("</tbody></table>")
        if v["tests"] or v["holds"]:
            parts.append('<h3 class="stage">TESTS AND HOLD POINTS:</h3>')
            parts += [f'<p class="cond">• {escape(a)}: {escape(b)}</p>' for a, b in v["tests"]]
            parts += [f'<p class="cond">• Hold point: {escape(h)}</p>' for h in v["holds"]]
        if v["checklist"]:
            parts.append(
                '<h3 class="stage">CHECKLIST:</h3><table class="spec"><thead><tr><th>Description</th><th>Yes</th><th>No</th></tr></thead><tbody>'
            )
            for c in v["checklist"]:
                parts.append(
                    f'<tr class="cl-stage"><td colspan="3">{escape(c["stage"].upper())}</td></tr>'
                )
                parts += [
                    f'<tr><td>{escape(q)}</td><td class="box">☐</td><td class="box">☐</td></tr>'
                    for q in c["items"]
                ]
            parts.append("</tbody></table>")
        parts.append(
            f'<h3 class="stage">GUARANTEE: {v["guarantee"]} years</h3>'
            if v["guarantee"]
            else '<h3 class="stage">GUARANTEE: as per the tender terms</h3>'
        )
        if v["client_scope"]:
            parts.append(
                '<h3 class="stage">CLIENT\'S SCOPE:</h3>'
                + "".join(f'<p class="cond">• {escape(t)}</p>' for t in v["client_scope"])
            )
        if internal and v["compliance"]:
            c = v["compliance"]
            parts.append(
                f'<h3 class="stage">COMPLIANCE (INTERNAL): {c.get("green", 0)} complies · {c.get("red", 0)} gaps · {c.get("confirm", 0)} to confirm</h3>'
            )
            parts.append(
                '<table class="spec"><thead><tr><th>Requirement</th><th>Client</th><th>Ours</th><th>Status</th><th>Gap</th></tr></thead><tbody>'
            )
            parts += [
                f'<tr><td>{escape(r["topic"])}</td><td>{escape(r["required"])}</td><td>{escape(r.get("ours") or "")}</td><td class="st-{r["status"]}">{STATUS_LABEL[r["status"]]}</td><td>{escape(r.get("gap") or "")}</td></tr>'
                for r in c["rows"]
            ]
            parts.append("</tbody></table>")
    if internal:
        qs = list(
            db.scalars(
                select(PreBidQuery)
                .where(PreBidQuery.methodology_id == m.id)
                .order_by(PreBidQuery.kind, PreBidQuery.id)
            )
        )
        parts.append(
            '<div class="page-break"></div><h2 class="item">PRE-BID QUERIES (INTERNAL)</h2><table class="spec"><thead><tr><th>Sr</th><th>Query</th><th>Reference</th><th>Kind</th></tr></thead><tbody>'
        )
        parts += [
            f"<tr><td>{i}</td><td>{escape(q.text)}</td><td>{escape(q.reference or '')}</td><td>{escape(q.kind)}</td></tr>"
            for i, q in enumerate(qs, 1)
        ]
        parts.append("</tbody></table>")
    body = "".join(parts)
    page = page.replace("</style>", EXTRA_CSS + "</style>", 1)
    return page.replace("</body>", body + "</body>", 1)


def _steps_html(stages: list[dict]) -> str:
    out = []
    for heading, steps in _numbered(stages):
        out.append(f'<h3 class="stage">{escape(heading)}:</h3>')
        out += [
            f'<div class="nitem"><span class="n">{num}.</span><div class="t">{escape(t)}</div></div>'
            for num, t in steps
        ]
    return "".join(out)


def pdf(db: Session, m: Methodology, internal: bool = False) -> bytes:
    from weasyprint import HTML

    return HTML(string=html(db, m, internal, for_pdf=True)).write_pdf()


# --- issued files ---------------------------------------------------------------------------------------


def issue(db: Session, m: Methodology, audience: str, user_id=None) -> list[MethodologyFile]:
    made = []
    for kind, data in (
        ("docx", docx(db, m, audience == "internal")),
        ("pdf", pdf(db, m, audience == "internal")),
    ):
        n = (
            db.scalar(
                select(func.count())
                .select_from(MethodologyFile)
                .where(
                    MethodologyFile.methodology_id == m.id,
                    MethodologyFile.kind == kind,
                    MethodologyFile.audience == audience,
                )
            )
            + 1
        )
        name = (
            f"{m.code}-R{m.revision}-v{m.version}"
            + ("-internal" if audience == "internal" else "")
            + (f"-{n}" if n > 1 else "")
            + f".{kind}"
        )
        rel = f"methodology/{m.code}/{name}"
        while media(rel).exists():  # never overwrite an issued file
            n += 1
            rel = f"methodology/{m.code}/{name.rsplit('.', 1)[0]}-{n}.{kind}"
        name = rel.rsplit("/", 1)[-1]
        media(rel).parent.mkdir(parents=True, exist_ok=True)
        media(rel).write_bytes(data)
        f = MethodologyFile(
            methodology_id=m.id,
            version=m.version,
            revision=m.revision,
            kind=kind,
            audience=audience,
            path=rel,
            file_name=name,
            size_bytes=len(data),
            quotation_id=m.quotation_id,
            created_by=user_id,
        )
        db.add(f)
        made.append(f)
    db.flush()
    return made


def share_with_client(db: Session, m: Methodology, site_id: int, user_id=None) -> None:
    """The client's copy (the latest client PDF) shared in the portal under the site's documents."""
    from app.portal.models import SiteDocument

    f = db.scalars(
        select(MethodologyFile)
        .where(
            MethodologyFile.methodology_id == m.id,
            MethodologyFile.audience == "client",
            MethodologyFile.kind == "pdf",
        )
        .order_by(MethodologyFile.id.desc())
    ).first()
    if f is None:
        f = next(x for x in issue(db, m, "client", user_id) if x.kind == "pdf")
    if db.scalar(
        select(SiteDocument.id).where(
            SiteDocument.site_id == site_id, SiteDocument.stored_path == f.path
        )
    ):
        return
    db.add(
        SiteDocument(
            site_id=site_id,
            title=f"Waterproofing methodology {m.code} R{m.revision}",
            kind="document",
            source="staff",
            stored_path=f.path,
            filename=f.file_name,
            share_with_client=True,
            created_by=user_id,
        )
    )
