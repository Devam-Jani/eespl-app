# ruff: noqa: E501  (the facts read better unwrapped)
"""Reading one old BOQ file (read only: the file is never changed).

Excel (xlsx, xlsm, xls): every sheet that is a BOQ (the same header detection and line rules as a
client BOQ import: app.tenders.boq_import) gives its priced lines: description, unit, quantity, the
rate we sold at, the item group, the make and remarks columns, QRO and NQ. For xlsx and xlsm the
workbook is also read with its formulas, and each rate cell is followed into the Working sheet
(app.knowledge.trace) for the build-up. The requirements come from the words
(app.knowledge.extract); the place from the folders and file name, else the sheet's first rows;
the date from the file name, a date written on the sheet, else the file's own date.

PDF: text-based PDFs that look like a BOQ (by name) are read with the PDF table reader; scanned
ones (no text) are listed as "scanned, not read" (no OCR yet); other PDFs (data sheets,
drawings) as "not a BOQ"."""

import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

from app.knowledge import extract
from app.knowledge.trace import Book, dec, trace

MAX_MB = 40
SKIP_SHEETS = re.compile(
    r"working|civil|summary|abstract|messurement|measurement|calculation|rate\s*analysis|^sheet\d*\s*\(\d+\)$",
    re.I,
)
BOQ_PDF = re.compile(r"boq|quot|offer|rate|inquiry|enquiry|tender|eespl|price|bill\s*of", re.I)
SQFT = Decimal("10.7639")


@dataclass
class FileFacts:
    path: str
    kind: str
    size: int
    status: str = "read"
    error: str | None = None
    sheets: list[dict] = field(default_factory=list)
    lines: list[dict] = field(default_factory=list)
    folder: str | None = None
    city: str | None = None
    region: str | None = None
    place_source: str | None = None
    on_date: Any = None
    date_source: str | None = None
    is_eespl: bool = False


def _text(v) -> str:
    return " ".join(str(v).split()) if v is not None else ""


def _grid_from_values(wb) -> dict[str, list[list[Any]]]:
    from app.tenders.boq_import import MAX_COLS, MAX_ROWS  # noqa: PLC0415

    out = {}
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        out[ws.title] = [
            list(r)
            for r in ws.iter_rows(
                max_row=min(ws.max_row or 0, MAX_ROWS), max_col=MAX_COLS, values_only=True
            )
        ]
    return out


def guess_layout(rows: list[list[Any]]) -> tuple[int, dict[str, int]] | None:
    """A BOQ sheet whose header does not say "Description" (our own offers: "500 | WATERPROOFING
    | Unit | Rate"): the long-text column, the unit column (cells that read as units) and the
    numbers to its right (quantity and rate, or the rate alone). (header row index, columns)."""
    from app.masters.units import normalise_unit  # noqa: PLC0415

    body = rows[:400]
    if len(body) < 3:
        return None
    width = max((len(r) for r in body), default=0)
    long_, units, nums = [0] * width, [0] * width, [0] * width
    for r in body:
        for c, v in enumerate(r):
            if isinstance(v, str) and len(v.strip()) > 40:
                long_[c] += 1
            elif (
                isinstance(v, str)
                and v.strip()
                and len(v.strip()) < 12
                and normalise_unit(v.strip())
            ):
                units[c] += 1
            elif isinstance(v, int | float) and not isinstance(v, bool):
                nums[c] += 1
    if not width or max(long_) < 3 or max(units) < 3:
        return None
    desc = long_.index(max(long_))
    unit = units.index(max(units))
    right = [c for c in range(unit + 1, width) if nums[c] >= 3]
    if unit <= desc or not right:
        return None
    first = next(
        (
            i
            for i, r in enumerate(body)
            if len(r) > unit and isinstance(r[unit], str) and normalise_unit(r[unit].strip())
        ),
        0,
    )
    cols = {"description": desc, "unit": unit}
    if desc > 0:
        cols["item_no"] = desc - 1
    if len(right) == 1:
        cols["rate"] = right[0]
    else:
        cols["qty"], cols["rate"] = right[0], right[1]
    return max(first - 1, 0), cols


def _lines_of(grid: dict, book: Book | None, facts: FileFacts) -> None:
    from app.tenders.boq_import import (  # noqa: PLC0415
        BoqFormatError,
        column_letter,
        detect_header,
        parse,
    )

    for name, rows in grid.items():
        if SKIP_SHEETS.search(name or "") or not rows:
            continue
        index, columns, score = detect_header(rows)
        if (
            score <= 0
            or "description" not in columns
            or not ({"rate", "qty", "amount"} & set(columns))
        ):
            guess = guess_layout(rows)
            if guess is None:
                facts.sheets.append({"name": name, "boq": False})
                continue
            index, column_map = guess
        else:
            column_map = {f: c["column"] for f, c in columns.items()}
        try:
            parsed = parse(rows, index + 1, column_map)
        except BoqFormatError:
            facts.sheets.append({"name": name, "boq": False})
            continue
        sections = parsed.sections
        found = 0
        for ln in parsed.lines:
            rate = ln.client_file_rate
            row = ln.source_row or (ln.rows[-1] if ln.rows else None)
            cell = (
                f"{column_letter(column_map['rate'])}{row}"
                if "rate" in column_map and row
                else None
            )
            if rate is None and book is not None and cell:
                rate = dec(book.value(name, cell))  # a formula saved without its value
            if rate is None and ln.qty_note not in ("QRO", "NQ"):
                continue
            working = None
            if book is not None and cell:
                try:
                    t = trace(book, name, cell)
                except Exception:  # noqa: BLE001  (a formula we cannot follow: no build-up)
                    t = None
                if t is not None and t.items:
                    working = t.summary()
            parent = (
                sections[ln.section].title
                if ln.section is not None and ln.section < len(sections)
                else None
            )
            facts.lines.append({
                "sheet": name, "row": row, "item_no": ln.item_no, "parent_item": parent, "description": ln.description,
                "unit": ln.unit, "unit_raw": ln.unit_raw, "qty": ln.qty, "qty_note": ln.qty_note, "rate": rate,
                "product_make": ln.client_product, "remarks": ln.client_remarks or ln.our_remarks, "working": working,
            })  # fmt: skip
            found += 1
        facts.sheets.append({"name": name, "boq": True, "lines": found, "header_row": index + 1})


def _first_cells(grid: dict, n: int = 15) -> list[str]:
    out = []
    for rows in list(grid.values())[:3]:
        for r in rows[:n]:
            t = " ".join(_text(v) for v in r if v not in (None, ""))
            if t:
                out.append(t)
    return out


def read_file(root: Path, path: Path, regions: dict[str, list[str]]) -> FileFacts:
    rel = path.relative_to(root).as_posix()
    kind = path.suffix.lower().lstrip(".")
    facts = FileFacts(path=rel, kind=kind, size=path.stat().st_size)
    parts = rel.split("/")
    facts.folder = parts[0] if len(parts) > 1 else None
    facts.is_eespl = path.name.upper().startswith("EESPL")
    grid: dict = {}
    saved = None  # the workbook's own saved (or created) date
    if facts.size > MAX_MB * 1024 * 1024:
        facts.status, facts.error = "skipped", f"larger than {MAX_MB} MB"
        return facts
    try:
        if kind == "xls" and path.read_bytes()[:2] == b"PK":
            kind = "xlsx"  # an .xlsx saved with the old name: read as what it is
        if kind in ("xlsx", "xlsm"):
            import io  # noqa: PLC0415

            from openpyxl import load_workbook  # noqa: PLC0415

            raw = path.read_bytes()  # a stream: openpyxl refuses a .xls name on an xlsx
            values = load_workbook(io.BytesIO(raw), data_only=True)
            formulas = load_workbook(io.BytesIO(raw), data_only=False)
            props = values.properties
            saved = (props.modified or props.created) if props else None
            grid = _grid_from_values(values)
            _lines_of(grid, Book(formulas, values), facts)
        elif kind == "xls":
            from app.tenders.boq_import import read_grid  # noqa: PLC0415

            grid = read_grid(path)
            _lines_of(grid, None, facts)  # old .xls: values only, no formulas to follow
        elif kind == "pdf":
            from pypdf import PdfReader  # noqa: PLC0415

            reader = PdfReader(str(path))
            text = "".join((p.extract_text() or "") for p in reader.pages[:3])
            if len(text.strip()) < 40:
                facts.status, facts.error = (
                    "scanned",
                    "scanned (no text layer): not read, no OCR yet",
                )
                return facts
            if not BOQ_PDF.search(path.name):
                facts.status, facts.error = (
                    "not_boq",
                    "text PDF, not a BOQ by its name (data sheet, drawing, letter)",
                )
                return facts
            from app.tenders.boq_import import BoqFormatError, read_grid  # noqa: PLC0415

            try:
                grid = read_grid(path)
            except BoqFormatError as exc:
                facts.status, facts.error = "not_boq", str(exc)
                return facts
            _lines_of(grid, None, facts)
        else:
            facts.status, facts.error = "skipped", "not a BOQ file type"
            return facts
    except Exception as exc:  # noqa: BLE001  (a broken file is reported, never fatal)
        facts.status, facts.error = "failed", f"{type(exc).__name__}: {str(exc)[:200]}"
        return facts
    if not facts.lines and kind == "pdf":
        facts.status, facts.error = "not_boq", "no priced BOQ lines found"
    texts = [("folder", " ".join(parts[:-1])), ("file name", path.stem)] + [
        ("sheet", t) for t in _first_cells(grid)
    ]
    facts.city, facts.region, facts.place_source = extract.place(texts, regions)
    facts.on_date, facts.date_source = extract.file_date(
        path.name, " ".join(parts[:-1]), _first_cells(grid, 25), saved
    )
    context = extract.requirements(path.stem, " ".join(parts[-2:-1]))
    for ln in facts.lines:
        _facts_of_line(ln, context)
    return facts


def _facts_of_line(ln: dict, context: extract.Requirements | None = None) -> None:
    """The requirements, the unit's rate per sqm, the cost and margin from the working. The
    area type comes from the item heading and the description's first sentence, else the
    file's name (a swimming pool BOQ), else the whole text."""
    from app.masters.units import normalise_unit  # noqa: PLC0415

    req = extract.requirements(
        ln["parent_item"], ln["description"], ln["remarks"], ln["product_make"]
    )
    head = extract.requirements(
        ln["parent_item"], re.split(r"(?<=[a-z0-9])[.:;]\s", ln["description"] or "", maxsplit=1)[0]
    )
    req.area_type = head.area_type or (context.area_type if context else None) or req.area_type
    ln["requirements"] = req.as_dict()
    unit = (ln.get("unit") or normalise_unit(ln.get("unit_raw") or "") or "").lower() or None
    ln["unit"] = unit
    rate = dec(ln["rate"])
    ln["rate"] = rate
    ln["rate_sqm"] = (
        rate
        if unit == "sqm"
        else (rate * SQFT).quantize(Decimal("0.01"))
        if unit in ("sqft", "sft") and rate is not None
        else None
    )
    w = ln.get("working")
    ln["cost_sqm"] = ln["margin_percent"] = None
    ln["to_check"], ln["check_note"] = False, None
    if w:
        cost = dec(w.get("cost"))
        ln["margin_percent"] = dec(w.get("margin_on_selling_percent"))
        ln["cost_sqm"] = (
            cost if unit != "sqft" or cost is None else (cost * SQFT).quantize(Decimal("0.01"))
        )
        worked = dec(w.get("rate"))
        if rate and worked and abs(worked - rate) > max(Decimal("1"), rate * Decimal("0.02")):
            ln["to_check"], ln["check_note"] = (
                True,
                f"the working gives {worked}, the BOQ says {rate}",
            )
    if rate is not None and (rate <= 0 or rate > 1_000_000):
        ln["to_check"], ln["check_note"] = True, "rate out of range"
