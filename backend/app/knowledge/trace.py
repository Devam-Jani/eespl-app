# ruff: noqa: E501  (the build-up rows read better unwrapped)
"""Following a BOQ rate cell's formula into the Working sheet: the build-up it adds up.

Our workings are blocks of rows: an item (surface preparation, a product, labour), its
consumption (coverage) and its price (often written price + freight: =120+5; labour as ₹ per sq
ft x 10.76), their product (C x D); a Total (=SUM); a margin row (a fraction, 0.25, times the
total); the rate (=SUM of the total and the margin), sometimes with items added after the margin
(protection, tiles). The rate cell on the BOQ points at it (=+Working!E13).

The trace walks the formulas from the rate cell (references, SUM ranges, sums of cells) and
classifies the cells it reaches:
- a product cell: two cells of one row multiplied, both numbers: (label, consumption, price);
- a margin: a fraction below 1 multiplied by a cell that is itself a total;
- anything else is followed into its references.
Every fact keeps its sheet and cell, so anyone can open the original."""

import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from openpyxl.formula import Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

SQFT = Decimal("10.76")
LABOUR = re.compile(r"labou?r|application\s*charge|applicator|manpower|workmanship", re.I)
PREP = re.compile(
    r"^\s*s\.?\s*p\.?\s*$|surface\s*prep|preparation|cleaning|grinding|chipping|hacking|\bprep\b",
    re.I,
)
MAX_CELLS = 400


def dec(v) -> Decimal | None:
    if v is None or isinstance(v, bool):
        return None
    try:
        d = Decimal(str(v))
    except (InvalidOperation, ValueError):
        return None
    return d if d.is_finite() else None


@dataclass
class Item:
    label: str
    kind: str  # product | labour | surface_prep | other
    consumption: Decimal | None
    price: Decimal | None
    freight: Decimal | None
    per_sqft: Decimal | None  # labour written as ₹ per sq ft x 10.76
    amount: Decimal | None
    ref: str
    after_margin: bool = False

    def as_dict(self) -> dict:
        return {k: (str(v) if isinstance(v, Decimal) else v) for k, v in self.__dict__.items()}


@dataclass
class BuildUp:
    sheet: str
    cell: str
    rate: Decimal | None
    items: list[Item] = field(default_factory=list)
    margins: list[tuple[Decimal, Decimal | None, str]] = field(
        default_factory=list
    )  # (fraction, base, ref)

    @property
    def margin_amount(self) -> Decimal:
        return sum((f * (b or 0) for f, b, _r in self.margins), Decimal(0))

    def summary(self) -> dict:
        rate = self.rate
        before = [i for i in self.items if not i.after_margin]
        cost = sum((i.amount or 0 for i in before), Decimal(0))
        extra = sum((i.amount or 0 for i in self.items if i.after_margin), Decimal(0))
        margin = self.margin_amount
        on_cost = (self.margins[0][0] * 100) if self.margins else None
        selling = (margin / (cost + margin) * 100) if margin and (cost + margin) else None
        labour = next((i for i in before if i.kind == "labour"), None)
        prep = sum((i.amount or 0 for i in before if i.kind == "surface_prep"), Decimal(0))
        return {
            "sheet": self.sheet, "cell": self.cell, "rate": str(rate) if rate is not None else None,
            "items": [i.as_dict() for i in self.items], "cost": str(round(cost, 2)), "after_margin": str(round(extra, 2)),
            "margin_amount": str(round(margin, 2)), "margin_on_cost_percent": str(round(on_cost, 2)) if on_cost is not None else None,
            "margin_on_selling_percent": str(round(selling, 2)) if selling is not None else None,
            "labour_sqm": str(labour.amount) if labour and labour.amount is not None else None,
            "labour_sqft": str(labour.per_sqft) if labour and labour.per_sqft is not None else None,
            "surface_prep_sqm": str(round(prep, 2)) if prep else None,
        }  # fmt: skip


class Book:
    """A workbook read twice: its formulas and its values (as Excel last saved them)."""

    def __init__(self, formulas, values):
        self.f, self.v = formulas, values

    def formula(self, sheet: str, cell: str):
        try:
            return self.f[sheet][cell].value
        except (KeyError, ValueError):
            return None

    def value(self, sheet: str, cell: str, depth: int = 0):
        """The value Excel saved; a formula saved without its value (never recalculated): worked
        out here when it is simple (numbers, + - * /, cell references, SUM)."""
        try:
            v = self.v[sheet][cell].value
        except (KeyError, ValueError):
            return None
        if v is None and depth < 30:
            f = self.formula(sheet, cell)
            if isinstance(f, str) and f.startswith("="):
                return self._eval(sheet, f, depth + 1)
        return v

    def _eval(self, sheet: str, formula: str, depth: int):
        try:
            tokens = Tokenizer(formula).items
        except Exception:  # noqa: BLE001
            return None
        expr = []
        stack_sum = 0
        for t in tokens:
            if t.type == "OPERAND" and t.subtype == "NUMBER":
                expr.append(t.value)
            elif t.type == "OPERAND" and t.subtype == "RANGE":
                ref = t.value
                sh = sheet
                if "!" in ref:
                    sh, ref = ref.rsplit("!", 1)
                    sh = sh.strip("'").replace("''", "'")
                try:
                    c1, r1, c2, r2 = range_boundaries(ref.replace("$", ""))
                except ValueError:
                    return None
                vals = []
                for r in range(r1, r2 + 1):
                    for c in range(c1, c2 + 1):
                        x = self.value(sh, f"{get_column_letter(c)}{r}", depth)
                        x = dec(x)
                        vals.append(x or Decimal(0))
                if len(vals) > 1 and not stack_sum:
                    return None
                expr.append(f"({'+'.join(str(x) for x in vals) or '0'})")
            elif t.type == "OPERATOR-INFIX" and t.value in "+-*/":
                expr.append(t.value)
            elif t.type == "OPERATOR-PREFIX" and t.value in "+-":
                expr.append(t.value)
            elif t.type == "FUNC" and t.subtype == "OPEN" and t.value.upper() == "SUM(":
                stack_sum += 1
                expr.append("(")
            elif t.type == "FUNC" and t.subtype == "CLOSE" and stack_sum:
                stack_sum -= 1
                expr.append(")")
            elif t.type == "SEP" and stack_sum:
                expr.append("+")
            elif t.type == "PAREN":
                expr.append(t.value)
            elif t.type == "WHITE-SPACE":
                continue
            else:
                return None
        text = "".join(expr)
        if not re.fullmatch(r"[\d.+\-*/()E]*", text):
            return None
        try:
            return Decimal(str(eval(text, {"__builtins__": {}}, {})))  # noqa: S307  (digits and operators only, checked above)
        except Exception:  # noqa: BLE001
            return None

    def has(self, sheet: str) -> bool:
        return sheet in self.f.sheetnames


def _refs(sheet: str, formula: str) -> tuple[list[tuple[str, str]], list[str]]:
    """The cells a formula reads, and its operators and functions (for the shape)."""
    cells, ops = [], []
    try:
        tokens = Tokenizer(formula).items
    except Exception:  # noqa: BLE001  (a formula openpyxl cannot read: no trace)
        return [], []
    for t in tokens:
        if t.type == "OPERAND" and t.subtype == "RANGE":
            ref = t.value
            sh = sheet
            if "!" in ref:
                sh, ref = ref.rsplit("!", 1)
                sh = sh.strip("'").replace("''", "'")
            ref = ref.replace("$", "")
            try:
                c1, r1, c2, r2 = range_boundaries(ref)
            except ValueError:
                continue
            if None in (c1, r1, c2, r2) or (c2 - c1 + 1) * (r2 - r1 + 1) > 300:
                continue
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    cells.append((sh, f"{get_column_letter(c)}{r}"))
        elif t.type in ("OPERATOR-INFIX", "FUNC"):
            ops.append(t.value.upper())
        elif t.type == "OPERAND" and t.subtype == "NUMBER":
            ops.append("NUM")
    return cells, ops


def _split(cell: str) -> tuple[str, int]:
    m = re.match(r"([A-Z]+)(\d+)", cell)
    return m.group(1), int(m.group(2))


def _label(book: Book, sheet: str, row: int, before_col: int) -> str:
    for c in range(1, before_col):
        v = book.value(sheet, f"{get_column_letter(c)}{row}")
        if isinstance(v, str) and v.strip() and not v.strip().startswith("="):
            if re.fullmatch(r"[\d.\s]+", v.strip()):
                continue
            return " ".join(v.split())[:120]
    return ""


def _price(
    book: Book, sheet: str, cell: str, value: Decimal | None
) -> tuple[Decimal | None, Decimal | None, Decimal | None]:
    """(price, freight, ₹ per sq ft) from how the price cell is written: =120+5, =11*10.76."""
    f = book.formula(sheet, cell)
    if isinstance(f, str) and f.startswith("="):
        body = f[1:].replace(" ", "").lstrip("+")
        m = re.fullmatch(r"(\d+(?:\.\d+)?)\+(\d+(?:\.\d+)?)", body)
        if m:
            return Decimal(m.group(1)), Decimal(m.group(2)), None
        m = re.fullmatch(r"(\d+(?:\.\d+)?)\*10\.76\d*", body) or re.fullmatch(
            r"10\.76\d*\*(\d+(?:\.\d+)?)", body
        )
        if m:
            return value, None, Decimal(m.group(1))
    return value, None, None


def trace(book: Book, sheet: str, cell: str) -> BuildUp | None:
    """The build-up behind a rate cell, or None when it is a typed number or leads nowhere."""
    start = book.formula(sheet, cell)
    if not (isinstance(start, str) and start.startswith("=")):
        return None
    out = BuildUp(sheet=sheet, cell=cell, rate=dec(book.value(sheet, cell)))
    seen: set[tuple[str, str]] = set()

    def is_total(sh: str, c: str) -> bool:
        f = book.formula(sh, c)
        return isinstance(f, str) and ("SUM(" in f.upper() or f.count("+") >= 2)

    def walk(sh: str, c: str, after_margin: bool) -> None:
        if (sh, c) in seen or len(seen) > MAX_CELLS or not book.has(sh):
            return
        seen.add((sh, c))
        f = book.formula(sh, c)
        if not (isinstance(f, str) and f.startswith("=")):
            return
        cells, ops = _refs(sh, f)
        col, row = _split(c)
        # a product of two cells of one row: an item of the build-up
        if len(cells) == 2 and ops.count("*") == 1 and "SUM" not in "".join(ops) and "/" not in ops:
            (s1, a), (s2, b) = cells
            ca, ra = _split(a)
            cb, rb = _split(b)
            va, vb = dec(book.value(s1, a)), dec(book.value(s2, b))
            # a margin: a fraction times a total
            for (_sx, _x, vx), (sy, y, vy) in (
                ((s1, a, va), (s2, b, vb)),
                ((s2, b, vb), (s1, a, va)),
            ):
                if vx is not None and Decimal(0) < vx < 1 and is_total(sy, y):
                    out.margins.append((vx, vy, f"{sh}!{c}"))
                    walk(sy, y, after_margin)
                    return
            if s1 == s2 == sh and ra == rb == row and va is not None and vb is not None:
                first, second = sorted(
                    [(column_index_from_string(ca), a, va), (column_index_from_string(cb), b, vb)]
                )
                label = _label(book, sh, row, first[0])
                price, freight, per_sqft = _price(book, sh, second[1], second[2])
                kind = (
                    "labour"
                    if LABOUR.search(label) or per_sqft is not None
                    else "surface_prep"
                    if PREP.search(label)
                    else "product"
                    if label
                    else "other"
                )
                out.items.append(Item(label=label or "(no label)", kind=kind, consumption=first[2], price=price, freight=freight,
                                      per_sqft=per_sqft, amount=dec(book.value(sh, c)), ref=f"{sh}!{c}", after_margin=after_margin))  # fmt: skip
                return
        # a sum of a total and a margin, then items added after it
        margin_seen = bool(out.margins)
        for sx, x in cells:
            walk(sx, x, after_margin or (margin_seen and bool(out.margins)))
            margin_seen = margin_seen or bool(out.margins)

    walk(sheet, cell, False)
    # items reached after the margin row (in the same SUM, after it): mark them
    if out.margins:
        mrow = _split(out.margins[0][2].split("!")[1])[1]
        for i in out.items:
            r = _split(i.ref.split("!")[1])[1]
            i.after_margin = r > mrow and i.ref.split("!")[0] == out.margins[0][2].split("!")[0]
    if not out.items and not out.margins:
        return None
    return out
