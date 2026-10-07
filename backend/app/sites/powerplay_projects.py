"""Past projects from a Powerplay export ("project Raw" sheet: _id, name) as closed sites.

- Demo / test / junk names and internal person names are skipped (listed for review).
- Exact duplicates (ignoring case and spaces) become one site; the other Powerplay ids are kept
  in the site's notes.
- Near-duplicates ("ANMAYA INFRABUILD" / "ANAMAYA INFRABUILD") stay separate but are flagged.
- Idempotent: a project already imported (by Powerplay id or by its normalised name among
  Powerplay sites) is not created again.
"""

import difflib
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.sites.models import Site
from app.sites.service import next_code

SHEET = "project Raw"
JUNK_WORDS = ("demo", "test", "abc", "abcd", "xyz")
JUNK_PHRASES = ("copy ", "return material")
INTERNAL_PEOPLE = ("devam jani", "devambhai site", "roshan bhai", "laxman bhai", "khusboo")
NEAR_DUPLICATE_RATIO = 0.9


@dataclass
class ProjectImport:
    rows: int = 0
    imported: list[str] = field(default_factory=list)
    already: int = 0
    skipped: list[tuple[str, str]] = field(default_factory=list)  # (name, reason)
    merged: list[tuple[str, int]] = field(default_factory=list)  # (name, how many rows)
    near_duplicates: list[tuple[str, str]] = field(default_factory=list)


def name_key(name: str) -> str:
    return re.sub(r"\s+", "", name).lower()


def skip_reason(name: str) -> str | None:
    text = " ".join(name.lower().split())
    words = set(re.findall(r"[a-z0-9]+", text))
    if len(re.sub(r"[^a-z0-9]", "", text)) <= 1:
        return "single letter"
    for w in JUNK_WORDS:
        if w in words:
            return f"junk word '{w}'"
    for p in JUNK_PHRASES:
        if p in f"{text} " and (p != "copy " or "copy" in words):
            return f"junk phrase '{p.strip()}'"
    for person in INTERNAL_PEOPLE:
        if person in text:
            return f"internal name '{person}'"
    return None


def read_projects(path: str | Path) -> list[tuple[str, str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET not in wb.sheetnames:
            raise ValueError(f"No '{SHEET}' sheet in {Path(path).name}")
        rows = list(wb[SHEET].iter_rows(values_only=True))
    finally:
        wb.close()
    header = [str(c or "").strip().lower() for c in rows[0]] if rows else []
    if "_id" not in header or "name" not in header:
        raise ValueError(f"'{SHEET}' needs the columns _id and name")
    i_id, i_name = header.index("_id"), header.index("name")
    out = []
    for r in rows[1:]:
        pid, name = r[i_id], r[i_name]
        if pid and name and str(name).strip():
            out.append((str(pid).strip(), " ".join(str(name).split())))
    return out


def plan(projects: list[tuple[str, str]], result: ProjectImport) -> list[tuple[str, list[str]]]:
    """[(name, [powerplay ids])] to import, after skips and duplicate merging."""
    groups: dict[str, tuple[str, list[str]]] = {}
    for pid, name in projects:
        result.rows += 1
        reason = skip_reason(name)
        if reason:
            result.skipped.append((name, reason))
            continue
        key = name_key(name)
        if key in groups:
            groups[key][1].append(pid)
        else:
            groups[key] = (name, [pid])
    for name, ids in groups.values():
        if len(ids) > 1:
            result.merged.append((name, len(ids)))
    names = [name for name, _ in groups.values()]
    for i, a in enumerate(names):
        for b in names[i + 1 :]:
            ratio = difflib.SequenceMatcher(None, a.lower(), b.lower()).ratio()
            if ratio >= NEAR_DUPLICATE_RATIO:
                result.near_duplicates.append((a, b))
    return list(groups.values())


def import_projects(db: Session, path: str | Path, user_id=None) -> ProjectImport:
    result = ProjectImport()
    groups = plan(read_projects(path), result)
    flagged: dict[str, str] = {}
    for a, b in result.near_duplicates:
        flagged.setdefault(a, b)
        flagged.setdefault(b, a)
    existing = db.execute(
        select(Site.source_ref, Site.name).where(Site.source == "powerplay")
    ).all()
    known_ids = {ref for ref, _ in existing}
    known_names = {name_key(name) for _, name in existing}
    for name, ids in groups:
        if known_ids & set(ids) or name_key(name) in known_names:
            result.already += 1
            continue
        notes = ["Imported from Powerplay (past project)."]
        if len(ids) > 1:
            notes.append(f"Merged duplicate Powerplay projects: {', '.join(ids)}.")
        if name in flagged:
            notes.append(f"Check: possibly the same project as '{flagged[name]}'.")
        db.add(
            Site(
                code=next_code(db),
                name=name[:300],
                status="closed",
                source="powerplay",
                source_ref=ids[0],
                notes=" ".join(notes),
                created_by=user_id,
            )
        )
        result.imported.append(name)
    db.flush()
    return result
