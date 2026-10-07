"""Reading a client's BOQ spreadsheet into sections and lines.

Every client lays out their BOQ differently, so this works in three steps:

1. read the file into a grid of cell values (xlsx/xls/csv; a merged range keeps its value in the
   top-left cell only, the other cells of the range read as empty),
2. find the header row (first 30 rows) and guess which column holds what, with a confidence,
3. walk the rows below the header and decide, for each, whether it is a section heading, a
   BOQ line, a continuation of the line above, a note, or something to skip.

Step 3's rules (see `parse`):
- A row is a *quantity row* when it has a unit, a quantity (number or QRO/NQ text) or a rate.
- Rows without quantities that come before a quantity row are resolved together with it:
  - a numbered title row directly followed by an unnumbered quantity row is one item
    ("1.2 PROTECTION SCREED" + the next row's text, unit and quantity);
  - lettered sub-items ("a)", "b)") under a numbered row without quantities get the parent's text
    as a prefix, and the parent becomes a section heading;
  - steps ("i)" ... "iv)") followed by a quantity row are one item with all the steps;
  - other numbered rows without quantities become section headings;
  - unnumbered rows right after a line continue that line's description; after a heading they
    are notes (skipped, counted).
- QRO / "Rate only" / RO in the quantity (or unit) column: qty empty, qty_note QRO.
  NQ / "Not quoted" / "By civil contractor" (in quantity or rate): qty_note NQ, not quoted.
- Blank rows, rows with no description, and Total / Sub total / Carried forward / Grand total
  rows are skipped and counted. Parsing stops at a "Terms & Conditions" block.
"""

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.masters.units import is_rate_only, normalise_unit

HEADER_SCAN_ROWS = 30
MAX_ROWS = 5000
MAX_COLS = 60
PREVIEW_ROWS = 30
QTY_PLACES = Decimal("0.001")  # boq_lines.qty is Numeric(14, 3)
RATE_PLACES = Decimal("0.01")

FIELDS = (
    "item_no",
    "description",
    "description2",
    "unit",
    "qty",
    "qty2",  # a second quantity column under the same heading (added to qty)
    "rate",
    "material_rate",
    "application_rate",
    "amount",
    "product",
    "remarks",
    "our_remarks",
)

# exact header names (normalised) and keywords the header may contain, per field
HEADER_RULES: dict[str, tuple[set[str], tuple[str, ...]]] = {
    "item_no": (
        {
            "sr no",
            "s no",
            "sn",
            "s n",
            "sl no",
            "s l no",
            "sr no s",
            "sno",
            "srno",
            "sr",
            "s r no",
            "item no",
            "item code",
            "serial no",
            "no",
            "boq no",
            "boq item no",
            "item number",
            "sr number",
            "sl",
        },
        ("sr no", "s no", "sl no", "item no", "serial"),
    ),
    "description": (
        {
            "description",
            "item description",
            "particulars",
            "description of work",
            "description of item",
            "work description",
            "item",
            "items",
            "item details",
            "scope of work",
            "nature of work",
        },
        ("description", "particular"),
    ),
    "description2": (
        {
            "detail description",
            "detailed description",
            "specification",
            "specifications",
            "detailed specification",
        },
        ("detail", "specification"),
    ),
    "unit": (
        {"unit", "uom", "units", "unit of measurement", "unit of measure", "u o m"},
        ("unit", "uom"),
    ),
    "qty": (
        {
            "qty",
            "quantity",
            "total quantity",
            "qnty",
            "quantities",
            "boq qty",
            "total qty",
            "est qty",
            "estimated quantity",
        },
        ("qty", "quantit"),
    ),
    "rate": (
        {
            "rate",
            "unit rate",
            "rate rs",
            "rate in rs",
            "rate inr",
            "rates",
            "total rate",
            "total rate m l",
            "total rate rs",
            "rate m l",
            "supply apply rate",
        },
        ("rate",),
    ),
    # split rates (material + application = total): kept for reference next to the total
    "material_rate": (
        {"material rate", "supply rate", "material", "material cost"},
        ("material rate", "supply rate", "material cost"),
    ),
    "application_rate": (
        {
            "application rate",
            "labour rate",
            "laying rate",
            "application",
            "labour",
            "application charges",
        },
        ("application rate", "labour rate", "laying rate", "application charge", "labour cost"),
    ),
    "amount": (
        {"amount", "amount rs", "total amount", "amount in rs", "value", "total", "amt"},
        ("amount", "amt"),
    ),
    "product": (
        {"product", "make", "brand", "approved make", "product brand", "make brand"},
        ("make", "brand", "product"),
    ),
    "remarks": (
        {"remarks", "remark", "client remarks", "query", "queries", "comments", "comment"},
        ("remark", "query", "comment"),
    ),
    "our_remarks": (
        {
            "ethios remarks",
            "eespl remarks",
            "consideration",
            "our remarks",
            "contractor remarks",
            "bidder remarks",
            "vendor remarks",
        },
        ("ethios", "eespl", "consideration", "contractor remark", "bidder remark"),
    ),
}
# Fields are claimed in this order, so "Ethios Remarks" goes to our_remarks before remarks
# sees it, and "Item No" goes to item_no before description could take "Item".
CLAIM_ORDER = (
    "item_no",
    "unit",
    "qty",
    "amount",
    "material_rate",
    "application_rate",
    "rate",
    "our_remarks",
    "remarks",
    "product",
    "description",
    "description2",
)
HEADER_WEIGHTS = {"description": 3, "qty": 2, "unit": 2, "rate": 1, "item_no": 1, "amount": 1}

QRO_TEXT = re.compile(
    r"^\s*(q\.?\s*r\.?\s*o\.?|rate\s*only|quote\s*rate\s*only|r\.?\s*o\.?|r/o)\s*$", re.I
)
NQ_TEXT = re.compile(
    r"^\s*(n\.?\s*q\.?|not\s*quoted|by\s*civil\s*contractor.*|.*civil\s*contractor"
    r"\s*scope.*|not\s*in\s*our\s*scope)\s*$",
    re.I,
)
TOTAL_ROW = re.compile(
    r"^\s*(grand\s*total|sub[\s-]*total|total|carried\s*forward|c/?f|b/?f|brought\s*forward)\b",
    re.I,
)
TC_MARKER = re.compile(r"\bterms\s*(&|and)\s*conditions\b", re.I)
# "GENERAL NOTES :", "Notes:", "General Points:" — a block of notes, not items
NOTES_HEAD = re.compile(r"^\s*(general\s+)?(notes?|points)\s*(:|-|$)", re.I)
MAIN_NO = re.compile(r"^\d+(\.\d+)*[a-z]?\.?$", re.I)
SUB_NO = re.compile(r"^\(?([a-h]|i{1,3}|iv|v|vi{0,3}|ix|x)[).]$", re.I)
BARE_SUB = re.compile(r"^[a-h]$")  # "a", "b" (lower case only: "A", "B" are sections)
DASHES = {"-", "--", "—", "–", "nil", "na", "n/a"}


class BoqFormatError(ValueError):
    pass


# --- step 1: reading ---------------------------------------------------------------------------


class Grid(dict):
    """{sheet name: rows}. For a PDF, `pages` gives each row's (first, last) page."""

    pages: dict[str, list[tuple[int, int]]]
    page_count: int | None

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.pages = {}
        self.page_count = None


PDF_SHEET = "PDF"


def read_grid(path: str | Path) -> dict[str, list[list[Any]]]:
    """{sheet name: rows of cell values}. Hidden sheets are left out. A PDF becomes one sheet
    (app.tenders.pdf_import), with the page of every row in `.pages`."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from app.tenders.pdf_import import read_pdf

        pdf = read_pdf(path)
        grid = Grid({PDF_SHEET: pdf.rows})
        grid.pages[PDF_SHEET] = pdf.pages
        grid.page_count = pdf.page_count
        return grid
    if suffix in (".xlsx", ".xlsm"):
        wb = load_workbook(path, data_only=True)  # not read-only: merged ranges are honoured
        sheets = {}
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            rows = []
            for row in ws.iter_rows(
                max_row=min(ws.max_row, MAX_ROWS), max_col=MAX_COLS, values_only=True
            ):
                rows.append(list(row))
            sheets[ws.title] = rows
        wb.close()
        return sheets
    if suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path))
        return {
            sh.name: [
                [
                    sh.cell_value(r, c) if sh.cell_type(r, c) != xlrd.XL_CELL_EMPTY else None
                    for c in range(min(sh.ncols, MAX_COLS))
                ]
                for r in range(min(sh.nrows, MAX_ROWS))
            ]
            for sh in book.sheets()
            if sh.visibility == 0
        }
    if suffix == ".csv":
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else csv.excel
        rows = [
            [c if c != "" else None for c in r][:MAX_COLS]
            for r in csv.reader(io.StringIO(text), dialect)
        ][:MAX_ROWS]
        return {path.stem[:31] or "CSV": rows}
    raise BoqFormatError("Upload an .xlsx, .xls, .csv or .pdf file")


# --- step 2: header row and columns -------------------------------------------------------------


def _norm_header(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def column_letter(index: int) -> str:
    letters, index = "", index + 1
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _match(field_name: str, header: str) -> float:
    exact, keywords = HEADER_RULES[field_name]
    if header in exact:
        return 0.95
    if any(k in header for k in keywords):
        return 0.75
    return 0.0


def _with_above(header_row: list[Any], above: list[Any]) -> tuple[list[Any], dict[int, str]]:
    """Header texts with the row above prefixed: "TOTAL QUANTITY" over "Up to Plinth" and
    "Above plinth" gives "TOTAL QUANTITY Up to Plinth" ... A heading spans to the right over
    cells that are None (merged); an empty string stops it. Returns (texts, {column: above})."""
    combined: list[Any] = []
    spans: dict[int, str] = {}
    current: str | None = None
    for i in range(max(len(header_row), len(above))):
        a = above[i] if i < len(above) else None
        if a is None:
            text = current
        else:
            current = _text(a)
            text = current
        if text:
            spans[i] = text
        own = header_row[i] if i < len(header_row) else None
        combined.append(" ".join(p for p in (text, _text(own)) if p) or own)
    return combined, spans


def guess_columns(
    header_row: list[Any], above: list[Any] | None = None
) -> dict[str, dict[str, Any]]:
    """{field: {column, letter, header, confidence, alternatives}} from one header row, or from
    a header row and the row above it (a heading over two sub-columns)."""
    spans: dict[int, str] = {}
    if above is not None:
        header_row, spans = _with_above(header_row, above)
    headers = {i: _norm_header(v) for i, v in enumerate(header_row) if _norm_header(v)}
    taken: set[int] = set()
    result: dict[str, dict[str, Any]] = {}
    for name in CLAIM_ORDER:
        scored = sorted(
            ((_match(name, h), i) for i, h in headers.items() if i not in taken),
            key=lambda x: (-x[0], x[1]),
        )
        scored = [(s, i) for s, i in scored if s > 0]
        if not scored:
            continue
        confidence, column = scored[0]
        alternatives = [i for s, i in scored[1:] if s == confidence]
        if name == "qty" and alternatives and spans.get(column):
            second = alternatives[0]
            if spans.get(second) == spans[column]:
                # "TOTAL QUANTITY" over "Up to Plinth" and "Above plinth": both are quantity
                taken.add(second)
                result["qty2"] = {
                    "column": second,
                    "letter": column_letter(second),
                    "header": str(header_row[second]).strip(),
                    "confidence": confidence,
                    "alternatives": [],
                }
                alternatives = alternatives[1:]
        if alternatives:  # e.g. "Sika Rate" and "Saint Gobain Rate": first one, less sure
            confidence = min(confidence, 0.6)
        if name == "description2" and "description" not in result:
            # only a "Specification" column: that is the description
            name, confidence = "description", min(confidence, 0.75)
        taken.add(column)
        result[name] = {
            "column": column,
            "letter": column_letter(column),
            "header": str(header_row[column]).strip(),
            "confidence": confidence,
            "alternatives": [column_letter(i) for i in alternatives],
        }
    return result


def _header_score(columns: dict[str, dict[str, Any]]) -> float:
    if "description" not in columns:
        return 0.0
    return sum(HEADER_WEIGHTS.get(f, 0.5) * c["confidence"] for f, c in columns.items())


def header_guess(rows: list[list[Any]], index: int) -> dict[str, dict[str, Any]]:
    """The column guess for the header on row `index` (0-based): the row alone, or with the row
    above when that reads better (a two-row header)."""
    plain = guess_columns(rows[index])
    if index == 0:
        return plain
    two = guess_columns(rows[index], rows[index - 1])
    return two if _header_score(two) > _header_score(plain) + 0.5 else plain


def detect_header(rows: list[list[Any]]) -> tuple[int, dict[str, dict[str, Any]], float]:
    """(0-based header row index, column guess, score) for the best row in the first 30."""
    best = (0, {}, 0.0)
    for index in range(min(len(rows), HEADER_SCAN_ROWS)):
        columns = header_guess(rows, index)
        score = _header_score(columns)
        if score > best[2] + 1e-9:
            best = (index, columns, score)
    return best


def choose_sheet(grid: dict[str, list[list[Any]]]) -> list[dict[str, Any]]:
    """Every sheet with its best header row, best first."""
    sheets = []
    for order, (name, rows) in enumerate(grid.items()):
        header_index, columns, score = detect_header(rows)
        sheets.append(
            {
                "name": name,
                "header_row": header_index + 1,
                "score": round(score, 2),
                "order": order,
                "columns": columns,
            }
        )
    return sorted(sheets, key=lambda s: (-s["score"], s["order"]))


# --- step 3: rows -------------------------------------------------------------------------------


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = " ".join(str(value).split())
    return text or None


def _number(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return Decimal(str(value))
    text = str(value).replace(",", "").replace("₹", "").replace(" ", "").strip()
    try:
        return Decimal(text) if text else None
    except InvalidOperation:
        return None


def _note(value: Any) -> str | None:
    text = _text(value)
    if not text:
        return None
    if QRO_TEXT.match(text):
        return "QRO"
    if NQ_TEXT.match(text):
        return "NQ"
    return None


def _kind(item_no: str | None) -> str | None:
    if not item_no:
        return None
    if SUB_NO.match(item_no) or BARE_SUB.match(item_no):
        return "sub"
    return "main"  # 1, 1.2, 653A, H, "Separate item (657)", "Additional item: 6.1" ...


@dataclass
class Row:
    number: int  # 1-based row in the sheet
    item_no: str | None
    kind: str | None
    description: str
    unit_raw: str | None
    qty: Decimal | None
    qty_note: str | None
    rate: Decimal | None
    product: str | None
    remarks: str | None
    our_remarks: str | None
    has_quantities: bool
    material_rate: Decimal | None = None
    application_rate: Decimal | None = None
    notes_head: bool = False


@dataclass
class ParsedSection:
    title: str
    rows: list[int]
    note: str | None = None  # a long general note that came with the heading


@dataclass
class ParsedLine:
    item_no: str | None
    description: str
    unit: str | None
    unit_raw: str | None
    qty: Decimal | None
    qty_note: str | None
    client_file_rate: Decimal | None
    client_product: str | None
    client_remarks: str | None
    our_remarks: str | None
    section: int | None  # index into sections
    rows: list[int]
    source_row: int | None = None  # the row with the quantity (client format export writes here)
    client_material_rate: Decimal | None = None
    client_application_rate: Decimal | None = None

    @property
    def status(self) -> str:
        return "not_quoted" if self.qty_note == "NQ" else "unpriced"


@dataclass
class ParsedBoq:
    sections: list[ParsedSection] = field(default_factory=list)
    lines: list[ParsedLine] = field(default_factory=list)
    skipped: list[tuple[int, str]] = field(default_factory=list)
    unrecognised_units: dict[str, int] = field(default_factory=dict)
    order: list[tuple[str, int]] = field(default_factory=list)  # ("section"|"line", index)

    def counts(self) -> dict[str, Any]:
        reasons: dict[str, int] = {}
        for _, reason in self.skipped:
            reasons[reason] = reasons.get(reason, 0) + 1
        return {
            "lines": len(self.lines),
            "sections": len(self.sections),
            "qro": sum(1 for ln in self.lines if ln.qty_note == "QRO"),
            "nq": sum(1 for ln in self.lines if ln.qty_note == "NQ"),
            "skipped": len(self.skipped),
            "skipped_by_reason": reasons,
            "unrecognised_units": self.unrecognised_units,
        }


def _money(value: Decimal | None) -> Decimal | None:
    return value.quantize(RATE_PLACES, ROUND_HALF_UP) if value is not None else None


def _join(*parts: str | None) -> str:
    return " — ".join(p for p in parts if p)


HEADING_MAX = 100
PARENT_NOTE_MAX = 400  # a heading note this short is the items' spec: it goes into their text


def split_heading(text: str) -> tuple[str, str | None]:
    """A heading that carries a general note ("600 WATER PROOFING — Unless otherwise specified,
    the Contractor shall ...") -> ("600 WATER PROOFING", "Unless otherwise specified, ...")."""
    text = text.strip()
    if len(text) <= HEADING_MAX:
        return text, None
    if " — " in text:
        head, rest = text.split(" — ", 1)
        if len(head) <= HEADING_MAX:
            return head.strip(), rest.strip()
    m = re.match(r"(.{8,100}?)(?::\s|\.\s|\s[-–]\s)\s*(.+)", text, re.S)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    cut = text[:HEADING_MAX].rsplit(" ", 1)[0]
    return f"{cut}…", text


def _extract(
    row: list[Any], number: int, cols: dict[str, int], labels: dict[str, str] | None = None
) -> tuple[Row | None, str | None]:
    def cell(name: str) -> Any:
        i = cols.get(name)
        return row[i] if i is not None and i < len(row) else None

    if all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
        return None, "blank row"
    if any(isinstance(v, str) and TC_MARKER.search(v) for v in row):
        return None, "terms & conditions"
    item_no = _text(cell("item_no"))
    description = _join(_text(cell("description")), _text(cell("description2")))
    if not description:
        if item_no and NOTES_HEAD.match(item_no):
            return None, "general notes"
        return None, "no description"
    if TOTAL_ROW.match(description):
        return None, "total row"

    unit_raw = _text(cell("unit"))
    qty_cell, rate_cell = cell("qty"), cell("rate")
    qty = _number(qty_cell)
    remarks = _text(cell("remarks"))
    if "qty2" in cols:
        # two quantity columns under one heading: qty is their sum, both kept in the remarks
        second = _number(cell("qty2"))
        if second is not None:
            qty = (qty or Decimal(0)) + second
        if _text(qty_cell) or _text(cell("qty2")):
            labels = labels or {}
            split = "; ".join(
                f"{labels.get(f) or f}: {_text(cell(f)) or '—'}" for f in ("qty", "qty2")
            )
            remarks = _join(split, remarks)
    qty_note = None if qty is not None else _note(qty_cell)
    rate = _number(rate_cell)
    if qty_note is None and rate is None and _note(rate_cell) == "NQ":
        qty_note = "NQ"  # some clients write NQ in the rate column
    if unit_raw and is_rate_only(unit_raw) and qty is None:
        qty_note = qty_note or "QRO"
    has_quantities = bool(unit_raw or qty is not None or qty_note or rate is not None)
    return (
        Row(
            number=number,
            item_no=item_no,
            kind=_kind(item_no),
            description=description,
            unit_raw=unit_raw[:50] if unit_raw else None,
            qty=qty,
            qty_note=qty_note,
            rate=rate,
            product=_text(cell("product")),
            remarks=remarks,
            our_remarks=_text(cell("our_remarks")),
            has_quantities=has_quantities,
            material_rate=_number(cell("material_rate")),
            application_rate=_number(cell("application_rate")),
            notes_head=len(description) <= 30
            and bool(NOTES_HEAD.match(description))
            and not has_quantities,
        ),
        None,
    )


def parse(
    rows: list[list[Any]],
    header_row: int,
    column_map: dict[str, int],
    aliases: dict[str, str] | None = None,
) -> ParsedBoq:
    """Parse the rows below `header_row` (1-based) using {field: 0-based column}."""
    result = ParsedBoq()
    if "description" not in column_map:
        raise BoqFormatError("Choose the description column")
    pending: list[Row] = []
    state = {"section": None, "parent": None, "last": None}  # last: "line" | "heading" | None
    head = rows[header_row - 1] if 0 < header_row <= len(rows) else []
    labels = {f: _text(head[c]) for f, c in column_map.items() if c is not None and c < len(head)}
    in_notes = False

    def add_section(r: Row) -> None:
        short, note = split_heading(r.description)
        title = f"{r.item_no} {short}" if r.item_no else short
        result.sections.append(ParsedSection(title=title[:500], rows=[r.number], note=note))
        result.order.append(("section", len(result.sections) - 1))
        state["section"] = len(result.sections) - 1
        state["parent"] = _join(short, note) if note and len(note) <= PARENT_NOTE_MAX else short
        state["siblings"] = False
        state["last"] = "heading"

    def continue_last_line(r: Row) -> None:
        line = result.lines[-1]
        line.description = f"{line.description} {r.description}"
        line.rows.append(r.number)
        line.client_remarks = _join(line.client_remarks, r.remarks) or None
        line.our_remarks = _join(line.our_remarks, r.our_remarks) or None
        line.client_product = line.client_product or r.product

    def flush_unclaimed(rows_: list[Row]) -> None:
        """Rows without quantities that do not belong to the next line."""
        for r in rows_:
            if r.kind == "main":
                add_section(r)
            elif state["last"] == "line" and r.kind is None:
                continue_last_line(r)
            else:
                result.skipped.append((r.number, "note without quantity"))

    def add_line(item_no: str | None, description: str, x: Row, extra_rows: list[Row]) -> None:
        unit = normalise_unit(x.unit_raw, aliases) if x.unit_raw else None
        if x.unit_raw and unit is None and not is_rate_only(x.unit_raw):
            result.unrecognised_units[x.unit_raw] = result.unrecognised_units.get(x.unit_raw, 0) + 1
        rows_ = [r.number for r in extra_rows] + [x.number]
        result.lines.append(
            ParsedLine(
                item_no=item_no[:50] if item_no else None,
                description=description,
                unit=unit,
                unit_raw=x.unit_raw,
                qty=x.qty.quantize(QTY_PLACES, ROUND_HALF_UP) if x.qty is not None else None,
                qty_note=x.qty_note,
                client_file_rate=(
                    x.rate.quantize(RATE_PLACES, ROUND_HALF_UP) if x.rate is not None else None
                ),
                client_product=x.product,
                client_material_rate=_money(x.material_rate),
                client_application_rate=_money(x.application_rate),
                client_remarks=_join(*(r.remarks for r in extra_rows), x.remarks) or None,
                our_remarks=_join(*(r.our_remarks for r in extra_rows), x.our_remarks) or None,
                section=state["section"],
                rows=rows_,
                source_row=x.number,
            )
        )
        result.order.append(("line", len(result.lines) - 1))
        state["last"] = "line"

    def resolve(x: Row) -> None:
        numbered = [i for i, r in enumerate(pending) if r.kind == "main"]
        if x.kind == "main" or not numbered:
            flush_unclaimed(pending)
            if x.kind == "main":
                state["parent"] = None
                state["siblings"] = False
                add_line(x.item_no, x.description, x, [])
            elif x.kind == "sub" and state["parent"]:
                add_line(x.item_no, _join(state["parent"], x.description), x, [])
            elif x.kind is None and state.get("siblings") and state["parent"] and not pending:
                # "6 Sealing around pipes" + "50mm dia pipe" (qty), then "75mm dia pipe" (qty)
                add_line(None, _join(state["parent"], x.description), x, [])
            else:
                add_line(x.item_no, x.description, x, [])
            return
        j = numbered[-1]
        before, title, after = pending[:j], pending[j], pending[j + 1 :]
        flush_unclaimed(before)
        if x.kind is None:
            # "1.2 PROTECTION SCREED" + next row with the quantities: one item
            add_line(
                title.item_no,
                _join(title.description, *(r.description for r in after), x.description),
                x,
                [title, *after],
            )
            state["parent"] = _join(title.description, *(r.description for r in after))
            state["siblings"] = True
        elif any(r.kind == "sub" for r in after):
            # "4 ..." then steps i) ... iv) and the quantity on v): one item with all the steps
            steps = [_join(r.item_no, r.description) for r in after] + [
                _join(x.item_no, x.description)
            ]
            add_line(title.item_no, _join(title.description, *steps), x, [title, *after])
            state["parent"] = None
        else:
            # "1 Surface preparation ..." then "a) For the terrace floor": a heading and sub-items
            add_section(title)
            state["parent"] = _join(state["parent"], *(r.description for r in after))
            add_line(x.item_no, _join(state["parent"], x.description), x, [])

    for index, values in enumerate(rows[header_row:], start=header_row + 1):
        row, reason = _extract(values, index, column_map, labels)
        if reason == "terms & conditions":
            result.skipped.append((index, "terms & conditions block"))
            remaining = sum(1 for v in rows[index:] if any(c not in (None, "") for c in v))
            result.skipped.extend(
                (index + k + 1, "terms & conditions block") for k in range(remaining)
            )
            break
        if row is None:
            if reason != "blank row":
                result.skipped.append((index, reason))
            else:
                result.skipped.append((index, "blank row"))
                in_notes = False  # a notes block ends at a blank row
            continue
        if row.notes_head:
            in_notes = True
        if in_notes and not row.has_quantities:
            result.skipped.append((index, "general notes"))
            continue
        in_notes = False
        if row.has_quantities:
            resolve(row)
            pending = []
        else:
            pending.append(row)
    flush_unclaimed(pending)

    _drop_empty_sections(result)
    return result


def _drop_empty_sections(result: ParsedBoq) -> None:
    used = {ln.section for ln in result.lines}
    keep = [i for i in range(len(result.sections)) if i in used]
    for i in range(len(result.sections)):
        if i not in used:
            for row in result.sections[i].rows:
                result.skipped.append((row, "heading with no lines under it"))
    remap = {old: new for new, old in enumerate(keep)}
    result.sections = [result.sections[i] for i in keep]
    for line in result.lines:
        line.section = remap.get(line.section) if line.section is not None else None
    result.order = [
        (kind, remap[i] if kind == "section" else i)
        for kind, i in result.order
        if kind == "line" or i in remap
    ]
    result.skipped.sort()


def page_label(pages: list[tuple[int, int]] | None, rows: list[int]) -> str | None:
    """ "p. 3" or "p. 3–4" for grid rows (1-based) of a PDF; None for spreadsheets."""
    if not pages or not rows:
        return None
    found = [pages[r - 1] for r in rows if 0 < r <= len(pages)]
    if not found:
        return None
    first, last = min(p[0] for p in found), max(p[1] for p in found)
    return f"p. {first}" if first == last else f"p. {first}–{last}"


def preview_rows(
    result: ParsedBoq, limit: int = PREVIEW_ROWS, pages: list[tuple[int, int]] | None = None
) -> list[dict[str, Any]]:
    out = []
    for kind, i in result.order[:limit]:
        if kind == "section":
            s = result.sections[i]
            out.append(
                {
                    "type": "section",
                    "title": s.title,
                    "note": s.note,
                    "rows": s.rows,
                    "page": page_label(pages, s.rows),
                }
            )
        else:
            ln = result.lines[i]
            out.append(
                {
                    "type": "line",
                    "item_no": ln.item_no,
                    "description": ln.description,
                    "unit": ln.unit,
                    "unit_raw": ln.unit_raw,
                    "qty": ln.qty,
                    "qty_note": ln.qty_note,
                    "client_file_rate": ln.client_file_rate,
                    "client_product": ln.client_product,
                    "client_remarks": ln.client_remarks,
                    "our_remarks": ln.our_remarks,
                    "status": ln.status,
                    "rows": ln.rows,
                    "page": page_label(pages, ln.rows),
                    "client_material_rate": ln.client_material_rate,
                    "client_application_rate": ln.client_application_rate,
                }
            )
    return out
