"""EESPL BOQ export: Excel and PDF in EESPL's standard layout, and "client format" (our rates
written into the client's own uploaded sheet).

Everything is built from a revision snapshot (app.tenders.revisions), live or frozen, which only
holds selling data: cost, margin, price source, suggestion scores and the client's own rates are
never part of an export.
"""

import base64
import html
import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from fastapi import HTTPException, status
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import settings
from app.masters.models import CompanyGstin, CompanyProfile
from app.tenders import revisions
from app.tenders.models import BoqImport, Tender

SUBJECT = "Waterproofing works"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
INR_FORMAT = '[>=10000000]"₹"##\\,##\\,##\\,##0.00;[>=100000]"₹"##\\,##\\,##0.00;"₹"#,##0.00'
QTY_FORMAT = "#,##0.###"


# --- the document -------------------------------------------------------------------------------


@dataclass
class Company:
    name: str
    gstin: str | None
    address: str | None
    logo: Path | None


@dataclass
class Document:
    company: Company
    snapshot: dict[str, Any]
    revision: str
    on: date
    gst_percent: Decimal

    @property
    def tender(self) -> dict[str, Any]:
        return self.snapshot["tender"]

    def groups(self) -> list[tuple[dict | None, list[dict]]]:
        """(section or None, its lines): lines without a section first, then each section."""
        by_section: dict[int | None, list[dict]] = {}
        for ln in self.snapshot["lines"]:
            by_section.setdefault(ln["section_id"], []).append(ln)
        out: list[tuple[dict | None, list[dict]]] = []
        if by_section.get(None):
            out.append((None, by_section[None]))
        out.extend((s, by_section.get(s["id"], [])) for s in self.snapshot["sections"])
        return out


def company(db: Session) -> Company:
    profile = db.get(CompanyProfile, 1)
    gstin = db.scalar(
        select(CompanyGstin).order_by(CompanyGstin.is_default.desc(), CompanyGstin.id).limit(1)
    )
    logo = None
    if profile and profile.logo_path:
        path = Path(settings.media_dir) / profile.logo_path
        logo = path if path.exists() else None
    name = (profile.legal_name or profile.trade_name) if profile else None
    return Company(
        name=name or "Ethios Enviro Solutions Pvt. Ltd.",
        gstin=gstin.gstin if gstin else None,
        address=gstin.address if gstin else None,
        logo=logo,
    )


def document(db: Session, tender: Tender, rev_no: int | None = None) -> Document:
    """The live tender (rev_no None) or a submitted revision, rebuilt from its snapshot."""
    if rev_no is None:
        snapshot = revisions.build_snapshot(db, tender)
        frozen = revisions.frozen(db, tender)
        label = revisions.revision_label(db, tender)
        on = frozen.submitted_at.date() if frozen else date.today()
    else:
        rev = revisions.get_revision(db, tender, rev_no)
        snapshot, label, on = rev.snapshot, revisions.label(rev.rev_no), rev.submitted_at.date()
    return Document(company(db), snapshot, label, on, Decimal(snapshot["totals"]["gst_percent"]))


def inr(value: Decimal | str | None) -> str:
    """₹12,34,567.89 (Indian grouping)."""
    if value is None or value == "":
        return ""
    v = Decimal(value).quantize(Decimal("0.01"))
    sign = "-" if v < 0 else ""
    whole, frac = f"{abs(v):.2f}".split(".")
    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        groups = []
        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]
        if head:
            groups.insert(0, head)
        whole = ",".join(groups + [tail])
    return f"{sign}₹{whole}.{frac}"


def qty_text(line: dict) -> str:
    if line["qty_note"]:
        return line["qty_note"]
    if line["qty"] is None:
        return ""
    return f"{Decimal(line['qty']).normalize():f}"


def _unit(line: dict) -> str:
    return line["unit"] or line["unit_raw"] or ""


def file_stem(doc: Document) -> str:
    return f"{doc.tender['code']}_{doc.revision.split(' ')[0]}_BOQ"


def gst_line(doc: Document) -> str:
    return f"GST @{doc.gst_percent.normalize():f}% extra at actual"


# --- Excel --------------------------------------------------------------------------------------

COLUMNS = [
    ("Sr no", 9),
    ("Description", 62),
    ("Unit", 8),
    ("Qty", 11),
    ("Rate", 15),
    ("Amount", 17),
    ("Product / brand", 20),
    ("Remarks", 26),
]
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="DCE6E0")
SECTION_FILL = PatternFill("solid", fgColor="EEF3F0")
WRAP = Alignment(wrap_text=True, vertical="top")


def _height(text: str, chars_per_line: int, line_points: float = 15) -> float:
    lines = sum(max(1, -(-len(part) // chars_per_line)) for part in text.split("\n"))
    return max(line_points, lines * line_points)


def xlsx(doc: Document) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"
    last_col = get_column_letter(len(COLUMNS))
    for i, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    def merged(
        row: int,
        value: str,
        *,
        start: int = 1,
        font: Font | None = None,
        height_chars: int | None = None,
        fill: PatternFill | None = None,
    ) -> None:
        ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=len(COLUMNS))
        cell = ws.cell(row=row, column=start, value=value)
        cell.alignment = WRAP
        if font:
            cell.font = font
        if fill:
            for col in range(start, len(COLUMNS) + 1):
                ws.cell(row=row, column=col).fill = fill
        if height_chars:
            ws.row_dimensions[row].height = _height(value, height_chars)

    # header
    if doc.company.logo:
        try:
            from openpyxl.drawing.image import Image as XlImage
            from PIL import Image as PilImage

            with PilImage.open(doc.company.logo) as im:
                im = im.convert("RGBA")
                ratio = 48 / im.height
                buf = io.BytesIO()
                im.resize((max(1, int(im.width * ratio)), 48)).save(buf, format="PNG")
            ws.add_image(XlImage(io.BytesIO(buf.getvalue())), "A1")
        except Exception:  # an unreadable logo never blocks the export
            pass
    merged(1, doc.company.name, start=2, font=Font(bold=True, size=14))
    ws.row_dimensions[1].height = 24
    merged(2, f"GSTIN: {doc.company.gstin}" if doc.company.gstin else "", start=2)
    merged(3, doc.company.address or "", start=2, height_chars=150)
    t = doc.tender
    site = ", ".join(p for p in (t["site_name"], t["site_city"], t["site_state"]) if p)
    rows = [
        f"Tender: {t['code']}  ·  Revision {doc.revision}  ·  Date: {doc.on:%d-%m-%Y}",
        f"Client: {t['client_name']}",
        f"Site: {site or '—'}",
        f"Subject: {SUBJECT}",
    ]
    for i, text in enumerate(rows, start=5):
        merged(i, text, font=Font(bold=i == 8))
    header_row = 10
    for i, (title, _) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=header_row, column=i, value=title)
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row = header_row + 1
    first_data = row
    for section, lines in doc.groups():
        if section is not None:
            merged(row, section["title"], font=Font(bold=True), fill=SECTION_FILL)
            row += 1
            if section["note"]:
                merged(row, section["note"], font=Font(italic=True, size=9), height_chars=230)
                row += 1
        start = row
        for ln in lines:
            values: list[Any] = [ln["client_item_no"], ln["description"], _unit(ln)]
            if ln["qty_note"]:
                values.append(ln["qty_note"])
            else:
                values.append(Decimal(ln["qty"]) if ln["qty"] is not None else None)
            if ln["qty_note"] == "NQ" or ln["status"] == "not_quoted":
                values += ["Not quoted", None]
            elif ln["rate"] is None:
                values += [None, None]
            else:
                values.append(Decimal(ln["rate"]))
                if ln["qty_note"] == "QRO":
                    values.append("Rate only")
                elif ln["qty"] is None:
                    values.append(None)
                else:
                    values.append(f"=ROUND(D{row}*E{row},2)")
            values += [ln["our_product"], ln["our_remarks"]]
            for col, value in enumerate(values, start=1):
                c = ws.cell(row=row, column=col, value=value)
                c.border = BOX
                c.alignment = WRAP
            ws.cell(row=row, column=4).number_format = QTY_FORMAT
            ws.cell(row=row, column=5).number_format = INR_FORMAT
            ws.cell(row=row, column=6).number_format = INR_FORMAT
            for col in (4, 5, 6):
                if isinstance(ws.cell(row=row, column=col).value, str):
                    ws.cell(row=row, column=col).alignment = Alignment(
                        horizontal="right", vertical="top"
                    )
            row += 1
        if section is not None:
            ws.cell(row=row, column=2, value=f"Total — {section['title']}").font = Font(bold=True)
            ws.cell(row=row, column=2).alignment = Alignment(horizontal="right", wrap_text=True)
            total = ws.cell(
                row=row, column=6, value=f"=SUBTOTAL(9,F{start}:F{max(start, row - 1)})"
            )
            total.font = Font(bold=True)
            total.number_format = INR_FORMAT
            total.border = BOX
            row += 1
    last_data = max(first_data, row - 1)
    row += 1
    ws.cell(row=row, column=2, value="Total (excl. GST)").font = Font(bold=True, size=12)
    ws.cell(row=row, column=2).alignment = Alignment(horizontal="right")
    grand = ws.cell(row=row, column=6, value=f"=SUBTOTAL(9,F{first_data}:F{last_data})")
    grand.font = Font(bold=True, size=12)
    grand.number_format = INR_FORMAT
    grand.border = BOX
    row += 1
    merged(row, gst_line(doc), start=2, font=Font(italic=True))
    row += 2

    if doc.snapshot["tc"]:
        merged(row, "Terms & Conditions", font=Font(bold=True))
        row += 1
        for i, text in enumerate(doc.snapshot["tc"], start=1):
            ws.cell(row=row, column=1, value=f"{i}.").alignment = Alignment(vertical="top")
            merged(row, text, start=2, height_chars=170)
            row += 1
        row += 1
    merged(row, f"For {doc.company.name}", font=Font(bold=True))
    row += 3
    merged(row, "Authorised signatory")

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.print_area = f"A1:{last_col}{row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.6)
    ws.oddFooter.center.text = "Page &P of &N"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# --- PDF ----------------------------------------------------------------------------------------

PDF_CSS = """
@page { size: A4 landscape; margin: 12mm 10mm 14mm;
  @bottom-center { content: "Page " counter(page) " of " counter(pages); font-family: "DejaVu Sans";
    font-size: 8pt; color: #555; } }
body { font-family: "DejaVu Sans", sans-serif; font-size: 8.5pt; color: #111; }
.head { display: flex; gap: 12px; align-items: center; border-bottom: 1.5px solid #2f6f4f;
  padding-bottom: 6px; }
.head img { height: 46px; }
.head h1 { font-size: 14pt; margin: 0; }
.meta { margin: 8px 0; display: grid; grid-template-columns: 1fr 1fr; gap: 2px 24px; }
.subject { font-weight: bold; margin: 4px 0 8px; }
table { width: 100%; border-collapse: collapse; }
thead { display: table-header-group; }
th { background: #dce6e0; }
th, td { border: 0.6px solid #999; padding: 3px 4px; vertical-align: top; }
td.num, th.num { text-align: right; white-space: nowrap; }
td.desc { width: 40%; }
tr { page-break-inside: avoid; }
tr.section td { background: #eef3f0; font-weight: bold; }
tr.note td { font-style: italic; font-size: 7.5pt; color: #333; }
tr.subtotal td { font-weight: bold; }
.totals { margin-top: 8px; width: 45%; margin-left: auto; }
.totals td { border: none; }
.totals .grand td { font-size: 10pt; font-weight: bold; border-top: 1px solid #333; }
h2 { font-size: 10pt; margin: 14px 0 4px; }
.tc { display: flex; gap: 6px; margin-bottom: 3px; page-break-inside: avoid; }
.tc .no { min-width: 18px; text-align: right; }
.tc .text { white-space: pre-line; }
.sign { margin-top: 26px; page-break-inside: avoid; }
.sign .line { margin-top: 36px; }
"""


def _e(value: Any) -> str:
    return html.escape("" if value is None else str(value))


def html_document(doc: Document) -> str:
    c, t = doc.company, doc.tender
    logo = ""
    if c.logo:
        mime = {
            ".png": "image/png",
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
            ".webp": "image/webp",
        }.get(c.logo.suffix.lower(), "image/png")
        logo = f'<img src="data:{mime};base64,{base64.b64encode(c.logo.read_bytes()).decode()}">'
    site = ", ".join(p for p in (t["site_name"], t["site_city"], t["site_state"]) if p) or "—"
    body = []
    for section, lines in doc.groups():
        if section is not None:
            body.append(f'<tr class="section"><td colspan="8">{_e(section["title"])}</td></tr>')
            if section["note"]:
                body.append(f'<tr class="note"><td colspan="8">{_e(section["note"])}</td></tr>')
        for ln in lines:
            nq = ln["qty_note"] == "NQ" or ln["status"] == "not_quoted"
            rate = "Not quoted" if nq else inr(ln["rate"])
            amount = "Rate only" if ln["qty_note"] == "QRO" and ln["rate"] else inr(ln["amount"])
            body.append(
                "<tr>"
                f'<td>{_e(ln["client_item_no"])}</td><td class="desc">{_e(ln["description"])}</td>'
                f"<td>{_e(_unit(ln))}</td><td class=\"num\">{_e(qty_text(ln))}</td>"
                f'<td class="num">{_e(rate)}</td><td class="num">{_e("" if nq else amount)}</td>'
                f'<td>{_e(ln["our_product"])}</td><td>{_e(ln["our_remarks"])}</td>'
                "</tr>"
            )
        if section is not None:
            body.append(
                f'<tr class="subtotal"><td></td><td colspan="4" class="num">Total — '
                f'{_e(section["title"])}</td><td class="num">{_e(inr(section["total"]))}</td>'
                "<td colspan=\"2\"></td></tr>"
            )
    tc = "".join(
        f'<div class="tc"><span class="no">{i}.</span><span class="text">{_e(text)}</span></div>'
        for i, text in enumerate(doc.snapshot["tc"], start=1)
    )
    totals = doc.snapshot["totals"]
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>{PDF_CSS}</style></head>
<body>
<div class="head">{logo}<div><h1>{_e(c.name)}</h1>
<div>{_e(f"GSTIN: {c.gstin}" if c.gstin else "")}</div><div>{_e(c.address)}</div></div></div>
<div class="meta">
<div><b>Tender:</b> {_e(t["code"])} · Revision {_e(doc.revision)}</div>
<div><b>Date:</b> {doc.on:%d-%m-%Y}</div>
<div><b>Client:</b> {_e(t["client_name"])}</div><div><b>Site:</b> {_e(site)}</div>
</div>
<div class="subject">Subject: {_e(SUBJECT)}</div>
<table><thead><tr><th>Sr no</th><th>Description</th><th>Unit</th><th class="num">Qty</th>
<th class="num">Rate</th><th class="num">Amount</th><th>Product / brand</th><th>Remarks</th>
</tr></thead><tbody>{"".join(body)}</tbody></table>
<table class="totals">
<tr class="grand"><td>Total (excl. GST)</td><td class="num">{_e(inr(totals["subtotal"]))}</td></tr>
<tr><td colspan="2">{_e(gst_line(doc))}</td></tr>
</table>
{f"<h2>Terms &amp; Conditions</h2>{tc}" if tc else ""}
<div class="sign"><b>For {_e(c.name)}</b><div class="line">Authorised signatory</div></div>
</body></html>"""


def pdf(doc: Document) -> bytes:
    from weasyprint import HTML

    return HTML(string=html_document(doc)).write_pdf()


# --- client format ------------------------------------------------------------------------------


def client_format(db: Session, tender: Tender, doc: Document) -> tuple[bytes, str, dict[str, int]]:
    """Our rates (and amounts, where the client's amount cell is not already a formula) written
    into the client's uploaded sheet, keeping its layout. Returns (bytes, filename, counts)."""
    imp = db.scalar(
        select(BoqImport)
        .where(BoqImport.tender_id == tender.id)
        .order_by(BoqImport.imported_at.desc(), BoqImport.id.desc())
        .limit(1)
    )
    if imp is None:
        raise HTTPException(status.HTTP_409_CONFLICT, "This tender has no imported client BOQ")
    path = Path(settings.media_dir) / imp.stored_path
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            f"The client's file is {path.suffix or 'not a spreadsheet'}: client format needs an "
            ".xlsx original. Export the EESPL format instead, or upload the BOQ again as .xlsx.",
        )
    rate_col = imp.column_map.get("rate")
    if rate_col is None:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            "No rate column was mapped when this BOQ was imported, so there is nowhere to put our "
            "rates. Import it again with the rate column chosen.",
        )
    amount_col = imp.column_map.get("amount")
    if not path.exists():
        raise HTTPException(status.HTTP_409_CONFLICT, "The uploaded client file is missing")
    wb = load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
    if imp.sheet not in wb.sheetnames:
        raise HTTPException(status.HTTP_409_CONFLICT, f"Sheet {imp.sheet!r} is not in the file")
    ws = wb[imp.sheet]
    counts = {"rates": 0, "amounts": 0, "skipped": 0}

    def put(row: int, col: int, value: Any) -> bool:
        cell = ws.cell(row=row, column=col + 1)
        if isinstance(cell, MergedCell):
            return False
        cell.value = value
        return True

    for ln in doc.snapshot["lines"]:
        row = ln.get("source_row")
        if not row or ln["rate"] is None or ln["qty_note"] == "NQ" or ln["status"] == "not_quoted":
            counts["skipped"] += bool(row is None or ln["rate"] is None)
            continue
        if put(row, rate_col, float(Decimal(ln["rate"]))):
            counts["rates"] += 1
        else:
            counts["skipped"] += 1
            continue
        if amount_col is not None and ln["amount"] is not None:
            existing = ws.cell(row=row, column=amount_col + 1).value
            if not (isinstance(existing, str) and existing.startswith("=")):
                if put(row, amount_col, float(Decimal(ln["amount"]))):
                    counts["amounts"] += 1
            else:
                counts["amounts"] += 1  # their formula computes it from our rate
    out = io.BytesIO()
    wb.save(out)
    data = out.getvalue()
    original = path.name.split("__", 1)[-1]
    name = f"{Path(original).stem}_{tender.code}_{doc.revision.split(' ')[0]}{path.suffix.lower()}"
    folder = Path(settings.media_dir) / "tenders" / str(tender.id) / "exports"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / f"{datetime.now():%Y%m%d-%H%M%S}__{name}").write_bytes(data)
    return data, name, counts
