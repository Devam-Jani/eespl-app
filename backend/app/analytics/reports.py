# ruff: noqa: E501  (HTML templates read better unwrapped)
"""Exports: Excel for every tile list and analytics table; PDFs in the EESPL format (management
summary, site status report for client meetings, receivables, sales pipeline); the weekly
management summary the worker files every Monday 8 am IST."""

from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from typing import Any

from fastapi.encoders import jsonable_encoder
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.analytics import kpi
from app.analytics.alerts import visible
from app.analytics.common import ZERO, inr_compact, money, now, pct, today
from app.analytics.models import RULE_LABELS, Alert, InvoiceSummary, SiteSummary, WeeklyReport
from app.analytics.summary import forecast_text
from app.execution.models import Dpr, Inspection, Mom, MomPoint
from app.execution.pdf import _doc, e
from app.execution.service import IST
from app.export import _cell_value
from app.masters.models import Client
from app.models import Role, RolePermission, User, UserRole
from app.portal import service as portal
from app.portal.models import Snag
from app.sites.models import Site, SiteNode
from app.tenders.export import inr

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# --- Excel ---------------------------------------------------------------------------------------


def xlsx_book(sheets: list[tuple], note: str | None = None) -> bytes:
    """[(title, columns [{key, label}], rows [dict])] -> one sheet each; a note (what the numbers
    count) goes above the header of every sheet."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, columns, rows in sheets:
        ws = wb.create_sheet(
            title=("".join(ch for ch in title if ch not in "[]:*?/\\")[:31]) or "Sheet"
        )
        if note:
            ws.append([note])
            ws.cell(row=1, column=1).font = Font(italic=True)
        head = 2 if note else 1
        ws.append([c["label"] for c in columns])
        for cell in ws[head]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F6B5C")
        for r in rows:
            ws.append([_cell_value(r.get(c["key"])) for c in columns])
        for i, c in enumerate(columns):
            width = (
                max([len(str(c["label"]))] + [len(str(r.get(c["key"]) or "")) for r in rows[:200]])
                + 2
            )
            ws.column_dimensions[ws.cell(row=head, column=i + 1).column_letter].width = min(
                60, width
            )
        ws.freeze_panes = f"A{head + 1}"
    if not wb.sheetnames:
        wb.create_sheet("Empty")
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def tiles_sheet(data: dict) -> tuple[str, list[dict], list[dict]]:
    """The tiles of a dashboard as one sheet."""
    rows = []
    for sec in data.get("sections") or [{"title": "", "tiles": data.get("tiles", [])}]:
        for t in sec["tiles"]:
            rows.append(
                {
                    "section": sec.get("title"),
                    "label": t["label"],
                    "value": t["value"],
                    "unit": t["unit"],
                    "prev": t.get("prev"),
                    "prev_label": t.get("prev_label"),
                    "note": t.get("note"),
                }
            )
    if data.get("margin"):
        for t in data["margin"]["tiles"]:
            rows.append(
                {
                    "section": "Margin",
                    "label": t["label"],
                    "value": t["value"],
                    "unit": t["unit"],
                    "note": t.get("note"),
                }
            )
    cols = [
        {"key": "section", "label": "Section"},
        {"key": "label", "label": "Tile"},
        {"key": "value", "label": "Value"},
        {"key": "unit", "label": "Unit"},
        {"key": "prev", "label": "Previous"},
        {"key": "prev_label", "label": "Previous = "},
        {"key": "note", "label": "Note"},
    ]
    return "Dashboard", cols, rows


# --- formatting ----------------------------------------------------------------------------------


def fmt(value: Any, unit: str = "inr") -> str:
    if value is None or value == "":
        return "—"
    if unit == "inr":
        return inr(Decimal(str(value)))
    if unit == "pct":
        return f"{Decimal(str(value)):.1f}%"
    if unit == "count":
        return f"{int(Decimal(str(value))):,}"
    return str(value)


def tile_fmt(value: Any, unit: str) -> str:
    """Tiles show compact money (₹29.49 Cr); tables keep exact rupees."""
    return inr_compact(value) if unit == "inr" and value not in (None, "") else fmt(value, unit)


def change(t: dict) -> str:
    if t.get("prev") in (None, "") or t["unit"] not in ("inr", "count", "pct"):
        return ""
    now_v, prev_v = Decimal(str(t["value"] or 0)), Decimal(str(t["prev"] or 0))
    diff = now_v - prev_v
    arrow = "▲" if diff > 0 else "▼" if diff < 0 else "="
    return f"{arrow} {tile_fmt(abs(diff), t['unit'])} vs {e(t.get('prev_label') or 'before')}"


def _when(as_of) -> str:
    if not as_of:
        return "not yet refreshed"
    dt = datetime.fromisoformat(str(as_of)) if not isinstance(as_of, datetime) else as_of
    return f"{dt.astimezone(IST):%d %b %Y %H:%M} IST"


# --- management summary --------------------------------------------------------------------------


def management_pdf(
    db: Session,
    data: dict,
    show_margin: bool,
    alerts: list[dict],
    delayed: list[dict],
    title_note: str = "",
) -> bytes:
    data = jsonable_encoder(data)
    blocks = []
    for sec in data["sections"]:
        cells = "".join(
            f"<td><div class='muted'>{e(t['label'])}</div><div class='big'>{tile_fmt(t['value'], t['unit'])}</div>"
            f"<div class='muted small'>{change(t)}{(' · ' + e(t['note'])) if t.get('note') else ''}</div></td>"
            for t in sec["tiles"]
        )
        blocks.append(f"<h2>{e(sec['title'])}</h2><table class='tiles'><tr>{cells}</tr></table>")
    lists = []
    if show_margin and data.get("margin"):
        m = data["margin"]
        cells = "".join(
            f"<td><div class='muted'>{e(t['label'])}</div><div class='big'>{tile_fmt(t['value'], t['unit'])}</div>"
            f"<div class='muted small'>{e(t.get('note') or '')}</div></td>"
            for t in m["tiles"]
        )
        blocks.append(f"<h2>Margin</h2><table class='tiles'><tr>{cells}</tr></table>")
        codes = dict(db.execute(select(Site.id, Site.code)).all())
        low = "".join(
            f"<tr><td>{e(codes.get(r['site_id']))}</td><td class='num'>{fmt(r['margin'], 'pct')}</td><td class='num'>{fmt(r['billed'])}</td></tr>"
            for r in m["lowest"]
        )
        lists.append(
            f"<div><h2>Lowest margin</h2><table class='mini'><thead><tr><th>Site</th><th class='num'>Margin</th><th class='num'>Billed</th></tr></thead>"
            f"<tbody>{low or '<tr><td colspan=3 class=muted>No billing yet.</td></tr>'}</tbody></table></div>"
        )
    rows = "".join(f"<tr><td>{e(a['rule_label'])}: {e(a['title'])}</td></tr>" for a in alerts[:5])
    late = "".join(
        f"<tr><td>{e(s['code'])}</td><td>{e(s['delay'])}</td><td class='num'>{fmt(s['progress'], 'pct')}</td><td class='num'>{fmt(s['elapsed'], 'pct')}</td></tr>"
        for s in delayed[:5]
    )
    more = (
        f"<tr><td colspan=4 class='muted'>and {len(delayed) - 5} more</td></tr>"
        if len(delayed) > 5
        else ""
    )
    lists.append(
        f"<div><h2>Open alerts</h2><table class='mini'><tbody>{rows or '<tr><td class=muted>None.</td></tr>'}</tbody></table></div>"
    )
    lists.append(
        f"<div><h2>Delayed sites</h2><table class='mini'><thead><tr><th>Site</th><th>Why</th><th class='num'>Done</th><th class='num'>Time gone</th></tr></thead>"
        f"<tbody>{late or '<tr><td colspan=4 class=muted>None.</td></tr>'}{more}</tbody></table></div>"
    )
    blocks.append(f"<div class='lists'>{''.join(lists)}</div>")
    # one landscape page: tight tiles, the three short lists side by side
    style = (
        "<style>@page { size: A4 landscape; margin: 7mm 9mm 9mm; } body { font-size: 7.5pt; } .head img { height: 34px; }"
        " .title { margin: 3px 0 2px; } .tiles td { width: 25%; padding: 2px 5px; } .big { font-size: 10.5pt; font-weight: bold; color: #2f6f4f; }"
        " .small { font-size: 6.5pt; } .mini td, .mini th { padding: 1px 3px; font-size: 7pt; } .lists { display: flex; gap: 8px; }"
        " .lists > div { flex: 1; } h2 { margin: 4px 0 2px; font-size: 8.5pt; } tr, table { page-break-inside: avoid; }</style>"
    )
    head = f"<div class='muted'>As of {e(_when(data.get('as_of')))}{(' · ' + e(title_note)) if title_note else ''}</div>"
    return _doc(db, "MANAGEMENT SUMMARY", style + head + "".join(blocks))


def delayed_sites(db: Session, demo: bool) -> list[dict]:
    rows = db.execute(
        select(SiteSummary, Site.code)
        .join(Site, Site.id == SiteSummary.site_id)
        .where(SiteSummary.is_demo == demo, SiteSummary.delayed)
        .order_by(SiteSummary.progress - func.coalesce(SiteSummary.elapsed, 0))
    ).all()
    return [
        {"code": code, "delay": s.delay_reason, "progress": s.progress, "elapsed": s.elapsed}
        for s, code in rows
    ]


# --- site status (for client meetings: no costs) -------------------------------------------------


def site_status_pdf(db: Session, site: Site) -> bytes:
    s = db.get(SiteSummary, site.id)
    client = db.get(Client, site.client_id) if site.client_id else None
    nodes = db.scalars(
        select(SiteNode)
        .where(
            SiteNode.site_id == site.id, SiteNode.kind.in_(("tower", "wing", "floor", "basement"))
        )
        .order_by(SiteNode.parent_id.nulls_first(), SiteNode.sort_order, SiteNode.id)
    ).all()
    parts = "".join(
        f"<tr><td>{'&nbsp;&nbsp;' if n.kind in ('floor', 'basement') else ''}{e(n.name)}</td><td class='num'>{Decimal(n.progress_percent):.0f}%</td></tr>"
        for n in nodes
    )
    dprs = db.scalars(
        select(Dpr)
        .where(Dpr.site_id == site.id, Dpr.status != "draft")
        .order_by(Dpr.on_date.desc())
        .limit(5)
    ).all()
    dpr_rows = "".join(
        f"<tr><td class='nowrap'>{d.on_date:%d %b}</td><td class='pre'>{e((d.work_done or '')[:400])}</td></tr>"
        for d in dprs
    )
    snags = db.scalars(
        select(Snag)
        .where(Snag.site_id == site.id, Snag.status.in_(("open", "in_progress", "fixed")))
        .order_by(Snag.id)
    ).all()
    snag_rows = "".join(
        f"<tr><td>{e(x.code)}</td><td>{e(x.title)}</td><td>{e(x.area or '')}</td><td>{e(x.status.replace('_', ' '))}</td></tr>"
        for x in snags
    )
    insp = db.scalars(
        select(Inspection)
        .where(Inspection.site_id == site.id)
        .order_by(Inspection.on_date.desc())
        .limit(8)
    ).all()
    insp_rows = "".join(
        f"<tr><td>{e(i.code)}</td><td>{i.on_date:%d %b}</td><td>{e(i.template.name)}</td><td class='{'pass' if i.result != 'fail' else 'fail'}'>{e(i.result.replace('_', ' '))}</td></tr>"
        for i in insp
    )
    points = db.execute(
        select(MomPoint, Mom.code)
        .join(Mom, Mom.id == MomPoint.mom_id)
        .where(Mom.site_id == site.id, MomPoint.status == "open")
    ).all()
    point_rows = "".join(
        f"<tr><td>{e(code)}</td><td>{e(p.text)}</td><td>{e(p.owner_name or 'EESPL')}</td><td>{p.due_date.strftime('%d %b') if p.due_date else ''}</td></tr>"
        for p, code in points
    )
    forecast = forecast_text(s.forecast_end, s.target_date, s.progress)[0] if s else "—"
    billing = (
        (
            f"<tr><td>Contract value</td><td class='num'>{fmt(s.contract_value)}</td></tr><tr><td>Certified to date</td><td class='num'>{fmt(s.certified)}</td></tr>"
            f"<tr><td>Billed to date (before GST)</td><td class='num'>{fmt(s.billed)}</td></tr><tr><td>Retention held</td><td class='num'>{fmt(s.retention_held)}</td></tr>"
        )
        if s
        else ""
    )
    body = f"""
<div class="meta"><div><b>Site:</b> {e(site.code)} · {e(site.name)}</div><div><b>Client:</b> {e(client.name if client else '')}</div>
<div><b>Start:</b> {site.start_date.strftime('%d %b %Y') if site.start_date else '—'} · <b>Planned end:</b> {site.target_date.strftime('%d %b %Y') if site.target_date else '—'}</div>
<div><b>Progress:</b> {Decimal(site.progress_percent):.0f}% · <b>Time gone:</b> {fmt(s.elapsed, 'pct') if s else '—'} · <b>Forecast end:</b> {forecast}</div></div>
<div style="display:flex;gap:10px"><div style="flex:1"><h2>Progress by tower / floor</h2><table><tbody>{parts or "<tr><td class='muted'>No structure yet.</td></tr>"}</tbody></table></div>
<div style="flex:1"><h2>Billing</h2><table><tbody>{billing or "<tr><td class='muted'>No contract yet.</td></tr>"}</tbody></table></div></div>
<h2>Recent daily reports</h2><table><tbody>{dpr_rows or "<tr><td class='muted'>No reports yet.</td></tr>"}</tbody></table>
<h2>Open snags</h2><table><thead><tr><th>Code</th><th>Snag</th><th>Area</th><th>Status</th></tr></thead><tbody>{snag_rows or "<tr><td colspan=4 class='muted'>None.</td></tr>"}</tbody></table>
<h2>Inspections</h2><table><thead><tr><th>Code</th><th>Date</th><th>Checklist</th><th>Result</th></tr></thead><tbody>{insp_rows or "<tr><td colspan=4 class='muted'>None yet.</td></tr>"}</tbody></table>
<h2>Open meeting points</h2><table><thead><tr><th>MOM</th><th>Point</th><th>Owner</th><th>Due</th></tr></thead><tbody>{point_rows or "<tr><td colspan=4 class='muted'>None.</td></tr>"}</tbody></table>
<p class="muted">Prepared {today():%d %b %Y}. Forecast end: from the progress rate of the last 30 days.</p>"""
    return _doc(db, "SITE STATUS REPORT", body)


# --- receivables ---------------------------------------------------------------------------------


def receivables_pdf(db: Session, demo: bool) -> bytes:
    day = today()
    names = dict(db.execute(select(Client.id, Client.name)).all())
    from app.finance.models import TaxInvoice

    numbers = dict(db.execute(select(TaxInvoice.id, TaxInvoice.number)).all())
    rows = db.scalars(
        select(InvoiceSummary)
        .where(InvoiceSummary.is_demo == demo, InvoiceSummary.outstanding > 0)
        .order_by(InvoiceSummary.client_id, InvoiceSummary.invoice_date)
    ).all()
    by_client: dict[int, list] = {}
    for r in rows:
        by_client.setdefault(r.client_id, []).append(r)
    buckets = {"0-30": ZERO, "31-60": ZERO, "61-90": ZERO, "90+": ZERO}
    body_rows = []
    for cid, invs in sorted(
        by_client.items(), key=lambda kv: -sum(Decimal(x.outstanding) for x in kv[1])
    ):
        tot = sum((Decimal(x.outstanding) for x in invs), ZERO)
        body_rows.append(
            f"<tr class='client'><td colspan='4'><b>{e(names.get(cid))}</b></td><td class='num'><b>{fmt(tot)}</b></td><td></td><td></td></tr>"
        )
        for x in invs:
            a = (day - x.invoice_date).days
            if Decimal(x.due):
                buckets[
                    "0-30" if a <= 30 else "31-60" if a <= 60 else "61-90" if a <= 90 else "90+"
                ] += Decimal(x.due)
            body_rows.append(
                f"<tr><td>{e(numbers.get(x.invoice_id))}</td><td>{x.invoice_date:%d %b %Y}</td><td>{x.due_date.strftime('%d %b %Y') if x.due_date else ''}</td>"
                f"<td class='num'>{a}</td><td class='num'>{fmt(x.outstanding)}</td><td class='num'>{fmt(x.retention_held)}</td><td class='num'>{fmt(x.due)}</td></tr>"
            )
    total = sum((Decimal(x.outstanding) for x in rows), ZERO)
    ageing = "".join(
        f"<td><div class='muted'>{k} days</div><b>{fmt(v)}</b></td>" for k, v in buckets.items()
    )
    body = (
        f"<p class='muted'>As of {day:%d %b %Y}. Total outstanding {fmt(total)}; the due part by age:</p><table><tr>{ageing}</tr></table>"
        "<h2>By client</h2><table><thead><tr><th>Invoice</th><th>Date</th><th>Due</th><th class='num'>Age</th><th class='num'>Outstanding</th>"
        f"<th class='num'>Retention</th><th class='num'>Due now</th></tr></thead><tbody>{''.join(body_rows) or '<tr><td colspan=7 class=muted>Nothing outstanding.</td></tr>'}</tbody></table>"
    )
    return _doc(db, "RECEIVABLES REPORT", body)


# --- sales pipeline ------------------------------------------------------------------------------


def pipeline_pdf(db: Session, principal, scope: str, demo: bool) -> bytes:
    data = jsonable_encoder(kpi.sales(db, principal, scope, demo))
    stages = "".join(
        f"<tr><td>{e(r['stage'].replace('_', ' '))}</td><td class='num'>{r['count']}</td><td class='num'>{fmt(r['value'])}</td></tr>"
        for r in data["pipeline"]
    )
    tiles = "".join(
        f"<td><div class='muted'>{e(t['label'])}</div><b>{fmt(t['value'], t['unit'])}</b><div class='muted'>{e(t.get('note') or '')}</div></td>"
        for t in data["tiles"]
    )
    from app.analytics.common import Filters
    from app.analytics.drill import drill as drill_list

    f = Filters(today() - timedelta(days=365), today(), demo=demo)
    due = drill_list(db, principal, "tenders", f, {"filter": "due_week"})["rows"]
    due_rows = "".join(
        f"<tr><td>{e(t['code'])}</td><td>{e(t['name'])}</td><td>{e(t['client'] or '')}</td><td>{e(t['owner'] or '')}</td><td>{t['due_on']:%d %b}</td><td class='num'>{fmt(t['quoted'])}</td></tr>"
        for t in due
    )
    late = drill_list(db, principal, "leads", f, {"filter": "follow_overdue"})["rows"]
    late_rows = "".join(
        f"<tr><td>{e(x['code'])}</td><td>{e(x['contact'])}</td><td>{e(x['company'] or '')}</td><td>{e(x['owner'] or '')}</td><td>{x['follow_up']:%d %b}</td></tr>"
        for x in late[:40]
    )
    who = "my leads and tenders" if scope == "own" else "all leads and tenders"
    body = (
        f"<p class='muted'>{today():%d %b %Y} · {who}</p><table><tr>{tiles}</tr></table>"
        f"<h2>Pipeline by stage</h2><table><thead><tr><th>Stage</th><th class='num'>Count</th><th class='num'>Value</th></tr></thead><tbody>{stages}</tbody></table>"
        f"<h2>Tenders due this week</h2><table><thead><tr><th>Tender</th><th>Name</th><th>Client</th><th>Owner</th><th>Due</th><th class='num'>Quoted</th></tr></thead><tbody>{due_rows or '<tr><td colspan=6 class=muted>None.</td></tr>'}</tbody></table>"
        f"<h2>Follow-ups overdue</h2><table><thead><tr><th>Lead</th><th>Contact</th><th>Company</th><th>Owner</th><th>Follow-up</th></tr></thead><tbody>{late_rows or '<tr><td colspan=5 class=muted>None.</td></tr>'}</tbody></table>"
    )
    return _doc(db, "SALES PIPELINE", body)


# --- weekly management summary -------------------------------------------------------------------

SYSTEM = SimpleNamespace(
    user=SimpleNamespace(id=None),
    permissions={"tender.margin": "all", "payroll.view": "all", "dashboard.company": "all"},
)


def company_users(db: Session) -> list:
    return list(
        db.scalars(
            select(User.id)
            .join(UserRole, UserRole.user_id == User.id)
            .join(RolePermission, RolePermission.role_id == UserRole.role_id)
            .where(
                RolePermission.permission_code == "dashboard.company",
                User.is_active,
                User.is_demo.is_(False),
            )
            .distinct()
        )
    )


def weekly(db: Session, at: datetime | None = None) -> WeeklyReport | None:
    """File this week's management summary (once per week) and tell management in the bell."""
    at = at or now()
    day = today()
    week_start = day - timedelta(days=day.weekday())
    if db.scalar(select(WeeklyReport.id).where(WeeklyReport.week_start == week_start)):
        return None
    data = kpi.management(db, SYSTEM, demo=False)
    open_alerts = db.execute(
        select(Alert.rule, func.count())
        .where(Alert.acknowledged_at.is_(None), Alert.closed_at.is_(None))
        .group_by(Alert.rule)
    ).all()
    report = WeeklyReport(
        week_start=week_start,
        generated_at=at,
        data=jsonable_encoder(
            {
                **data,
                "delayed": delayed_sites(db, False),
                "alerts": [
                    {"rule_label": RULE_LABELS[r], "title": f"{n} open"} for r, n in open_alerts
                ],
            }
        ),
    )
    db.add(report)
    db.flush()
    portal.notify(
        db,
        company_users(db),
        "weekly_report",
        f"Management summary for the week of {week_start:%d %b %Y}",
        "/reports",
    )
    db.commit()
    return report


def weekly_pdf(db: Session, report: WeeklyReport, show_margin: bool) -> bytes:
    d = report.data
    return management_pdf(
        db,
        d,
        show_margin,
        d.get("alerts", []),
        d.get("delayed", []),
        f"week of {report.week_start:%d %b %Y}",
    )


def open_alerts_for(db: Session, principal) -> list[dict]:
    return [
        {"rule_label": RULE_LABELS[a.rule], "title": a.title}
        for a in visible(db, principal, limit=8)
    ]


__all__ = ["XLSX", "date", "money", "pct", "Role"]
