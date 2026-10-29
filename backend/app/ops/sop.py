# ruff: noqa: E501
"""The role SOPs: one button-by-button guide per role, from docs/sop (see its README). The role's own
text (docs/sop/<role>.md) names its pages in the order of a working day; each page's section
(docs/sop/pages/<page>.md) is shared by the roles that open it, its lines tagged {perm: ...} kept only
for the roles holding one of those permissions. Read in the app (Help > How to, the "How to" link of
each page: a user reads their own roles' SOPs, the super admin and the director all of them), and
printed to PDF and Word with the DEMO screenshots of the training company:

    python -m app.cli sops --screens data/sop/screens --out data/sop
"""

import json
import os
import re
from dataclasses import dataclass
from datetime import date
from html import escape
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Role

VERSION = "1.0"
SEE_ALL = {"super_admin", "director"}
INDEX = "index"  # the super admin's SOP on users and logins (docs/sop/index.md)
OPS = Path(__file__).parent
PERM_TAG = re.compile(r"\s*\{perm:\s*([^}]*)\}\s*$")


def sop_dir() -> Path:
    """docs/sop: next to the backend in the repo and in CI, /repo/docs/sop in the dev container,
    /app/docs/sop in the image."""
    if os.environ.get("SOP_DIR"):
        return Path(os.environ["SOP_DIR"])
    for d in (OPS.parents[2] / "docs" / "sop", Path("/repo/docs/sop"), OPS.parents[1] / "docs" / "sop"):
        if (d / "pages").is_dir():
            return d
    return OPS.parents[2] / "docs" / "sop"


def slug(page: str) -> str:
    return page.strip("/").replace("/", "-").replace(":", "") or "home"


def _front(text: str) -> tuple[dict, str]:
    text = text.replace("\r\n", "\n")
    if not text.startswith("---\n"):
        return {}, text
    head, _, body = text[4:].partition("\n---\n")
    meta = {}
    for line in head.splitlines():
        k, sep, v = line.partition(":")
        if sep:
            meta[k.strip()] = v.strip()
    return meta, body.strip("\n")


@dataclass
class Page:
    page: str
    title: str
    routes: list[str]
    mark: str
    body: str


@dataclass
class RoleDoc:
    role: str
    title: str
    width: int
    body: str


def pages() -> dict[str, Page]:
    out = {}
    for f in sorted((sop_dir() / "pages").glob("*.md")):
        meta, body = _front(f.read_text(encoding="utf-8"))
        page = meta.get("page") or f.stem
        routes = [r.strip() for r in meta.get("routes", "").split(",") if r.strip()]
        out[page] = Page(page, meta.get("title", page), routes or ([page] if page.startswith("/") else []), meta.get("mark", ""), body)
    return out


def role_docs() -> dict[str, RoleDoc]:
    out = {}
    for f in sorted(sop_dir().glob("*.md")):
        meta, body = _front(f.read_text(encoding="utf-8"))
        if meta.get("role"):
            out[meta["role"]] = RoleDoc(meta["role"], meta.get("title", meta["role"]), int(meta.get("width", 1600)), body)
    return out


# --- who opens what (the menu's rules, menu.ts) -------------------------------------------------------


def _menu() -> list[dict]:
    return json.loads((OPS / "menu.json").read_text(encoding="utf-8"))


def _routes() -> dict:
    return json.loads((OPS / "routes.json").read_text(encoding="utf-8"))


def in_menu(item: dict, have: set[str]) -> bool:
    return any(p in have for p in item["perms"]) and all(p in have for p in item.get("all", [])) and not any(p in have for p in item.get("none", []))


def may_open(item: dict, have: set[str]) -> bool:
    return any(p in have for p in item.get("view", item["perms"])) and all(p in have for p in item.get("all", [])) and not any(p in have for p in item.get("none", []))


def opens(have: set[str]) -> list[str]:
    """Every route a role with these permissions may open: menu pages and detail pages."""
    out = [i["to"] for i in _menu() if may_open(i, have)]
    out += [r["path"] for r in _routes()["detail"] if any(p in have for p in r["perms"])]
    return out


def route_pages() -> dict[str, str]:
    """route -> the page section that covers it."""
    return {r: p.page for p in pages().values() for r in p.routes}


def perms_of(db: Session, role: str) -> set[str]:
    if role == INDEX:
        role = "super_admin"
    r = db.scalar(select(Role).where(Role.code == role))
    return {rp.permission_code for rp in r.permissions} if r else set()


# --- assembling a role's SOP --------------------------------------------------------------------------


_ITEM = re.compile(r"\s*([-*]|\d+[.)]) ")


def keep(body: str, have: set[str]) -> str:
    """The body without the list items (with their wrapped lines and sub-items) and lines tagged
    for permissions the role does not hold; the tags taken out."""
    lines = body.splitlines()
    out: list[str] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        j = i + 1
        if _ITEM.match(line):  # the item: its wrapped lines and deeper lines up to the next item at its level
            depth = len(line) - len(line.lstrip())
            while j < len(lines) and lines[j].strip() and (len(lines[j]) - len(lines[j].lstrip()) > depth) and not _ITEM.match(lines[j]):
                j += 1
        unit = lines[i:j]
        tags = [m for m in (PERM_TAG.search(x) for x in unit) if m]
        if tags and not any(c.strip() in have for m in tags for c in m.group(1).split(",")):
            # drop the item and the sub-items under it
            depth = len(line) - len(line.lstrip())
            while j < len(lines) and lines[j].strip() and len(lines[j]) - len(lines[j].lstrip()) > depth:
                j += 1
            i = j
            continue
        out += [PERM_TAG.sub("", x) if PERM_TAG.search(x) else x for x in unit]
        i = j
    return "\n".join(out)


def _access(db: Session, role: str, have: set[str]) -> str:
    from app.ops.roles_report import NEVER, SCOPE  # noqa: PLC0415

    groups: dict[str, list[str]] = {}
    for item in _menu():
        if in_menu(item, have):
            groups.setdefault(item["section"], []).append(item["label"])
    r = db.scalar(select(Role).where(Role.code == ("super_admin" if role == INDEX else role)))
    scopes = {rp.scope for rp in r.permissions} if r else set()
    lines = ["## What you can and cannot see", ""]
    if groups:
        lines += ["Your menu, by group (the app shows only these):", ""]
        lines += [f"- **{g}**: {', '.join(items)}" for g, items in groups.items()]
    else:
        lines += ["You have no staff menu: you see the client portal of your own sites only."]
    if "assigned" in scopes:
        lines += ["", f"Some of your work is limited to {SCOPE['assigned']}: you see a site when you are a member of it (Settings > Team, or the site's team)."]
    never = [label for label, codes in NEVER if not any(c in have for c in codes)]
    if never:
        lines += ["", "You never see:", ""] + [f"- {n}" for n in never]
    return "\n".join(lines)


def assemble(db: Session, role: str) -> dict:
    """The role's SOP as sections: {id, title, page, markdown}, in reading order."""
    doc = role_docs()[role]
    have = perms_of(db, role)
    all_pages = pages()
    sections: list[dict] = []

    def add(sid: str, title: str, md: str, page: str = "") -> None:
        sections.append({"id": sid, "title": title, "page": page, "markdown": md.strip("\n")})

    cur: dict | None = None
    for line in keep(doc.body, have).splitlines():
        s = line.strip()
        if s.startswith("## "):
            cur = {"id": slug(re.sub(r"[^a-z0-9]+", "-", s[3:].lower()).strip("-")), "title": s[3:], "page": "", "markdown": ""}
            sections.append(cur)
        elif s == "{{access}}":
            add("access", "What you can and cannot see", _access(db, role, have).split("\n", 2)[2])
            cur = None
        elif s == "{{first-login}}":
            p = all_pages["_first-login"]
            add("first-login", p.title, keep(p.body, have))
            cur = None
        elif s.startswith("@page "):
            p = all_pages[s[6:].strip()]
            add(slug(p.page), p.title, keep(p.body, have), p.page)
            cur = None
        elif cur is not None:
            cur["markdown"] += line + "\n"
        elif s:
            cur = {"id": f"part-{len(sections)}", "title": "", "page": "", "markdown": line + "\n"}
            sections.append(cur)
    for x in sections:
        x["markdown"] = x["markdown"].strip("\n")
    return {"role": role, "title": doc.title, "width": doc.width, "version": VERSION, "sections": sections}


def included(role: str) -> set[str]:
    """The pages the role's SOP file includes (its @page lines)."""
    return {line.strip()[6:].strip() for line in role_docs()[role].body.splitlines() if line.strip().startswith("@page ")}


def readable(user_roles: list[str]) -> list[str]:
    """The SOPs a person may read: their own roles', or all for the super admin and the director."""
    docs = role_docs()
    if SEE_ALL & set(user_roles):
        return list(docs)
    return [r for r in user_roles if r in docs]


# --- markdown (the subset docs/sop uses) to blocks, HTML and Word -------------------------------------

_INLINE = re.compile(r"(\*\*(.+?)\*\*|`([^`]+)`|\[([^\]]+)\]\(([^)]+)\)|(?<![*\w])\*([^*]+)\*(?!\*))")


def blocks(md: str) -> list[tuple]:
    """("h", level, text) ("p", text) ("list", ordered, [(depth, text)]) ("table", rows) ("hr",)"""
    out: list[tuple] = []
    para: list[str] = []
    items: list[tuple[int, str]] = []
    ordered = False
    table: list[list[str]] = []

    def flush() -> None:
        nonlocal para, items, table
        if para:
            out.append(("p", " ".join(para)))
        if items:
            out.append(("list", ordered, items))
        if table:
            out.append(("table", table))
        para, items, table = [], [], []

    for raw in md.splitlines():
        line = raw.rstrip()
        s = line.lstrip()
        indent = len(line) - len(s)
        if not s:
            flush()
            continue
        if s.startswith("|"):
            if para or items:
                flush()
            cells = [c.strip() for c in s.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                table.append(cells)
            continue
        h = re.match(r"(#{1,4}) (.*)", s)
        if h and not indent:
            flush()
            out.append(("h", len(h.group(1)), h.group(2)))
            continue
        if s in ("---", "***"):
            flush()
            out.append(("hr",))
            continue
        li = re.match(r"([-*]|\d+[.)]) (.*)", s)
        if li:
            if para or table:
                flush()
            if not items:
                ordered = li.group(1)[0].isdigit()
            items.append((indent // 2, li.group(2)))
            continue
        if items and indent:
            d, t = items[-1]
            items[-1] = (d, f"{t} {s}")
            continue
        if items or table:
            flush()
        para.append(s)
    flush()
    return out


def inline_html(text: str) -> str:
    out, pos = [], 0
    for m in _INLINE.finditer(text):
        out.append(escape(text[pos : m.start()]))
        if m.group(2) is not None:
            out.append(f"<b>{inline_html(m.group(2))}</b>")
        elif m.group(3) is not None:
            out.append(f"<code>{escape(m.group(3))}</code>")
        elif m.group(4) is not None:
            href = m.group(5)
            safe = href if href.startswith(("/", "#", "https://")) else "#"
            out.append(f'<a href="{escape(safe)}">{escape(m.group(4))}</a>')
        else:
            out.append(f"<i>{escape(m.group(6))}</i>")
        pos = m.end()
    out.append(escape(text[pos:]))
    return "".join(out)


def to_html(md: str, shift: int = 0) -> str:
    parts = []
    for b in blocks(md):
        if b[0] == "h":
            lv = min(6, b[1] + shift)
            parts.append(f"<h{lv}>{inline_html(b[2])}</h{lv}>")
        elif b[0] == "p":
            parts.append(f"<p>{inline_html(b[1])}</p>")
        elif b[0] == "hr":
            parts.append("<hr>")
        elif b[0] == "table":
            head, *rows = b[1]
            parts.append("<table><thead><tr>" + "".join(f"<th>{inline_html(c)}</th>" for c in head) + "</tr></thead><tbody>")
            parts += ["<tr>" + "".join(f"<td>{inline_html(c)}</td>" for c in r) + "</tr>" for r in rows]
            parts.append("</tbody></table>")
        else:
            tag = "ol" if b[1] else "ul"
            depth = -1
            for d, t in b[2]:
                while depth < d:
                    parts.append(f"<{tag}>" if depth < 0 else f"<{'ul'}>")
                    depth += 1
                while depth > d:
                    parts.append("</ul>")
                    depth -= 1
                parts.append(f"<li>{inline_html(t)}</li>")
            while depth > 0:
                parts.append("</ul>")
                depth -= 1
            parts.append(f"</{tag}>")
    return "\n".join(parts)


def for_app(db: Session, role: str) -> dict:
    doc = assemble(db, role)
    for s in doc["sections"]:
        s["html"] = to_html(s.pop("markdown"), shift=1)
    return doc


# --- PDF and Word ------------------------------------------------------------------------------------


def file_name(title: str, ext: str) -> str:
    return f"EESPL SOP - {title}.{ext}"


def _shot(screens: Path | None, role: str, page: str) -> Path | None:
    f = screens / role / f"{slug(page)}.png" if screens and page else None
    return f if f and f.exists() else None


def html(db: Session, role: str, screens: Path | None) -> str:
    doc = assemble(db, role)
    title, today = doc["title"], date.today()
    narrow = doc["width"] < 800
    parts = [f"""<html><head><meta charset="utf-8"><style>
@page {{ size: A4; margin: 14mm 14mm 18mm;
  @bottom-left {{ content: "EESPL App · SOP: {escape(title)} · version {VERSION} · {today:%d.%m.%Y}"; font-size: 7.5pt; color: #64736f; }}
  @bottom-right {{ content: "Page " counter(page) " of " counter(pages); font-size: 7.5pt; color: #64736f; }} }}
body {{ font-family: DejaVu Sans, sans-serif; font-size: 9.5pt; color: #1c2b28; line-height: 1.4; }}
h1 {{ font-size: 24pt; margin: 50mm 0 4px; color: #0f6b5c; }}
h2 {{ font-size: 14pt; margin: 0 0 6px; border-bottom: 2px solid #0f6b5c; padding-bottom: 3px; page-break-after: avoid; }}
h3 {{ font-size: 11pt; margin: 10px 0 4px; color: #0f6b5c; page-break-after: avoid; }} h4 {{ font-size: 10pt; margin: 8px 0 2px; }}
.sec {{ page-break-before: always; }} .muted {{ color: #64736f; }} ul, ol {{ margin: 3px 0 6px 16px; padding: 0; }} li {{ margin: 2px 0; }}
p {{ margin: 4px 0 6px; }} code {{ font-size: 8.5pt; background: #eef2f1; padding: 0 2px; }}
table {{ border-collapse: collapse; width: 100%; margin: 4px 0 8px; }} td, th {{ border: 1px solid #d5dcda; padding: 3px 5px; vertical-align: top; text-align: left; }}
th {{ background: #eef6f4; }} img.shot {{ border: 1px solid #d5dcda; margin: 4px 0 8px; }}
img.wide {{ width: 100%; }} img.narrow {{ width: 62mm; float: right; margin: 0 0 6px 10px; }}
.card {{ border: 2px solid #0f6b5c; padding: 10px 14px; }} .toc li {{ margin: 1px 0; }}
</style></head><body>
<h1>{escape(title)}</h1>
<p style="font-size:13pt">EESPL App: standard operating procedure</p>
<p class="muted">Version {VERSION}, {today:%d.%m.%Y}. Written from the app itself; the pictures come from the training company's
DEMO data (names and numbers are invented). Practise in the training company first: it is reset every night.</p>
<h3>Contents</h3><ol class="toc">"""]
    parts += [f"<li>{escape(s['title'])}</li>" for s in doc["sections"] if s["title"]]
    parts.append("</ol>")
    for s in doc["sections"]:
        card = s["id"] == "quick-card"
        parts.append('<div class="sec">' + (f"<h2>{escape(s['title'])}</h2>" if s["title"] else ""))
        if s["page"]:
            parts.append(f'<p class="muted">Page: {escape(s["page"])}</p>')
            shot = _shot(screens, role, s["page"])
            if shot:
                parts.append(f'<img class="shot {"narrow" if narrow else "wide"}" src="{shot.resolve().as_uri()}">')
        parts.append(('<div class="card">' if card else "") + to_html(s["markdown"], shift=1) + ("</div>" if card else ""))
        parts.append("</div>")
    parts.append("</body></html>")
    return "".join(parts)


def pdf(db: Session, role: str, screens: Path | None) -> tuple[bytes, int]:
    from weasyprint import HTML  # noqa: PLC0415

    rendered = HTML(string=html(db, role, screens)).render()
    return rendered.write_pdf(), len(rendered.pages)


def _runs(par, text: str) -> None:
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos : m.start()])
        if m.group(2) is not None:
            par.add_run(m.group(2).replace("**", "")).bold = True
        elif m.group(3) is not None:
            par.add_run(m.group(3)).font.name = "Consolas"
        elif m.group(4) is not None:
            par.add_run(m.group(4))
        else:
            par.add_run(m.group(6)).italic = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def _page_field(par) -> None:
    from docx.oxml import OxmlElement  # noqa: PLC0415
    from docx.oxml.ns import qn  # noqa: PLC0415

    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        run = par.add_run()
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def docx(db: Session, role: str, screens: Path | None) -> bytes:
    from io import BytesIO  # noqa: PLC0415

    from docx import Document  # noqa: PLC0415
    from docx.enum.text import WD_BREAK  # noqa: PLC0415
    from docx.shared import Cm, Pt, RGBColor  # noqa: PLC0415

    doc = assemble(db, role)
    title, today = doc["title"], date.today()
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.6)
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10)
    foot = sec.footer.paragraphs[0]
    foot.text = f"EESPL App · SOP: {title} · version {VERSION} · {today:%d.%m.%Y} · Page "
    foot.runs[0].font.size = Pt(8)
    _page_field(foot)
    h = d.add_heading(title, 0)
    h.runs[0].font.color.rgb = RGBColor(0x0F, 0x6B, 0x5C)
    d.add_paragraph("EESPL App: standard operating procedure")
    d.add_paragraph(f"Version {VERSION}, {today:%d.%m.%Y}. Written from the app itself; the pictures come from the training company's DEMO data "
                    "(names and numbers are invented). Practise in the training company first: it is reset every night.")
    narrow = doc["width"] < 800
    for s in doc["sections"]:
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        if s["title"]:
            d.add_heading(s["title"], 1)
        if s["page"]:
            d.add_paragraph(f"Page: {s['page']}").runs[0].italic = True
            shot = _shot(screens, role, s["page"])
            if shot:
                d.add_picture(str(shot), width=Cm(6.5) if narrow else Cm(17.5))
        for b in blocks(s["markdown"]):
            if b[0] == "h":
                d.add_heading(b[2].replace("**", ""), min(4, b[1]))
            elif b[0] == "p":
                _runs(d.add_paragraph(), b[1])
            elif b[0] == "table":
                t = d.add_table(rows=0, cols=len(b[1][0]))
                t.style = "Table Grid"
                for row in b[1]:
                    cells = t.add_row().cells
                    for i, c in enumerate(row[: len(cells)]):
                        _runs(cells[i].paragraphs[0], c)
            elif b[0] == "list":
                for depth, text in b[2]:
                    style = ("List Number" if b[1] and depth == 0 else "List Bullet") + (" 2" if depth else "")
                    _runs(d.add_paragraph(style=style), text)
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def build_all(db: Session, screens: Path | None, out: Path, only: list[str] | None = None) -> list[dict]:
    out.mkdir(parents=True, exist_ok=True)
    made = []
    for role, doc in role_docs().items():
        if only and role not in only:
            continue
        data, n = pdf(db, role, screens)
        (out / file_name(doc.title, "pdf")).write_bytes(data)
        (out / file_name(doc.title, "docx")).write_bytes(docx(db, role, screens))
        shots = sum(1 for p in included(role) if _shot(screens, role, p))
        made.append({"role": role, "title": doc.title, "pages": n, "sections": len(included(role)), "screenshots": shots})
    return made
