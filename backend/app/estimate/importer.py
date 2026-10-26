# ruff: noqa: E501
"""The planning team's old estimate sheets read back: each sheet's areas and coverages, listed
where they differ from the measured drawing or from the library, and every figure typed as text
instead of a number or a formula ("300 TO 340", "1.8 kg/sqm"), which no total can add.

The sheets are not one layout, so they are read by what a row says: a row naming a measured area
(by its name or code) with a number is that area's figure; a row naming a library product is that
product's line, its coverage the number in a "coverage" / "consumption" column, or else the
constant multiplied in its quantity formula (=B5*1.8*1.05: 1.8, the wastage 1.05 apart). The same
area given two figures across the sheets (1,693 and 1,695 sqm) and one material at two thicknesses
(140 mm and 165 mm) are listed too."""

import io
import re
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.estimate import library as lib
from app.estimate.models import Coverage
from app.masters.models import Product

D = Decimal
TEXT_FIGURE = re.compile(r"^\s*\d+(?:\.\d+)?\s*(?:(?:to|-|–|—)\s*\d+(?:\.\d+)?|(?:kg|ltr|l|sqm|sqft|mm|nos?|%)\b.*)", re.I)
MM = re.compile(r"(\d{2,3})\s*mm\b", re.I)
CONST = re.compile(r"(?<![A-Z$\d.])(\d+(?:\.\d+)?)(?![\d.]*[A-Z(])", re.I)


def _num(v) -> Decimal | None:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, int | float | Decimal):
        return D(str(v))
    return None


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def read(db: Session, data: bytes, name: str, measured: list | None = None) -> dict:
    """measured: the site's areas ([service.Measured]) to compare with; None: areas are not judged."""
    wb_f = load_workbook(io.BytesIO(data), data_only=False)
    wb_v = load_workbook(io.BytesIO(data), data_only=True)
    products = []
    for p in db.scalars(select(Product).where(Product.is_active.is_(True))):
        cov = db.scalar(select(Coverage).where(Coverage.product_id == p.id, Coverage.system_id.is_(None)))
        cov = cov or db.scalar(select(Coverage).where(Coverage.product_id == p.id).limit(1))
        if cov is None:
            continue
        for nm in {p.name, *(p.aliases or [])}:
            if nm and len(nm) >= 4:
                products.append((_norm(nm), p, cov))
    products.sort(key=lambda t: -len(t[0]))
    areas = []
    for m in measured or []:
        areas.append((m, [_norm(m.name)] + ([_norm(m.code)] if m.code and len(m.code) >= 2 else [])))
    out = {"file": name, "sheets": [], "findings": []}
    seen_area: dict[str, list[tuple[str, Decimal]]] = {}
    seen_mm: dict[str, set[int]] = {}
    for ws in wb_f.worksheets:
        wv = wb_v[ws.title]
        sh = {"name": ws.title, "areas": [], "coverages": [], "typed_as_text": []}
        cov_col = None
        for row in ws.iter_rows():
            texts = [(c, str(c.value)) for c in row if isinstance(c.value, str) and not str(c.value).startswith("=")]
            heads = {c.column: _norm(t) for c, t in texts}
            for col, t in heads.items():
                if t.startswith(("coverage", "consumption")):
                    cov_col = col
            for c, t in texts:
                if TEXT_FIGURE.match(t):
                    sh["typed_as_text"].append({"cell": c.coordinate, "text": t.strip()})
            label = " ".join(t for _, t in texts)
            low = _norm(label)
            if not low:
                continue
            nums = [(c, _num(wv[c.coordinate].value)) for c in row if _num(wv[c.coordinate].value) is not None]
            for mm in MM.findall(label):
                for key, _p, _c in products:
                    if key in low:
                        seen_mm.setdefault(_p.name, set()).add(int(mm))
                        break
            hit = next(((m, k) for m, keys in areas for k in keys if k and re.search(rf"(?<!\w){re.escape(k)}(?!\w)", low)), None)
            if hit and nums:
                m, _k = hit
                value = nums[0][1]
                drawn = m.floor_sqm
                sh["areas"].append({"cell": nums[0][0].coordinate, "label": label.strip()[:120], "value": value, "measured": drawn, "diff": value - drawn})
                seen_area.setdefault(m.name, []).append((ws.title, value))
                continue
            prod = next(((p, cov) for key, p, cov in products if re.search(rf"(?<!\w){re.escape(key)}(?!\w)", low)), None)
            if prod is None:
                continue
            p, cov = prod
            value, where = None, None
            if cov_col is not None:
                v = _num(wv.cell(row[0].row, cov_col).value)
                if v is not None:
                    value, where = v, wv.cell(row[0].row, cov_col).coordinate
            if value is None:
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        consts = [D(x) for x in CONST.findall(c.value[1:]) if D(x) not in (D(0), D(1))]
                        consts = [x for x in consts if not (D("1") < x <= D("1.25") and len(consts) > 1)]  # the wastage factor apart
                        if consts:
                            value, where = consts[0], c.coordinate
                            break
            if value is None:
                continue
            lib_v = D(cov.value)
            sh["coverages"].append({"cell": where, "product": p.name, "value": value, "library": lib_v, "library_unit": cov.unit, "differs": value != lib_v})
        out["sheets"].append(sh)
        for a in sh["areas"]:
            if abs(a["diff"]) >= D("0.01"):
                out["findings"].append(f"{ws.title} {a['cell']}: {a['label']} is {lib.n(a['value'])} sqm; the drawing measures {lib.n(a['measured'])} sqm ({'+' if a['diff'] > 0 else ''}{lib.n(a['diff'])})")
        for cv in sh["coverages"]:
            if cv["differs"]:
                out["findings"].append(f"{ws.title} {cv['cell']}: {cv['product']} at {lib.n(cv['value'], 4)}; the library says {lib.n(cv['library'], 4)} {cv['library_unit']}")
        for t in sh["typed_as_text"]:
            out["findings"].append(f"{ws.title} {t['cell']}: '{t['text']}' is typed as text, not a number or a formula (no total adds it)")
    for area, vals in seen_area.items():
        distinct = sorted({v for _, v in vals})
        if len(distinct) > 1:
            out["findings"].append(f"{area} has {len(distinct)} figures across the sheets: " + ", ".join(f"{lib.grouped(v, 3)} sqm ({s})" for s, v in vals))
    for prod, mms in seen_mm.items():
        if len(mms) > 1:
            out["findings"].append(f"{prod} is worked at {' and '.join(f'{x} mm' for x in sorted(mms))}: one thickness is wrong")
    return out


def safe_decimal(text: str) -> Decimal | None:
    try:
        return D(text)
    except (InvalidOperation, TypeError):
        return None
