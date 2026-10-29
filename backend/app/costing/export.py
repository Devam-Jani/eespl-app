# ruff: noqa: E501  (the build-up rows and queries read better unwrapped)
"""The rate working as Excel, in the layout of our own "Working" sheets (the 5 Acre BOQ's
"WORKING (2)", the Shela inquiry's working): per main item, SR NO | ITEM | COVRAGE | COST OF
MATERIAL | COST PER AREA (SQMT) | COST PER AREA (SQFT) | RATE REF, with live formulas: the product
rows (coverage with wastage × price + freight), LABOUR (₹ per sq ft × 10.7639), CTC, MARGIN, RATE.
Cushions follow as visible rows. Internal only: never in what goes to the client."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.costing import prices, working
from app.costing.models import RateWorking
from app.masters.models import Product, ProductPrice
from app.masters.rate import SQFT_PER_SQM
from app.timefmt import ist_today

HEADERS = [
    "SR NO",
    "ITEM",
    "COVRAGE",
    "COST OF MATERIAL",
    "COST PER AREA (SQMT)",
    "COST PER AREA (SQFT)",
    "RATE REF",
]
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
CUSHION_FILL = PatternFill("solid", fgColor="FFF2CC")
F = f"{SQFT_PER_SQM}"


def _label(db: Session, w: RateWorking) -> str:
    if w.tender_id:
        from app.tenders.models import Tender  # noqa: PLC0415

        t = db.get(Tender, w.tender_id)
        return f"{t.code} R{w.rev_no} {t.name}"
    from app.quotations.models import Quotation  # noqa: PLC0415

    q = db.get(Quotation, w.quotation_id)
    return f"{q.code} R{q.revision} {q.client_firm}"


def add_working_sheet(
    wb: Workbook, db: Session, w: RateWorking, title: str = "WORKING"
) -> Worksheet:
    ws = wb.create_sheet(title)
    lines = {x["key"]: x for x in working.client_lines(db, w)}
    basis = working.margin_basis(db, w)
    ws.append([f"Rate working: {_label(db, w)} (internal: costs and margins)"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([f"Margin basis: {working.BASIS_LABEL[basis]}. Labour: ₹ per sq ft × {F} per sqm."])
    ws.append(HEADERS + ["USED / ACTUAL (cushion)"])
    for c in ws[3]:
        c.font, c.fill = BOLD, HEAD_FILL
    for n, g in enumerate(w.groups, start=1):
        result = working.work(db, g)
        ws.append([])
        r = ws.max_row + 1
        ws.append([n, g.name.upper(), None, None, None, None,
                   f"wastage {g.wastage_percent.normalize():f} %" + (f" ({g.wastage_source})" if g.wastage_source else "")
                   + (f", {g.coats} coats" if g.coats else "")])  # fmt: skip
        ws.cell(r, 2).font = BOLD
        first = ws.max_row + 1
        ws.append([None, "SURFACE PREPARATION", 1, float(g.surface_prep or 0), None])
        ws.cell(ws.max_row, 5).value = f"=C{ws.max_row}*D{ws.max_row}"
        w_pct = float(g.wastage_percent or 0) / 100
        for p in g.products:
            product = db.get(Product, p.product_id)
            row = ws.max_row + 1
            if p.consumption is None or p.price is None:
                ws.append([None, product.name, None, None, None, None,
                           "consumption to confirm" if p.consumption is None else "no price"])  # fmt: skip
                continue
            ref = "price list"
            if p.price_id:
                info = prices.price_out(db, db.get(ProductPrice, p.price_id))
                if info:
                    ref = f"{info['source']} {info['date']:%d-%m-%Y}" + (
                        f" {info['by']}" if info["by"] else ""
                    )
            if p.consumption_reason:
                ref += f"; consumption changed: {p.consumption_reason}"
            ws.append([None, product.name,
                       f"={float(p.consumption)}*(1+{w_pct})",
                       f"={float(p.price)}+{float(p.freight_used or 0)}",
                       f"=C{row}*D{row}", None, ref])  # fmt: skip
            if p.freight_actual is not None:
                ws.cell(row, 8).value = f"freight {p.freight_used} used / {p.freight_actual} actual"
        lab = ws.max_row + 1
        ws.append([None, "LABOUR", 1, f"={float(g.labour_used_sqft or 0)}*{F}", f"=C{lab}*D{lab}", None,
                   f"₹{g.labour_used_sqft} per sq ft"])  # fmt: skip
        if g.labour_paid_sqft is not None:
            ws.cell(
                lab, 8
            ).value = f"labour ₹{g.labour_used_sqft} used / ₹{g.labour_paid_sqft} paid per sq ft"
        ctc = ws.max_row + 1
        ws.append([None, "CTC", None, None, f"=SUM(E{first}:E{lab})", f"=E{ctc}/{F}"])
        mar = ctc + 1
        # on cost: margin = m × CTC; on the selling price: margin = CTC ÷ (1 − m) − CTC
        margin_formula = f"=D{mar}*E{ctc}" if basis == "cost" else f"=E{ctc}/(1-D{mar})-E{ctc}"
        ws.append(
            [
                None,
                "MARGIN" + (" (on cost)" if basis == "cost" else " (on selling price)"),
                None,
                float(g.margin_percent or 0) / 100,
                margin_formula,
                f"=E{mar}/{F}",
            ]
        )
        rate = mar + 1
        ws.append([None, "RATE", None, None, f"=SUM(E{ctc}:E{mar})", f"=E{rate}/{F}",
                   f"app: ₹{result['rate_sqm']} / sqm, ₹{result['rate_sqft']} / sq ft"
                   if result["rate_sqm"] is not None else "; ".join(result["errors"])])  # fmt: skip
        for col in (2, 5, 6):
            ws.cell(rate, col).font = BOLD
        for c in result.get("cushions", []):
            ws.append([None, f"CUSHION: {c['what']}", None, None, float(c["per_sqm"]),
                       f"=E{ws.max_row + 1}/{F}", None,
                       f"used {c['used']} / actual {c['actual']}"])  # fmt: skip
            for cell in ws[ws.max_row]:
                cell.fill = CUSHION_FILL
        if result.get("library_median_sqm") is not None:
            ws.append([None, "RATE LIBRARY MEDIAN", None, None, float(result["library_median_sqm"]), None,
                       f"{result['library_diff_percent']} % off" + (" (flagged)" if result["library_flagged"] else "")])  # fmt: skip
        nos = [
            lines[k]["item_no"] or str(lines[k]["id"])
            for k in (working._line_key(x) for x in g.lines)
            if k in lines
        ]
        ws.append([None, f"Client lines ({len(nos)}): " + ", ".join(str(x) for x in nos)])
        ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True)
    for col, width in zip("ABCDEFGH", (7, 44, 16, 18, 20, 20, 36, 34), strict=True):
        ws.column_dimensions[col].width = width
    return ws


def working_xlsx(db: Session, w: RateWorking) -> tuple[bytes, str]:
    wb = Workbook()
    wb.remove(wb.active)
    add_working_sheet(wb, db, w)
    buf = BytesIO()
    wb.save(buf)
    stem = _label(db, w).split(" ")[0].replace("/", "-")
    return buf.getvalue(), f"{stem}-rate-working-{ist_today():%Y%m%d}.xlsx"
