"""Excel exports for list pages.

Each list endpoint has an `/export` twin that runs the same query with the same permission and
the same field hiding, and returns an .xlsx built here.
"""

import uuid
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from fastapi import Response
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXPORT_ROW_LIMIT = 20000
LOCAL_TZ = ZoneInfo("Asia/Kolkata")  # Excel has no time zones; show Indian time

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="0F6B5C")


def _cell_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.astimezone(LOCAL_TZ).replace(tzinfo=None)
        return value
    if isinstance(value, date | int | float | Decimal):
        return value
    if isinstance(value, uuid.UUID):
        return str(value)
    if isinstance(value, list | tuple | set):
        return ", ".join(str(v) for v in value)
    if isinstance(value, dict):
        return ", ".join(f"{k}: {v}" for k, v in value.items())
    return str(value)


def xlsx_bytes(sheet_title: str, columns: Sequence[str], rows: Iterable[Sequence[Any]]) -> bytes:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title=sheet_title[:31] or "Export")
    ws.freeze_panes = "A2"
    widths = [len(c) + 2 for c in columns]
    header = []
    for title in columns:
        cell = WriteOnlyCell(ws, value=title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        header.append(cell)
    body = []
    for row in rows:
        values = [_cell_value(v) for v in row]
        for i, v in enumerate(values[: len(widths)]):
            widths[i] = min(60, max(widths[i], len(str(v)) + 2 if v is not None else 0))
        body.append(values)
    for i, width in enumerate(widths):
        ws.column_dimensions[_column_letter(i)].width = width
    ws.append(header)
    for values in body:
        ws.append(values)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _column_letter(index: int) -> str:
    letters = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def xlsx_response(
    name: str, columns: Sequence[str], rows: Iterable[Sequence[Any]], sheet_title: str | None = None
) -> Response:
    stamp = datetime.now(LOCAL_TZ).strftime("%Y%m%d")
    return Response(
        content=xlsx_bytes(sheet_title or name.replace("-", " ").title(), columns, rows),
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="eespl-{name}-{stamp}.xlsx"'},
    )
