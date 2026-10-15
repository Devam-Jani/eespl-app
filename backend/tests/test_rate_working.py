# ruff: noqa: E501, F811  (invented data reads better unwrapped; shared fixtures are imported)
"""Rate working, the master price list, labour paid back to back, material in lots and the
handover schedule (the director's review, 12 Oct). Invented data only."""

from datetime import UTC, date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.costing.models import LabourBill, RateGroup, RateWorking
from app.execution.models import Attendance, Labour
from app.finance.models import ClientContract, ContractLine, RaBill, RaBillLine, VendorBill
from app.masters.models import LibraryItem, Product, System, SystemComponent, Vendor
from app.masters.rate import SQFT_PER_SQM
from app.material.models import Indent
from app.portal.models import Notification
from app.process import engine
from app.schedule import service as schedule
from app.sites.models import AreaScope, Site, SiteMember, SiteNode, StageTemplate, StageTemplateStep
from app.tenders.models import BoqLine, Tender
from tests.conftest import login
from tests.test_finance import boss, contract, world  # noqa: F401  (fixtures)


def d(v) -> Decimal:
    return Decimal(str(v))


def cents(v: Decimal) -> Decimal:
    return v.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


# --- the master price list -----------------------------------------------------------------------


def test_the_last_price_used_is_shown_and_a_change_over_10_percent_is_flagged(boss, db, login_as):
    c, h = boss
    p = Product(code="T-PU", name="Invented PU coating", unit="kg")
    db.add(p)
    db.commit()
    r = c.post(
        f"/api/prices/products/{p.id}",
        json={
            "rate": "250",
            "source": "vendor_quote",
            "ref": "Q-17",
            "effective_from": "2026-09-01",
        },
        headers=h,
    )
    assert r.status_code == 201 and r.json()["check"]["last"] is None
    hint = c.get(f"/api/prices/last?product_id={p.id}&rate=280", headers=h).json()
    assert d(hint["last"]["rate"]) == 250 and hint["last"]["date"] == "2026-09-01"
    assert hint["last"]["by"] == "boss" and hint["last"]["source"] == "vendor_quote"
    assert d(hint["change_percent"]) == d("12.0") and hint["flagged"]
    small = c.get(f"/api/prices/last?product_id={p.id}&rate=270", headers=h).json()
    assert d(small["change_percent"]) == 8 and not small["flagged"]
    r = c.post(f"/api/prices/products/{p.id}", json={"rate": "280"}, headers=h)
    assert r.json()["check"]["flagged"]
    page = c.get(f"/api/prices/products/{p.id}", headers=h).json()
    assert d(page["current"]["rate"]) == 280 and page["history"][0]["flagged"]
    assert not page["history"][1]["flagged"]  # the first price has nothing to compare with
    # freight per unit, per product or per vendor and godown, with its own history
    v = Vendor(name="Invented Polymers", type="material_supplier")
    db.add(v)
    db.commit()
    c.post(f"/api/prices/products/{p.id}/freight", json={"per_unit": "5"}, headers=h)
    c.post(
        f"/api/prices/products/{p.id}/freight", json={"per_unit": "3", "vendor_id": v.id}, headers=h
    )
    from app.costing.prices import freight_for

    assert freight_for(db, p.id).per_unit == 5 and freight_for(db, p.id, v.id).per_unit == 3
    # purchase enters prices (POs) and sees the hint, not the cost history
    pc, ph = login_as("store_purchase")
    assert pc.get(f"/api/prices/last?product_id={p.id}&rate=300", headers=ph).status_code == 200
    assert pc.get(f"/api/prices/products/{p.id}", headers=ph).status_code == 403
    # the limit is a setting
    assert (
        c.put(
            "/api/prices/settings", json={"price_change_flag_percent": "15"}, headers=h
        ).status_code
        == 200
    )
    assert not c.get(f"/api/prices/last?product_id={p.id}&rate=320", headers=h).json()["flagged"]


# --- the rate working ----------------------------------------------------------------------------

TEXTS = {
    "toilet": "Supplying and applying Ultra Flex elastomeric coating in toilet of tower",
    "raft": "Raft waterproofing with HDPE membrane below the raft",
    "terrace": "Terrace hybrid PU coating over the protection screed",
}


@pytest.fixture
def boq(boss, db):
    """A client BOQ of 12 lines that are three main items (toilets, raft, terrace), a raft system
    with two products, their prices and freight, and the raft lines' library median."""
    c, h = boss
    t = Tender(code="T-2026-9101", name="Invented annex", status="draft")
    db.add(t)
    db.flush()
    lib = LibraryItem(
        source_key="t" * 40,
        match_key="m" * 40,
        description="Invented raft HDPE",
        unit="sqm",
        median_rate=500,
    )
    db.add(lib)
    db.flush()
    lines = []
    for i in range(12):
        kind = ("toilet", "raft", "terrace")[i % 3]
        unit = "sqft" if (kind == "terrace" or i == 4) else "sqm"
        ln = BoqLine(tender_id=t.id, sort_order=i, client_item_no=str(600 + i), description=TEXTS[kind], unit=unit, qty=100 + i,
                     library_item_id=lib.id if kind == "raft" else None)  # fmt: skip
        lines.append(ln)
    db.add_all(lines)
    p1 = Product(code="T-HDPE", name="Invented HDPE membrane", unit="sqm")
    p2 = Product(code="T-PRIMER", name="Invented primer", unit="ltr")
    db.add_all([p1, p2])
    db.flush()
    s = System(
        code="T-RAFT",
        name="Invented raft HDPE system",
        surface_prep_per_unit=30,
        labour_rate=12,
        labour_unit="sqft",
        default_margin_percent=25,
    )
    s.components = [
        SystemComponent(product_id=p1.id, consumption_per_unit=d("1.2")),
        SystemComponent(product_id=p2.id, consumption_per_unit=d("0.3")),
    ]
    db.add(s)
    db.commit()
    for p, rate, freight in ((p1, "250", "5"), (p2, "80", "2")):
        c.post(
            f"/api/prices/products/{p.id}",
            json={"rate": rate, "effective_from": "2026-09-01"},
            headers=h,
        )
        c.post(f"/api/prices/products/{p.id}/freight", json={"per_unit": freight}, headers=h)
    return {
        "tender": t.id,
        "lines": [x.id for x in lines],
        "system": s.id,
        "products": (p1.id, p2.id),
    }


def _working(c, h, boq):
    r = c.get(f"/api/rate-working/tenders/{boq['tender']}", headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def _group(w, word):
    return next(g for g in w["groups"] if word.lower() in g["name"].lower())


def test_grouping_suggests_main_items_and_lines_move_between_them(boss, boq):
    c, h = boss
    w = _working(c, h, boq)
    assert w["line_count"] == 12 and len(w["ungrouped"]) == 12
    w = c.post(f"/api/rate-working/{w['id']}/suggest", headers=h).json()
    assert len(w["groups"]) == 3 and w["ungrouped"] == []
    assert sorted(len(g["lines"]) for g in w["groups"]) == [4, 4, 4]  # sqm and sq ft lines together
    raft, toilet = _group(w, "raft"), _group(w, "toilet")
    moved = raft["lines"][0]["key"]
    w = c.post(
        f"/api/rate-working/{w['id']}/move",
        json={"line_key": moved, "to_group_id": toilet["id"]},
        headers=h,
    ).json()
    assert (
        moved in [x["key"] for x in _group(w, "toilet")["lines"]]
        and len(_group(w, "raft")["lines"]) == 3
    )
    w = c.post(
        f"/api/rate-working/{w['id']}/move",
        json={"line_key": moved, "new_group_name": "Invented lift pit"},
        headers=h,
    ).json()
    assert len(w["groups"]) == 4 and [x["key"] for x in _group(w, "lift pit")["lines"]] == [moved]
    w = c.post(f"/api/rate-working/{w['id']}/move", json={"line_key": moved}, headers=h).json()
    assert [x["key"] for x in w["ungrouped"]] == [moved]


def test_the_rate_formula_cushions_reasons_flags_and_the_rate_to_every_line(
    boss, boq, db, login_as
):
    c, h = boss
    w = c.post(f"/api/rate-working/{_working(c, h, boq)['id']}/suggest", headers=h).json()
    raft = _group(w, "raft")
    g = c.put(
        f"/api/rate-working/groups/{raft['id']}", json={"system_id": boq["system"]}, headers=h
    ).json()
    assert (
        g["application"] == "raft" and d(g["wastage_percent"]) == 7
    )  # raft 5-7 %: the upper figure
    # consumption × (price + freight) × (1 + wastage), + surface prep + labour ₹/sq ft × 10.7639, + margin
    material = (d("1.2") * (250 + 5) + d("0.3") * (80 + 2)) * d("1.07")
    cost = material + 30 + 12 * SQFT_PER_SQM
    # a new working: the margin on the selling price, rate = cost ÷ (1 − 25 %)
    assert d(g["rate_sqm"]) == cents(cost / d("0.75"))
    assert d(g["rate_sqft"]) == cents(cost / d("0.75") / SQFT_PER_SQM)
    labour_row = next(r for r in g["rows"] if r["item"] == "Labour")
    assert d(labour_row["per_sqm"]) == cents(12 * SQFT_PER_SQM)
    # cushions are visible lines: labour used vs paid, freight used vs actual
    prods = g["products"]
    body = {
        "labour_paid_sqft": "10",
        "products": [
            {
                "id": prods[0]["id"],
                "product_id": prods[0]["product_id"],
                "consumption": "1.2",
                "freight_used": "5",
                "freight_actual": "3",
            },
            {
                "id": prods[1]["id"],
                "product_id": prods[1]["product_id"],
                "consumption": "0.3",
                "freight_used": "2",
            },
        ],
    }
    g = c.put(f"/api/rate-working/groups/{raft['id']}", json=body, headers=h).json()
    cushions = {x["what"]: d(x["per_sqm"]) for x in g["cushions"]}
    assert cushions["Labour"] == cents(2 * SQFT_PER_SQM)
    assert cushions["Freight on Invented HDPE membrane"] == cents(2 * d("1.2") * d("1.07"))
    # a consumption other than the system's needs a reason
    body["products"][0]["consumption"] = "1.4"
    r = c.put(f"/api/rate-working/groups/{raft['id']}", json=body, headers=h)
    assert r.status_code == 422 and "say why" in r.json()["detail"]
    body["products"][0]["consumption_reason"] = "Invented: deeper laps on this raft"
    g = c.put(f"/api/rate-working/groups/{raft['id']}", json=body, headers=h).json()
    assert g["products"][0]["consumption_changed"] and g["products"][0][
        "consumption_reason"
    ].startswith("Invented")
    # more than 15 % off the rate library's median (₹500 per sqm) is flagged
    assert (
        d(g["library_median_sqm"]) == 500
        and g["library_flagged"]
        and d(g["library_diff_percent"]) > 15
    )
    # the worked rate flows to every client line of the group, per sqm or per sq ft
    out = c.post(f"/api/rate-working/{w['id']}/apply", headers=h).json()
    assert out["applied"] == 4
    db.expire_all()
    rows = {
        ln.id: ln
        for ln in db.scalars(
            select(BoqLine).where(BoqLine.id.in_([x["id"] for x in raft["lines"]]))
        )
    }
    for ln in rows.values():
        assert ln.status == "priced" and ln.source == "working"
        assert ln.rate == (d(g["rate_sqm"]) if ln.unit == "sqm" else d(g["rate_sqft"]))
    assert any(ln.unit == "sqft" for ln in rows.values())
    assert db.get(Tender, boq["tender"]).quoted_total > 0
    # sales sees the main items and their rates, never costs, cushions or margins
    sales_client, sh = login_as("sales")
    t = db.get(Tender, boq["tender"])
    from app.models import User

    t.created_by = db.scalar(select(User.id).where(User.email == "sales@example.com"))
    db.commit()
    seen = sales_client.get(f"/api/rate-working/tenders/{boq['tender']}", headers=sh).json()
    seen_raft = _group(seen, "raft")
    assert d(seen_raft["rate_sqm"]) == d(g["rate_sqm"])
    for key in (
        "products",
        "rows",
        "cost_sqm",
        "margin_percent",
        "cushions",
        "labour_used_sqft",
        "surface_prep",
    ):
        assert key not in seen_raft
    assert not seen["sees_cost"]
    assert (
        sales_client.put(
            f"/api/rate-working/groups/{raft['id']}", json={"margin_percent": "5"}, headers=sh
        ).status_code
        == 403
    )
    # the client's Excel never has the working; the internal one has it as its own sheet
    client_file = load_workbook(
        BytesIO(c.get(f"/api/tenders/{boq['tender']}/export.xlsx", headers=h).content)
    )
    assert "WORKING" not in client_file.sheetnames
    internal = c.get(f"/api/tenders/{boq['tender']}/export.xlsx?internal=true", headers=h)
    book = load_workbook(BytesIO(internal.content))
    assert "WORKING" in book.sheetnames
    cells = [
        str(x) for row in book["WORKING"].iter_rows(values_only=True) for x in row if x is not None
    ]
    assert (
        "COVRAGE" in cells
        and "CTC" in cells
        and "MARGIN (on selling price)" in cells
        and any(x.startswith("CUSHION") for x in cells)
    )
    assert (
        sales_client.get(
            f"/api/tenders/{boq['tender']}/export.xlsx?internal=true", headers=sh
        ).status_code
        == 403
    )
    alone = c.get(f"/api/rate-working/{w['id']}/export.xlsx", headers=h)
    assert (
        alone.status_code == 200 and "WORKING" in load_workbook(BytesIO(alone.content)).sheetnames
    )


def test_a_new_revision_copies_the_working_and_keeps_the_submitted_one(boss, boq, db):
    c, h = boss
    w = c.post(f"/api/rate-working/{_working(c, h, boq)['id']}/suggest", headers=h).json()
    raft = _group(w, "raft")
    c.put(
        f"/api/rate-working/groups/{raft['id']}",
        json={"system_id": boq["system"], "margin_percent": "20"},
        headers=h,
    )
    t = db.get(Tender, boq["tender"])
    t.revision = 1  # R0 submitted: the next change starts R1
    db.commit()
    w1 = _working(c, h, boq)
    assert w1["id"] != w["id"] and w1["rev_no"] == 1 and not w1["locked"]
    assert [g["name"] for g in w1["groups"]] == [g["name"] for g in w["groups"]]
    assert d(_group(w1, "raft")["margin_percent"]) == 20
    db.expire_all()
    assert db.get(RateWorking, w["id"]).locked
    r = c.put(f"/api/rate-working/groups/{raft['id']}", json={"margin_percent": "15"}, headers=h)
    assert r.status_code == 409  # R0's working is kept as it was


# --- labour back to back -------------------------------------------------------------------------


def _ra(db, contract_id, site_id, seq, period_to, status, line_id, qty, certified=None):
    b = RaBill(
        code=f"RA-T-{seq}",
        contract_id=contract_id,
        site_id=site_id,
        seq=seq,
        period_to=period_to,
        status=status,
    )
    db.add(b)
    db.flush()
    db.add(
        RaBillLine(
            ra_bill_id=b.id,
            contract_line_id=line_id,
            qty=qty,
            certified_qty=certified,
            rate=500,
            amount=qty * 500,
        )
    )
    db.commit()


def test_piece_rate_labour_is_the_billed_quantity_or_the_certified_one_when_lower(boss, world, db):
    c, h = boss
    contract(c, h, world)
    k = db.scalar(select(ClientContract).where(ClientContract.site_id == world["site"]))
    coating = db.scalar(
        select(ContractLine).where(ContractLine.contract_id == k.id, ContractLine.unit == "sqm")
    )
    gang = Vendor(
        name="Invented Applicators", type="labour_contractor", pan="ABCPA1234A"
    )  # an individual: 1 %
    db.add(gang)
    db.commit()
    card = c.post(
        "/api/labour/rate-cards",
        json={
            "contractor_id": gang.id,
            "job": "Coating",
            "covers": "surface preparation, cleaning, crack and joint treatment, vata, primer and coat",
            "unit": "sqft",
            "rate": "10",
            "valid_from": "2026-04-01",
        },
        headers=h,
    ).json()
    _ra(
        db,
        k.id,
        world["site"],
        1,
        date(2026, 9, 30),
        "certified",
        coating.id,
        d(4000),
        certified=d(3500),
    )
    _ra(db, k.id, world["site"], 2, date(2026, 10, 31), "submitted", coating.id, d(300))
    pick = {
        "contractor_id": gang.id,
        "site_id": world["site"],
        "lines": [{"contract_line_id": coating.id, "rate_card_id": card["id"]}],
    }
    b = c.post(
        "/api/labour/bills/piece",
        json=pick | {"period_from": "2026-09-01", "period_to": "2026-09-30"},
        headers=h,
    ).json()
    ln = b["lines"][0]
    assert (
        d(ln["qty"]) == 3500 and ln["basis"] == "certified"
    )  # the client certified less than billed
    assert d(ln["rate"]) == (10 * SQFT_PER_SQM).quantize(
        Decimal("0.0001")
    )  # ₹10 per sq ft on a sqm item
    assert d(b["gross"]) == cents(d(3500) * (10 * SQFT_PER_SQM).quantize(Decimal("0.0001")))
    assert (
        b["tds_section"] == "194C"
        and d(b["tds_percent"]) == 1
        and d(b["tds"]) == cents(d(b["gross"]) / 100)
    )
    assert d(b["net"]) == d(b["gross"]) - d(b["tds"])
    # the next month: only what was billed since (300, not certified yet), never typed
    b2 = c.post(
        "/api/labour/bills/piece",
        json=pick | {"period_from": "2026-10-01", "period_to": "2026-10-31"},
        headers=h,
    ).json()
    assert d(b2["lines"][0]["qty"]) == 300 and b2["lines"][0]["basis"] == "billed"
    r = c.post(
        "/api/labour/bills/piece",
        json=pick | {"period_from": "2026-10-01", "period_to": "2026-10-31"},
        headers=h,
    )
    assert r.status_code == 422  # nothing more to pay
    # approved: a vendor bill, and the site's labour cost
    from app.execution.service import actuals

    before = actuals(db, world["site"])["labour"]
    a = c.post(f"/api/labour/bills/{b['id']}/approve", headers=h).json()
    assert a["status"] == "approved"
    db.expire_all()
    assert db.scalar(select(VendorBill).where(VendorBill.bill_no == b["number"])).tds_amount == d(
        b["tds"]
    )
    assert actuals(db, world["site"])["labour"] == before + d(b["gross"])


def test_a_gang_on_attendance_is_billed_by_its_hours_on_each_site(boss, world, db):
    c, h = boss
    other = Site(code="S-2026-9811", name="Invented second site", status="active")
    gang = Vendor(name="Invented Laxman gang", type="labour_contractor")
    db.add_all([other, gang])
    db.flush()
    w1 = Labour(
        name="Invented worker one", type="subcontractor", subcontractor_id=gang.id, daily_wage=800
    )
    w2 = Labour(
        name="Invented worker two", type="subcontractor", subcontractor_id=gang.id, daily_wage=800
    )
    db.add_all([w1, w2])
    db.flush()
    rows = [
        (w1, world["site"], date(2026, 9, 1), "present", d(8)),
        (w1, other.id, date(2026, 9, 2), "present", d(6)),
        (w2, world["site"], date(2026, 9, 1), "half_day", None),  # half a standard day: 4 hours
        (w2, other.id, date(2026, 9, 3), "absent", None),
    ]
    for lab, site_id, on, st, hours in rows:
        db.add(
            Attendance(
                labour_id=lab.id,
                site_id=site_id,
                on_date=on,
                status=st,
                hours=hours,
                daily_wage=800,
            )
        )
    db.commit()
    b = c.post(
        "/api/labour/bills/attendance",
        json={"contractor_id": gang.id, "month": "2026-09-15"},
        headers=h,
    ).json()
    by_site = {ln["site_id"]: ln for ln in b["lines"]}
    assert set(by_site) == {world["site"], other.id}
    assert d(by_site[world["site"]]["qty"]) == 12 and d(by_site[world["site"]]["amount"]) == 1200
    assert d(by_site[other.id]["qty"]) == 6 and d(by_site[other.id]["amount"]) == 600
    assert d(b["gross"]) == 1800 and d(b["tds"]) == 0  # under the 194C limits
    r = c.post(
        "/api/labour/bills/attendance",
        json={"contractor_id": gang.id, "month": "2026-09-01"},
        headers=h,
    )
    assert r.status_code == 409  # the month is billed once
    c.post(f"/api/labour/bills/{b['id']}/approve", headers=h)
    from app.costing.labour import site_labour

    assert site_labour(db, other.id)["labour_cost"] == 600


# --- material in lots ----------------------------------------------------------------------------


def test_lots_are_called_off_against_the_plan_and_over_the_balance_needs_approval(
    boss, world, db, login_as
):
    c, h = boss
    p = Product(code="T-CRYS", name="Invented crystalline powder", unit="kg")
    db.add(p)
    db.commit()
    site = world["site"]
    r = c.post(
        f"/api/schedule/sites/{site}/plan/fix",
        json={"lines": [{"product_id": p.id, "planned_qty": "1000"}]},
        headers=h,
    )
    assert r.status_code == 422  # differs from the computed quantity: needs a note
    plan = c.post(
        f"/api/schedule/sites/{site}/plan/fix",
        json={
            "lines": [
                {"product_id": p.id, "planned_qty": "1000", "note": "Invented: from the survey"}
            ]
        },
        headers=h,
    ).json()
    assert d(plan["products"][0]["planned"]) == 1000
    node = db.scalar(select(SiteNode).where(SiteNode.site_id == site))
    lot1 = c.post(
        f"/api/schedule/sites/{site}/lots",
        json={"node_ids": [node.id], "days": 4, "lines": [{"product_id": p.id, "qty": "100"}]},
        headers=h,
    ).json()
    assert lot1["status"] == "approved" and lot1["indent"]
    lot2 = c.post(
        f"/api/schedule/sites/{site}/lots",
        json={"lines": [{"product_id": p.id, "qty": "940"}]},
        headers=h,
    ).json()
    assert lot2["status"] == "approved"  # 1,040: within the plan plus 5 %
    r = c.post(
        f"/api/schedule/sites/{site}/lots",
        json={"lines": [{"product_id": p.id, "qty": "60"}]},
        headers=h,
    )
    assert r.status_code == 422 and "reason" in r.json()["detail"]  # 1,100 > 1,050
    lot3 = c.post(
        f"/api/schedule/sites/{site}/lots",
        json={
            "lines": [{"product_id": p.id, "qty": "60"}],
            "reason": "Invented: the raft grew by 40 sqm",
        },
        headers=h,
    ).json()
    assert lot3["status"] == "pending" and lot3["over_plan"] and lot3["indent"] is None
    plan = c.get(f"/api/schedule/sites/{site}/plan", headers=h).json()
    row = plan["products"][0]
    assert d(row["called_off"]) == 1040 and d(row["pending"]) == 60 and d(row["balance"]) == -40
    # the site cannot approve its own call-off; planning does
    sc, sh = login_as("site_supervisor")
    assert (
        sc.post(
            f"/api/schedule/lots/{lot3['id']}/decide", json={"approve": True}, headers=sh
        ).status_code
        == 403
    )
    pc, ph = login_as("planning")
    done = pc.post(
        f"/api/schedule/lots/{lot3['id']}/decide",
        json={"approve": True, "note": "Invented: agreed"},
        headers=ph,
    ).json()
    assert done["status"] == "approved" and done["indent"]
    assert db.scalar(select(Indent).where(Indent.id == done["indent_id"])).status == "submitted"


# --- handover dates and the schedule -------------------------------------------------------------


@pytest.fixture
def places(world, db, make_user):
    """Two places on the site with a 5 + 2 + 8 day stage template, the site's engineer and
    supervisor, planning and the director."""
    site = db.get(Site, world["site"])
    eng = make_user("eng@example.com", "site_engineer", name="Invented Engineer")
    tpl = StageTemplate(name="Invented coating stages", created_by=eng.id)  # removed after the test
    tpl.steps = [StageTemplateStep(sort_order=i, name=n, weight_percent=w, typical_days=days)
                 for i, (n, w, days) in enumerate((("Surface", 20, 5), ("Primer", 20, 2), ("Coats", 60, 8)))]  # fmt: skip
    db.add(tpl)
    db.flush()
    nodes = []
    for kind, name in (("raft", "Invented block A raft"), ("terrace", "Invented block B terrace")):
        n = SiteNode(site_id=site.id, kind=kind, name=name)
        db.add(n)
        db.flush()
        db.add(
            AreaScope(
                site_id=site.id, node_id=n.id, stage_template_id=tpl.id, qty=100, progress_percent=0
            )
        )
        nodes.append(n)
    sup = make_user("sup@example.com", "site_supervisor", name="Invented Supervisor")
    plan = make_user("plan@example.com", "planning", name="Invented Planner")
    director = make_user("dir@example.com", "director", name="Invented Director")
    db.add_all([SiteMember(site_id=site.id, user_id=eng.id, role_on_site="incharge"),
                SiteMember(site_id=site.id, user_id=sup.id, role_on_site="supervisor"),
                SiteMember(site_id=site.id, user_id=plan.id, role_on_site="office")])  # fmt: skip
    site.site_incharge_id = eng.id
    db.commit()
    return {
        "site": site.id,
        "nodes": [n.id for n in nodes],
        "eng": eng,
        "sup": sup,
        "plan": plan,
        "director": director,
    }


def test_handover_dates_target_finish_and_the_two_delays(places, db, client):
    site, (a, b) = places["site"], places["nodes"]
    today = schedule.today()
    eh = login(client, "eng@example.com")
    # the engineer enters what the client promised; the client's word alone starts nothing
    r = client.put(
        f"/api/schedule/sites/{site}/nodes/{a}/promise",
        json={
            "promised_on": (today - timedelta(days=10)).isoformat(),
            "promised_by": "Invented PMC, Mr. Shah",
            "client_says_on": (today - timedelta(days=9)).isoformat(),
        },
        headers=eh,
    ).json()
    assert r["status"] == "awaited" and r["received_on"] is None and r["target_finish"] is None
    assert (
        r["overdue_handover"] and r["client_delay_days"] == 10
    )  # not handed over yet: still counting
    # a promised date passed and not confirmed: the engineer and planning are told
    engine.ensure_rules(db)
    schedule.timed_events(db, datetime.now(UTC))
    db.commit()
    told = {
        n.user_id
        for n in db.scalars(
            select(Notification).where(Notification.rule_code == "handover_overdue")
        )
    }
    assert told == {places["eng"].id, places["plan"].id}
    # only the supervisor confirms the real handover
    assert (
        client.post(
            f"/api/schedule/sites/{site}/nodes/{a}/handover", json={"received": True}, headers=eh
        ).status_code
        == 403
    )
    sh = login(client, "sup@example.com")
    r = client.post(
        f"/api/schedule/sites/{site}/nodes/{a}/handover",
        json={"received": True, "on_date": (today - timedelta(days=6)).isoformat()},
        headers=sh,
    ).json()
    assert r["received_on"] == (today - timedelta(days=6)).isoformat() and r["stage_days"] == 15
    assert r["target_finish"] == (today + timedelta(days=9)).isoformat()  # received + 5 + 2 + 8
    assert r["client_delay_days"] == 4 and r["our_delay_days"] == 0
    db.expire_all()
    assert db.get(SiteNode, a).front_ready  # the handover is the work-front-ready tick
    assert all(
        n.closed_at
        for n in db.scalars(
            select(Notification).where(Notification.rule_code == "handover_overdue")
        )
    )
    # B: handed over on time, not finished by its target: our delay
    client.put(
        f"/api/schedule/sites/{site}/nodes/{b}/promise",
        json={
            "promised_on": (today - timedelta(days=30)).isoformat(),
            "promised_by": "Invented PMC",
        },
        headers=eh,
    )
    r = client.post(
        f"/api/schedule/sites/{site}/nodes/{b}/handover",
        json={"received": True, "on_date": (today - timedelta(days=30)).isoformat()},
        headers=sh,
    ).json()
    assert r["client_delay_days"] == 0 and r["our_delay_days"] == 15 and r["overdue_finish"]
    whole = client.get(f"/api/schedule/sites/{site}/handover", headers=eh).json()
    assert whole["client_delay_days"] == 4 and whole["our_delay_days"] == 15
    # stage days changed for this site move the target
    days = client.get(f"/api/schedule/sites/{site}/stage-days", headers=eh).json()
    steps = next(t for t in days if t["template"] == "Invented coating stages")["steps"]
    client.put(
        f"/api/schedule/sites/{site}/stage-days",
        json=[{"step_id": steps[2]["step_id"], "days": 10}],
        headers=eh,
    )
    assert (
        client.get(f"/api/schedule/sites/{site}/handover", headers=eh).json()["our_delay_days"]
        == 13
    )  # B: 30 - 17
    # target finish passed: the engineer, then the director after 3 days
    schedule.timed_events(db, datetime.now(UTC))
    db.commit()
    first = db.scalar(
        select(Notification).where(
            Notification.rule_code == "finish_overdue", Notification.escalation_level == 0
        )
    )
    assert first.user_id == places["eng"].id
    engine.escalate(db, first.created_at + timedelta(days=2))
    db.commit()
    assert not db.scalar(
        select(Notification).where(
            Notification.rule_code == "finish_overdue",
            Notification.user_id == places["director"].id,
        )
    )
    engine.escalate(db, first.created_at + timedelta(days=3, hours=1))
    db.commit()
    assert db.scalar(
        select(Notification).where(
            Notification.rule_code == "finish_overdue",
            Notification.user_id == places["director"].id,
        )
    ).escalated
    # the director's view across sites
    dh = login(client, "dir@example.com")
    delays = client.get("/api/schedule/delays", headers=dh).json()
    assert delays["client_delay_days"] == 4 and delays["our_delay_days"] == 13
    # finishing the area closes the to-do
    scope = db.scalar(select(AreaScope).where(AreaScope.node_id == b))
    scope.dpr_done_qty = 100
    from app.sites import work

    work.refresh_scope(db, scope)
    db.commit()
    assert all(
        n.closed_at
        for n in db.scalars(select(Notification).where(Notification.rule_code == "finish_overdue"))
    )
    assert (
        db.scalar(select(LabourBill)) is None and db.scalar(select(RateGroup)) is None
    )  # nothing else made
