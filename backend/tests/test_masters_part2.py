"""M1 part 2: vendors, bank-detail hiding, categories, unit conversions, the RO rule, Powerplay
imports, Excel exports and company settings."""

from decimal import Decimal as D
from io import BytesIO
from pathlib import Path

import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, text

from app.config import settings
from app.db import engine
from app.main import app
from app.masters.conversions import ConversionError, Rule, convert_with
from app.masters.importers import import_library
from app.masters.models import Category, LibraryLine, Product, Vendor
from app.masters.powerplay import import_materials, import_team, import_vendors
from app.masters.units import is_rate_only, normalise_unit
from app.models import AuditLog, User
from tests.conftest import PASSWORD, login

FIXTURES = Path(__file__).parent / "fixtures" / "powerplay"
GSTIN = "27AAPFU0939F1ZV"
BANK = {
    "account_name": "Ramesh Traders",
    "account_number": "123456789012",
    "ifsc": "HDFC0001234",
    "bank": "HDFC Bank",
    "branch": "Navrangpura",
}


def _client(make_user, role, email=None):
    email = email or f"{role.replace('_', '-')}@example.com"
    make_user(email, role)
    c = TestClient(app)
    return c, login(c, email)


# --- vendors -----------------------------------------------------------------------------------


def test_vendor_gstin_and_ifsc_validation(login_as):
    c, h = login_as("office_admin")
    assert (
        c.post(
            "/api/vendors", json={"name": "V", "gstin": "27AAPFU0939F1ZX"}, headers=h
        ).status_code
        == 422
    )
    r = c.post(
        "/api/vendors", json={"name": "V", "gstin": GSTIN.lower(), "type": "transporter"}, headers=h
    )
    assert r.status_code == 201 and r.json()["gstin"] == GSTIN
    vid = r.json()["id"]
    for bad in (
        {"ifsc": "HDFC1234567"},
        {"ifsc": "HDF0001234"},
        {"account_number": "12AB34"},
        {"account_number": "123"},
    ):
        r = c.post(f"/api/vendors/{vid}/bank-accounts", json={**BANK, **bad}, headers=h)
        assert r.status_code == 422, bad
    r = c.post(f"/api/vendors/{vid}/bank-accounts", json={**BANK, "ifsc": "hdfc0001234"}, headers=h)
    assert r.status_code == 201
    account = r.json()["bank_accounts"][0]
    assert (account["ifsc"], account["is_primary"], account["masked"]) == (
        "HDFC0001234",
        True,
        False,
    )


def test_bank_details_only_with_settings_company_or_finance_view(make_user):
    admin, ah = _client(make_user, "office_admin")
    vid = admin.post("/api/vendors", json={"name": "Ramesh Traders"}, headers=ah).json()["id"]
    admin.post(f"/api/vendors/{vid}/bank-accounts", json=BANK, headers=ah)

    for role, visible in (("store_purchase", False), ("estimator", False), ("accounts", True)):
        c, h = _client(make_user, role)
        bank = c.get(f"/api/vendors/{vid}", headers=h).json()["bank_accounts"][0]
        assert bank["account_name"] == "Ramesh Traders"
        if visible:
            assert (bank["account_number"], bank["ifsc"], bank["masked"]) == (
                "123456789012",
                "HDFC0001234",
                False,
            )
        else:
            assert (bank["account_number"], bank["ifsc"], bank["masked"]) == (None, None, True)
        listed = str(c.get("/api/vendors", headers=h).json())
        assert ("123456789012" in listed) is visible
        sheet = load_workbook(BytesIO(c.get("/api/vendors/export", headers=h).content)).active
        assert ("Account number" in [cell.value for cell in sheet[1]]) is visible

    # store/purchase can edit vendors but not bank details they cannot see
    sp, sh = _client(make_user, "store_purchase", "sp2@example.com")
    assert sp.patch(f"/api/vendors/{vid}", json={"city": "Surat"}, headers=sh).status_code == 200
    assert sp.post(f"/api/vendors/{vid}/bank-accounts", json=BANK, headers=sh).status_code == 403


def test_vendor_contacts_products_and_audit(login_as, db):
    c, h = login_as("super_admin")
    pid = c.post(
        "/api/products",
        json={"code": "P1", "name": "Primer", "unit": "kg", "category": "coating"},
        headers=h,
    ).json()["id"]
    r = c.post(
        "/api/vendors",
        json={"name": "Acme", "contacts": [{"name": "Ravi", "is_primary": True}]},
        headers=h,
    )
    vid = r.json()["id"]
    r = c.put(
        f"/api/vendors/{vid}/products",
        json=[{"product_id": pid, "last_rate": "310", "lead_time_days": 3}],
        headers=h,
    )
    assert r.json()["products"][0]["product_name"] == "Primer"
    assert (
        c.put(f"/api/vendors/{vid}/products", json=[{"product_id": 999999}], headers=h).status_code
        == 422
    )
    actions = {a.action for a in db.scalars(select(AuditLog).where(AuditLog.entity == "vendor"))}
    assert {"vendor.create", "vendor.products.update"} <= actions


# --- categories ----------------------------------------------------------------------------------


def test_category_migration_keeps_every_product_category(db):
    """Downgrade to 0008 (products.category text), add products, upgrade: every product keeps its
    category through category_id."""
    cfg = Config(str(Path(__file__).resolve().parent.parent / "alembic.ini"))
    cfg.set_main_option(
        "script_location", str(Path(__file__).resolve().parent.parent / "migrations")
    )
    db.close()
    engine.dispose()
    command.downgrade(cfg, "0008")
    try:
        values = [
            "membrane",
            "coating",
            "chemical",
            "admixture",
            "sealant",
            "waterstop",
            "accessory",
            "other",
        ]
        with engine.begin() as conn:
            for i, category in enumerate(values):
                conn.execute(
                    text(
                        "INSERT INTO products (code, name, category, unit) "
                        "VALUES (:c, :n, :k, 'kg')"
                    ),
                    {"c": f"MIG-{i}", "n": f"Product {category}", "k": category},
                )
    finally:
        engine.dispose()
        command.upgrade(cfg, "head")
    engine.dispose()
    with engine.connect() as conn:
        rows = conn.execute(
            text(
                "SELECT p.code, lower(c.name) FROM products p "
                "LEFT JOIN categories c ON c.id = p.category_id "
                "WHERE p.code LIKE 'MIG-%' ORDER BY p.code"
            )
        ).all()
    assert [r[1] for r in rows] == values
    assert all(r[1] is not None for r in rows)


def test_categories_crud_and_seed(login_as, db):
    c, h = login_as("estimator")  # library.edit
    work = c.get("/api/categories", params={"kind": "work"}, headers=h).json()["items"]
    assert [w["name"] for w in work][:3] == ["Raft / basement", "Retaining wall", "Lift pit"]
    assert len(work) == 12
    r = c.post("/api/categories", json={"kind": "material", "name": "Primer"}, headers=h)
    assert r.status_code == 201
    assert (
        c.post(
            "/api/categories", json={"kind": "material", "name": "primer"}, headers=h
        ).status_code
        == 409
    )
    pid = c.post(
        "/api/products",
        json={"code": "X", "name": "X", "unit": "kg", "category_id": r.json()["id"]},
        headers=h,
    ).json()
    assert pid["category"] == "Primer"
    assert c.delete(f"/api/categories/{r.json()['id']}", headers=h).status_code == 409  # in use
    sales, sh = login_as("sales", email="s@example.com")
    assert sales.get("/api/categories", headers=sh).status_code == 200
    assert (
        sales.post("/api/categories", json={"kind": "work", "name": "Z"}, headers=sh).status_code
        == 403
    )


# --- unit conversions ----------------------------------------------------------------------------

RULES = [
    Rule("sqm", "sqft", D("10.76391042")),
    Rule("rmt", "rft", D("3.28083990")),
    Rule("roll", "sqm", D("10"), product_id=7),  # 1 roll of product 7 covers 10 sqm
    Rule("set", "nos", D("4")),
    Rule("set", "nos", D("6"), product_id=8),  # product 8 comes in sets of 6
]


def test_convert_generic_reverse_and_chained():
    assert convert_with(RULES, D("2"), "sqm", "sqft") == D("21.52782084")
    assert convert_with(RULES, D("10.76391042"), "sqft", "sqm") == D("1")
    assert convert_with(RULES, D("5"), "kg", "kg") == D("5")
    # product-specific, chained through a generic rule
    assert convert_with(RULES, D("3"), "roll", "sqft", product_id=7) == D("322.9173126")


def test_convert_product_specific_wins_over_generic():
    assert convert_with(RULES, D("2"), "set", "nos") == D("8")
    assert convert_with(RULES, D("2"), "set", "nos", product_id=8) == D("12")


def test_convert_raises_when_no_path():
    with pytest.raises(ConversionError):
        convert_with(RULES, D("1"), "kg", "sqm")
    with pytest.raises(ConversionError):
        convert_with(RULES, D("1"), "roll", "sqm")  # the roll rule is for product 7 only


def test_convert_endpoint_with_stored_conversions(login_as):
    c, h = login_as("estimator")
    r = c.get(
        "/api/units/convert", params={"qty": "2", "from_unit": "sqm", "to_unit": "sqft"}, headers=h
    )
    assert D(r.json()["result"]) == D("21.52782084")
    pid = c.post(
        "/api/products", json={"code": "M1", "name": "Membrane roll", "unit": "roll"}, headers=h
    ).json()["id"]
    r = c.post(
        "/api/unit-conversions",
        json={"from_unit": "roll", "to_unit": "sqm", "factor": "10", "product_id": pid},
        headers=h,
    )
    assert r.status_code == 201 and r.json()["product_name"] == "Membrane roll"
    r = c.get(
        "/api/units/convert",
        headers=h,
        params={"qty": "3", "from_unit": "roll", "to_unit": "sqft", "product_id": pid},
    )
    assert D(r.json()["result"]) == D("322.9173126")
    r = c.get(
        "/api/units/convert", params={"qty": "1", "from_unit": "kg", "to_unit": "sqm"}, headers=h
    )
    assert r.status_code == 422


# --- RO = rate only ------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("text", "rate_only", "unit"),
    [
        ("RO", True, None),
        ("R.O", True, None),
        ("Rate only", True, None),
        ("sqm (RO)", True, "sqm"),
        ("Sqm", False, "sqm"),
        ("Rmt", False, "rmt"),
    ],
)
def test_rate_only_spellings(text, rate_only, unit):
    assert is_rate_only(text) is rate_only
    assert normalise_unit(text) == unit


def test_import_stores_ro_lines_as_blank_unit_with_qro(db, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rate library"
    ws.append(
        [
            "Description",
            "Unit",
            "No. of BOQs",
            "Latest rate (₹)",
            "Min",
            "Median",
            "Max",
            "Latest client folder",
            "Latest source file",
            "Product / make",
            "Remarks",
            "From EESPL-priced file",
            "Check",
        ]
    )
    ws.append(
        ["Pipe sleeve packing", "RO", 1, 300, 300, 300, 300, "A", "a.xlsx", None, None, "Yes", None]
    )
    ws = wb.create_sheet("All lines")
    ws.append(
        [
            "Client folder",
            "File",
            "Sheet",
            "Row",
            "Item No",
            "Parent item",
            "Description",
            "Unit (as written)",
            "Unit",
            "Qty",
            "Qty note",
            "Rate (₹)",
            "Product / make",
            "Remarks",
            "EESPL file",
            "Check",
        ]
    )
    ws.append(
        [
            "A",
            "A/a.xlsx",
            "BOQ",
            1,
            "1",
            None,
            "Pipe sleeve packing",
            "RO",
            "RO",
            None,
            None,
            300,
            None,
            None,
            "Yes",
            None,
        ]
    )
    ws.append(
        [
            "A",
            "A/a.xlsx",
            "BOQ",
            2,
            "2",
            None,
            "Coving",
            "sqm (RO)",
            "sqm (RO)",
            None,
            None,
            90,
            None,
            None,
            "Yes",
            None,
        ]
    )
    path = tmp_path / "ro.xlsx"
    wb.save(path)

    result = import_library(db, path)
    assert "RO" not in result.unrecognised_units
    lines = {ln.row: ln for ln in db.scalars(select(LibraryLine))}
    assert (lines[1].unit, lines[1].unit_raw, lines[1].qty_note) == (None, "RO", "QRO")
    assert (lines[2].unit, lines[2].qty_note) == ("sqm", "QRO")
    assert lines[1].library_item_id is not None


# --- Powerplay imports ---------------------------------------------------------------------------


def test_powerplay_materials_import(db):
    first = import_materials(db, FIXTURES / "materials.xlsx")
    assert first.mapping["name"] == "Material Name" and first.mapping["unit"] == "UOM"
    assert (first.rows_in_file, first.created, first.duplicates_merged) == (8, 6, 1)
    assert sorted(r for _, r in first.skipped) == ["no material name", "unknown unit 'bundle'"]
    assert sorted(first.categories_created) == ["CONSUMABLE", "MASTERSEAL", "Primer", "SIKA"]

    products = {p.name: p for p in db.scalars(select(Product))}
    assert products["MASTER SEAL M645"].brand == "MASTERSEAL"
    assert products["Zz Soft Wiper"].brand == "3M"
    assert products["Acme Primer"].brand == "Acme Coatings"
    assert products["Acme Primer"].category.name == "Primer"
    assert products["Non Woven Geotextile - 120 GSM"].category.name == "Other"  # OTHERS -> Other
    assert products["Non Woven Geotextile - 120 GSM"].unit == "sqm"
    assert products["Breaker Panu"].category_id is None
    assert all(p.code.startswith("PP-") for p in products.values())

    second = import_materials(db, FIXTURES / "materials.xlsx")
    assert (second.created, second.already_present, second.categories_created) == (0, 6, [])
    assert db.scalar(select(text("count(*)")).select_from(Product)) == 6


def test_powerplay_vendors_import(db):
    first = import_vendors(db, FIXTURES / "vendors.xlsx")
    assert (first.rows_in_file, first.created, first.duplicates_merged) == (4, 3, 1)
    assert [r for _, r in first.skipped] == ["no vendor name"]
    warnings = " | ".join(w for _, w in first.warnings)
    assert "GSTIN '27ABC' ignored" in warnings and "IFSC 'HDFC1234' ignored" in warnings
    vendors = {v.name: v for v in db.scalars(select(Vendor))}
    ramesh = vendors["Ramesh Traders"]
    assert (ramesh.type, ramesh.gstin, ramesh.city) == ("material_supplier", GSTIN, "Ahmedabad")
    assert ramesh.contacts[0].phone == "9800000001"
    assert ramesh.bank_accounts[0].ifsc == "HDFC0001234"
    assert vendors["Sample Waterproofing Contractor"].type == "labour_contractor"
    assert vendors["Sample Waterproofing Contractor"].gstin is None
    assert vendors["Swift Transport"].type == "transporter"
    assert vendors["Swift Transport"].bank_accounts == []

    second = import_vendors(db, FIXTURES / "vendors.xlsx")
    assert (second.created, second.already_present) == (0, 3)
    assert len(db.scalars(select(Vendor)).all()) == 3


def test_powerplay_team_import_creates_inactive_users_without_password(db, client, login_as):
    first = import_team(db, FIXTURES / "team.xlsx")
    assert (first.rows_in_file, first.created, first.duplicates_merged) == (4, 3, 1)
    assert [r for _, r in first.skipped] == ["Powerplay's own support account"]
    users = {u.full_name: u for u in db.scalars(select(User))}
    asha, ravi, neha = users["Asha Patel"], users["Ravi Kumar"], users["Neha Shah"]
    assert (asha.email, asha.job_title, asha.phone) == (
        "asha.patel@example.com",
        "Project Manager",
        "9000000001",
    )
    assert ravi.email is None and neha.email is None
    assert all(not u.is_active and u.password_hash is None for u in (asha, ravi, neha))

    # cannot log in, and cannot be activated without an email and a password
    r = client.post(
        "/api/auth/login", json={"email": "asha.patel@example.com", "password": "x" * 12}
    )
    assert r.status_code == 401
    admin, h = login_as("super_admin", email="boss@example.com")
    assert (
        admin.patch(f"/api/users/{asha.id}", json={"is_active": True}, headers=h).status_code == 422
    )
    admin.post(f"/api/users/{asha.id}/reset-password", json={"password": PASSWORD}, headers=h)
    assert (
        admin.patch(f"/api/users/{asha.id}", json={"is_active": True}, headers=h).status_code == 200
    )
    assert login(TestClient(app), "asha.patel@example.com")
    listed = {u["full_name"]: u for u in admin.get("/api/users", headers=h).json()}
    assert listed["Ravi Kumar"]["has_password"] is False

    second = import_team(db, FIXTURES / "team.xlsx")
    assert (second.created, second.already_present) == (0, 3)


# --- Excel export --------------------------------------------------------------------------------


def _sheet(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response.headers["content-disposition"]
    ws = load_workbook(BytesIO(response.content)).active
    return [[cell.value for cell in row] for row in ws.iter_rows()]


def test_exports_are_valid_xlsx_with_the_right_columns(login_as, make_user):
    c, h = login_as("super_admin", email="boss@example.com")
    c.post("/api/products", json={"code": "P1", "name": "Primer", "unit": "kg"}, headers=h)
    pid = c.get("/api/products", headers=h).json()["items"][0]["id"]
    c.post(f"/api/products/{pid}/prices", json={"purchase_rate": "310"}, headers=h)
    c.post("/api/clients", json={"name": "Acme Builders"}, headers=h)

    rows = _sheet(c.get("/api/products/export", headers=h))
    assert rows[0][:5] == ["Code", "Name", "Brand", "Category", "Unit"]
    assert "Purchase rate" in rows[0] and rows[1][rows[0].index("Purchase rate")] == 310

    sales, sh = _client(make_user, "sales")
    rows = _sheet(sales.get("/api/products/export", headers=sh))
    assert "Purchase rate" not in rows[0] and rows[1][1] == "Primer"  # cost hidden, as in the list

    assert _sheet(c.get("/api/clients/export", headers=h))[1][0] == "Acme Builders"
    assert _sheet(c.get("/api/categories/export", params={"kind": "work"}, headers=h))[0] == [
        "Kind",
        "Name",
        "Parent",
        "Order",
        "Products",
        "Active",
    ]
    assert _sheet(c.get("/api/users/export", headers=h))[0][:3] == ["Name", "Email", "Phone"]
    assert _sheet(c.get("/api/audit/export", headers=h))[0][:3] == ["When", "User", "Action"]
    for path in (
        "/api/vendors/export",
        "/api/tags/export",
        "/api/unit-conversions/export",
        "/api/systems/export",
        "/api/tc/clauses/export",
        "/api/settings/gstins/export",
        "/api/settings/bank-accounts/export",
        "/api/library/search/export",
    ):
        assert _sheet(c.get(path, headers=h))[0], path
    # same permission as the list
    assert sales.get("/api/users/export", headers=sh).status_code == 403
    assert sales.get("/api/vendors/export", headers=sh).status_code == 403


# --- company settings ----------------------------------------------------------------------------


def test_company_settings_permissions_profile_gstin_and_bank(make_user):
    accounts, ah = _client(make_user, "accounts")  # settings.company + finance.view
    estimator, eh = _client(make_user, "estimator")
    assert estimator.get("/api/settings/company", headers=eh).status_code == 403
    assert estimator.get("/api/settings/bank-accounts", headers=eh).status_code == 403

    r = accounts.patch(
        "/api/settings/company",
        json={
            "legal_name": "Ethios Enviro Solutions Pvt Ltd",
            "pan": "aapfu0939f",
            "tan": "AHMA12345B",
            "cin": "U74999GJ2019PTC123456",
            "tds_percent": "2",
            "email": "info@example.com",
        },
        headers=ah,
    )
    assert r.status_code == 200 and r.json()["pan"] == "AAPFU0939F"
    assert r.json()["default_gst_percent"] == "18.00"
    assert (
        accounts.patch("/api/settings/company", json={"tan": "BAD"}, headers=ah).status_code == 422
    )

    g1 = accounts.post(
        "/api/settings/gstins",
        json={"gstin": GSTIN, "state": "Maharashtra", "address": "Mumbai"},
        headers=ah,
    ).json()
    assert g1["is_default"]  # the first one becomes the default
    g2 = accounts.post(
        "/api/settings/gstins",
        json={"gstin": "24AAPFU0939F1ZZ", "state": "Gujarat", "address": "Ahmedabad"},
        headers=ah,
    )
    assert g2.status_code == 422  # bad check digit
    assert (
        accounts.post(
            "/api/settings/gstins", json={"gstin": GSTIN, "state": "X", "address": "Y"}, headers=ah
        ).status_code
        == 409
    )

    b = accounts.post(
        "/api/settings/bank-accounts", json={**BANK, "account_name": "EESPL"}, headers=ah
    )
    assert b.status_code == 201 and b.json()["is_default"]
    assert accounts.get("/api/settings/bank-accounts", headers=ah).json()[0]["account_number"] == (
        "123456789012"
    )


def test_logo_upload_validates_and_is_served(make_user, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "media_dir", str(tmp_path))
    c, h = _client(make_user, "office_admin")
    png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
    r = c.post(
        "/api/settings/company/logo", files={"file": ("logo.png", png, "image/png")}, headers=h
    )
    assert r.status_code == 200 and r.json()["has_logo"]
    assert (tmp_path / "company" / "logo.png").read_bytes() == png
    assert c.get("/api/settings/company/logo", headers=h).content == png
    fake = c.post(
        "/api/settings/company/logo",
        headers=h,
        files={"file": ("logo.png", b"<svg/>", "image/png")},
    )
    assert fake.status_code == 422
    svg = c.post(
        "/api/settings/company/logo",
        headers=h,
        files={"file": ("logo.svg", b"<svg/>", "image/svg+xml")},
    )
    assert svg.status_code == 422
    assert c.delete("/api/settings/company/logo", headers=h).status_code == 204
    assert not (tmp_path / "company" / "logo.png").exists()


def test_tags_crud(login_as):
    c, h = login_as("estimator")
    r = c.post("/api/tags", json={"module": "issue", "name": "Leakage"}, headers=h)
    assert r.status_code == 201
    assert (
        c.post("/api/tags", json={"module": "issue", "name": "leakage"}, headers=h).status_code
        == 409
    )
    assert (
        c.post("/api/tags", json={"module": "material", "name": "Leakage"}, headers=h).status_code
        == 201
    )
    assert c.post("/api/tags", json={"module": "bogus", "name": "x"}, headers=h).status_code == 422
    r = c.patch(f"/api/tags/{r.json()['id']}", json={"is_archived": True}, headers=h)
    assert r.json()["is_archived"]
    assert len(c.get("/api/tags", params={"archived": "false"}, headers=h).json()["items"]) == 1


def test_categories_seeded_as_rows(db):
    kinds = dict(db.execute(select(Category.kind, text("count(*)")).group_by(Category.kind)).all())
    assert kinds == {"material": 8, "work": 12}
