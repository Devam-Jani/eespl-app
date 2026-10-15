# ruff: noqa: E501, F811  (invented data reads better unwrapped; shared fixtures are imported)
"""From the rate working review: Samgard is Bronco Cemguard, the margin basis (on cost or on the
selling price), the application's wastage from the settings, and 10.7639 shown as used.
Invented data only."""

from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select, text

from app.db import engine as db_engine
from app.masters.models import Product, System
from app.masters.named_systems import cemguard
from app.masters.rate import SQFT_PER_SQM
from app.survey.models import AreaType
from tests.test_finance import boss  # noqa: F401  (fixtures)
from tests.test_rate_working import _group, _working, boq  # noqa: F401  (fixtures)


def d(v) -> Decimal:
    return Decimal(str(v))


def cents(v: Decimal) -> Decimal:
    return v.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def test_samgard_is_bronco_cemguard_and_samgard_still_finds_it(boss, db):
    c, h = boss
    with db_engine.begin() as conn:
        conn.execute(
            text("DELETE FROM systems WHERE code LIKE 'CEMGUARD%' OR code LIKE 'SAMGARD%'")
        )
        # an older database: the systems still called Samgard
        conn.execute(
            text(
                "INSERT INTO systems (code, name, unit, to_confirm) VALUES ('SAMGARD', 'Samgard', 'sqm', true), ('SAMGARD-PLUS', 'Samgard Plus', 'sqm', true)"
            )
        )
        ids = cemguard(conn)
        assert cemguard(conn) == ids  # safe to run again
    systems = {s.code: s for s in db.scalars(select(System).where(System.code.like("CEMGUARD%")))}
    assert {s.name for s in systems.values()} == {"Cemguard", "Cemguard Plus"}
    assert not db.scalar(select(System.id).where(System.code.like("SAMGARD%")))
    plus = systems["CEMGUARD-PLUS"]
    assert [c.product.name for c in plus.components] == ["BRONCO CEMGUARD PLUS"]
    assert plus.components[0].consumption_per_unit is None and plus.components[0].to_confirm
    assert plus.components[0].coats == 2  # "1st coat ... 2nd & final coat" in our offers
    assert systems["CEMGUARD"].components[0].coats is None
    cem = db.scalar(select(Product).where(Product.code == "BRONCO-CEMGUARD"))
    assert "SAMGARD" in cem.aliases and cem.brand == "Bronco"
    assert db.scalar(select(Product).where(Product.code == "BRONCO-CEMCRYSTAL-PLUS")) is not None
    from app.masters.models import ProductPrice

    assert not db.scalar(
        select(ProductPrice.id).where(ProductPrice.product_id == cem.id)
    )  # no invented price
    # "Samgard" still finds them
    found = [p["name"] for p in c.get("/api/products?q=samgard", headers=h).json()["items"]]
    assert {"BRONCO CEMGUARD", "BRONCO CEMGUARD PLUS"} <= set(found)
    found = [s["name"] for s in c.get("/api/systems?q=samgard", headers=h).json()["items"]]
    assert {"Cemguard", "Cemguard Plus"} <= set(found)
    sunken = db.scalar(select(AreaType).where(AreaType.name == "Sunken slab"))
    assert sunken.default_system_id == systems["CEMGUARD-PLUS"].id
    for name in ("Sunken slab", "Toilet / bathroom", "Terrace"):
        a = db.scalar(select(AreaType).where(AreaType.name == name))
        a.default_system_id, a.suitable_system_ids = None, []
    db.commit()


def test_margin_on_cost_or_on_the_selling_price_shown_on_the_sheet_and_in_the_excel(boss, boq):
    c, h = boss
    w = c.post(f"/api/rate-working/{_working(c, h, boq)['id']}/suggest", headers=h).json()
    raft = _group(w, "raft")
    g = c.put(
        f"/api/rate-working/groups/{raft['id']}", json={"system_id": boq["system"]}, headers=h
    ).json()
    cost = d(g["cost_sqm"])
    assert g["margin_basis"] == "cost" and d(g["rate_sqm"]) == cents(cost * d("1.25"))
    assert "rate = cost × (1 + margin)" in g["margin_basis_label"]
    assert (
        c.put("/api/prices/settings", json={"margin_basis": "selling"}, headers=h).json()[
            "margin_basis"
        ]
        == "selling"
    )
    g = c.get(f"/api/rate-working/{w['id']}", headers=h).json()
    g = next(x for x in g["groups"] if x["id"] == raft["id"])
    assert g["margin_basis"] == "selling" and "cost ÷ (1 − margin)" in g["margin_basis_label"]
    assert d(g["rate_sqm"]) == cents(cost / d("0.75"))
    assert d(g["rate_sqft"]) == cents(cost / d("0.75") / SQFT_PER_SQM)
    assert d(g["margin_sqm"]) == cents(cost / d("0.75") - cost)
    book = load_workbook(
        BytesIO(c.get(f"/api/rate-working/{w['id']}/export.xlsx", headers=h).content)
    )
    cells = [x for row in book["WORKING"].iter_rows(values_only=True) for x in row if x is not None]
    assert any(
        isinstance(x, str) and x.startswith("Margin basis: margin on selling price") for x in cells
    )
    assert "MARGIN (on selling price)" in cells
    assert any(
        isinstance(x, str) and "/(1-D" in x for x in cells
    )  # the Excel works it the same way
    assert (
        c.put("/api/prices/settings", json={"margin_basis": "sideways"}, headers=h).status_code
        == 422
    )


def test_the_application_fills_the_wastage_from_the_settings_and_says_so(boss, boq):
    c, h = boss
    w = c.post(f"/api/rate-working/{_working(c, h, boq)['id']}/suggest", headers=h).json()
    toilet = _group(w, "toilet")
    g = c.put(
        f"/api/rate-working/groups/{toilet['id']}", json={"application": "footing"}, headers=h
    ).json()
    assert d(g["wastage_percent"]) == 15 and "footing 12-15 %" in g["wastage_source"]
    assert {a["name"] for a in g["applications"]} >= {"raft", "footing"}
    g = c.put(
        f"/api/rate-working/groups/{toilet['id']}", json={"wastage_percent": "9"}, headers=h
    ).json()
    assert d(g["wastage_percent"]) == 9 and g["wastage_source"] == "typed for this site"
    # the system's coats are shown (blank ones as to confirm)
    raft = _group(w, "raft")
    g = c.put(
        f"/api/rate-working/groups/{raft['id']}", json={"system_id": boq["system"]}, headers=h
    ).json()
    assert (
        g["application"] == "raft" and "Settings > Process rules: raft 5-7 %" in g["wastage_source"]
    )
    assert [x["coats"] for x in g["system_coats"]] == [None, None]


def test_the_labour_line_shows_10_7639_as_used(boss, boq):
    c, h = boss
    w = c.post(f"/api/rate-working/{_working(c, h, boq)['id']}/suggest", headers=h).json()
    raft = _group(w, "raft")
    g = c.put(
        f"/api/rate-working/groups/{raft['id']}", json={"system_id": boq["system"]}, headers=h
    ).json()
    labour = next(r for r in g["rows"] if r["item"] == "Labour")
    assert "× 10.7639" in labour["ref"] and d(labour["per_sqm"]) == cents(12 * SQFT_PER_SQM)
    book = load_workbook(
        BytesIO(c.get(f"/api/rate-working/{w['id']}/export.xlsx", headers=h).content)
    )
    texts = [
        x for row in book["WORKING"].iter_rows(values_only=True) for x in row if isinstance(x, str)
    ]
    assert any("× 10.7639 per sqm" in x for x in texts) and any("*10.7639" in x for x in texts)
    assert not any("10.764)" in x or "10.764 " in x for x in texts)
