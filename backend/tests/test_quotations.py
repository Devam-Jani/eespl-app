# ruff: noqa: E501  (invented offer wording reads better unwrapped)
"""Quotations: the offer import, libraries with versions, placeholders, rates, revisions, status
changes with follow-ups, won hand-over, Word and PDF output. Invented data only."""

import io
from datetime import date, timedelta
from decimal import Decimal

import pdfplumber
import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from sqlalchemy import select

from app.analytics import alerts
from app.analytics.models import Alert
from app.config import settings
from app.crm.models import Lead, LeadActivity
from app.masters.models import LibraryItem, Product, ProductPrice, System, SystemComponent
from app.masters.rate import system_rate
from app.models import User
from app.quotations.importer import import_offer
from app.quotations.models import (
    LibraryVersion,
    OfferItem,
    OfferLine,
    Quotation,
    QuotationFile,
    QuotationFollowUp,
    Reference,
    SpecBlock,
)
from app.sites.models import Site
from app.tenders.models import BoqLine, Tender


def _png() -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (40, 20), (15, 110, 90)).save(buf, "PNG")
    return buf.getvalue()


PNG = _png()


def D(x) -> Decimal:
    return Decimal(str(x))


@pytest.fixture(autouse=True)
def media(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "media_dir", str(tmp_path))
    return tmp_path


@pytest.fixture
def boss(login_as):
    return login_as("super_admin", email="boss@example.com")


# --- the invented offer for the import -----------------------------------------------------------


def _num(p, num_id: int) -> None:
    ppr = p._p.get_or_add_pPr()
    num = OxmlElement("w:numPr")
    lvl = OxmlElement("w:ilvl")
    lvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    num.append(lvl)
    num.append(nid)
    ppr.append(num)


def _p(d, *parts, num: int | None = None):
    """parts: plain strings, ("b", text) bold, ("h", text) highlighted."""
    p = d.add_paragraph()
    for part in parts:
        kind, text = part if isinstance(part, tuple) else ("", part)
        r = p.add_run(text)
        if "b" in kind:
            r.bold = True
        if "h" in kind:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    if num:
        _num(p, num)
    return p


def _row(t, cells):
    row = t.add_row()
    for i, c in enumerate(cells):
        paras = c if isinstance(c, list) else [c]
        row.cells[i].text = paras[0]
        for extra in paras[1:]:
            row.cells[i].add_paragraph(extra)
    return row


def invented_offer(path):
    d = Document()
    d.sections[0].header.paragraphs[0].add_run().add_picture(io.BytesIO(PNG))
    _p(d, "Date: ", ("b", "01.08.2026"))
    _p(d, "Firm Name: ", ("b", "DEMO BUILDERS"))
    _p(d, ("b", "Ahmedabad"))
    _p(d, "Kind Attn.: ", ("b", "Mr. Test Person"))
    _p(d, "Dear Sir,")
    _p(d, "SUB: OFFER FOR ", ("h", "TERRACE"), " WATERPROOFING WORK.")
    _p(
        d,
        "Reference to subject, we offer waterproofing by using ",
        ("b", "ACME"),
        " products at ",
        ("b", "Demo Builders"),
        " with the methodology.",
    )
    _p(d, "Thanking you.")
    _p(d, ("b", "For ACME COATINGS PVT. LTD."))
    _p(d, ("b", "Test Signer | Sales Engineer"))
    _p(d, "Encl.:")
    _p(d, "Technical Specification", num=5)
    _p(d, "Budgetary Offer.", num=5)
    # specifications
    _p(
        d,
        ("bh", "OPTION :1 - "),
        (
            "b",
            "ITEM NO: 1 - TECHNICAL SPECIFICATION & METHEDOLOGY FOR TERRACE WATERPROOFING USING PU COATING",
        ),
    )
    _p(d, ("b", "SURFACE PREPARATION:"))
    _p(d, "Clean the slab free from dust.", num=11)
    _p(d, "Remove the old tiles. ", ("h", "(Client’s scope)"), num=11)
    _p(d, ("b", "PROTECTIVE COATING"))
    _p(d, "Apply two coats of ", ("b", "ACME PU"), " done on on the slab.", num=11)
    _p(d, ("bh", "OR"))
    _p(
        d,
        ("bh", "OPTION :2 - "),
        (
            "b",
            "ITEM NO: 1 - TECHNICAL SPECIFICATION FOR TERRACE WATERPROOFING USING CEMENTITIOUS COATING",
        ),
    )
    _p(d, ("b", "SURFACE PREPARATION:"))
    _p(d, "Wet the slab. Al low to dry for 2 hours.", num=12)
    _p(d, ("b", "ITEM NO: 2 - TECHNICAL SPECIFICATION FOR TANK AREA WATERPROOFING WORKS:"))
    _p(d, ("b", "LAYING MEMBRANE (HORIZINTAL AREA):"))
    _p(d, "Lay the membrane.", num=13)
    _p(d, ("b", "PROTECTIVE COATING"))
    _p(d, "Coating shall be done entirely over the Roof.", num=13)
    # budgetary offer
    _p(d, ("b", "ITEM NO: 1 - BUDGETARY OFFER FOR TERRACE AREA WATERPROOFING WORK:"))
    t = d.add_table(rows=1, cols=4)
    for i, h in enumerate(["Sr No.", "Item Description", "UoM", "Rate"]):
        t.rows[0].cells[i].text = h
    _row(t, ["Opt.1.", "Two coats of ACME PU membrane", "SQ FT", "60.00"])
    _row(t, ["Opt.2.", "Two coats of ACME cementitious coating", "SQ FT", "45.00"])
    _row(t, ["2.", "Pipe sleeve packing", "NOS", "400.00"])
    _row(t, ["", "Pressure grouting (If required)", "NOS", "350.00"])
    _p(d, ("b", "ITEM NO: 2 - BUDGETARY OFFER FOR TANK AREA WATERPROOFING WORK:"))
    t = d.add_table(rows=1, cols=4)
    for i, h in enumerate(["Sr No.", "Item Description", "UoM", "Rate"]):
        t.rows[0].cells[i].text = h
    _row(
        t,
        [
            "1.",
            ["Surface preparation.", "Option: 1 - coat A", "OR", "Option: 2 - coat B"],
            ["SQ FT", "SQ FT"],
            ["45.00", "47.00"],
        ],
    )
    r = _row(t, ["", "Protective plaster", "Client’s Scope", ""])
    r.cells[2].merge(r.cells[3])
    # terms
    _p(d, ("b", "GENERAL TERMS AND CONDITIONS"))
    _p(d, ("b", "CLIENTS OBLIGATION"))
    _p(d, "Water shall be provided free of charge.", num=1)
    _p(d, ("b", "Force Majeure: "), "prices may change.", num=1)
    _p(d, ("b", "TAXES: "), num=1)
    _p(d, "GST is extra.")
    _p(d, ("b", "VALIDITY: "), num=1)
    _p(d, "30 days.")
    _p(
        d,
        (
            "b",
            "Work shall be carried out by our Authorized Applicator. And the Work Order shall be issued in favor of our Authorized Applicator - ETHIOS ENVIRO SOLUTIONS PVT LTD.",
        ),
        num=1,
    )
    _p(d, ("b", "TERMS OF PAYMENT: "), "Payment as follows:", num=1)
    _p(d, "RA bills in 7 days.", num=7)
    _p(d, "Final bill in 15 days.", num=7)
    _p(d, ("b", "Our Esteemed Clients for Waterproofing Projects in Gujarat"))
    t = d.add_table(rows=1, cols=5)
    for i, h in enumerate(["Sr. No.", "Client Name", "Project", "Application", "Area Covered"]):
        t.rows[0].cells[i].text = h
    _row(t, ["1", "Demo Developers", "Demo Heights", "Terrace - PU coating", "12,500 SQFT"])
    _row(t, ["2", "Sample Infra", "Sample Park", "Expansion joint", "355 RMT"])
    d.save(str(path))
    return path


def test_import_of_an_invented_offer(db, tmp_path, media):
    res = import_offer(db, invented_offer(tmp_path / "offer.docx"))
    db.commit()
    assert res.counts == {
        "letterheads": 1,
        "letter_templates": 1,
        "items": 2,
        "spec_blocks": 3,
        "options": 4,
        "offer_lines": 7,
        "tc_clauses": 6,
        "references": 2,
    }
    for fix in (
        "METHEDOLOGY -> METHODOLOGY",
        "on on -> on",
        "Al low -> Allow",
        "HORIZINTAL -> HORIZONTAL",
        "over the Roof -> over the surface (tanks, walls, pools)",
    ):
        assert res.fixes.get(fix) == 1, (fix, res.fixes)
    assert any(k.startswith("duplicated UoM") for k in res.fixes)
    blocks = list(db.scalars(select(SpecBlock).order_by(SpecBlock.id)))
    assert [b.option_label for b in blocks] == ["1", "2", None]
    assert blocks[0].heading_prefix == "TECHNICAL SPECIFICATION & METHODOLOGY FOR"
    steps = blocks[0].sections[0]["steps"]
    assert steps[1]["client_scope"] and "Client" not in steps[1]["text"]
    assert "**ACME PU**" in blocks[0].sections[1]["steps"][0]["text"]
    assert (
        "(raft" not in blocks[2].sections[0]["heading"]
        and "horizontal" in blocks[2].sections[0]["heading"]
    )
    assert "over the surface" in blocks[2].sections[1]["steps"][0]["text"]
    lines = list(db.scalars(select(OfferLine).order_by(OfferLine.id)))
    assert [(ln.uom, ln.default_rate) for ln in lines[:4]] == [
        ("sqft", D("60.00")),
        ("sqft", D("45.00")),
        ("nos", D("400.00")),
        ("nos", D("350.00")),
    ]
    assert lines[3].if_required and "If required" not in lines[3].description
    split = lines[4:6]
    assert [ln.default_rate for ln in split] == [D("45.00"), D("47.00")]
    assert "coat A" in split[0].description and "coat B" not in split[0].description
    assert lines[6].client_scope and lines[6].default_rate is None
    tank = db.scalar(select(OfferItem).where(OfferItem.name == "Tank area"))
    assert [ln.option for ln in tank.lines] == ["1", "2", None]
    assert all(r.needs_check for r in blocks + lines)
    refs = list(db.scalars(select(Reference).order_by(Reference.sort_order)))
    assert [(r.area_value, r.area_unit, r.state) for r in refs] == [
        (D("12500"), "SQFT", "Gujarat"),
        (D("355"), "RMT", "Gujarat"),
    ]
    # every imported row has its first version
    assert db.scalar(select(LibraryVersion).where(LibraryVersion.action == "import")) is not None
    with pytest.raises(ValueError, match="Already imported"):
        import_offer(db, tmp_path / "offer.docx")


# --- a library built through the API -------------------------------------------------------------


@pytest.fixture
def lib(boss, db):
    client, h = boss
    pu = Product(code="Q-PU", name="Invented PU coating", unit="kg")
    db.add(pu)
    db.flush()
    db.add(
        ProductPrice(
            product_id=pu.id,
            purchase_rate=D(300),
            freight_per_unit=D(0),
            effective_from=date(2026, 1, 1),
        )
    )
    system = System(code="Q-SYS", name="Invented terrace PU system", unit="sqm", labour_rate=50)
    db.add(system)
    db.flush()
    db.add(
        SystemComponent(
            system_id=system.id, product_id=pu.id, consumption_per_unit=D("1.5"), wastage_percent=5
        )
    )
    item = LibraryItem(
        source_key="Q-LIB-1",
        match_key="q-lib-1",
        description="Invented brick bat screed 75 mm",
        unit="sqft",
        median_rate=D(27),
        latest_rate=D(28),
        boq_count=3,
    )
    db.add(item)
    db.commit()

    def post(kind, body):
        r = client.post(f"/api/quotations/library/{kind}", json=body, headers=h)
        assert r.status_code == 201, r.text
        return r.json()["id"]

    letter = post(
        "letter",
        {
            "name": "Invented letter",
            "opening": "Date: **{date}**\nFirm Name: **{client_firm}**\n**{client_city}**\nKind Attn.: **{attention}**",
            "subject": "SUB: OFFER FOR {areas_list} WATERPROOFING WORK.",
            "body": "We offer **{brand}** products at **{project}**.\nRegards, {salesperson} ({designation}). Enclosed: {enclosures}.",
            "enclosures": "Technical Specification\nBudgetary Offer",
        },
    )
    head = post(
        "letterhead",
        {
            "name": "Invented Co",
            "company_name": "INVENTED COATINGS PVT. LTD.",
            "signatory_firm": "For INVENTED COATINGS PVT. LTD.",
            "brand": "INVENTO",
            "letter_template_id": letter,
        },
    )
    s1 = post(
        "spec",
        {
            "title": "Terrace waterproofing using invented PU",
            "option_label": "1",
            "sections": [
                {
                    "heading": "Surface preparation",
                    "steps": [
                        {"text": "Clean the slab."},
                        {"text": "Remove tiles.", "client_scope": True},
                    ],
                },
                {
                    "heading": "Protective coating",
                    "steps": [{"text": "Two coats of **Invented PU coating**."}],
                },
            ],
        },
    )
    s2 = post(
        "spec",
        {
            "title": "Terrace waterproofing using invented cementitious coat",
            "option_label": "2",
            "sections": [{"heading": "Protective coating", "steps": [{"text": "Two coats."}]}],
        },
    )
    s3 = post(
        "spec",
        {
            "title": "Tank waterproofing",
            "sections": [
                {"heading": "Surface preparation", "steps": [{"text": "Clean the walls."}]},
                {"heading": "Protective coating", "option": "1", "steps": [{"text": "Coat A."}]},
                {"heading": "Protective coating", "option": "2", "steps": [{"text": "Coat B."}]},
            ],
        },
    )
    l_sys = post(
        "line",
        {
            "description": "Invented PU coating system, two coats",
            "uom": "sqm",
            "rate_source": "system",
            "system_id": system.id,
        },
    )
    l_fix = post(
        "line",
        {
            "description": "Cementitious coating, two coats ==(Rate shall change for higher thickness)==",
            "uom": "sqft",
            "rate_source": "fixed",
            "default_rate": "45.00",
        },
    )
    l_lib = post(
        "line",
        {
            "description": "Brick bat screed 75 mm",
            "uom": "sqft",
            "rate_source": "library",
            "library_item_id": item.id,
        },
    )
    l_req = post(
        "line",
        {
            "description": "Pressure grouting",
            "uom": "nos",
            "rate_source": "fixed",
            "default_rate": "350",
            "if_required": True,
        },
    )
    l_cli = post(
        "line",
        {
            "description": "Protective plaster",
            "uom": "sqft",
            "rate_source": "fixed",
            "client_scope": True,
        },
    )
    l_tank = post(
        "line",
        {
            "description": "Tank coating",
            "uom": "sqft",
            "rate_source": "fixed",
            "default_rate": "49",
        },
    )
    terrace = post(
        "item",
        {
            "name": "Terrace area",
            "budget_title": "Terrace area waterproofing work",
            "specs": [{"spec_block_id": s1}, {"spec_block_id": s2}],
            "lines": [
                {"offer_line_id": l_sys, "option": "1"},
                {"offer_line_id": l_fix, "option": "2"},
                {"offer_line_id": l_lib},
                {"offer_line_id": l_req},
                {"offer_line_id": l_cli},
            ],
        },
    )
    tank = post(
        "item",
        {
            "name": "Tank area",
            "budget_title": "Tank area waterproofing work",
            "specs": [{"spec_block_id": s3}],
            "lines": [{"offer_line_id": l_tank}],
        },
    )
    return {
        "head": head,
        "letter": letter,
        "s1": s1,
        "terrace": terrace,
        "tank": tank,
        "system": system.id,
        "l_sys": l_sys,
        "l_fix": l_fix,
        "l_lib": l_lib,
    }


def new_quotation(client, h, lib, items=None, **extra):
    body = {
        "client_firm": "Invented Builders",
        "client_city": "Surat",
        "attention": "Mr. Invented",
        "project": "Invented Towers",
        "letterhead_id": lib["head"],
        "items": items or [{"offer_item_id": lib["terrace"]}, {"offer_item_id": lib["tank"]}],
        **extra,
    }
    r = client.post("/api/quotations", json=body, headers=h)
    assert r.status_code == 201, r.text
    return r.json()


def lines_of(q):
    return [ln for it in q["items"] for ln in it["lines"]]


def test_placeholders_filled_and_unknown_ones_refused(boss, lib):
    client, h = boss
    r = client.post(
        "/api/quotations/library/letter",
        json={"name": "Bad", "subject": "For {client_name}", "body": "x"},
        headers=h,
    )
    assert r.status_code == 422 and "{client_name}" in r.json()["detail"]
    q = new_quotation(client, h, lib)
    assert q["code"].startswith("QTN-") and q["revision"] == 0
    html = client.get(f"/api/quotations/{q['id']}/preview", headers=h).json()["html"]
    for text in (
        "Invented Builders",
        "Surat",
        "Mr. Invented",
        "INVENTO",
        "Invented Towers",
        "TERRACE AREA &amp; TANK AREA",
        date.today().strftime("%d.%m.%Y"),
    ):
        assert text in html, text
    assert "{" not in html.split("<body>")[1].split("<style>")[0]
    r = client.put(f"/api/quotations/{q['id']}", json={"subject": "Offer {unknown}"}, headers=h)
    assert r.status_code == 422


def test_rates_from_system_library_fixed_override_and_margin(boss, lib, login_as, db):
    client, h = boss
    q = new_quotation(client, h, lib)
    by_src = {ln["offer_line_id"]: ln for ln in lines_of(q)}
    expected = system_rate(db, db.get(System, lib["system"]))
    sys_line = by_src[lib["l_sys"]]
    assert (
        D(sys_line["rate"]) == expected.rate.quantize(D("0.01"))
        and sys_line["rate_source"] == "system"
    )
    assert D(sys_line["cost_rate"]) == expected.base_cost.quantize(D("0.01"))
    assert (
        D(by_src[lib["l_lib"]]["rate"]) == D("27.00")
        and by_src[lib["l_lib"]]["rate_source"] == "library"
    )
    assert D(by_src[lib["l_fix"]]["rate"]) == D("45.00")
    # override for this quotation only
    r = client.put(
        f"/api/quotations/{q['id']}/lines/{by_src[lib['l_fix']]['id']}",
        json={"rate": "52.50"},
        headers=h,
    )
    line = next(ln for ln in lines_of(r.json()) if ln["offer_line_id"] == lib["l_fix"])
    assert line["overridden"] and D(line["rate"]) == D("52.50") and line["rate_source"] == "manual"
    assert db.get(OfferLine, lib["l_fix"]).default_rate == D("45.00")
    # a salesperson without tender.margin sees selling rates only
    sales, sh = login_as("sales", email="sales@example.com")
    sq = new_quotation(sales, sh, lib)
    assert all("cost_rate" not in ln and "margin_percent" not in ln for ln in lines_of(sq))
    # the offer never prints cost or margin
    pdf = client.get(f"/api/quotations/{q['id']}/pdf", headers=h).content
    with pdfplumber.open(io.BytesIO(pdf)) as doc:
        text = " ".join(p.extract_text() or "" for p in doc.pages)
    cost = f"{expected.base_cost:.2f}"
    assert cost not in text and "margin" not in text.lower()
    assert "52.50" in text and "Client's Scope" in text and "(If required)" in text


def test_editing_a_quotation_leaves_the_library_and_save_back_versions_it(boss, lib, db):
    client, h = boss
    old = new_quotation(client, h, lib)
    q = new_quotation(client, h, lib)
    item = q["items"][0]
    specs = item["specs"]
    specs[0]["sections"][0]["steps"][0]["text"] = "Clean the slab with a wire brush."
    r = client.put(
        f"/api/quotations/{q['id']}/items/{item['id']}", json={"specs": specs}, headers=h
    )
    assert r.status_code == 200
    block = client.get(f"/api/quotations/library/spec/{lib['s1']}", headers=h).json()
    assert block["version"] == 1 and block["sections"][0]["steps"][0]["text"] == "Clean the slab."
    r = client.post(
        f"/api/quotations/{q['id']}/save-back",
        json={"kind": "spec", "item_id": item["id"], "spec_index": 0},
        headers=h,
    )
    assert r.status_code == 200, r.text
    assert r.json()["saved_back"]["version"] == 2
    block = client.get(f"/api/quotations/library/spec/{lib['s1']}", headers=h).json()
    assert block["sections"][0]["steps"][0]["text"] == "Clean the slab with a wire brush."
    versions = client.get(f"/api/quotations/library/spec/{lib['s1']}/versions", headers=h).json()
    assert [v["action"] for v in versions] == ["save_back", "create"]
    # the older quotation keeps its own copy; a new one gets the new text
    kept = client.get(f"/api/quotations/{old['id']}", headers=h).json()
    assert kept["items"][0]["specs"][0]["sections"][0]["steps"][0]["text"] == "Clean the slab."
    fresh = new_quotation(client, h, lib)
    assert (
        fresh["items"][0]["specs"][0]["sections"][0]["steps"][0]["text"]
        == "Clean the slab with a wire brush."
    )
    # restore the first version (a third version)
    r = client.post(f"/api/quotations/library/spec/{lib['s1']}/restore/1", headers=h)
    assert (
        r.json()["version"] == 3
        and r.json()["sections"][0]["steps"][0]["text"] == "Clean the slab."
    )


def test_revision_diff_lost_needs_a_reason_and_follow_ups(boss, lib, db):
    client, h = boss
    lead = Lead(
        code="L-Q-1",
        contact_name="Mr. Invented",
        company="Invented Builders",
        status="contacted",
        owner_id=db.scalar(select(User.id).where(User.email == "boss@example.com")),
    )
    db.add(lead)
    db.commit()
    q = new_quotation(client, h, lib, lead_id=lead.id)
    r = client.post(f"/api/quotations/{q['id']}/status", json={"status": "sent"}, headers=h)
    assert r.status_code == 200, r.text
    sent = r.json()
    assert sent["status"] == "sent" and {f["kind"] for f in sent["files"]} == {"docx", "pdf"}
    assert sorted(f["day"] for f in sent["followups"]) == [2, 7, 15, 30]
    assert sent["followups"][0]["due_on"] == (date.today() + timedelta(days=2)).isoformat()
    acts = list(
        db.scalars(
            select(LeadActivity).where(
                LeadActivity.lead_id == lead.id, LeadActivity.type == "quotation"
            )
        )
    )
    assert any(a.quotation_file_id for a in acts)
    db.refresh(lead)
    assert lead.status == "quoted"
    # R1 with two rates changed
    r1 = client.post(f"/api/quotations/{q['id']}/revision", headers=h).json()
    assert r1["revision"] == 1 and r1["code"] == q["code"]
    for ln in lines_of(r1)[:2]:
        client.put(
            f"/api/quotations/{r1['id']}/lines/{ln['id']}",
            json={"rate": str(D(ln["rate"]) - 5)},
            headers=h,
        )
    diff = client.get(f"/api/quotations/{r1['id']}/diff", headers=h).json()
    assert diff["from"] == "R0" and diff["to"] == "R1" and len(diff["changed"]) == 2
    assert all(D(c["new_rate"]) == D(c["old_rate"]) - 5 for c in diff["changed"])
    # the old revision cannot be edited
    assert (
        client.put(f"/api/quotations/{q['id']}", json={"notes": "x"}, headers=h).status_code == 409
    )
    # sending R1 replaces R0's follow-ups
    client.post(f"/api/quotations/{r1['id']}/status", json={"status": "sent"}, headers=h)
    open_ = db.scalars(select(QuotationFollowUp).where(QuotationFollowUp.status == "open")).all()
    assert len(open_) == 4 and {f.quotation_id for f in open_} == {r1["id"]}
    # an overdue follow-up raises an alert for the salesperson
    open_[0].due_on = date.today() - timedelta(days=1)
    db.commit()
    alerts.run(db)
    db.commit()
    assert db.scalar(select(Alert).where(Alert.rule == "followup_overdue")) is not None
    # lost needs a reason, then the follow-ups stop
    assert (
        client.post(
            f"/api/quotations/{r1['id']}/status", json={"status": "lost"}, headers=h
        ).status_code
        == 422
    )
    r = client.post(
        f"/api/quotations/{r1['id']}/status",
        json={"status": "lost", "lost_reason": "price", "lost_note": "Invented competitor"},
        headers=h,
    )
    assert r.json()["status"] == "lost" and r.json()["lost_reason"] == "price"
    assert not db.scalars(select(QuotationFollowUp).where(QuotationFollowUp.status == "open")).all()


def test_won_creates_the_tender_and_site_and_stops_follow_ups(boss, lib, db):
    client, h = boss
    q = new_quotation(client, h, lib)
    client.post(f"/api/quotations/{q['id']}/status", json={"status": "sent"}, headers=h)
    terrace = q["items"][0]
    r = client.post(
        f"/api/quotations/{q['id']}/status",
        json={"status": "won", "choices": {str(terrace["id"]): "2"}},
        headers=h,
    )
    assert r.status_code == 200, r.text
    res = r.json()["result"]
    tender = db.get(Tender, res["tender_id"])
    site = db.get(Site, res["site_id"])
    assert tender.status == "won" and site.tender_id == tender.id
    lines = list(db.scalars(select(BoqLine).where(BoqLine.tender_id == tender.id)))
    texts = " | ".join(ln.description for ln in lines)
    assert (
        "Cementitious coating" in texts and "Invented PU coating system" not in texts
    )  # option 2 chosen
    assert "Protective plaster" not in texts  # client's scope never
    assert (
        len(lines) == 4
    )  # opt 2, screed, grouting (if required), tank coating ... and nothing else
    assert not db.scalars(select(QuotationFollowUp).where(QuotationFollowUp.status == "open")).all()
    # a second quotation won onto an existing site links it
    q2 = new_quotation(client, h, lib)
    r = client.post(
        f"/api/quotations/{q2['id']}/status", json={"status": "won", "site_id": site.id}, headers=h
    )
    assert r.json()["result"]["site_id"] == site.id


def test_expiry_after_validity(boss, lib, db):
    client, h = boss
    q = new_quotation(
        client, h, lib, validity_days=10, quote_date=(date.today() - timedelta(days=11)).isoformat()
    )
    client.post(f"/api/quotations/{q['id']}/status", json={"status": "sent"}, headers=h)
    client.get("/api/quotations", headers=h)
    assert db.get(Quotation, q["id"]).status == "expired"
    assert not db.scalars(select(QuotationFollowUp).where(QuotationFollowUp.status == "open")).all()


def test_issued_files_are_kept_per_revision_and_never_overwritten(boss, lib, db, media):
    client, h = boss
    q = new_quotation(client, h, lib)
    client.post(f"/api/quotations/{q['id']}/issue", headers=h)
    client.post(f"/api/quotations/{q['id']}/issue", headers=h)
    files = list(db.scalars(select(QuotationFile).where(QuotationFile.quotation_id == q["id"])))
    names = sorted(f.file_name for f in files)
    assert len(names) == 4 and len(set(names)) == 4 and f"{q['code']}-R0.pdf" in names
    assert all((media / f.path).exists() for f in files)
    r = client.get(f"/api/quotations/files/{files[0].id}", headers=h)
    assert r.status_code == 200


def test_docx_and_pdf_keep_headings_with_their_tables(boss, lib, db):
    client, h = boss
    picks = [{"offer_item_id": lib["terrace"]}, {"offer_item_id": lib["tank"]}] * 6
    q = new_quotation(client, h, lib, items=picks)
    for it in q["items"]:  # long descriptions so headings fall near page ends
        for ln in it["lines"]:
            client.put(
                f"/api/quotations/{q['id']}/lines/{ln['id']}",
                json={"description": ln["description"] + "\n" + "Long invented wording. " * 12},
                headers=h,
            )
    r = client.get(f"/api/quotations/{q['id']}/docx", headers=h)
    assert r.status_code == 200
    d = Document(io.BytesIO(r.content))
    heads = [p for p in d.paragraphs if "BUDGETARY OFFER FOR" in p.text]
    assert len(heads) == 12 and all(p.paragraph_format.keep_with_next for p in heads)
    for t in d.tables:
        first = t.rows[0]._tr.trPr
        assert first.find(qn("w:tblHeader")) is not None
        assert all(row._tr.trPr.find(qn("w:cantSplit")) is not None for row in t.rows)
    ors = [p for p in d.paragraphs if p.text.strip() == "OR"]
    assert ors and all(p.paragraph_format.keep_with_next for p in ors)
    footer = d.sections[0].footer._element.xml
    assert "NUMPAGES" in footer and f"{q['code']} R0" in footer
    pdf = client.get(f"/api/quotations/{q['id']}/pdf", headers=h).content
    with pdfplumber.open(io.BytesIO(pdf)) as doc:
        pages = [(p.extract_text() or "").split("\n") for p in doc.pages]
    assert len(pages) > 4
    seen = 0
    for lines in pages:
        assert any(f"Page {pages.index(lines) + 1} of {len(pages)}" in ln for ln in lines)
        for i, ln in enumerate(lines):
            if "BUDGETARY OFFER FOR" in ln or "TECHNICAL SPECIFICATION FOR" in ln:
                seen += 1
                after = [
                    x
                    for x in lines[i + 1 :]
                    if x.strip() and "Page " not in x and q["code"] not in x
                ]
                assert len(after) >= 2, f"heading alone at the bottom of a page: {ln}"
                if "BUDGETARY" in ln:
                    assert any(x.startswith("Sr No.") for x in after[:3]), ln
    assert seen >= 24


def test_permissions_client_403_and_sales_own_scope(boss, lib, login_as):
    client, h = boss
    q = new_quotation(client, h, lib)
    portal, ph = login_as("client", email="client@example.com")
    for path in (
        "/api/quotations",
        f"/api/quotations/{q['id']}",
        "/api/quotations/library/spec",
        "/api/quotations/lookups",
    ):
        assert portal.get(path, headers=ph).status_code == 403, path
    sales, sh = login_as("sales", email="sales2@example.com")
    assert sales.get(f"/api/quotations/{q['id']}", headers=sh).status_code == 404  # not theirs
    own = new_quotation(sales, sh, lib)
    assert [x["id"] for x in sales.get("/api/quotations", headers=sh).json()] == [own["id"]]
    assert (
        sales.post(
            "/api/quotations/library/letter",
            json={"name": "x", "subject": "x", "body": "x"},
            headers=sh,
        ).status_code
        == 403
    )
    r = sales.post(f"/api/quotations/{own['id']}/save-back", json={"kind": "letter"}, headers=sh)
    assert r.status_code == 403
    est, eh = login_as("estimator", email="est@example.com")
    assert est.get(f"/api/quotations/{q['id']}", headers=eh).status_code == 200
    assert (
        est.post(
            f"/api/quotations/{q['id']}/status", json={"status": "sent"}, headers=eh
        ).status_code
        == 403
    )


def test_survey_quantities_and_totals(boss, lib, db):
    from app.survey.models import AreaType, Survey, SurveyArea

    client, h = boss
    terrace_type = db.scalar(select(AreaType).where(AreaType.name == "Terrace"))
    db.get(OfferItem, lib["terrace"]).area_type_id = terrace_type.id
    site = Site(code="S-Q-1", name="Invented site", status="active")
    db.add(site)
    db.flush()
    s = Survey(code="SUR-Q-1", title="Invented survey", status="approved", site_id=site.id)
    db.add(s)
    db.flush()
    db.add(
        SurveyArea(
            survey_id=s.id,
            name="Roof",
            area_type_id=terrace_type.id,
            shape="direct",
            direct_area_sqm=D(100),
            treated_area_sqm=D(100),
            count=1,
        )
    )
    db.commit()
    q = new_quotation(client, h, lib)
    r = client.post(
        f"/api/quotations/{q['id']}/survey-quantities", json={"survey_id": s.id}, headers=h
    )
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["show_amounts"] and out["survey_fill"]["lines_filled"] >= 3
    by = {ln["offer_line_id"]: ln for ln in lines_of(out)}
    assert D(by[lib["l_sys"]]["qty"]) == D("100.000")
    assert D(by[lib["l_lib"]]["qty"]) == D("1076.390")
    # total: option 1 counts where options are offered; "if required" and client's scope never
    expected = D(by[lib["l_sys"]]["amount"]) + D(by[lib["l_lib"]]["amount"])
    assert D(out["total"]) == expected
