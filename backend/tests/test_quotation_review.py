# ruff: noqa: E501, F811  (invented offer wording unwrapped; the shared `lib` fixture is imported)
"""Quotations review (part 0): re-issue keeps both files, work-front-ready limits the indent
draft, "area type missing", the bungalow-style import (header / footer / watermark images, mixed
styles, split item titles, a sub-heading, new stage headings, unnumbered "if required" rows),
empty letter lines dropped, presets, consistency warnings. Invented data only."""

import io
from decimal import Decimal

import pytest
from docx import Document
from docx.shared import Cm
from sqlalchemy import select

from app.config import settings
from app.crm.models import Lead, LeadActivity
from app.masters.models import Client, Product, System, SystemComponent
from app.material.models import Indent
from app.models import User
from app.quotations.checks import check
from app.quotations.docx_out import _watermark
from app.quotations.importer import import_offer
from app.quotations.models import Letterhead, OfferItem, OfferPreset, QuotationFile, SpecBlock
from app.quotations.render import fill_lines
from app.sites.models import Site, SiteNode
from app.survey.models import AreaType
from app.tenders.models import Tender
from tests.test_quotations import _num, _p, _png, _row, lib, new_quotation  # noqa: F401  (fixtures)


@pytest.fixture(autouse=True)
def media(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "media_dir", str(tmp_path))
    return tmp_path


@pytest.fixture
def boss(login_as):
    return login_as("super_admin", email="boss@example.com")


def _png_size(w: int, h: int) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (w, h), (20, 90, 160)).save(buf, "PNG")
    return buf.getvalue()


def bungalow_offer(path, tmp_path):
    """An invented offer laid out like the EESPL bungalow one."""
    d = Document()
    sec = d.sections[0]
    sec.header.paragraphs[0].add_run().add_picture(io.BytesIO(_png_size(2480, 600)), width=Cm(21))
    mark = tmp_path / "mark.png"
    mark.write_bytes(_png_size(800, 800))
    _watermark(sec.header, mark)
    sec.footer.paragraphs[0].add_run().add_picture(io.BytesIO(_png_size(2480, 420)), width=Cm(21))
    _p(d, "Date: ", ("b", "02.03.2026"))
    _p(d, "Client Name: ", ("b", "Demo Villa Owners"))
    _p(d, ("b", "SUB: OFFER FOR ALL AREA WATERPROOFING WORK."))
    _p(d, "Dear Sir,")
    _p(
        d,
        "Reference to subject, we offer waterproofing by Using ",
        ("b", "ACME Products"),
        " at ",
        ("b", "Demo Villa Owners"),
        ".",
    )
    _p(d, "Thanking you.")
    _p(d, "For ", ("b", "ETHIOS ENVIRO SOLUTIONS PVT. LTD.,"))
    _p(d, ("b", "Test Signer"))
    _p(d, "Encl.:")
    _p(d, "Technical Specification", num=5)
    t = d.add_paragraph(style="Title")
    t.add_run("ITEM NO – 1").bold = True
    t = d.add_paragraph(style="Title")
    t.add_run(
        "TECHNICAL SPECIFICATION FOR RETAINING WALL AREA WATERPROOFING WORK BY USING ACME COATING SYSTEMS"
    ).bold = True
    _p(d, ("b", "LAYING WATERPROOFING MEMBRANE (WALL / VERTICAL AREA)"))
    _p(d, ("b", "SURFACE PREPARATION:"))
    p = d.add_paragraph(style="List Paragraph")
    p.add_run("Clean the wall; west water pipes fixed.")
    _num(p, 25)
    _p(d, ("b", "TERMINATION"))
    p = d.add_paragraph(style="Header")
    p.add_run("Termination at 300 mm by light temping.")
    _num(p, 25)
    _p(d, ("b", "COVING"))
    _p(d, "Corner tape with ", ("b", "ACME TAPE"), " fixed.", num=25)
    _p(d, ("b", "PROTECTIVE COATING"))
    _p(
        d,
        "Apply ",
        ("b", "ACME SBS MEMBRANE"),
        " after pond testing lying 45gsm mesh, overall 1.2 kg in 2 coats.",
        num=25,
    )
    _p(d, ("b", "BUDGETARY OFFER FOR RETAINING WALL AREA WATERPROOFING WORK:"))
    tb = d.add_table(rows=1, cols=4)
    for i, h in enumerate(["Sr No.", "Item Description", "UoM", "Rate"]):
        tb.rows[0].cells[i].text = h
    r = _row(tb, ["1.", "", "SQ FT", "80.00"])
    r.cells[1].paragraphs[0].add_run("Laying ")
    r.cells[1].paragraphs[0].add_run("ACME SBS MEMBRANE").bold = True
    r.cells[1].paragraphs[0].add_run(" with 40gsm mesh in 3 coats and ")
    r.cells[1].paragraphs[0].add_run("ACME PRIMER X").bold = True
    _row(tb, ["", "Pressure grouting (if required)", "NOS", "350.00"])
    _p(d, ("b", "GENERAL TERMS AND CONDITIONS"))
    _p(d, ("b", "CLIENTS OBLIGATION"))
    _p(d, "Water shall be provided free of charge.", num=1)
    _p(d, ("b", "TAXES. "), "GST is extra.", num=1)
    d.save(str(path))
    return path


def test_bungalow_import_images_mixed_styles_and_new_stages(db, tmp_path, media):
    res = import_offer(
        db, bungalow_offer(tmp_path / "bungalow.docx", tmp_path), preset="Bungalow - test"
    )
    db.commit()
    assert (
        res.counts["items"] == 1
        and res.counts["spec_blocks"] == 1
        and res.counts["offer_lines"] == 2
    )
    assert res.counts["presets"] == 1 and res.counts["tc_clauses"] == 2
    for fix in (
        "west water -> waste water",
        "temping -> tamping",
        "lying -> laying",
        "CLIENTS OBLIGATION -> CLIENT'S OBLIGATIONS",
    ):
        assert res.fixes.get(fix) == 1, (fix, res.fixes)
    # our own company: the seeded EESPL letterhead is filled in with the three images
    head = db.scalar(select(Letterhead).where(Letterhead.name == "EESPL"))
    assert head.header_image_path and head.footer_image_path and head.watermark_path
    assert (
        head.signatory_firm == "For ETHIOS ENVIRO SOLUTIONS PVT. LTD."
        and head.signatory_name == "Test Signer"
    )
    assert not head.signatory_designation
    block = db.scalar(select(SpecBlock))
    assert block.subtitle == "LAYING WATERPROOFING MEMBRANE (WALL / VERTICAL AREA)"
    assert [s["heading"] for s in block.sections] == [
        "Surface preparation",
        "Termination",
        "Coving",
        "Protective coating",
    ]
    assert all(
        len(s["steps"]) == 1 for s in block.sections
    )  # List Paragraph / Header styles: all steps
    item = db.scalar(select(OfferItem))
    assert item.budget_title.lower().startswith("retaining wall area")
    assert [ln.line.if_required for ln in item.lines] == [False, True]
    # the letter keeps only what this offer has: no city, no attention line
    from app.quotations.models import LetterTemplate

    lt = db.get(LetterTemplate, head.letter_template_id)
    assert lt.opening == "Date: **{date}**\nClient Name: **{client_firm}**"
    assert "{areas_list}" in lt.subject and "**{brand} Products**" in lt.body


def test_items_of_an_area_that_already_has_one_are_named_by_system(db, tmp_path, media):
    at = db.scalar(select(AreaType).where(AreaType.name == "Basement retaining wall"))
    db.add(OfferItem(name="Retaining wall area", budget_title="x", area_type_id=at.id))
    db.commit()
    import_offer(db, bungalow_offer(tmp_path / "b.docx", tmp_path))
    names = sorted(db.scalars(select(OfferItem.name)))
    assert "Retaining wall - SBS self-adhesive membrane" in names


def test_empty_placeholder_lines_are_dropped():
    lines = fill_lines(
        "Date: **{date}**\nFirm: **{client_firm}**\n**{client_city}**\nKind Attn.: **{attention}**\n\nDear Sir,",
        {"date": "01.01.2026", "client_firm": "X", "client_city": "", "attention": ""},
    )
    assert lines == ["Date: **01.01.2026**", "Firm: **X**", "", "Dear Sir,"]


def test_consistency_warnings():
    spec = [
        {
            "sections": [
                {
                    "steps": [
                        {
                            "text": "Apply **ACME HYBRID PU** with 45gsm mesh; 1st coat, 2nd coat; primer coverage of 5 – 6 sq. mtr. per Kg"
                        }
                    ]
                }
            ]
        }
    ]
    warn = check(
        spec,
        ["Apply **ACME CEMSHIELD HYBRID PU** with 40gsm mesh in 3 coats; primer at 0.25 kg / m2"],
    )
    text = " | ".join(warn)
    assert "ACME CEMSHIELD HYBRID PU is in the budgetary offer but not in the specification" in text
    assert "ACME HYBRID PU is in the specification but not in the budgetary offer" in text
    assert "Different gsm (mesh): specification 45 gsm, budgetary offer 40 gsm" in text
    assert "Different coat count: specification 2 coats, budgetary offer 3 coats" in text
    assert "Different consumption" in text
    assert (
        check(spec, ["Apply **ACME HYBRID PU** with 45 gsm mesh, two coats, 0.18 kg/m2 primer"])
        == []
    )


def test_reissue_keeps_both_files_and_marks_the_old_one(boss, lib, db, media):
    client, h = boss
    lead = Lead(
        code="L-R-1", contact_name="Mr. Invented", company="Invented Builders", status="contacted"
    )
    db.add(lead)
    db.commit()
    q = new_quotation(client, h, lib, lead_id=lead.id)
    client.post(f"/api/quotations/{q['id']}/issue", headers=h)
    r = client.post(f"/api/quotations/{q['id']}/issue", headers=h).json()
    files = {f["file_name"]: f for f in r["files"]}
    assert set(files) == {
        f"{q['code']}-R0.docx",
        f"{q['code']}-R0.pdf",
        f"{q['code']}-R0-2.docx",
        f"{q['code']}-R0-2.pdf",
    }
    assert (
        files[f"{q['code']}-R0.pdf"]["replaced"] and not files[f"{q['code']}-R0-2.pdf"]["replaced"]
    )
    assert all((media / f.path).exists() for f in db.scalars(select(QuotationFile)))
    acts = [
        a.text
        for a in db.scalars(
            select(LeadActivity).where(LeadActivity.lead_id == lead.id).order_by(LeadActivity.id)
        )
    ]
    assert any(" issued by " in a for a in acts) and any(
        " re-issued by " in a and "replaces" in a for a in acts
    )


def test_area_type_missing_and_picked_on_the_quotation(boss, lib, db):
    client, h = boss
    rows = client.get("/api/quotations/library/item", headers=h).json()
    assert all(r["area_type_missing"] for r in rows)  # the invented items have none
    q = new_quotation(client, h, lib)
    terrace = db.scalar(select(AreaType).where(AreaType.name == "Terrace"))
    item = q["items"][0]
    r = client.put(
        f"/api/quotations/{q['id']}/items/{item['id']}",
        json={"area_type_id": terrace.id},
        headers=h,
    )
    assert r.json()["items"][0]["area_type_id"] == terrace.id
    assert (
        db.get(OfferItem, lib["terrace"]).area_type_id is None
    )  # the library is not guessed or changed


def test_a_preset_builds_a_quotation_with_its_items_in_order(boss, lib, db):
    client, h = boss
    r = client.post(
        "/api/quotations/library/preset",
        json={
            "name": "Bungalow - test",
            "letterhead_id": lib["head"],
            "items": [
                {"offer_item_id": lib["tank"]},
                {"offer_item_id": lib["terrace"], "options": ["2"]},
            ],
            "include_references": False,
        },
        headers=h,
    )
    assert r.status_code == 201, r.text
    preset = r.json()
    assert preset["version"] == 1 and preset["item_names"] == ["Tank area", "Terrace area"]
    q = client.post(
        "/api/quotations",
        json={
            "client_firm": "Invented Villa",
            "project": "Invented Villa",
            "preset_id": preset["id"],
        },
        headers=h,
    ).json()
    assert [i["name"] for i in q["items"]] == ["Tank area", "Terrace area"]
    assert (
        q["items"][1]["options"] == ["2"]
        and q["references"] == []
        and q["letterhead_id"] == lib["head"]
    )
    r = client.put(
        f"/api/quotations/library/preset/{preset['id']}",
        json={"include_references": True},
        headers=h,
    )
    assert r.json()["version"] == 2
    assert db.scalar(select(OfferPreset)).include_references


def test_work_front_ready_limits_the_indent_draft(boss, db):
    client, h = boss
    me = db.scalar(select(User.id).where(User.email == "boss@example.com"))
    raft = db.scalar(select(AreaType).where(AreaType.name == "Raft"))
    mem = Product(
        code="Q-RM", name="Invented raft membrane", unit="sqm", pack_size=10, pack_unit="roll"
    )
    system = System(code="Q-RAFT", name="Invented raft system", unit="sqm", labour_rate=10)
    c = Client(name="Invented Builders")
    db.add_all([mem, system, c])
    db.flush()
    db.add(
        SystemComponent(
            system_id=system.id,
            product_id=mem.id,
            consumption_per_unit=Decimal(1),
            wastage_percent=0,
        )
    )
    t = Tender(code="T-F-1", name="Invented", client_id=c.id, status="won")
    db.add(t)
    db.flush()
    site = Site(code="S-F-1", name="Invented site", client_id=c.id, tender_id=t.id, status="active")
    db.add(site)
    db.flush()
    blocks = []
    for name in ("Block A", "Block B", "Block C"):
        n = SiteNode(site_id=site.id, kind="tower", name=name)
        db.add(n)
        db.flush()
        blocks.append(n)
    db.commit()
    s = client.post(
        "/api/surveys", json={"site_id": site.id, "title": "Invented rafts"}, headers=h
    ).json()
    for b in blocks:
        r = client.post(
            f"/api/surveys/{s['id']}/areas",
            json={
                "name": f"{b.name} raft",
                "tower": b.name,
                "node_id": b.id,
                "area_type_id": raft.id,
                "shape": "direct",
                "direct_area_sqm": 2000,
                "system_id": system.id,
            },
            headers=h,
        )
        assert r.status_code == 201, r.text
    assert (
        client.post(f"/api/surveys/{s['id']}/indent", json={}, headers=h).status_code == 422
    )  # nothing ready yet
    r = client.post(
        f"/api/sites/{site.id}/nodes/{blocks[0].id}/front-ready",
        data={"ready": "true"},
        files={"photo": ("p.png", _png(), "image/png")},
        headers=h,
    )
    assert r.status_code == 200, r.text
    node = next(n for n in r.json() if n["id"] == blocks[0].id)
    assert node["front_ready"] and node["front_ready_photo"] and node["front_ready_by_name"]
    fronts = client.get(f"/api/surveys/{s['id']}/fronts", headers=h).json()["floors"]
    assert sum(f["ready_areas"] for f in fronts) == 1
    r = client.post(f"/api/surveys/{s['id']}/indent", json={}, headers=h)
    assert r.status_code == 201, r.text
    ind = db.scalars(select(Indent).order_by(Indent.id.desc())).first()
    qty = sum(ln.qty for ln in ind.lines)
    assert qty == Decimal(
        "2140"
    )  # Block A only: 2,000 sqm + the Raft 7 % wastage = 2,140 sqm (214 rolls)
    r = client.post(f"/api/surveys/{s['id']}/indent", json={"include_not_ready": True}, headers=h)
    ind2 = db.scalars(select(Indent).order_by(Indent.id.desc())).first()
    assert sum(ln.qty for ln in ind2.lines) == Decimal("6420")
    assert me is not None
    sm = client.get(f"/api/surveys/site-map/{site.id}", headers=h).json()
    assert blocks[0].id in sm["ready"] and blocks[1].id not in sm["ready"]
