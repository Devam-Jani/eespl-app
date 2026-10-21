# ruff: noqa: E501, F811  (invented data reads better unwrapped; shared fixtures are imported)
"""The work order, billing and close-out milestone on invented data: a tender quoted to three
civil contractors, the work order check, client roles in the portal, client-approved
measurements, RFI zones and the billing checklist, the billing cycle and payment reminders on a
frozen clock (never an email to a client), the work completion certificate, a discontinued
product, material reconciliation and a labour claim."""

import io
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.closeout.models import ClientWorkOrder, SiteIssueLog
from app.main import app
from app.masters.models import Client, Product
from app.portal.models import NotificationOutbox
from app.process.models import ProcessEvent
from app.quotations.models import Quotation
from app.sites.models import Site, SiteMember, SiteNode
from app.team.models import Measurement
from app.tenders.models import BoqLine, BoqSection, Tender
from tests.conftest import login
from tests.test_finance import boss, contract, invoiced, world  # noqa: F401  (fixtures)
from tests.test_portal import PASSWORD, accept

D = Decimal


def _client(db, name):
    c = Client(name=name, state="Gujarat", city="Ahmedabad")
    db.add(c)
    db.flush()
    return c


def _tender(db, name="Invented Phase-2 towers"):
    dev = _client(db, "Invented Developer")
    t = Tender(code="T-2026-9601", name=name, status="draft", client_id=dev.id, guarantee_years=10)
    db.add(t)
    db.flush()
    s = BoqSection(tender_id=t.id, title="Toilets", sort_order=1)
    db.add(s)
    db.flush()
    db.add_all([BoqLine(tender_id=t.id, section_id=s.id, sort_order=1, client_item_no="1", description="Toilet waterproofing with acrylic cementitious coating and 45 GSM mesh", unit="sqm", qty=500, rate=Decimal("620")),
                BoqLine(tender_id=t.id, section_id=s.id, sort_order=2, client_item_no="2", description="Brick bat coba in sunken toilets", unit="sqm", qty=300, rate=Decimal("310"))])  # fmt: skip
    db.commit()
    return t


# --- 1. one tender, several civil contractors -------------------------------------------------------


def test_a_tender_quoted_to_three_contractors_one_wins_then_we_win_as_l1(boss, db):
    c, h = boss
    t = _tender(db)
    ids = [_client(db, f"Invented Civil Contractor {x}").id for x in "ABC"]
    db.commit()
    for cid in ids:
        r = c.post(f"/api/closeout/tenders/{t.id}/customers", json={"client_id": cid}, headers=h)
        assert r.status_code == 201, r.text
    custs = r.json()["customers"]
    assert len(custs) == 3 and all(
        D(str(x["quotation"]["total"])) == D("403000") for x in custs
    )  # 500 x 620 + 300 x 310, from the same BOQ
    assert len({x["quotation"]["code"] for x in custs}) == 3  # each its own quotation
    winner = next(x for x in custs if x["client_id"] == ids[1])
    out = c.post(
        f"/api/closeout/tenders/{t.id}/customers/{winner['id']}/main-won", headers=h
    ).json()
    assert out["main_contractor"] == "Invented Civil Contractor B"
    states = {
        x["client"]: (x["main_status"], x["quotation"]["status"], x["quotation"]["lost_reason"])
        for x in out["customers"]
    }
    assert states["Invented Civil Contractor A"] == ("lost", "lost", "main_job_lost")
    assert states["Invented Civil Contractor B"] == ("won", "draft", None)
    db.expire_all()
    assert db.get(Tender, t.id).status == "draft"  # not our loss
    # our stage with the winner: L1, then won
    for name, rank, us in (
        ("EESPL", 1, True),
        ("Invented Applicator Two", 2, False),
        ("Invented Applicator Three", 3, False),
    ):
        assert (
            c.post(
                f"/api/team/tenders/{t.id}/bidders",
                json={"name": name, "rank": rank, "is_us": us},
                headers=h,
            ).status_code
            == 201
        )
    assert c.patch(f"/api/tenders/{t.id}", json={"status": "won"}, headers=h).status_code == 200
    award = c.get(f"/api/team/tenders/{t.id}/award", headers=h).json()
    assert any(b["is_us"] and b.get("label", "L1") == "L1" for b in award["bidders"])
    assert (
        c.post(
            f"/api/closeout/tenders/{t.id}/customers", json={"client_id": ids[0]}, headers=h
        ).status_code
        == 422
    )  # decided


# --- 2. the work order check -------------------------------------------------------------------------


def _wo_quotation(db):
    from app.quotations.models import QuotationItem, QuotationLine

    q = Quotation(code="QTN-2026-9601", revision=0, is_latest=True, client_firm="Invented Builders", project="Invented tower", quote_date=date(2026, 10, 1),
                  status="won", guarantee_years=10,
                  terms=[{"category": "client_scope", "text": "Lifting and shifting of material to all floors shall be provided by the client free of cost."},
                         {"category": "client_scope", "text": "Water and electricity shall be provided by the client."},
                         {"category": "payment", "text": "Our R.A. bills shall be released within 7 (seven) days from the date of submission."}])  # fmt: skip
    item = QuotationItem(sort_order=0, name="Toilets", budget_title="Toilets", options=[], specs=[])
    item.lines.append(
        QuotationLine(
            sort_order=0,
            description="Toilet waterproofing with INVENTA COAT acrylic coating",
            uom="sqm",
            rate=D("620"),
            qty=D("500"),
        )
    )
    item.lines.append(
        QuotationLine(
            sort_order=1,
            description="Brick bat coba in sunken toilets",
            uom="sqm",
            rate=D("310"),
            qty=D("300"),
        )
    )
    q.items.append(item)
    db.add(q)
    db.add(
        Product(code="INV-COAT", name="INVENTA COAT", unit="kg", discontinued_on=date(2026, 9, 1))
    )
    db.commit()
    return q


def test_work_order_deviations_block_acceptance_until_resolved_then_it_is_locked(boss, db):
    c, h = boss
    q = _wo_quotation(db)
    wo = c.post(
        "/api/closeout/work-orders",
        json={"quotation_id": q.id, "client_ref": "WO/INV/77"},
        headers=h,
    ).json()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Sr No", "Description", "Unit", "Qty", "Rate"])
    ws.append(
        ["1", "Toilet waterproofing with INVENTA COAT acrylic coating", "cum", 500, 620]
    )  # per cum: a unit change
    ws.append(["2", "Brick bat coba in sunken toilets", "sqm", 300, 310])
    ws.append([])
    ws.append(
        [
            "Terms: Lifting and shifting of material shall be provided by the contractor. Payment within 45 days of certification. Retention 5%."
        ]
    )
    buf = io.BytesIO()
    wb.save(buf)
    r = c.post(
        f"/api/closeout/work-orders/{wo['id']}/file",
        files={
            "file": (
                "wo.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=h,
    )
    assert r.status_code == 200, r.text
    got = r.json()
    assert (
        got["read"]["lines"] == 2
        and got["terms"]["payment_days"] == 45
        and got["terms"]["payment_basis"] == "certification"
    )
    assert "material lifting and shifting" in got["terms"]["scope_on_us"]
    out = c.post(f"/api/closeout/work-orders/{wo['id']}/compare", headers=h).json()
    kinds = {d["kind"] for d in out["deviations"]}
    assert {"unit", "scope", "discontinued", "payment", "retention"} <= kinds
    assert any(d["kind"] == "scope" and "lifting" in d["topic"] for d in out["deviations"])
    r = c.post(f"/api/closeout/work-orders/{wo['id']}/accept", headers=h)
    assert r.status_code == 422 and "still open" in r.json()["detail"]
    for d in out["deviations"]:
        body = (
            {"action": "amend", "note": "Ask the client", "amend_text": None}
            if d["kind"] == "scope"
            else {"action": "accept", "note": "Agreed with the client"}
        )
        assert (
            c.post(
                f"/api/closeout/work-orders/{wo['id']}/deviations/{d['id']}", json=body, headers=h
            ).status_code
            == 200
        )
    assert (
        db.scalar(select(ProcessEvent).where(ProcessEvent.type == "wo_amend")) is not None
    )  # the salesperson's to-do
    assert c.post(f"/api/closeout/work-orders/{wo['id']}/accept", headers=h).json()["locked"]
    for call in (lambda: c.put(f"/api/closeout/work-orders/{wo['id']}", json={"payment_days": 30}, headers=h),
                 lambda: c.put(f"/api/closeout/work-orders/{wo['id']}/lines", json={"lines": []}, headers=h),
                 lambda: c.post(f"/api/closeout/work-orders/{wo['id']}/compare", headers=h),
                 lambda: c.delete(f"/api/closeout/work-orders/{wo['id']}", headers=h)):  # fmt: skip
        assert call().status_code == 409
    assert db.get(ClientWorkOrder, wo["id"]).status == "accepted"


# --- 3. client roles ----------------------------------------------------------------------------------


def _client_login(c, h, client_id, site_ids, email, roles):
    r = c.post(
        "/api/portal-admin/invites",
        json={
            "full_name": f"Invented {email.split('@')[0]}",
            "email": email,
            "client_id": client_id,
            "site_ids": site_ids,
            "roles": roles,
        },
        headers=h,
    )
    assert r.status_code == 201, r.text
    assert accept(r.json()["link"]).status_code == 200
    return login(TestClient(app), email, PASSWORD)


@pytest.fixture
def portal_site(boss, world, db):
    c, h = boss
    db.get(Site, world["site"]).site_incharge_id = None
    db.commit()
    assert (
        c.put(
            f"/api/portal-admin/sites/{world['site']}",
            json={"visible": True, "sections": {}},
            headers=h,
        ).status_code
        == 200
    )
    return world


ENDPOINTS = [
    ("get", "/api/portal/sites/{s}/dprs", "progress.view"),
    ("get", "/api/portal/sites/{s}/billing", "bill.view"),
    ("get", "/api/portal/sites/{s}/overdue", "payment.view"),
    ("get", "/api/portal/sites/{s}/measurements", "measurement.approve"),
    ("get", "/api/portal/sites/{s}/handover", "progress.view"),
    ("get", "/api/portal/sites/{s}/documents", "documents.view"),
    ("get", "/api/portal/sites/{s}/snags", "snag.raise"),
]
ALLOWED = {  # role -> the endpoints above it may call
    "project_manager": {0, 1, 2, 3, 4, 5, 6},
    "planning": {0, 4, 5, 6},  # progress, handover, documents; snags through progress
    "site": {0, 4, 5, 6},
    "billing": {0, 1, 3, 4, 5, 6},
    "accounts": {1, 2, 5},
}


def _terrace(db, w) -> int:
    return db.scalar(
        select(SiteNode.id).where(SiteNode.site_id == w["site"], SiteNode.kind == "terrace")
    )


def test_each_client_role_on_each_portal_endpoint_and_another_clients_site(boss, portal_site, db):
    c, h = boss
    w = portal_site
    other = _client(db, "Other Invented Client")
    db.commit()
    s2 = c.post(
        "/api/sites",
        json={"name": "Other Invented Site", "client_id": other.id, "state": "Gujarat"},
        headers=h,
    ).json()["id"]
    c.put(f"/api/portal-admin/sites/{s2}", json={"visible": True, "sections": {}}, headers=h)
    for role, allowed in ALLOWED.items():
        hh = _client_login(c, h, w["client"], [w["site"]], f"{role}-client@example.com", [role])
        cc = TestClient(app)
        for i, (method, path, _right) in enumerate(ENDPOINTS):
            r = getattr(cc, method)(path.format(s=w["site"]), headers=hh)
            assert (r.status_code == 200) == (i in allowed), (
                role,
                path,
                r.status_code,
                r.text[:200],
            )
            if i not in allowed:
                assert r.status_code == 403
            assert (
                getattr(cc, method)(path.format(s=s2), headers=hh).status_code == 404
            )  # never another client's site
        if role in (
            "planning",
            "site",
        ):  # only planning (and the project manager) changes the dates
            r = cc.put(
                f"/api/portal/sites/{w['site']}/nodes/{_terrace(db, w)}/handover-promise",
                json={"promised_on": "2026-12-01"},
                headers=hh,
            )
            assert r.status_code == (200 if role == "planning" else 403), r.text
    me = TestClient(app).get("/api/portal/me/roles", headers=hh).json()
    assert [r["code"] for r in me["roles"]] == ["accounts"] and "payment.view" in me["rights"]
    # a role that may not sign: 403 on the sign endpoints
    hh = _client_login(c, h, w["client"], [w["site"]], "accounts2-client@example.com", ["accounts"])
    assert (
        TestClient(app)
        .post(f"/api/portal/sites/{w['site']}/measurements/1/approve", headers=hh)
        .status_code
        == 403
    )


# --- 5. measurements approved by the client ------------------------------------------------------------


@pytest.fixture
def measured(boss, portal_site, db):
    c, h = boss
    w = portal_site
    k = contract(
        c, h, w
    )  # its own (laser) measurements are approved: written before approval was needed
    for m in db.scalars(select(Measurement).where(Measurement.site_id == w["site"])):
        db.delete(m)
    site = db.get(Site, w["site"])
    site.measurements_need_approval = True
    tower = SiteNode(site_id=site.id, kind="tower", name="Tower A")
    db.add(tower)
    db.flush()
    flats = []
    for n in ("101", "102"):
        f = SiteNode(site_id=site.id, parent_id=tower.id, kind="flat", name=f"Flat {n}")
        db.add(f)
        db.flush()
        flats.append(
            (
                f,
                SiteNode(site_id=site.id, parent_id=f.id, kind="toilet", name="Toilet"),
                SiteNode(site_id=site.id, parent_id=f.id, kind="balcony", name="Balcony"),
            )
        )
        db.add_all(flats[-1][1:])
    db.commit()
    pu = next(ln for ln in k["lines"] if ln["description"] == "PU coating")
    return {
        **w,
        "contract": k,
        "pu": pu["id"],
        "tower": tower.id,
        "flats": [(f.id, t.id, b.id) for f, t, b in flats],
    }


def _line(c, h, w, node, length="20", breadth="1.5", **kw):
    r = c.post(f"/api/closeout/sites/{w['site']}/measurements", json={"description": "PU coating, toilet floor", "contract_line_id": w["pu"], "node_id": node,
                                                                     "length_m": length, "breadth_m": breadth, **kw}, headers=h)  # fmt: skip
    return r


def test_the_client_corrects_a_line_staff_record_another_approved_lines_lock_and_bills_take_only_approved(
    boss, measured, db
):
    c, h = boss
    w = measured
    a = _line(c, h, w, w["flats"][0][1]).json()
    b = _line(c, h, w, w["flats"][1][1], length="10", breadth="2").json()
    assert a["status"] == "draft" and D(str(a["qty"])) == D("30.000")
    assert (
        c.post(
            f"/api/closeout/sites/{w['site']}/measurement-sets",
            json={"ids": [a["id"], b["id"]]},
            headers=h,
        ).status_code
        == 200
    )
    hh = _client_login(c, h, w["client"], [w["site"]], "billing-eng@example.com", ["billing"])
    cc = TestClient(app)
    r = cc.post(
        f"/api/portal/sites/{w['site']}/measurements/{a['id']}/correct",
        json={"field": "length_m", "new_value": "19.42", "remark": "20 ft is 19 ft 5 in on site"},
        headers=hh,
    )
    assert r.status_code == 200, r.text
    ch = r.json()["changes"][0]
    assert (
        ch["old"] == D("20.000")
        or str(ch["old"]).startswith("20")
        and str(ch["new"]).startswith("19.42")
    )
    assert ch["by"] == "Invented billing-eng" and ch["at"] and ch["status"] == "pending"
    assert (
        cc.post(
            f"/api/portal/sites/{w['site']}/measurements/{a['id']}/approve", headers=hh
        ).status_code
        == 422
    )  # our answer first
    out = c.post(
        f"/api/closeout/measurement-changes/{ch['id']}/answer", json={"accept": True}, headers=h
    ).json()
    assert (
        out["status"] == "approved"
        and D(str(out["qty"])) == D("29.130")
        and out["approved_name"] == "Invented billing-eng"
    )
    # the other line: the client signed on paper; staff record it, with the signed copy
    r = c.post(
        f"/api/closeout/measurements/{b['id']}/record-approval",
        data={"approved_name": "Invented Client Engineer"},
        headers=h,
    )
    assert r.status_code == 422 and "signed sheet" in r.json()["detail"]
    r = c.post(
        f"/api/closeout/measurements/{b['id']}/record-approval",
        data={"approved_name": "Invented Client Engineer"},
        files={"file": ("signed.pdf", b"%PDF-1.4 invented", "application/pdf")},
        headers=h,
    )
    assert (
        r.status_code == 200 and r.json()["recorded_by_staff"] and r.json()["status"] == "approved"
    )
    # approved: locked for everyone (the super admin included)
    assert (
        c.put(
            f"/api/closeout/measurements/{a['id']}",
            json={"description": "x", "length_m": "25"},
            headers=h,
        ).status_code
        == 409
    )
    assert c.delete(f"/api/closeout/measurements/{a['id']}", headers=h).status_code == 409
    d = _line(
        c, h, w, w["flats"][0][2], length="5", breadth="1"
    ).json()  # a draft line, not approved
    ra = c.post(
        f"/api/finance/contracts/{w['contract']['id']}/ra-bills",
        json={"period_to": "2026-10-31"},
        headers=h,
    ).json()
    pu = next(ln for ln in ra["lines"] if ln["contract_line_id"] == w["pu"])
    assert D(str(pu["suggested_qty"])) == D("49.130")  # 29.13 + 20, never the draft 5
    db.expire_all()
    assert (
        db.get(Measurement, d["id"]).ra_bill_id is None
        and db.get(Measurement, a["id"]).ra_bill_id == ra["id"]
    )
    x = c.get(f"/api/closeout/sites/{w['site']}/measurement-sheet.xlsx", headers=h)
    sheet = openpyxl.load_workbook(io.BytesIO(x.content)).active
    text = " ".join(
        str(v) for row in sheet.iter_rows(values_only=True) for v in row if v is not None
    )
    assert "19.42" in text and "Invented billing-eng" in text and "recorded by staff" in text


# --- 6. RFI zones and the billing checklist ------------------------------------------------------------


def test_outside_our_zones_a_measurement_needs_a_reason_and_the_checklist_shows_the_bill(
    boss, measured, db
):
    c, h = boss
    w = measured
    (f1, t1, b1), (f2, t2, b2) = w["flats"]
    assert (
        c.post(
            f"/api/closeout/sites/{w['site']}/zones",
            json={"name": "Flat 101", "node_id": f1},
            headers=h,
        ).status_code
        == 201
    )
    assert (
        c.post(
            f"/api/closeout/sites/{w['site']}/zones",
            json={
                "name": "Flat 102",
                "node_id": f2,
                "ours": False,
                "other_applicator": "Invented Other Applicator",
            },
            headers=h,
        ).status_code
        == 201
    )
    r = _line(c, h, w, t2)
    assert r.status_code == 422 and "outside our RFI zones" in r.json()["detail"]
    r = _line(
        c, h, w, t2, zone_reason="The client asked us to redo it after the other applicator left"
    )
    assert r.status_code == 201 and "outside our RFI zones" in r.json()["warning"]
    m = _line(c, h, w, t1).json()
    db.get(Measurement, m["id"]).status = "approved"
    db.get(Measurement, r.json()["id"]).status = "approved"
    db.commit()
    ra = c.post(
        f"/api/finance/contracts/{w['contract']['id']}/ra-bills",
        json={"period_to": "2026-10-31"},
        headers=h,
    ).json()
    g = c.get(
        f"/api/closeout/sites/{w['site']}/billing-checklist",
        params={"tower_id": w["tower"]},
        headers=h,
    ).json()
    rows = {r["flat"]: r["cells"] for r in g["rows"]}
    assert (
        rows["Flat 101"]["Toilet"]["status"] == "billed"
        and rows["Flat 101"]["Toilet"]["bill"] == ra["code"]
    )
    assert rows["Flat 101"]["Balcony"]["status"] == "waiting_handover"
    assert (
        rows["Flat 102"]["Toilet"]["status"] == "not_ours"
        and "Invented Other Applicator" in rows["Flat 102"]["Toilet"]["label"]
    )
    assert (
        c.get(
            f"/api/closeout/sites/{w['site']}/billing-checklist",
            params={"tower_id": w["tower"], "month": "2026-09"},
            headers=h,
        ).json()["rows"][0]["cells"]["Toilet"]["status"]
        != "billed"
    )


# --- 7-8. billing cycle, the stuck bill, payment reminders, never an email to a client ----------------


def test_a_stuck_bill_escalates_and_payment_reminders_come_on_due_plus_5_and_plus_30(
    boss, portal_site, db, monkeypatch, make_user
):
    c, h = boss
    w = portal_site
    k, ra, inv = invoiced(c, h, w, on="2026-10-07")
    site = db.get(Site, w["site"])
    from app.models import User

    sales = make_user("sales-inv@example.com", "sales")
    db.add(SiteMember(site_id=site.id, user_id=sales.id, role_on_site="sales"))
    wo = ClientWorkOrder(
        code="CWO-2026-9001",
        site_id=site.id,
        client_id=w["client"],
        status="accepted",
        payment_basis="invoice",
        payment_days=15,
    )
    db.add(wo)
    db.commit()
    hh = _client_login(c, h, w["client"], [w["site"]], "accounts-client@example.com", ["accounts"])
    r = c.put(
        f"/api/closeout/clients/{w['client']}/billing-cycle",
        json={
            "submit_from_day": 25,
            "submit_to_day": 30,
            "ho_from_day": 1,
            "ho_to_day": 7,
            "payment_days": 30,
        },
        headers=h,
    )
    assert r.status_code == 200
    assert (
        c.post(
            f"/api/closeout/bills/{ra['id']}/steps",
            data={"step": "signed_pm", "on_date": "2026-10-08"},
            headers=h,
        ).status_code
        == 422
    )  # the signed copy
    assert (
        c.post(
            f"/api/closeout/bills/{ra['id']}/steps",
            data={
                "step": "submitted_site",
                "on_date": "2026-10-08",
                "by_name": "Invented site office",
            },
            headers=h,
        ).status_code
        == 200
    )
    bill = c.post(
        f"/api/closeout/bills/{ra['id']}/steps",
        data={"step": "signed_pm", "on_date": "2026-10-09"},
        files={"file": ("signed.pdf", b"%PDF-1.4", "application/pdf")},
        headers=h,
    ).json()
    assert (
        bill["due_on"] == "2026-10-22" and "15 days after the invoice" in bill["due_how"]
    )  # the work order's term
    from app.process import hooks

    def at(day):
        return datetime.combine(day, datetime.min.time(), tzinfo=UTC) + timedelta(hours=4)

    def events(kind):
        return {
            e.dedupe_key for e in db.scalars(select(ProcessEvent).where(ProcessEvent.type == kind))
        }

    hooks.timed_events(db, at(date(2026, 10, 15)))  # 6 days at "signed by the PM"
    db.commit()
    assert not events("bill_stuck")
    hooks.timed_events(db, at(date(2026, 10, 16)))  # 7 days: the salesperson
    db.commit()
    assert events("bill_stuck") and not events("bill_stuck_long")
    hooks.timed_events(db, at(date(2026, 10, 19)))  # 10 days: escalated to the director
    db.commit()
    assert events("bill_stuck_long") and not events("payment_due")
    hooks.timed_events(db, at(date(2026, 10, 22)))  # the due date: accounts
    db.commit()
    assert events("payment_due") and not events("payment_overdue")
    hooks.timed_events(db, at(date(2026, 10, 27)))  # +5: the salesperson
    db.commit()
    assert events("payment_overdue") and not events("payment_overdue_long")
    c.post(
        f"/api/closeout/bills/{ra['id']}/followups",
        json={
            "spoke_to": "Invented client accountant",
            "said": "Cheque next week",
            "promised_on": "2026-11-03",
        },
        headers=h,
    )
    hooks.timed_events(db, at(date(2026, 11, 3)))
    db.commit()
    assert events("payment_promised")
    hooks.timed_events(db, at(date(2026, 11, 21)))  # +30: director, salesperson, site engineer
    db.commit()
    assert events("payment_overdue_long")
    from app.closeout import portal_routes

    monkeypatch.setattr(portal_routes, "today", lambda: date(2026, 11, 21))
    od = TestClient(app).get(f"/api/portal/sites/{w['site']}/overdue", headers=hh).json()
    assert od and od[0]["overdue_days"] == 30
    client_ids = {
        u for (u,) in db.execute(select(User.id).where(User.email.like("%-client@example.com")))
    }
    from app.portal.models import Notification

    notes = list(db.scalars(select(Notification).where(Notification.user_id.in_(client_ids))))
    assert any(n.kind == "bill_overdue" for n in notes)  # in the portal
    emails = db.scalars(
        select(NotificationOutbox).where(
            NotificationOutbox.channel == "email",
            NotificationOutbox.notification_id.in_([n.id for n in notes] or [0]),
        )
    ).all()
    assert emails == []  # never an email to a client


# --- 9. work completion --------------------------------------------------------------------------------


def _give_system(db, w):
    """The PU coating line's system (on its BOQ line), with one product."""
    from app.finance.models import ContractLine
    from app.masters.models import System, SystemComponent

    s = System(code="INV-PU", name="Invented PU system", unit="sqm")
    p = Product(code="INV-PU-1", name="Invented PU coat", unit="kg")
    db.add_all([s, p])
    db.flush()
    db.add(SystemComponent(system_id=s.id, product_id=p.id, consumption_per_unit=D("1.5")))
    cl = db.get(ContractLine, w["pu"])
    db.get(BoqLine, cl.boq_line_id).system_id = s.id
    db.commit()
    return s


def test_the_work_completion_takes_approved_quantities_and_is_locked_once_signed(
    boss, measured, db
):
    c, h = boss
    w = measured
    (f1, t1, b1), _ = w["flats"]
    a = _line(c, h, w, t1).json()
    _line(c, h, w, b1, length="3", breadth="1")  # a draft: not in the certificate
    db.get(Measurement, a["id"]).status = "approved"
    db.commit()
    r = c.post(
        f"/api/closeout/sites/{w['site']}/work-completions",
        json={"completion_date": "2026-11-15", "guarantee_years": 10},
        headers=h,
    )
    assert r.status_code == 422 and "No system for" in r.json()["detail"]  # every line its system
    _give_system(db, w)
    r = c.post(
        f"/api/closeout/sites/{w['site']}/work-completions",
        json={"completion_date": "2026-11-15", "guarantee_years": 10},
        headers=h,
    )
    assert r.status_code == 201, r.text
    wc = r.json()
    assert [(x["area"].split(" / ")[-1], D(x["qty"])) for x in wc["areas"]] == [
        ("Toilet", D("30.000"))
    ]
    assert wc["guarantee_expiry"] == "2036-11-15"
    assert c.get(f"/api/closeout/work-completions/{wc['id']}/pdf", headers=h).content[:4] == b"%PDF"
    assert (
        b"<svg" in c.get(f"/api/closeout/work-completions/{wc['id']}/plan.svg", headers=h).content
    )
    hh = _client_login(c, h, w["client"], [w["site"]], "pm-client@example.com", ["project_manager"])
    assert (
        TestClient(app)
        .post(f"/api/portal/sites/{w['site']}/work-completions/{wc['id']}/sign", headers=hh)
        .status_code
        == 200
    )
    assert (
        c.put(
            f"/api/closeout/work-completions/{wc['id']}", json={"note": "changed"}, headers=h
        ).status_code
        == 409
    )
    assert c.delete(f"/api/closeout/work-completions/{wc['id']}", headers=h).status_code == 409
    add = c.post(
        f"/api/closeout/work-completions/{wc['id']}/addendum",
        json={
            "completion_date": "2026-12-01",
            "note": "The toilet re-measured after the client's request",
        },
        headers=h,
    )
    assert (
        add.status_code == 201
        and add.json()["addendum_of_id"] == wc["id"]
        and add.json()["status"] == "draft"
    )
    reg = c.get("/api/closeout/guarantees", headers=h).json()
    assert reg[0]["code"] == wc["code"] and reg[0]["expiry"] == "2036-11-15"
    leak = c.post(
        "/api/closeout/issues",
        json={
            "site_id": w["site"],
            "category": "leakage",
            "title": "Leak below flat 101 toilet",
            "work_completion_id": wc["id"],
            "node_id": t1,
        },
        headers=h,
    )
    assert leak.status_code == 201
    not_ours = c.post(
        "/api/closeout/issues",
        json={
            "site_id": w["site"],
            "category": "leakage",
            "title": "Leak at flat 102",
            "work_completion_id": wc["id"],
            "node_id": w["flats"][1][1],
        },
        headers=h,
    )
    assert not_ours.status_code == 422 and "not our work" in not_ours.json()["detail"]
    # no guarantee given: found from the place (a damp patch in the toilet of flat 101)
    auto = c.post("/api/closeout/issues", json={"site_id": w["site"], "category": "leakage", "title": "Damp patch in flat 101", "node_id": f1}, headers=h).json()
    assert auto["work_completion_id"] == wc["id"]
    elsewhere = c.post("/api/closeout/issues", json={"site_id": w["site"], "category": "leakage", "title": "Damp patch in flat 102", "node_id": w["flats"][1][0]}, headers=h).json()
    assert elsewhere["work_completion_id"] is None and "check whether it is our work" in elsewhere["description"]
    # a whole flat chosen for a certificate brings its toilet and balcony
    db.get(Measurement, _line(c, h, w, b1, length="2", breadth="1").json()["id"]).status = "approved"
    db.commit()
    flat = c.post(f"/api/closeout/sites/{w['site']}/work-completions", json={"completion_date": "2026-12-05", "node_ids": [f1], "guarantee_years": 5}, headers=h)
    assert flat.status_code == 201, flat.text
    flat = flat.json()
    assert sorted(x["area"].split(" / ")[-1] for x in flat["areas"]) == ["Balcony", "Toilet"]


def test_a_site_engineer_works_on_their_own_site_only(boss, measured, db, login_as):
    from app.models import User

    c, h = boss
    w = measured
    ce, he = login_as("site_engineer", email="se-closeout@example.com")
    uid = db.scalar(select(User.id).where(User.email == "se-closeout@example.com"))
    assert ce.get(f"/api/closeout/sites/{w['site']}/zones", headers=he).status_code == 404  # not allocated
    db.add(SiteMember(site_id=w["site"], user_id=uid, role_on_site="incharge"))
    db.commit()
    assert ce.get(f"/api/closeout/sites/{w['site']}/zones", headers=he).status_code == 200
    r = ce.post(f"/api/closeout/sites/{w['site']}/zones", json={"name": "Flat 101", "node_id": w["flats"][0][0]}, headers=he)
    assert r.status_code == 201, r.text
    r = ce.post("/api/closeout/issues", json={"site_id": w["site"], "category": "material_not_arrived", "title": "Mesh not delivered"}, headers=he)
    assert r.status_code == 201, r.text
    assert ce.get(f"/api/closeout/sites/{w['site']}/billing-checklist", headers=he).status_code == 403  # billing's


# --- 10. a product discontinued --------------------------------------------------------------------------


def test_a_discontinued_product_raises_issues_on_open_methodologies_and_sites(boss, world, db):
    from app.finance.models import ClientContract, ContractLine
    from app.masters.models import System, SystemComponent
    from app.methodology.models import Methodology, MethodologyItem

    c, h = boss
    p = Product(code="INV-GUARD", name="INVENTA GUARD", unit="kg")
    db.add(p)
    db.flush()
    s = System(code="INV-SYS", name="Invented guard system")
    s.components = [SystemComponent(product_id=p.id, consumption_per_unit=D("2"))]
    db.add(s)
    db.flush()
    line = db.scalar(select(BoqLine).where(BoqLine.description == "PU coating"))
    line.system_id = s.id
    k = ClientContract(site_id=world["site"], client_id=world["client"])
    db.add(k)
    db.flush()
    db.add(
        ContractLine(
            contract_id=k.id,
            boq_line_id=line.id,
            description="PU coating",
            unit="sqm",
            qty=100,
            rate=500,
        )
    )
    m = Methodology(code="MS-2026-9901", title="Invented methodology", version=1)
    db.add(m)
    db.flush()
    db.add(
        MethodologyItem(
            methodology_id=m.id,
            name="Toilets",
            stages=[
                {
                    "heading": "PROTECTIVE COATING",
                    "steps": [{"text": "2 coats of INVENTA GUARD.", "product": "INVENTA GUARD"}],
                }
            ],
        )
    )
    db.commit()
    r = c.patch(f"/api/products/{p.id}", json={"discontinued_on": "2026-10-20"}, headers=h)
    assert r.status_code == 200, r.text
    issues = list(
        db.scalars(select(SiteIssueLog).where(SiteIssueLog.category == "product_discontinued"))
    )
    assert {(i.methodology_id, i.site_id) for i in issues} == {(m.id, None), (None, world["site"])}
    assert all("new method statement and client approval needed" in i.title for i in issues)
    # the methodology's revision without the product closes its issue
    from app.methodology import build

    it = db.scalar(select(MethodologyItem).where(MethodologyItem.methodology_id == m.id))
    it.stages = [
        {
            "heading": "PROTECTIVE COATING",
            "steps": [{"text": "2 coats of INVENTA FLEX.", "product": "INVENTA FLEX"}],
        }
    ]
    build.save_version(
        db, db.get(Methodology, m.id), "INVENTA GUARD replaced by INVENTA FLEX", None
    )
    db.commit()
    assert (
        db.scalar(select(SiteIssueLog.status).where(SiteIssueLog.methodology_id == m.id))
        == "resolved"
    )


# --- 11. material reconciliation and a labour claim -------------------------------------------------------


def test_reconciliation_excess_and_short_and_a_claim_against_the_measurement(boss, measured, db):
    from app.masters.models import System, SystemComponent
    from app.material.models import SiteIssue, SiteIssueLine, Store

    c, h = boss
    w = measured
    coat, primer = (
        Product(code="INV-C1", name="INVENTA COAT 2", unit="kg"),
        Product(code="INV-P1", name="INVENTA PRIMER", unit="kg"),
    )
    db.add_all([coat, primer])
    db.flush()
    s = System(code="INV-S2", name="Invented PU system")
    s.components = [
        SystemComponent(product_id=coat.id, consumption_per_unit=D("1.5")),
        SystemComponent(product_id=primer.id, consumption_per_unit=D("0.2")),
    ]
    db.add(s)
    db.flush()
    db.scalar(select(BoqLine).where(BoqLine.description == "PU coating")).system_id = s.id
    store = db.scalar(select(Store).where(Store.site_id == w["site"]))
    iss = SiteIssue(
        code="ISS-2026-9901",
        kind="issue",
        store_id=store.id,
        site_id=w["site"],
        issued_on=date(2026, 10, 10),
    )
    db.add(iss)
    db.flush()
    db.add_all([SiteIssueLine(issue_id=iss.id, product_id=coat.id, qty=D("60"), rate=D("100"), value=D("6000")),
                SiteIssueLine(issue_id=iss.id, product_id=primer.id, qty=D("4"), rate=D("80"), value=D("320"))])  # fmt: skip
    db.commit()
    a = _line(c, h, w, w["flats"][0][1], measured_on="2026-10-12").json()  # 30 sqm
    db.get(Measurement, a["id"]).status = "approved"
    db.commit()
    rec = {
        r["product"]: r
        for r in c.get(
            f"/api/closeout/sites/{w['site']}/reconciliation",
            params={"month": "2026-10"},
            headers=h,
        ).json()["rows"]
    }
    assert rec["INVENTA COAT 2"]["needed"] == "45.000" or D(
        str(rec["INVENTA COAT 2"]["needed"])
    ) == D("45")
    assert rec["INVENTA COAT 2"]["state"] == "excess" and D(
        str(rec["INVENTA COAT 2"]["difference"])
    ) == D("15")
    assert rec["INVENTA PRIMER"]["state"] == "short" and D(
        str(rec["INVENTA PRIMER"]["difference"])
    ) == D("-2")
    r = c.post(
        f"/api/closeout/sites/{w['site']}/labour-claims",
        json={
            "contractor": "Invented Labour Gang",
            "node_id": w["flats"][0][1],
            "claimed_qty": "36",
            "period": "2026-10",
        },
        headers=h,
    )
    claim = r.json()[0]
    assert (
        D(str(claim["claimed"])) == D("36")
        and D(str(claim["approved"])) == D("30")
        and D(str(claim["excess_claim"])) == D("6")
        and D(str(claim["pays"])) == D("30")
    )
