# ruff: noqa: E501
"""Material estimates from the measured drawing, invented data: the consumption library (one
coverage per product per system; a BOQ or method statement figure outside its min and max warns),
wastage on its own, packs rounded up, mix designs against the planning team's figures, one area set
for every line, cut-outs into Misc, the missed-area warning, the Excel equal to the page, an old
planning sheet read back, the margin check, free-issue materials, the T&C wording, and the site's
request turned into an indent."""

import io
from datetime import date, datetime, time
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import select

from app.estimate import demo, terms
from app.estimate import library as lib
from app.estimate import mix as mixes
from app.estimate.models import Coverage, MixDesign
from app.masters.models import Product, System
from app.models import User
from app.sites.models import Site, SiteMember
from tests.test_finance import boss  # noqa: F401  (fixture)

D = Decimal


def _rect(x, y, w, h):
    return [[x, y, 0], [x + w, y, 0], [x + w, y + h, 0], [x, y + h, 0]]


@pytest.fixture
def terrace(boss, db):  # noqa: F811
    """A site with an upper terrace (two shafts inside it), a lower terrace, the DEMO library."""
    c, h = boss
    ids = demo.library(db)
    pu = System(code="S-PU-T", name="Invented PU coating", unit="sqm", upturn_mm=D(300))
    db.add(pu)
    site = Site(code="S-TER-1", name="Invented terraces", status="active")
    db.add(site)
    db.commit()
    for code, name, o in (("UT", "Upper terrace", _rect(0, 0, 24, 14)), ("S1", "Shaft 1", _rect(3, 3, 2, 1.5)), ("S2", "Shaft 2", _rect(17, 9, 1.5, 1.2)),
                          ("LT", "Lower terrace", _rect(0, -12, 25, 8))):  # fmt: skip
        r = c.post(f"/api/measure/sites/{site.id}/areas", json={"outline": o, "code": code, "name": name}, headers=h)
        assert r.status_code == 201, r.text
    nod = db.scalar(select(MixDesign.id).where(MixDesign.code == "DEMO-NODULAR-114"))
    return {"c": c, "h": h, "site": site, "ids": ids, "pu": pu, "nod": nod}


def _estimate(t, lines=()):
    c, h = t["c"], t["h"]
    e = c.post("/api/estimate/estimates", json={"site_id": t["site"].id, "title": "Terraces"}, headers=h)
    assert e.status_code == 201, e.text
    e = e.json()
    for body in lines:
        r = c.post(f"/api/estimate/estimates/{e['id']}/lines", json=body, headers=h)
        assert r.status_code == 201, r.text
        e = r.json()
    return e


# --- the library ----------------------------------------------------------------------------------


def test_juno_a_method_statement_figure_outside_the_library_warns(terrace, db):
    """Juno: the method statement says 1.2-1.5 kg/sqm for Ultrathane, the BOQ 1.8-1.9, the library
    1.8 (1.8 to 2.0). The method statement warns; the BOQ does not."""
    c, h = terrace["c"], terrace["h"]
    ms = "Apply two coats of Ultrathane at 1.2–1.5 kg/sqm over the primer. Primer: Epoxy Primer 021 at 0.2 kg/sqm."
    boq = "PU coating with Ultrathane 1.8-1.9 kg/sqm, Epoxy Primer 021 at 5 sqm per kg."
    out = lib.check_text(db, ms, None, "Method statement")
    assert [w["product"] for w in out] == ["Ultrathane"] and "1.2–1.5 kg/sqm is outside the library's 1.8–2 kg/sqm" in out[0]["warning"]
    assert lib.check_text(db, boq) == []  # 1.8-1.9 is inside; 5 sqm per kg = 0.2 kg/sqm, inside 0.15-0.25
    assert [w["product"] for w in lib.check_text(db, "Ultrathane Logic Foam at 0.5 kg/sqm")] == ["Ultrathane Logic Foam"]  # the longer name, not Ultrathane
    assert lib.check_text(db, "Ultrathane, 2.5 m wide strips at 1.2 m c/c") == []  # no coverage written
    r = c.post("/api/estimate/check-text", json={"text": ms}, headers=h).json()
    assert len(r["warnings"]) == 1
    # the BOQ before it is sent: a tender's lines
    from app.tenders.models import BoqLine, Tender

    t = Tender(code="T-INV-1", name="Invented", status="draft")
    db.add(t)
    db.flush()
    db.add_all([BoqLine(tender_id=t.id, sort_order=1, client_item_no="iv", description=boq), BoqLine(tender_id=t.id, sort_order=2, client_item_no="v", description="Ultrathane at 1.0 kg/sqm")])
    db.commit()
    chk = c.get("/api/estimate/checks", params={"tender_id": t.id}, headers=h).json()
    assert [w["where"] for w in chk["coverage"]] == ["BOQ item v"]
    # the method statement before it is sent: the export waits until the warning is seen
    from app.methodology.models import Methodology, MethodologyItem

    m = Methodology(code="MS-INV-1", title="Invented")
    db.add(m)
    db.flush()
    db.add(MethodologyItem(methodology_id=m.id, name="Terrace", stages=[{"name": "Coating", "steps": [{"text": ms}]}]))
    db.commit()
    r = c.post(f"/api/methodology/methodologies/{m.id}/export", json={"audience": "client"}, headers=h)
    assert r.status_code == 409 and "Ultrathane at 1.2–1.5 kg/sqm" in r.json()["detail"]


def test_wastage_is_its_own_factor_and_packs_round_up(terrace, db):
    p = db.get(Product, terrace["ids"]["Ultrathane"])
    assert p.wastage_percent == 5 and p.pack_size == 20 and p.pack_unit == "drum"
    pk = lib.packs(D("100.01"), p)
    assert pk.count == 6 and pk.text == "6 drums of 20 kg"  # 5.0005 drums: always up
    assert lib.packs(D(100), p).count == 5
    e = _estimate(terrace, [{"product_id": p.id, "system_id": terrace["pu"].id}])
    ln = e["lines"][0]
    assert ln["coverage"].startswith("1.8 kg/sqm") and "%" not in ln["coverage"] and D(ln["wastage_percent"]) == 5
    base = sum(D(r["base_sqm"]) for r in ln["rows"])
    assert D(ln["qty"]) == base * D("1.8") * D("1.05") and ln["packs"] == lib.packs(D(ln["qty"]), p).text
    assert all("× (1 + 5%)" in r["formula"] for r in ln["rows"])
    # every change to the library is logged
    cov = db.scalar(select(Coverage).where(Coverage.product_id == p.id, Coverage.system_id.is_(None)))
    c, h = terrace["c"], terrace["h"]
    r = c.put(f"/api/estimate/coverages/{cov.id}", json={"product_id": p.id, "value": "1.9", "min_value": "1.8", "max_value": "2", "unit": "kg/sqm", "source": "tds"}, headers=h)
    assert r.status_code == 200 and r.json()["source_label"] == "data sheet (TDS)"
    assert c.put(f"/api/estimate/coverages/{cov.id}", json={"product_id": p.id, "value": "2.5", "min_value": "1.8", "max_value": "2", "unit": "kg/sqm"}, headers=h).status_code == 422
    changes = c.get("/api/estimate/library/changes", headers=h).json()
    assert changes[0]["action"] == "coverage.update" and changes[0]["before"]["value"] in ("1.8", "1.8000") and changes[0]["after"]["value"] in ("1.9", "1.9000")


def test_mix_designs_match_the_planning_sheet(terrace, db):
    """Nodular concrete 1:1:4: 3.52 kg/sqft of nodular and of cement at 165 mm, sand × 4, 2.77 at
    130 mm; the admixture per bag of cement (0.5 L SBR latex per 50 kg; Technomix 200 ml)."""
    m = db.scalar(select(MixDesign).where(MixDesign.code == "DEMO-NODULAR-114"))
    at = {p.role: p for p in mixes.parts(db, m, D(165))}
    assert at["nodular"].per_sqft == D("3.52") and at["cement"].per_sqft == D("3.52") and at["sand"].per_sqft == D("14.08")
    assert at["cement"].per_sqm == D("3.52") * lib.SQFT_PER_SQM
    assert at["admixture"].per_sqm == at["cement"].per_sqm / 50 * D("0.5") and at["admixture"].unit == "ltr"
    at130 = {p.role: p for p in mixes.parts(db, m, D(130))}
    assert at130["nodular"].per_sqft == D("2.77") and at130["cement"].per_sqft == D("2.77") and at130["sand"].per_sqft == D("11.09")
    plaster = db.scalar(select(MixDesign).where(MixDesign.code == "DEMO-PLASTER-14"))
    pp = {p.role: p for p in mixes.parts(db, plaster, D(20))}
    assert pp["cement"].per_sqm == D("0.020") * D("1.27") / 5 * 1440  # 20 mm × 1.27 × 1/5 × 1440
    assert pp["admixture"].per_sqm == pp["cement"].per_sqm / 50 * D("0.2")  # Technomix WL 200 ml per bag
    c, h = terrace["c"], terrace["h"]
    got = c.get(f"/api/estimate/mixes/{m.id}", params={"thickness_mm": "130"}, headers=h).json()
    assert {p["role"]: p["per_sqft"] for p in got["parts"]}["nodular"] in ("2.77", 2.77)


# --- the estimate ---------------------------------------------------------------------------------


def test_one_area_set_and_cut_outs_go_to_misc(terrace, db):
    t = terrace
    e = _estimate(t, [{"product_id": t["ids"]["Ultrathane"], "system_id": t["pu"].id}, {"product_id": t["ids"]["120 gsm geotextile"], "system_id": t["pu"].id, "part": "floor"},
                      {"product_id": t["ids"]["45 gsm fibre mesh"], "system_id": t["pu"].id, "part": "upturn"}])  # fmt: skip
    by = {a["code"]: a for a in e["areas"]}
    assert set(by) == {"UT", "S1", "S2", "LT"} and not e["missed"]
    ut = by["UT"]
    assert D(ut["floor_sqm"]) == 336 and D(ut["deduction_sqm"]) == D("4.8") and by["S1"]["source"] == "misc" and by["S1"]["parent_id"] == ut["id"]
    assert by["S1"]["use_floor"] is False and by["S1"]["use_upturn"] is True and by["S1"]["name"].startswith("Misc: Shaft 1")
    # one area set: every line reads the same floor for the same area
    floors = {ln["product"]: {r["area_id"]: D(r["floor_sqm"]) for r in ln["rows"]} for ln in e["lines"]}
    assert floors["Ultrathane"][ut["id"]] == floors["120 gsm geotextile"][ut["id"]] == D("331.2")  # 336 less the two shafts, counted once
    assert by["S1"]["id"] not in floors["120 gsm geotextile"]  # a shaft's floor is not laid
    geo = next(ln for ln in e["lines"] if ln["product"] == "120 gsm geotextile")
    assert D(geo["qty"]) == (D("331.2") + 200) * D("1.1")
    mesh = next(ln for ln in e["lines"] if ln["product"] == "45 gsm fibre mesh")
    per = 76 + D(7) + D("5.4") + 66  # the terraces' and the shafts' peripheries
    assert D(mesh["qty"]) == per * D("0.3") * D("1.1") and mesh["packs"] == "2 rolls of 50 sqm"


def test_a_measured_area_left_out_warns_until_it_is_in_or_excluded(terrace, db):
    """Dholera: a 29.99 sqm area measured on the drawing never reached the estimate."""
    t = terrace
    c, h = t["c"], t["h"]
    e = _estimate(t, [{"product_id": t["ids"]["Ultrathane"], "system_id": t["pu"].id}])
    assert not e["warnings"]
    r = c.post(f"/api/measure/sites/{t['site'].id}/areas", json={"outline": _rect(30, -12, 5.45, 5.5028), "code": "CN", "name": "Canopy"}, headers=h)
    assert r.status_code == 201
    e = c.get(f"/api/estimate/estimates/{e['id']}", headers=h).json()
    assert [m["name"] for m in e["missed"]] == ["Canopy"] and "Canopy (29.99 sqm)" in e["warnings"][0]
    assert c.post(f"/api/estimate/estimates/{e['id']}/issue", headers=h).status_code == 409
    assert c.post(f"/api/estimate/estimates/{e['id']}/exclude", json={"key": e["missed"][0]["key"], "reason": ""}, headers=h).status_code == 422
    e = c.post(f"/api/estimate/estimates/{e['id']}/exclude", json={"key": e["missed"][0]["key"], "reason": "Done by the civil contractor"}, headers=h).json()
    assert not e["missed"] and not e["warnings"]
    cn = next(a for a in e["areas"] if a["code"] == "CN")
    assert cn["included"] is False and cn["excluded_reason"] == "Done by the civil contractor"
    assert all(r["area_id"] != cn["id"] for r in e["lines"][0]["rows"])
    assert c.post(f"/api/estimate/estimates/{e['id']}/issue", headers=h).json()["status"] == "issued"


def test_the_excel_totals_equal_the_page(terrace, db):
    t = terrace
    c, h = t["c"], t["h"]
    e = _estimate(t, [{"product_id": t["ids"]["Nodular"], "system_id": None, "mix_design_id": t["nod"], "thickness_mm": "165", "part": "floor"},
                      {"product_id": t["ids"]["Ultrathane"], "system_id": t["pu"].id}])  # fmt: skip
    r = c.get(f"/api/estimate/estimates/{e['id']}/excel", headers=h)
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[-2:] == ["Per area", "Summary"] and wb.sheetnames[0] == "Nodular"  # per material, per area, the summary
    s = wb["Summary"]
    rows = {row[0]: row for row in s.iter_rows(min_row=5, values_only=True) if row[0] and row[1] is not None}
    for line in e["summary"]:
        assert D(str(rows[line["product"]][1])) == D(line["qty"]).quantize(D("0.001")) and rows[line["product"]][3] == line["packs"]
    assert any(row[0] == "Prepared by: Planning" for row in s.iter_rows(values_only=True))
    nod = wb["Nodular"]
    total = next(row for row in nod.iter_rows(values_only=True) if row[0] == "Total")
    assert D(str(total[7])) == D(e["lines"][0]["qty"]).quantize(D("0.001"))
    assert c.get(f"/api/estimate/estimates/{e['id']}/pdf", headers=h).content[:4] == b"%PDF"


def test_an_old_planning_sheet_is_read_against_the_drawing_and_the_library(terrace, db):
    """What the old sheets did: the same floor as 1,693 and 1,695, a coverage typed differently,
    "300 TO 340" as text, the latex at 140 mm for a 165 mm screed."""
    t = terrace
    c, h = t["c"], t["h"]
    e = _estimate(t)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PU"
    ws.append(["Area", "Sqm", "Coverage", "Qty"])
    ws.append(["Upper terrace", 336, None, None])
    ws.append(["Ultrathane", 336, 1.5, "=B3*C3*1.05"])
    ws.append(["Upturn", "300 TO 340", None, None])
    ws2 = wb.create_sheet("Screed")
    ws2.append(["Area", "Sqm"])
    ws2.append(["Upper terrace", 338])
    ws2.append(["SBR Latex for 140 mm screed", 338, None, "=B3*0.38"])
    ws2.append(["Nodular 165 mm", 338])
    ws2.append(["SBR Latex at 165 mm", 338])
    buf = io.BytesIO()
    wb.save(buf)
    r = c.post(f"/api/estimate/estimates/{e['id']}/import", files={"file": ("old.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=h)
    assert r.status_code == 200, r.text
    f = "\n".join(r.json()["findings"])
    assert "Screed B2: Upper terrace is 338 sqm; the drawing measures 336 sqm (+2)" in f
    assert "PU C3: Ultrathane at 1.5; the library says 1.8 kg/sqm" in f
    assert "'300 TO 340' is typed as text" in f
    assert "Upper terrace has 2 figures across the sheets: 336 sqm (PU), 338 sqm (Screed)" in f
    assert "SBR Latex is worked at 140 mm and 165 mm" in f


# --- the margin check, free issue, the T&C wording ----------------------------------------------------


def _working(t, db, rate="1072", free=()):
    from app.costing.models import RateGroup, RateGroupLine, RateGroupProduct, RateWorking
    from app.masters.models import ProductPrice
    from app.tenders.models import BoqLine, Tender

    for name, price in (("Ultrathane", "480"), ("Epoxy Primer 021", "350"), ("Cement", "8")):
        db.add(ProductPrice(product_id=t["ids"][name], purchase_rate=D(price), effective_from=date(2026, 1, 1), source="manual"))
    tender = Tender(code="T-INV-2", name="Invented terraces", status="draft")
    db.add(tender)
    db.flush()
    b = BoqLine(tender_id=tender.id, sort_order=1, client_item_no="iv", unit="sqm", description="PU coating", rate=D(rate))
    db.add(b)
    w = RateWorking(tender_id=tender.id, rev_no=0, free_issue_product_ids=[t["ids"][n] for n in free])
    db.add(w)
    db.flush()
    g = RateGroup(working_id=w.id, name="PU coating", system_id=t["pu"].id, rate_sqm=D(rate))
    db.add(g)
    db.flush()
    for k, name in enumerate(("Epoxy Primer 021", "Ultrathane", "Cement", "Cemkrete")):
        db.add(RateGroupProduct(group_id=g.id, product_id=t["ids"][name], sort_order=k))
    db.add(RateGroupLine(group_id=g.id, boq_line_id=b.id))
    db.commit()
    return tender


def test_margin_warns_above_the_directors_share_and_free_issue_is_out_of_cost(terrace, db, login_as):
    """Juno item iv: material at the client sheet's figures was about 91% of ₹1,072/sqm."""
    t = terrace
    c, h = t["c"], t["h"]
    tender = _working(t, db, free=("Cement",))
    m = c.get("/api/estimate/margin", params={"tender_id": tender.id}, headers=h).json()
    row = m["lines"][0]
    material = (D("0.2") * D("1.05") * 350 + D("1.8") * D("1.05") * 480).quantize(D("0.01"))
    assert D(row["material_sqm"]) == material and D(row["material_share_percent"]) == (material / 1072 * 100).quantize(D("0.1"))  # 91.5%
    assert D(row["margin_percent"]) == ((1072 - material) / 1072 * 100).quantize(D("0.1"))  # on the selling price
    assert any("Material alone is 91.5% of the rate ₹1,072.00/sqm (over 70%)" in w for w in row["warnings"])
    assert row["free_issue"] == ["Cement"] and row["no_price"] == ["Cemkrete"]  # free issue out of the cost; no price: listed, not complete
    # the director raises the share: no warning; another role cannot change it
    pc, ph = login_as("planning", email="planner@example.com")
    assert pc.put("/api/estimate/margin/settings", json={"material_share_warn_percent": "95"}, headers=ph).status_code == 403
    assert c.put("/api/estimate/margin/settings", json={"material_share_warn_percent": "95"}, headers=h).status_code == 200
    row = c.get("/api/estimate/margin", params={"tender_id": tender.id}, headers=h).json()["lines"][0]
    assert not any("Material alone" in w for w in row["warnings"])
    # free issue: the clause is written in the tender's T&C, the cost drops
    out = c.put("/api/estimate/margin/free-issue", json={"tender_id": tender.id, "product_ids": [t["ids"]["Cement"], t["ids"]["Sand"]]}, headers=h).json()
    assert out["free_issue_clause"] == "Cement and Sand shall be supplied by the client free of cost at site; they are not in our rates."
    from app.tenders.models import TenderTc

    tcs = list(db.scalars(select(TenderTc).where(TenderTc.tender_id == tender.id)))
    assert len(tcs) == 1 and tcs[0].text_override.endswith("supplied by the client free of cost at site; they are not in our rates.")
    assert c.get("/api/estimate/checks", params={"tender_id": tender.id}, headers=h).json()["terms"] == []  # the app's own clause reads right


def test_tc_wording_that_gives_civil_materials_free_is_flagged(terrace, db):
    c, h = terrace["c"], terrace["h"]
    wrong = "Cement, sand and water shall be provide by us at free of cost."
    f = terms.check(wrong, "T&C 3")
    assert f["corrected"] == "Cement, sand and water shall be provide by the client at free of cost." and "T&C 3" in f["warning"]
    assert terms.check("Water and electricity shall be provided by the client free of cost.") is None
    assert terms.check("Lifting of material free of cost by EESPL.") is not None
    from app.masters.models import TcClause

    cl = TcClause(text=wrong, category="client_scope")
    db.add(cl)
    db.commit()
    lib_f = c.get("/api/estimate/terms/library", headers=h).json()
    assert [x["clause_id"] for x in lib_f] == [cl.id]
    assert c.post("/api/estimate/terms/fix", json={"clause_id": cl.id}, headers=h).json() == []
    db.refresh(cl)
    assert cl.text == "Cement, sand and water shall be provide by the client at free of cost."


# --- the site's request -------------------------------------------------------------------------


def _engineer(t, db, login_as):
    ec, eh = login_as("site_engineer", email="engineer@example.com")
    uid = db.scalar(select(User.id).where(User.email == "engineer@example.com"))
    db.add(SiteMember(site_id=t["site"].id, user_id=uid, role_on_site="incharge"))
    db.commit()
    return ec, eh


def test_a_site_request_by_area_and_thickness_becomes_an_indent(terrace, db, login_as):
    t = terrace
    c, h = t["c"], t["h"]
    e = _estimate(t, [{"product_id": t["ids"]["Nodular"], "mix_design_id": t["nod"], "thickness_mm": "165", "part": "floor"}])
    ut = next(a for a in e["areas"] if a["code"] == "UT")
    ec, eh = _engineer(t, db, login_as)
    body = {"site_id": t["site"].id, "product_id": t["ids"]["Nodular"], "area_sqm": "280", "thickness_mm": "130", "mix_design_id": t["nod"], "estimate_area_id": ut["id"]}
    pre = ec.post("/api/estimate/requests", json=body | {"preview": True}, headers=eh).json()
    assert pre["formula"] == "280 sqm × 10.7639 × 2.77 kg/sqft (Nodular concrete 1:1:4, 130 mm) = 8,349 kg = 334 bags of 25 kg"
    assert D(pre["qty"]) == 8349 and pre["packs"] == 334 and pre["over"] is False
    r = ec.post("/api/estimate/requests", json=body, headers=eh)
    assert r.status_code == 200, r.text
    r = r.json()
    assert r["status"] == "requested" and D(r["left_before"]) == (D("331.2") * D("3.52") * lib.SQFT_PER_SQM)
    assert ec.post(f"/api/estimate/requests/{r['id']}/decide", json={"approve": True}, headers=eh).status_code == 403  # planning decides
    # a changed quantity needs a note; approved, the indent is raised (the indent → PO flow)
    assert c.post(f"/api/estimate/requests/{r['id']}/decide", json={"approve": True, "qty": "8400"}, headers=h).status_code == 422
    d = c.post(f"/api/estimate/requests/{r['id']}/decide", json={"approve": True, "qty": "8400", "note": "Round to whole bags plus the ramp"}, headers=h).json()
    assert d["status"] == "indented" and d["indent"]["status"] == "submitted"
    from app.material.models import Indent

    ind = db.get(Indent, d["indent"]["id"])
    assert ind.site_id == t["site"].id and ind.lines[0].qty == 8400 and ind.lines[0].unit == "kg" and "336 bags of 25 kg" in ind.lines[0].remark and r["code"] in ind.remark


def test_asking_for_more_than_the_estimate_has_left_needs_a_reason(terrace, db, login_as):
    t = terrace
    e = _estimate(t, [{"product_id": t["ids"]["Nodular"], "mix_design_id": t["nod"], "thickness_mm": "165", "part": "floor"}])
    lt = next(a for a in e["areas"] if a["code"] == "LT")  # 200 sqm × 37.89 kg = 7,578 kg planned
    ec, eh = _engineer(t, db, login_as)
    body = {"site_id": t["site"].id, "product_id": t["ids"]["Nodular"], "area_sqm": "200", "thickness_mm": "165", "mix_design_id": t["nod"], "estimate_area_id": lt["id"]}
    assert ec.post("/api/estimate/requests", json=body, headers=eh).status_code == 200  # all of it
    again = ec.post("/api/estimate/requests", json=body | {"area_sqm": "20"}, headers=eh)
    assert again.status_code == 422 and "more than the estimate has left" in again.json()["detail"]
    ok = ec.post("/api/estimate/requests", json=body | {"area_sqm": "20", "reason": "Second layer at the drain falls"}, headers=eh)
    assert ok.status_code == 200 and ok.json()["over_reason"] == "Second layer at the drain falls"
    bal = {b["product"]: b for b in t["c"].get(f"/api/estimate/estimates/{e['id']}/balance", headers=t["h"]).json()}
    assert D(bal["Nodular"]["requested"]) == D(ok.json()["qty"]) + 7578


# --- the bill's material check ----------------------------------------------------------------------


def test_the_material_check_lists_billed_items_with_no_coverage(db, boss):  # noqa: F811
    from app.closeout.models import ClientWorkOrderLine
    from app.masters.models import SystemComponent
    from app.material.models import StockLedger, Store
    from tests.test_measure_billing import pool as _pool  # noqa: F401
    from tests.test_slips_register import _slip

    pytest.importorskip("app.measure.register")
    # built inline (the pool fixture is not requested here): a site, its work order and contract
    from tests.test_measure_billing import pool as fixture_fn

    pool = fixture_fn.__wrapped__(boss, db)
    c, h = pool["c"], pool["h"]
    ids = demo.library(db)  # Cemguard Plus: 1.75 kg/sqm for any system
    sy = System(code="S-INV-2", name="Invented membrane system", unit="sqm")
    bare = System(code="S-INV-3", name="Invented system, product not in the library", unit="sqm")
    db.add_all([sy, bare])
    db.flush()
    other = Product(code="P-INV-9", name="Invented sealant", unit="kg")
    db.add(other)
    db.flush()
    db.add_all([SystemComponent(system_id=sy.id, product_id=ids["Cemguard Plus"]), SystemComponent(system_id=bare.id, product_id=other.id)])
    w1, w3 = db.get(ClientWorkOrderLine, pool["wl"]["1"]), db.get(ClientWorkOrderLine, pool["wl"]["3"])
    w1.system_id, w3.system_id = sy.id, bare.id
    db.get(ClientWorkOrderLine, pool["wl"]["2"]).system_id = sy.id  # a count: no area to cover
    store = Store(name="Invented site store", kind="site", site_id=pool["site"].id)
    db.add(store)
    db.flush()
    db.add(StockLedger(store_id=store.id, product_id=ids["Cemguard Plus"], qty=D(800), unit="kg", rate=D(0), value=D(0), ref_type="transfer_in", at=datetime.combine(date(2026, 8, 1), time(12)).astimezone()))
    db.commit()
    s = _slip(pool)
    for item, n, length, width in (("1", "1", "100", "10"), ("2", "3", None, None), ("3", "1", "10", "10")):
        c.post(f"/api/measure/sheets/{s['id']}/slip-lines", json={"particulars": f"Item {item}", "no": n, "length": length, "width": width, "wo_line_id": pool["wl"][item]}, headers=h)
    c.post(f"/api/measure/sheets/{s['id']}/send", headers=h)
    b = c.post(f"/api/measure/sheets/{s['id']}/bill", json={"basis": "verbal", "verbal_by": "Client's site engineer", "verbal_on": "2026-05-23", "proforma_on": "2026-09-01"}, headers=h).json()
    m = c.get(f"/api/measure/bills/{b['id']}/material", headers=h).json()
    row = next(r for r in m["rows"] if r["product"] == "Cemguard Plus")
    assert D(row["needed"]) == (D(1000) / D("10.7639") * D("1.75")).quantize(D("0.01")) and "(library)" in row["from"][0]
    why = {u["item"]: u["why"] for u in m["uncovered"]}
    assert m["complete"] is False and why["2"] == "measured in nos, not by area" and why["3"] == "no coverage for Invented sealant"
