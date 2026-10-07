"""Rate policies and their backtest, EESPL BOQ export (Excel, PDF, client format), revisions."""

import io
import re
import zipfile
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select

from app.config import settings
from app.masters.models import LibraryItem, LibraryLine, TcClause, TcTemplate, TcTemplateClause
from app.masters.rate_policy import History, Obs, backtest, choose, file_order, winner
from app.models import AuditLog
from app.tenders.boq_import import parse, split_heading
from app.tenders.models import BoqImport
from tests.test_tenders import (
    boq,
    confirm,
    imported_tender,
    line,
    membrane_system,
    new_client,
    new_tender,
    upload,
)

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture(autouse=True)
def media(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "media_dir", str(tmp_path))


@pytest.fixture
def boss(login_as):
    return login_as("super_admin", email="boss@example.com")


# --- backtest maths ------------------------------------------------------------------------------


def _history() -> dict[int, History]:
    one = [
        Obs(Decimal(100), "f1", "X", 0),
        Obs(Decimal(110), "f2", "Y", 1),
        Obs(Decimal(130), "f3", "X", 2),
    ]
    return {
        1: History(1, "sqm", "f3", one),
        2: History(2, "sqm", "f1", [Obs(Decimal(50), "f1", "X", 0)]),  # quoted once: no pool
    }


def test_policies_on_a_known_pool():
    pool = [Obs(Decimal(100), "f1", "X", 0), Obs(Decimal(110), "f2", "Y", 1)]
    assert choose("latest", pool).rate == Decimal("110.00")
    assert choose("latest", pool, preferred_file="f1").rate == Decimal(
        "100.00"
    )  # item's own latest
    assert choose("median", pool).rate == Decimal("105.00")
    assert choose("channel_last", pool, "Y").rate == Decimal("110.00")
    assert choose("channel_last", pool, "Z").policy == "median"  # no same-channel rate: fallback
    assert choose("channel_median", pool, "x").rate == Decimal("100.00")  # case-insensitive channel
    rates = [Obs(Decimal(v), f"f{v}", None, 0) for v in (1, 50, 51, 52, 53, 54, 55, 56, 57, 1000)]
    assert choose("trimmed_mean", rates).rate == Decimal("53.50")  # 1 and 1000 dropped
    assert choose("median", []) is None


def test_leave_one_boq_out_backtest():
    rows = {r["policy"]: r for r in backtest(_history(), unlinked=1)}
    # targets: f1 100 <- {110, 130}, f2 110 <- {100, 130}, f3 130 <- {100, 110}; the f1 line of
    # item 2 has nothing to learn from; 1 unlinked line. Coverage = 3 of 5.
    latest = rows["latest"]  # 130 (+30%), 130 (+18.2%), 110 (-15.4%: f3 hidden, f2 is next)
    assert (latest["within_10"], latest["within_20"], latest["above_by_10"]) == (0.0, 66.7, 66.7)
    assert (latest["median_abs_error"], latest["coverage"], latest["lines"]) == (18.2, 60.0, 5)
    median = rows["median"]  # 120 (+20%), 115 (+4.5%), 105 (-19.2%)
    assert (median["within_5"], median["within_10"], median["within_20"]) == (33.3, 33.3, 100.0)
    assert (median["above_by_10"], median["median_abs_error"]) == (33.3, 19.2)
    client = rows["channel_median"]  # X: 130 (+30%), Y: none -> 115 (+4.5%), X: 100 (-23.1%)
    assert (client["within_10"], client["median_abs_error"]) == (33.3, 23.1)
    assert winner(list(rows.values())) == "median"


def test_boq_order_from_latest_files():
    order = file_order({1: {"a", "b", "c"}, 2: {"b", "d"}}, {1: "c", 2: "d"})
    assert order["a"] < order["c"] and order["b"] < order["c"] and order["b"] < order["d"]


# --- policy in suggest ---------------------------------------------------------------------------


@pytest.fixture
def crystal_history(db):
    item = LibraryItem(
        source_key="t-crystal",
        match_key="crystal",
        unit="sqm",
        unit_raw="Sqm",
        description="Crystalline coating to raft with two coats",
        boq_count=3,
        latest_rate=Decimal(600),
        median_rate=Decimal(500),
        latest_source_file="c.xlsx",
    )
    db.add(item)
    db.flush()
    for i, (folder, rate) in enumerate(
        (("EXAMPLE DEVELOPERS", 400), ("OTHER CO", 500), ("OTHER CO", 600))
    ):
        db.add(
            LibraryLine(
                source_key=f"t-crystal-{i}",
                library_item_id=item.id,
                channel=folder,
                file=f"{folder}/{'abc'[i]}.xlsx",
                row=5,
                description=item.description,
                unit="sqm",
                rate=Decimal(rate),
            )
        )
    db.commit()
    return item


def _suggested(client, headers, tid):
    r = client.post(f"/api/tenders/{tid}/boq/suggest", headers=headers)
    assert r.status_code == 200, r.text
    crystal = line(boq(client, headers, tid), "Crystalline coating")
    detail = client.get(f"/api/tenders/{tid}/lines/{crystal['id']}", headers=headers).json()
    return crystal, detail["candidates"][0]


def test_rate_policy_setting_changes_the_suggestion(boss, crystal_history):
    client, headers = boss
    tid = imported_tender(client, headers)["id"]
    # the tender came through the channel the library calls "EXAMPLE DEVELOPERS"
    r = client.post("/api/channels", json={"name": "Example Developers", "type": "partner"},
                    headers=headers)  # fmt: skip
    assert r.status_code == 201, r.text
    r = client.patch(f"/api/tenders/{tid}", json={"channel_id": r.json()["id"]}, headers=headers)
    assert r.status_code == 200, r.text

    def policy(name):
        r = client.patch("/api/settings/company", json={"rate_policy": name}, headers=headers)
        assert r.status_code == 200, r.text

    policy("latest")
    crystal, cand = _suggested(client, headers, tid)
    assert crystal["rate"] == "600.00"
    h = cand["details"]
    assert (h["used"], h["latest_rate"], h["median_rate"], h["n_boqs"]) == (
        "latest",
        "600.00",
        "500.00",
        3,
    )
    assert (h["min_rate"], h["max_rate"], h["channel_last_rate"]) == ("400.00", "600.00", "400.00")
    assert h["warning"] is True and h["above_median_percent"] == "20.0"  # > 15% above median
    assert [s["rate"] for s in h["sources"]] == ["600.00"]

    policy("channel_median")  # the backtest default
    crystal, cand = _suggested(client, headers, tid)
    assert crystal["rate"] == "400.00" and cand["details"]["used"] == "channel_median"
    assert cand["details"]["warning"] is False
    assert cand["details"]["sources"][0]["channel"] == "EXAMPLE DEVELOPERS"

    policy("median")
    crystal, _ = _suggested(client, headers, tid)
    assert crystal["rate"] == "500.00"
    assert (
        client.patch(
            "/api/settings/company", json={"rate_policy": "guess"}, headers=headers
        ).status_code
        == 422
    )


def test_rate_history_min_max_need_tender_margin(login_as, crystal_history):
    client, headers = login_as("sales", email="sales@example.com")
    tid = new_tender(client, headers, new_client(client, headers))["id"]
    assert confirm(client, headers, tid, upload(client, headers, tid)).status_code == 200
    _, cand = _suggested(client, headers, tid)
    h = cand["details"]
    assert h["median_rate"] == "500.00" and h["latest_rate"] == "600.00"  # selling rates: shown
    assert "min_rate" not in h and "max_rate" not in h  # cost-type fields: absent


# --- export --------------------------------------------------------------------------------------


@pytest.fixture
def priced(boss, db):
    """The fixture BOQ, priced: a system line (cost 320, rate 384), typed rates, QRO and NQ
    lines, T&C with a per-tender override, products and remarks."""
    client, headers = boss
    membrane_system(client, headers)
    clauses = [TcClause(text=t, category="general") for t in ("Water by client", "GST extra")]
    db.add_all(clauses)
    db.flush()
    db.add(
        TcTemplate(
            name="Standard",
            is_default=True,
            clauses=[TcTemplateClause(clause_id=c.id, sort_order=i) for i, c in enumerate(clauses)],
        )
    )
    db.commit()
    tid = imported_tender(client, headers)["id"]
    client.post(f"/api/tenders/{tid}/boq/suggest", headers=headers)
    data = boq(client, headers, tid)
    assert line(data, "APP membrane")["source"] == "system"
    edits = [
        {"id": line(data, "Crystalline coating")["id"], "rate": "470", "our_product": "Acme Seal"},
        {"id": line(data, "Pipe sleeve")["id"], "rate": "355.55", "our_remarks": "Per opening"},
        {"id": line(data, "Protective screed")["id"], "rate": "275"},  # QRO
        {"id": line(data, "Lift pit")["id"], "rate": "1499.99"},
        {"id": line(data, "Planter box")["id"], "rate": "640"},
    ]
    r = client.patch(f"/api/tenders/{tid}/lines", json={"lines": edits}, headers=headers)
    assert r.status_code == 200, r.text
    terms = client.get(f"/api/tenders/{tid}/tc", headers=headers).json()
    r = client.put(
        f"/api/tenders/{tid}/tc",
        json={
            "items": [
                {
                    "clause_id": terms[1]["clause_id"],
                    "text_override": "GST extra at 18%, as applicable",
                },
                {"clause_id": terms[0]["clause_id"]},
            ]
        },
        headers=headers,
    )
    assert r.status_code == 200
    return client, headers, tid


FORBIDDEN_TEXT = re.compile(r"cost|margin|score|client_file_rate|suggest|\bsource\b", re.I)


def _evaluate(ws) -> dict[str, Decimal | str | None]:
    """A tiny evaluator for the two formula shapes the export writes."""
    cache: dict[str, Decimal | str | None] = {}

    def value(ref: str):
        if ref in cache:
            return cache[ref]
        v = ws[ref].value
        if isinstance(v, str) and v.startswith("=ROUND("):
            a, b = re.match(r"=ROUND\(([A-Z]+\d+)\*([A-Z]+\d+),2\)", v).groups()
            v = round(Decimal(str(value(a))) * Decimal(str(value(b))), 2)
        elif isinstance(v, str) and v.startswith("=SUBTOTAL(9,"):
            col, lo, hi = re.match(r"=SUBTOTAL\(9,([A-Z]+)(\d+):[A-Z]+(\d+)\)", v).groups()
            total = Decimal(0)
            for r in range(int(lo), int(hi) + 1):
                raw = ws[f"{col}{r}"].value
                if isinstance(raw, str) and raw.startswith("=SUBTOTAL"):
                    continue  # SUBTOTAL ignores nested subtotals
                x = value(f"{col}{r}")
                if isinstance(x, int | float | Decimal):
                    total += Decimal(str(x))
            v = total
        cache[ref] = v
        return v

    return value


def test_export_xlsx(priced):
    client, headers, tid = priced
    data = boq(client, headers, tid)
    r = client.get(f"/api/tenders/{tid}/export.xlsx", headers=headers)
    assert r.status_code == 200, r.text
    assert (
        "T-" in r.headers["content-disposition"]
        and "_R0_BOQ.xlsx" in r.headers["content-disposition"]
    )
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    cells = {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
    text = "\n".join(str(v) for v in cells.values())
    assert "Revision R0 (draft)" in text and "Subject: Waterproofing works" in text
    value = _evaluate(ws)
    total_ref = next(k for k, v in cells.items() if v == "Total (excl. GST)").replace("B", "F")
    assert value(total_ref) == Decimal(data["totals"]["subtotal"])
    rows = {ws[f"B{r}"].value: r for r in range(1, ws.max_row + 1) if ws[f"B{r}"].value}
    crystal = next(r for d, r in rows.items() if "Crystalline coating" in str(d))
    assert ws[f"F{crystal}"].value == f"=ROUND(D{crystal}*E{crystal},2)"
    assert value(f"F{crystal}") == Decimal("56635.00") and ws[f"G{crystal}"].value == "Acme Seal"
    screed = next(r for d, r in rows.items() if "Protective screed" in str(d))
    assert (ws[f"D{screed}"].value, ws[f"E{screed}"].value, ws[f"F{screed}"].value) == (
        "QRO",
        Decimal("275") if isinstance(ws[f"E{screed}"].value, Decimal) else 275,
        "Rate only",
    )
    coba = next(r for d, r in rows.items() if "Brick bat coba" in str(d))
    assert (ws[f"D{coba}"].value, ws[f"E{coba}"].value) == ("NQ", "Not quoted")
    for heading in ("Total — 1 Basement waterproofing", "Total — 2 Terrace waterproofing"):
        sub = rows[heading]
        expected = sum(
            Decimal(ln["amount"])
            for ln in data["lines"]
            if ln["amount"]
            and ln["section_id"]
            == next(s["id"] for s in data["sections"] if heading.endswith(s["title"]))
        )
        assert value(f"F{sub}") == expected
    start = next(r for r in range(1, ws.max_row + 1) if ws[f"A{r}"].value == "Terms & Conditions")
    tc = [
        ws[f"B{r}"].value
        for r in range(start + 1, ws.max_row + 1)
        if re.fullmatch(r"\d+\.", str(ws[f"A{r}"].value))
    ]
    assert tc == ["GST extra at 18%, as applicable", "Water by client"]
    assert "GST @18% extra at actual" in text and "Authorised signatory" in text
    # print setup
    assert ws.page_setup.orientation == "landscape" and ws.print_title_rows == "$10:$10"
    assert ws.freeze_panes == "A11" and ws.print_area
    # nothing cost-related anywhere in the file
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        blob = "\n".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist())
    assert not FORBIDDEN_TEXT.search(text)
    numbers = {str(v) for v in cells.values() if isinstance(v, int | float | Decimal)}
    assert not numbers & {"320", "320.0", "450", "450.0", "520", "1500.0", "600"}  # cost, client
    assert "<v>320</v>" not in blob and "<v>450</v>" not in blob


def test_export_pdf(priced):
    client, headers, tid = priced
    data = boq(client, headers, tid)
    r = client.get(f"/api/tenders/{tid}/export.pdf", headers=headers)
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    reader = PdfReader(io.BytesIO(r.content))
    text = "\n".join(p.extract_text() for p in reader.pages)
    flat = " ".join(text.split())
    assert "Total (excl. GST)" in flat and "GST @18% extra at actual" in flat
    subtotal = Decimal(data["totals"]["subtotal"])
    assert f"{subtotal:,.2f}".replace(",", "")[:4] in flat.replace(",", "")
    assert "1. GST extra at 18%, as applicable" in flat and "2. Water by client" in flat
    assert "Page 1 of" in flat and "Authorised signatory" in flat
    assert "QRO" in flat and "Rate only" in flat and "Not quoted" in flat
    assert not FORBIDDEN_TEXT.search(flat)
    assert "₹320.00" not in flat and "₹450.00" not in flat


def test_export_client_format(priced, db):
    client, headers, tid = priced
    r = client.get(f"/api/tenders/{tid}/export/client", headers=headers)
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content))["BOQ"]
    # the fixture's header is on row 4; "a) Crystalline coating" is on row 6, "b) Pipe sleeve" 7
    assert ws["B6"].value.startswith("Crystalline coating")
    assert (ws["E6"].value, ws["F6"].value) == (470, 56635)
    assert (ws["E7"].value, ws["F7"].value) == (355.55, 4266.6)
    assert ws["A1"].value.startswith("Name of work")  # their layout is kept
    assert "B17:G17" in {str(m) for m in ws.merged_cells.ranges}
    assert int(r.headers["x-rates-written"]) >= 5
    # an .xls original cannot be written back
    imp = db.scalar(select(BoqImport).where(BoqImport.tender_id == tid))
    xls = Path(settings.media_dir) / imp.stored_path
    xls.with_suffix(".xls").write_bytes(xls.read_bytes())
    imp.stored_path = str(Path(imp.stored_path).with_suffix(".xls"))
    db.commit()
    r = client.get(f"/api/tenders/{tid}/export/client", headers=headers)
    assert r.status_code == 409 and ".xlsx" in r.json()["detail"]


def test_client_format_needs_a_rate_column(boss):
    client, headers = boss
    tid = new_tender(client, headers, new_client(client, headers))["id"]
    preview = upload(client, headers, tid)
    preview["column_map"].pop("rate")
    assert confirm(client, headers, tid, preview).status_code == 200
    r = client.get(f"/api/tenders/{tid}/export/client", headers=headers)
    assert r.status_code == 409 and "rate column" in r.json()["detail"]


# --- revisions -----------------------------------------------------------------------------------


def test_revisions_submit_edit_compare_and_export(priced, db):
    client, headers, tid = priced
    data = boq(client, headers, tid)
    crystal = line(data, "Crystalline coating")
    r = client.post(f"/api/tenders/{tid}/submit", json={"note": "First offer"}, headers=headers)
    assert r.status_code == 200, r.text
    assert (r.json()["status"], r.json()["revision_label"]) == ("submitted", "R0")
    assert client.post(f"/api/tenders/{tid}/submit", json={}, headers=headers).status_code == 409

    # an edit after submitting starts R1 as a draft; R0 stays as it was
    r = client.patch(
        f"/api/tenders/{tid}/lines",
        json={"lines": [{"id": crystal["id"], "rate": "450"}]},
        headers=headers,
    )
    assert r.status_code == 200
    tender = client.get(f"/api/tenders/{tid}", headers=headers).json()
    assert (tender["status"], tender["revision"], tender["revision_label"]) == (
        "draft",
        1,
        "R1 (draft)",
    )
    actions = [a.action for a in db.scalars(select(AuditLog).order_by(AuditLog.id))]
    assert "tender.submit" in actions and "tender.revision.start" in actions
    assert client.post(f"/api/tenders/{tid}/submit", json={}, headers=headers).status_code == 200

    revs = client.get(f"/api/tenders/{tid}/revisions", headers=headers).json()
    assert [(x["label"], x["note"]) for x in revs] == [("R0", "First offer"), ("R1", None)]
    cmp = client.get(f"/api/tenders/{tid}/revisions/compare?a=0&b=1", headers=headers).json()
    changed = [x for x in cmp["lines"] if x["change"] != "same"]
    assert len(changed) == 1
    assert (changed[0]["rate_a"], changed[0]["rate_b"], changed[0]["rate_delta"]) == (
        "470.00",
        "450.00",
        "-20.00",
    )
    assert changed[0]["amount_delta"] == "-2410.00"  # 120.5 × -20
    assert Decimal(cmp["subtotal_delta"]) == Decimal("-2410.00")
    assert Decimal(cmp["grand_total_delta"]) == Decimal("-2843.80")  # with 18% GST

    # the old revision exports from its snapshot, not the live lines
    old = load_workbook(
        io.BytesIO(client.get(f"/api/tenders/{tid}/export.xlsx?rev=0", headers=headers).content)
    ).active
    new = load_workbook(
        io.BytesIO(client.get(f"/api/tenders/{tid}/export.xlsx", headers=headers).content)
    ).active

    def rate_of(ws):
        row = next(
            r for r in range(1, ws.max_row + 1) if "Crystalline coating" in str(ws[f"B{r}"].value)
        )
        return ws[f"E{row}"].value

    assert (rate_of(old), rate_of(new)) == (470, 450)
    assert "Revision R0" in str(old["A5"].value) and "Revision R1" in str(new["A5"].value)
    assert client.get(f"/api/tenders/{tid}/export.xlsx?rev=7", headers=headers).status_code == 404


def test_export_permissions(boss, login_as):
    boss_client, boss_headers = boss
    tid = new_tender(boss_client, boss_headers, new_client(boss_client, boss_headers))["id"]
    sales, sales_headers = login_as("sales", email="sales@example.com")  # tender.view: own
    for path in ("export.xlsx", "export.pdf", "export/client", "revisions"):
        assert sales.get(f"/api/tenders/{tid}/{path}", headers=sales_headers).status_code == 404
    customer, customer_headers = login_as("client", email="customer@example.com")
    for path in ("export.xlsx", "export.pdf", "export/client", "revisions"):
        assert (
            customer.get(f"/api/tenders/{tid}/{path}", headers=customer_headers).status_code == 403
        )
    own = new_tender(sales, sales_headers, new_client(sales, sales_headers, "Own Co"))["id"]
    assert sales.get(f"/api/tenders/{own}/export.xlsx", headers=sales_headers).status_code == 200


# --- addendum: section notes, cost totals ------------------------------------------------------


def test_heading_with_a_general_note_is_split(db):
    note = (
        "Unless otherwise specified, the Contractor shall carry out waterproofing in "
        "basements, terraces and sunken portions through an approved firm."
    )
    assert split_heading(f"WATER PROOFING — {note}") == ("WATER PROOFING", note)
    assert split_heading("Terrace waterproofing") == ("Terrace waterproofing", None)
    rows = [
        ["Sr", "Description", "Unit", "Qty"],
        ["600", "WATER PROOFING", None, None],
        ["", None, None, None],
        ["601", "Brick bat coba", "sqm", 10],
    ]
    rows[1][1] = f"WATER PROOFING — {note}"
    parsed = parse(rows, 1, {"item_no": 0, "description": 1, "unit": 2, "qty": 3})
    assert parsed.sections[0].title == "600 WATER PROOFING" and parsed.sections[0].note == note


def test_cost_totals_are_none_without_system_lines(boss):
    client, headers = boss
    tid = imported_tender(client, headers)["id"]
    data = boq(client, headers, tid)
    assert data["totals"]["cost_total"] is None and data["totals"]["margin_amount"] is None
    assert client.get(f"/api/tenders/{tid}", headers=headers).json()["cost_total"] is None
