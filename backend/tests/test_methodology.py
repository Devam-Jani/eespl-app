# ruff: noqa: E501, F811  (invented data reads better unwrapped; shared fixtures are imported)
"""The methodology builder, on invented data only (no client SOP, no real offer): seeding the
library from an invented offer (fixes, "to confirm", nothing invented), edits need a reason and
are versioned, invented old- and new-format SOP PDFs import into rules with their SOP, revision and
page, the rule is chosen by the toilet's size and the refuge's exposure, the compliance check
(a DFT below the minimum is red with the gap, a make not on the list is red with two choices, a
missing value is to confirm), the Word file (stages, numbered steps, spec table, checklist; the
client's copy without internals), the won BOQ on site (stage template, blocking hold point, pond
test records) and who sees what."""

import io
from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document
from sqlalchemy import select

from app.methodology import build as bld
from app.methodology import seed, sop
from app.methodology.models import (
    ClientStandard,
    Methodology,
    MethodProduct,
    MethodSystem,
    PondTest,
    StandardContradiction,
)
from app.sites.models import AreaScope, Site, SiteMember, SiteNode, StageTemplate, Task
from app.tenders.models import BoqLine, BoqSection, Tender
from tests.test_finance import boss  # noqa: F401  (fixtures)

# --- invented offers -----------------------------------------------------------------------------


def _bold(d, text):
    d.add_paragraph().add_run(text).bold = True


def offer_docx(path: Path, items: list[dict]) -> Path:
    d = Document()
    d.add_paragraph("Date: 01.01.2026")
    _bold(d, "SUB: OFFER FOR INVENTED WATERPROOFING WORK.")
    d.add_paragraph("Reference to subject, we submit our offer by Using BRONCO Products.")
    for it in items:
        _bold(d, it["title"])
        for heading, steps in it["stages"]:
            if heading == "OR":
                _bold(d, "OR")
                continue
            _bold(d, heading)
            for s in steps:
                d.add_paragraph(s)
    for i, it in enumerate(items, 1):
        _bold(d, f"ITEM NO: {i} - BUDGETARY OFFER FOR INVENTED AREA:")
        t = d.add_table(rows=2, cols=4)
        for c, v in zip(
            t.rows[0].cells, ("Sr No.", "Item Description", "UoM", "Rate"), strict=True
        ):
            c.text = v
        for c, v in zip(
            t.rows[1].cells,
            ("1.", it.get("boq", "Providing invented waterproofing"), "SQ FT", "55.00"),
            strict=True,
        ):
            c.text = v
    _bold(d, "GENERAL TERMS AND CONDITIONS")
    d.add_paragraph("Measurement shall be in Sq. Ft of overall surface to be treated by us.")
    d.save(str(path))
    return path


SUNKEN = {
    "title": "ITEM NO: 1 - TECHNICAL SPECIFICATION & METHEDOLOGY FOR SUNKEN AREA WATERPROOFING USING BRONCO SYSTEM",
    "stages": [
        ("SURFACE PREPARATION", ["Cleaning the invented RCC surface. (Client's scope)", "Opening the cracks and filling with BRONCO INVENTMIX."]),
        ("PROTECTIVE COATING", ["Providing & applying 1st coat BRONCO INVENTGUARD – two component acrylic cementitious coating shall be done entirely over the Roof. Allow to dry the coat for 3 to 4 hrs.",
                                "On application of 2nd & final coat BRONCO INVENTGUARD – two component acrylic cementitious coating shall be done on on the surface. Allow to dry the coat for 8 hrs."]),
        ("PROTECTION LAYER", ["Laying invented brick bats of 75 mm over the coating.", "Do water curing at least 3 days."]),
    ],
}  # fmt: skip
TERRACE_A = {
    "title": "ITEM NO: 2 - TECHNICAL SPECIFICATION FOR TERRACE AREA WATERPROOFING WORKS BY USING BRONCO SYSTEM:",
    "stages": [
        ("SURFACE PREPARATION", ["Cleaning the invented roof."]),
        ("PROTECTIVE COATING", ["Applying BRONCO INVENTPU with dilution of water (1:1) as primer at a coverage of 5 – 6 sq. mtr. per Kg.",
                                "The 1st coat of BRONCO INVENTPU shall be applied over the primer by brush.",
                                "Laying 40 GSM non-woven geotextile / 50 GSM fiberglass mesh reinforcement.",
                                "The 2nd coat (final coat) of BRONCO INVENTPU shall be applied within 24 hrs, overall consumption of material is 1.4 kg in 2-coat."]),
    ],
}  # fmt: skip
TERRACE_B = {
    "title": "ITEM NO: 1 - TECHNICAL SPECIFICATION FOR TERRACE AREA WATERPROOFING WORK BY USING BRONCO COATING SYSTEMS",
    "stages": [
        ("SURFACE PREPARATION", ["Cleaning the invented roof."]),
        ("PROTECTIVE COATING", ["Applying BRONCO INVENTPU with dilution of water (1:1) as primer.", "Laying 45gsm non-woven fabric / fiber glass mesh reinforcement.",
                                "The 1st coat of BRONCO INVENTPU shall be applied over the primer.", "The 2nd coat (final coat) of BRONCO INVENTPU, overall consumption of material is 1.2 kg."]),
    ],
}  # fmt: skip


def test_seeding_from_an_invented_offer_fixes_errors_keeps_unknowns_to_confirm(db, tmp_path):
    a = offer_docx(tmp_path / "Invented Offer @ Invented One.docx", [SUNKEN, TERRACE_A])
    b = offer_docx(tmp_path / "Invented Offer @ Invented Two.docx", [TERRACE_B])
    report = seed.seed(db, [a, b], data_root=tmp_path / "no-datasheets")
    db.commit()
    assert report["items"] == 3 and report["systems"] == 2 and report["merged"] == 1
    whys = {f["why"] for f in report["fixes"]}
    assert "spelling" in whys and "repeated word" in whys  # METHEDOLOGY, "on on"
    roof = [f for f in report["fixes"] if "over the Roof" in f["why"]]
    assert roof and "over the surface" in roof[0]["after"]  # sunken is not a roof
    sunk = db.scalar(select(MethodSystem).where(MethodSystem.elements.contains(["sunken"])))
    steps = [x for s in sunk.stages for x in s["steps"]]
    assert [s["heading"] for s in sunk.stages] == [
        "SURFACE PREPARATION",
        "PROTECTIVE COATING",
        "PROTECTION LAYER",
    ]
    assert (
        steps[0]["scope"] == "client"
        and steps[2]["wait"] == "3 to 4 hrs"
        and steps[2]["product"] == "BRONCO INVENTGUARD"
    )
    assert all("over the Roof" not in x["text"] for x in steps) and "on on" not in steps[3]["text"]
    assert sunk.guarantee_years is None and any(
        c["field"] == "guarantee_years" and c["value"] == "to confirm" for c in sunk.to_confirm
    )
    assert not steps[2].get("consumption")  # nothing read, nothing invented
    p = db.scalar(select(MethodProduct).where(MethodProduct.name == "BRONCO INVENTGUARD"))
    assert p.values == {} and p.datasheet_name is None
    terrace = db.scalar(select(MethodSystem).where(MethodSystem.elements.contains(["terrace"])))
    labels = {c["label"]: c["note"] for c in terrace.to_confirm}
    assert (
        "40, 50 GSM" in labels["Mesh / fabric reinforcement"]
        and "45 GSM" in labels["Mesh / fabric reinforcement"]
    )
    assert "1.4" in labels["Consumption"] and "1.2" in labels["Consumption"]
    assert (
        len(terrace.sources) == 2 and terrace.boq_description == "Providing invented waterproofing"
    )


def test_library_edits_need_a_reason_and_are_versioned(boss, db, tmp_path):
    c, h = boss
    seed.seed(
        db, [offer_docx(tmp_path / "Invented @ Invented One.docx", [SUNKEN])], data_root=tmp_path
    )
    db.commit()
    s = db.scalar(select(MethodSystem))
    assert (
        c.put(
            f"/api/methodology/systems/{s.id}", json={"change": {"guarantee_years": 10}}, headers=h
        ).status_code
        == 422
    )
    r = c.put(
        f"/api/methodology/systems/{s.id}",
        json={"change": {"guarantee_years": 10}, "reason": "Our guarantee letter says 10 years"},
        headers=h,
    )
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["version"] == 2 and out["guarantee_years"] == 10
    assert out["changes"][0]["reason"] == "Our guarantee letter says 10 years" and out["changes"][
        0
    ]["changes"] == [{"field": "guarantee_years"}]
    assert not any(x["field"] == "guarantee_years" for x in out["to_confirm_list"])


# --- invented SOPs ---------------------------------------------------------------------------------

CSS = "<style>table{border-collapse:collapse} td{border:1px solid #000;padding:3px} body{font:10pt sans-serif} .pb{page-break-after:always}</style>"


def _pdf(path: Path, html: str) -> Path:
    from weasyprint import HTML

    path.write_bytes(HTML(string=f"<html><head>{CSS}</head><body>{html}</body></html>").write_pdf())
    return path


def old_sop(path: Path) -> Path:
    return _pdf(
        path,
        """
<table><tr><td></td><td>INVENTED PROPERTIES LIMITED</td><td>SOP No.: INV-WP-RWCONFINED-003</td></tr><tr><td></td><td>STANDARD OPERATING PROCEDURE</td><td>Rev No.: 2</td></tr>
<tr><td></td><td>RETAINING WALL CONFINED WATERPROOFING</td><td>Date: 01-01-2021</td></tr></table>
<p>• Element : Retaining wall</p><p>• Condition: 1. Confined Area</p><p>• Options Evaluated:</p>
<table><tr><td></td><td>1.2 mm HDPE Membrane</td><td>Crystalline</td></tr><tr><td>Recommended</td><td>Yes</td><td>No</td></tr><tr><td>Reason</td><td>Bonds to the concrete poured on it</td><td>Inflexible</td></tr></table>
<p>• Recommended System for all regions: HDPE Membrane (1.2 mm thick).</p><p>Sketch of Approved System:-</p>
<div class="pb"></div>
<p>Proposed Recommended Makes:</p>
<table><tr><td>Recommended System</td><td>Vendors</td><td></td><td></td></tr><tr><td></td><td>BASF</td><td>FOSROC</td><td>SIKA</td></tr>
<tr><td>Pre-Applied HDPE Membrane, min 1.2 mm thick</td><td>INVENTSEAL 731</td><td>INVENT HDPE-P (approved as per TPT)</td><td>INVENTPROOF T-10 (not approved)</td></tr></table>
<p>Method Statement: Step 1:- Lay the membrane with an overlap of 70-80 mm between two sheets.</p>
<p>This overlap between retaining wall membrane and raft membrane should be minimum 100 mm.</p>
<p>Material Specification:</p><p>1.2 mm HDPE Membrane</p>
<table><tr><td></td><td>Standard/Specification</td><td>Expected value</td></tr><tr><td>Elongation</td><td>ASTM D 412</td><td>&gt;300 %</td></tr><tr><td>Tensile Strength</td><td>ASTM D 412</td><td>Min. 25 MPa</td></tr></table>
<p>Sampling Frequency Once per site if less than 1000 m2. Mode of testing Third party.</p>
<div class="pb"></div>
<p>Check List</p><p>1. Installation</p><p>Description Yes No</p><p>Is the surface clean and free from any loose material?</p><p>Is there overlap of 70-80 mm between two sheets of HDPE membrane?</p>
<p>2. Post Concreting</p><p>Description Yes No</p><p>Any visible defects in concrete?</p>
<p>BOQ Description:</p><p>Supplying and installing 1.2mm HDPE membrane, providing guarantee of water proofing work on stamp paper in approved format for a period of 10 years from completion.</p>
<p>Mode of measurement: Measurement as per drawing dimensions shall be paid.</p>""",
    )


def new_sop(
    path: Path, no: str, title: str, systems: list[tuple[str, str]], extra: str = ""
) -> Path:
    subs = "".join(
        f"<p>1.1.{i} {name}</p><p>• {text}</p>" for i, (name, text) in enumerate(systems, 1)
    )
    return _pdf(
        path,
        f"""
<div class="pb"><p>STANDARD OPERATING PROCEDURE</p><p>{title}</p><p>{no}</p></div>
<p>{no} | Issue 01 (1st Jan. 2026) | Rev 00</p><p>1. Scope</p><p>The invented scope.</p><p>1.1 Approved System</p>{subs}<p>2. Responsibility</p>
<div class="pb"></div><p>{no} | Issue 01 (1st Jan. 2026) | Rev 00</p>
<p>Table 2: Performance Criteria for Acrylic Cementitious Coating</p>
<table><tr><td>Parameter</td><td>Tensile Strength</td><td>Elongation</td></tr><tr><td>Reference standard</td><td>ASTM D 412</td><td>ASTM D 412</td></tr><tr><td>Minimum acceptable value</td><td>min. 1.5 MPa</td><td>min. 120%</td></tr></table>
<div class="pb"></div><p>{no} | Issue 01 (1st Jan. 2026) | Rev 00</p>
<p>8. Procedure</p><p>8.1 Surface Preparation</p><p>• Pre-ponding to be done for at least 24 hrs to analyse the points of leakage.</p>
<p>8.2 Application of Coatings</p>{extra}<p>8.3 2nd Post-Ponding Water Test</p><p>• Complete the water-ponding on the floor with 75 mm of water for 3 days (72 hours).</p>
<p>9. Environment, Health, and Safety</p>""",
    )


SMALL = "Acrylic cementitious coating with a minimum elongation of 150%, at a total consumption rate of 2 kg/m² and minimum Dry Film Thickness (DFT) 1 mm, reinforced with a 45 GSM glass fibre mesh for floors and walls."
LARGE = "1.5mm Dry Film Thickness (DFT) 2 component polyurethane on the floor and acrylic cementitious coating with min. 150% elongation on walls. This polyurethane (PU) coat will terminate on the wall at 100 mm above SFL."


@pytest.fixture
def standard(db, tmp_path) -> ClientStandard:
    folder = tmp_path / "invented-sops"
    folder.mkdir()
    old_sop(folder / "03_Invented retaining wall confined.pdf")
    new_sop(folder / "INV-QUALITY-SOP-WATERPROOFING-096 TOILET UNDERSLUNG.pdf", "INV-QUALITY-SOP-WATERPROOFING-096", "WATERPROOFING OF WET AREAS - TOILET UNDERSLUNG",
            [("Small Toilet (for residential toilet/ less than 4.5 m2)", SMALL), ("Large Toilet (for common bathroom/ more than 4.5 m2)", LARGE)],
            "<p>• The polyurethane coat will terminate on the wall at 50 mm above the slab level.</p>")  # fmt: skip
    new_sop(folder / "INV-QUALITY-SOP-WATERPROOFING-097 REFUGE.pdf", "INV-QUALITY-SOP-WATERPROOFING-097", "WATERPROOFING OF WET AREAS - REFUGE AND SERVICE FLOOR",
            [("Open to Sky Refuge", "1.5 mm Dry Film Thickness (DFT) 2 Component Polyurethane (PU) coating on floor and wall."), ("Covered Refuge, Service and Fire Check Floor", SMALL)])  # fmt: skip
    std = sop.import_standard(db, folder, "Invented client standard", "Invented Client")
    db.commit()
    return std


def test_invented_old_and_new_format_sops_import_into_rules(db, standard):
    rep = standard.report
    assert rep["read"] == 3 and rep["formats"] == {"old": 1, "new": 2}
    rules = {(r.element, r.condition_label): r for r in standard.rules}
    rw = rules[("retaining_wall", "Confined Area")]
    assert rw.system_type == "hdpe" and rw.condition == {"confined": True}
    assert (
        rw.document.sop_no == "INV-WP-RWCONFINED-003"
        and rw.document.revision == "2"
        and rw.page == 1
    )
    assert {(m["brand"], m["status"]) for m in rw.makes} == {
        ("BASF", "approved"),
        ("Fosroc", "tpt"),
        ("Sika", "not_approved"),
    }
    assert any(
        s["property"] == "elongation" and s["num"] == 300 and s["op"] == "min" for s in rw.min_specs
    )
    assert rw.guarantee_years == 10 and [c["stage"] for c in rw.checklist] == [
        "Installation",
        "Post Concreting",
    ]
    assert rw.options_evaluated[0]["recommended"] and "Bonds" in rw.options_evaluated[0]["reason"]
    small = rules[("toilet", "Small Toilet (for residential toilet/ less than 4.5 m2)")]
    assert (
        small.condition == {"area_sqm": {"lt": 4.5}, "toilet_kind": "underslung"}
        and small.system_type == "acrylic_cementitious"
    )
    assert (
        small.document.sop_no == "INV-QUALITY-SOP-WATERPROOFING-096"
        and small.document.revision == "I01 R00"
        and small.page == 2
    )
    assert {(t["kind"], t["hours"]) for t in small.tests} >= {("pre_pond", 24), ("pond", 72)}
    assert (
        rules[("toilet", "Large Toilet (for common bathroom/ more than 4.5 m2)")].system_type
        == "pu"
    )
    assert rules[("refuge", "Open to Sky Refuge")].condition == {"exposure": "open"}
    topics = {
        c.topic
        for c in db.scalars(
            select(StandardContradiction).where(StandardContradiction.standard_id == standard.id)
        )
    }
    assert (
        "Overlap of the membrane" in topics and "Termination of the PU coat on the wall" in topics
    )
    assert any(
        t.startswith("Elongation") for t in topics
    )  # 150 % in the system, 120 % in the table
    assert standard.approved_makes == ["BASF", "Fosroc", "Sika"]


# --- the methodology ---------------------------------------------------------------------------------


def _systems(db):
    acr = MethodProduct(
        name="INVENTED ACRYL COAT",
        brand="Bronco",
        kind="coating",
        values={
            "elongation": {
                "value": "min. 150%",
                "num": 150,
                "unit": "%",
                "op": "min",
                "source": "invented TDS",
            }
        },
    )
    db.add(acr)
    coat = [{"text": "1st coat of INVENTED ACRYL COAT.", "product": "INVENTED ACRYL COAT", "products": ["INVENTED ACRYL COAT"], "scope": "ours", "to_confirm": []},
            {"text": "2nd coat of INVENTED ACRYL COAT.", "product": "INVENTED ACRYL COAT", "products": ["INVENTED ACRYL COAT"], "scope": "ours", "to_confirm": []}]  # fmt: skip
    a = MethodSystem(code="T-ACR", name="Invented acrylic system", brand="Bronco", system_type="acrylic_cementitious", elements=["sunken", "toilet", "balcony"],
                     stages=[{"heading": "SURFACE PREPARATION", "steps": [{"text": "Clean the slab.", "scope": "client", "products": [], "to_confirm": []}]},
                             {"heading": "PROTECTIVE COATING", "steps": coat},
                             {"heading": "PROTECTION LAYER", "steps": [{"text": "Brick bat coba 75 mm.", "scope": "ours", "products": [], "to_confirm": []}]}],
                     checklist=[{"stage": "PROTECTIVE COATING", "items": ["Is the coating applied in two coats?"]}],
                     specs={"dft": {"value": "0.8 mm", "num": 0.8, "unit": "mm", "source": "invented TDS"}, "consumption": {"value": "2 kg/m²", "num": 2, "unit": "kg/m²"}})  # fmt: skip
    p = MethodSystem(code="T-PU", name="Invented PU system", brand="Bronco", system_type="pu", elements=["terrace", "refuge"],
                     stages=[{"heading": "PROTECTIVE COATING", "steps": [{"text": "2 coats of INVENTED PU.", "product": "INVENTED PU", "products": ["INVENTED PU"], "scope": "ours", "to_confirm": []}]}])  # fmt: skip
    db.add_all([a, p])
    db.commit()
    return a, p


def _tender(
    db,
    standard,
    owner=None,
    lines=(
        ("Toilets", "Toilet waterproofing, underslung toilets"),
        ("Refuge", "Refuge floor waterproofing, open to sky"),
    ),
):
    from app.masters.models import Client

    cl = Client(name="Invented Client Builders", city="Ahmedabad", state="Gujarat")
    db.add(cl)
    db.flush()
    n = db.scalar(select(Tender.id).order_by(Tender.id.desc()).limit(1)) or 0
    t = Tender(
        code=f"T-2026-97{n + 1:02d}",
        name="Invented methodology tender",
        status="draft",
        client_id=cl.id,
        client_standard_id=standard.id,
        owner_id=owner,
    )
    db.add(t)
    db.flush()
    for i, (sec, text) in enumerate(lines, 1):
        s = BoqSection(tender_id=t.id, title=sec, sort_order=i)
        db.add(s)
        db.flush()
        db.add(
            BoqLine(
                tender_id=t.id,
                section_id=s.id,
                sort_order=1,
                client_item_no=str(i),
                description=f"{sec} — {text}",
                unit="sqm",
                qty=Decimal("40"),
            )
        )
    db.commit()
    return t


def test_the_rule_is_chosen_by_toilet_size_and_refuge_exposure(boss, db, standard):
    places = bld.places_from_texts(
        [[0, 0, 100, "0", "TOILET 1480 X 2135"], [9000, 0, 100, "0", "TOILET 2135 X 2440"]],
        0.001,
        "invented.dxf",
    )
    assert sorted(p["area_sqm"] for p in places) == [3.16, 5.21]
    _systems(db)
    c, h = boss
    t = _tender(db, standard)
    m = c.post(
        "/api/methodology/methodologies", json={"tender_id": t.id, "brand": "Bronco"}, headers=h
    ).json()
    mo = db.get(Methodology, m["id"])
    toilet = next(i for i in mo.items if i.element == "toilet")
    bld.set_places(toilet, places)
    items = bld.split_toilets(db, mo, toilet)
    for it in items:
        db.flush()
        bld.choose(db, mo, it, standard)
    db.commit()
    small, large = items
    assert small.places[0]["area_sqm"] == 3.16 and "Small Toilet" in small.rule_note
    assert large.places[0]["area_sqm"] == 5.21 and "Large Toilet" in large.rule_note
    refuge = next(i for i in mo.items if i.element == "refuge")
    assert (
        "Open to Sky Refuge" in refuge.rule_note
        and db.get(MethodSystem, refuge.system_id).system_type == "pu"
    )


def test_compliance_red_with_the_gap_make_choices_and_to_confirm(boss, db, standard):
    _systems(db)
    c, h = boss
    t = _tender(db, standard, lines=(("Toilets", "Toilet waterproofing, underslung toilets"),))
    m = c.post(
        "/api/methodology/methodologies",
        json={"tender_id": t.id, "brand": "Bronco", "by_hand": {"Toilets": {"area_sqm": 3.16}}},
        headers=h,
    ).json()
    it = m["items"][0]
    rows = {r["topic"]: r for r in it["compliance"]["rows"]}
    assert (
        rows["Dry film thickness (DFT)"]["status"] == "red"
        and rows["Dry film thickness (DFT)"]["gap"] == "0.2 mm short (0.8 vs 1)"
    )
    assert rows["Elongation"]["status"] == "green" and rows["Elongation"]["required"].startswith(
        "min. 150%"
    )  # the stricter of 150 / 120
    assert (
        rows["Tensile strength"]["status"] == "confirm"
        and rows["Tensile strength"]["ours"] == "to confirm"
    )
    make = rows["Make"]
    assert make["status"] == "red" and [ch["key"] for ch in make["choices"]] == ["approved", "keep"]
    assert (
        make["choices"][1]["query"]
        == "Request approval of INVENTED ACRYL COAT (Bronco) as an equivalent make (TPT)"
    )
    assert rows["Mesh / fabric"]["status"] == "red" and "45 GSM" in rows["Mesh / fabric"]["gap"]
    assert any(q["kind"] == "make" for q in m["prebid"]) and any(
        q["kind"] == "contradiction" for q in m["prebid"]
    )
    # the estimator edits a step: a reason is needed, the version moves on
    r = c.post(
        f"/api/methodology/items/{it['id']}/steps",
        json={"stage": 1, "step": None, "text": "Embed a 45 GSM glass fibre mesh in the 1st coat."},
        headers=h,
    )
    assert r.status_code == 422
    r = c.post(
        f"/api/methodology/items/{it['id']}/steps",
        json={
            "stage": 1,
            "step": 1,
            "text": "Embed a 45 GSM glass fibre mesh in the 1st coat.",
            "reason": "The client asks for the mesh",
        },
        headers=h,
    )
    assert r.status_code == 200, r.text
    assert (
        r.json()["version"] == m["version"] + 1
        and r.json()["items"][0]["edits"][0]["reason"] == "The client asks for the mesh"
    )
    rows = {
        x["topic"]: x
        for x in c.post(f"/api/methodology/methodologies/{m['id']}/recheck", headers=h).json()[
            "items"
        ][0]["compliance"]["rows"]
    }
    assert rows["Mesh / fabric"]["status"] == "green"


def _texts(data: bytes) -> list[str]:
    d = Document(io.BytesIO(data))
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            out.append(" | ".join(c.text for c in r.cells))
    return out


def test_the_word_file_has_stages_steps_spec_table_checklist_and_the_client_copy_no_internals(
    boss, db, standard
):
    from app.methodology import export

    _systems(db)
    c, h = boss
    t = _tender(db, standard, lines=(("Toilets", "Toilet waterproofing, underslung toilets"),))
    m = c.post(
        "/api/methodology/methodologies",
        json={"tender_id": t.id, "brand": "Bronco", "by_hand": {"Toilets": {"area_sqm": 3.16}}},
        headers=h,
    ).json()
    mo = db.get(Methodology, m["id"])
    client = _texts(export.docx(db, mo, internal=False))
    text = "\n".join(client)
    assert (
        "ITEM NO: 1 - TECHNICAL SPECIFICATION & METHODOLOGY FOR TOILETS USING BRONCO SYSTEM" in text
    )
    assert (
        "SURFACE PREPARATION:" in client
        and "PROTECTIVE COATING:" in client
        and "TESTING:" in client
    )
    assert any(x.startswith("1.\tClean the slab.") for x in client) and any(
        x.startswith("2.\t1st coat of INVENTED ACRYL COAT") for x in client
    )
    assert (
        "Property | Standard | Required | Ours" in text
        and "Dry film thickness (DFT) |  | min. 1 mm | 0.8 mm" in text
    )
    assert (
        "CHECKLIST:" in client
        and "Description | Yes | No" in text
        and "GUARANTEE: 10 years" in text
        and "CLIENT'S SCOPE:" in client
    )
    assert (
        "COMPLIANCE" not in text
        and "PRE-BID" not in text
        and "Gap" not in text
        and "to confirm" not in text.lower()
        and "Why:" not in text
    )
    internal = "\n".join(_texts(export.docx(db, mo, internal=True)))
    assert (
        "COMPLIANCE (INTERNAL)" in internal
        and "PRE-BID QUERIES (INTERNAL)" in internal
        and "Request approval of INVENTED ACRYL COAT" in internal
    )
    files = c.post(
        f"/api/methodology/methodologies/{m['id']}/export", json={"audience": "client"}, headers=h
    ).json()
    assert {f["kind"] for f in files} == {"docx", "pdf"} and all(
        f["file_name"].startswith(m["code"]) for f in files
    )


def test_a_won_boq_gives_the_site_stage_templates_blocking_hold_points_and_pond_tests(
    boss, db, standard
):
    _systems(db)
    c, h = boss
    t = _tender(db, standard, lines=(("Toilets", "Toilet waterproofing, underslung toilets"),))
    m = c.post(
        "/api/methodology/methodologies",
        json={"tender_id": t.id, "brand": "Bronco", "by_hand": {"Toilets": {"area_sqm": 3.16}}},
        headers=h,
    ).json()
    assert c.patch(f"/api/tenders/{t.id}", json={"status": "won"}, headers=h).status_code == 200
    site = c.post(
        "/api/sites/from-tender", json={"tender_id": t.id, "start_date": "2026-11-02"}, headers=h
    ).json()
    db.expire_all()
    item = db.get(Methodology, m["id"]).items[0]
    tpl = db.get(StageTemplate, item.stage_template_id)
    steps = {s.name: s for s in tpl.steps}
    # the pond tests on the coating, before the protection covers it
    assert list(steps) == [
        "Surface Preparation",
        "Protective Coating",
        "Testing",
        "Protection Layer",
    ]
    assert steps["Protective Coating"].checklist_template_id  # the client's checklist as the form
    assert steps["Testing"].hold_point and not steps["Protection Layer"].hold_point
    ponds = list(db.scalars(select(PondTest).where(PondTest.site_id == site["id"])))
    assert {(p.kind, p.required_hours) for p in ponds} >= {("pre_pond", 24), ("pond", 72)}
    # the hold point blocks the next stage until it is certified
    node = SiteNode(site_id=site["id"], kind="tower", name="Invented tower")
    db.add(node)
    db.flush()
    line = db.scalar(select(BoqLine).where(BoqLine.tender_id == t.id))
    db.add(
        AreaScope(
            site_id=site["id"],
            node_id=node.id,
            boq_line_id=line.id,
            stage_template_id=tpl.id,
            qty=Decimal("40"),
            unit="sqm",
        )
    )
    db.commit()
    assert c.post(f"/api/sites/{site['id']}/tasks/generate", headers=h).status_code in (200, 201)
    tasks = {x.name: x for x in db.scalars(select(Task).where(Task.site_id == site["id"]))}
    for name in ("Surface Preparation",):
        assert (
            c.patch(
                f"/api/sites/{site['id']}/tasks/{tasks[name].id}",
                json={"status": "in_progress"},
                headers=h,
            ).status_code
            == 200
        )
    db.query(Task).filter(
        Task.id.in_([tasks[x].id for x in ("Surface Preparation", "Protective Coating", "Testing")])
    ).update({"status": "done"}, synchronize_session=False)
    db.commit()
    r = c.patch(
        f"/api/sites/{site['id']}/tasks/{tasks['Protection Layer'].id}",
        json={"status": "in_progress"},
        headers=h,
    )
    assert r.status_code == 409 and "'Testing' is a hold point" in r.json()["detail"]
    # the pond test needs its hours and photos to pass
    p = next(x for x in ponds if x.kind == "pond")
    r = c.put(
        f"/api/methodology/pond-tests/{p.id}",
        json={
            "started_at": "2026-11-10T08:00:00+05:30",
            "ended_at": "2026-11-11T08:00:00+05:30",
            "result": "pass",
        },
        headers=h,
    )
    assert r.status_code == 422 and "72 hours" in r.json()["detail"]
    # the client's copy is shared in the portal
    from app.portal.models import SiteDocument

    assert db.scalar(
        select(SiteDocument).where(
            SiteDocument.site_id == site["id"], SiteDocument.share_with_client
        )
    ).filename.endswith(".pdf")


def test_who_sees_what(boss, db, standard, login_as, make_user):
    _systems(db)
    c, h = boss
    sc, sh = login_as("sales", email="sales-m@example.com")
    from app.models import User

    sales = db.scalar(select(User).where(User.email == "sales-m@example.com"))
    own = _tender(db, standard, owner=sales.id, lines=(("Toilets", "Toilet waterproofing"),))
    other = _tender(db, standard, lines=(("Refuge", "Refuge open to sky"),))
    m_own = c.post("/api/methodology/methodologies", json={"tender_id": own.id}, headers=h).json()
    m_other = c.post(
        "/api/methodology/methodologies", json={"tender_id": other.id}, headers=h
    ).json()
    got = sc.get(f"/api/methodology/methodologies/{m_own['id']}", headers=sh)
    assert (
        got.status_code == 200
        and "compliance" not in got.json()["items"][0]
        and "prebid" not in got.json()
    )
    assert sc.get(f"/api/methodology/methodologies/{m_other['id']}", headers=sh).status_code == 404
    assert (
        sc.post(
            f"/api/methodology/methodologies/{m_own['id']}/export",
            json={"audience": "internal"},
            headers=sh,
        ).status_code
        == 403
    )
    assert (
        sc.post(
            f"/api/methodology/methodologies/{m_own['id']}/export",
            json={"audience": "client"},
            headers=sh,
        ).status_code
        == 200
    )
    assert (
        sc.post(
            "/api/methodology/methodologies", json={"tender_id": own.id}, headers=sh
        ).status_code
        == 403
    )
    for role in ("planning", "estimator", "director"):
        rc, rh = login_as(role, email=f"{role}-m@example.com")
        assert (
            "compliance"
            in rc.get(f"/api/methodology/methodologies/{m_other['id']}", headers=rh).json()[
                "items"
            ][0]
        )
    ec, eh = login_as("site_engineer", email="se-m@example.com")
    assert ec.get("/api/methodology/systems", headers=eh).status_code == 403
    site = Site(code="S-2026-9971", name="Invented methodology site", status="active")
    db.add(site)
    db.flush()
    db.add(
        SiteMember(
            site_id=site.id,
            user_id=db.scalar(select(User.id).where(User.email == "se-m@example.com")),
            role_on_site="incharge",
        )
    )
    db.commit()
    assert ec.get(f"/api/methodology/sites/{site.id}", headers=eh).status_code == 200
    other_site = Site(code="S-2026-9972", name="Not theirs", status="active")
    db.add(other_site)
    db.commit()
    assert ec.get(f"/api/methodology/sites/{other_site.id}", headers=eh).status_code == 404
    cc, ch = login_as("client", email="client-m@example.com")
    assert cc.get(f"/api/methodology/methodologies/{m_own['id']}", headers=ch).status_code == 403
