"""Writes invented client BOQ files for the tender import tests (no company data).
Run from backend/:  python tests/fixtures/make_boq_fixtures.py

client_boq.xlsx, sheet "BOQ" (after a "Cover" sheet):
- title rows, a blank row, and the header on row 4 (not row 1)
- a title merged across the sheet, and a section heading merged across B:G
- numbered sections (1, 2, 5) with lettered sub-items, and a main item with its own quantities
- a description that runs on into a second, unnumbered row
- QRO in the qty column, "Rate only" in the qty column, "By civil contractor" (NQ)
- a sub-total row and a grand total row, a fully blank row
- odd units: "Sq.Mt", "Nos.", "R.Mtr" (recognised) and "Pair" (not recognised)
- a Terms & Conditions block after the items

client_boq.csv: the same kind of content with the header on row 1 and other column names.
"""

import csv
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent

HEADER = ["Sr. No.", "Description of Item", "Unit", "Qty", "Rate (Rs.)", "Amount (Rs.)", "Remarks"]

ROWS = [
    ["1", "Basement waterproofing", None, None, None, None, None],
    ["a)", "Crystalline coating to raft with two coats", "Sq.Mt", 120.5, 450, 54225, None],
    ["b)", "Pipe sleeve sealing with polymer modified mortar", "Nos.", 12, 350, 4200, None],
    [None, None, None, None, None, None, None],
    ["2", "Terrace waterproofing", None, None, None, None, None],
    ["a)", "APP membrane 4 mm thick laid over terrace slab,", "sqm", 300, 520, 156000, None],
    [None, "including primer and 100 mm laps, complete.", None, None, None, None, None],
    ["b)", "Protective screed 50 mm thick over membrane", "Sq.Mt", "QRO", 260, None, None],
    ["c)", "Expansion joint sealing with PU sealant", "R.Mtr", "Rate only", None, None, None],
    [
        "d)",
        "Brick bat coba with average 100 mm thickness",
        "sqm",
        "By civil contractor",
        None,
        None,
        "By others",
    ],  # fmt: skip
    [None, "Sub Total of Terrace", None, None, None, 156000, None],
    ["3", "Lift pit injection grouting", "Pair", 2, 1500.004, 3000.01, None],
    ["5", "External areas", None, None, None, None, None],  # merged B:G below
    ["a)", "Planter box waterproofing with two coats", "sqm", 40.1234, 600, 24074, "Approved make"],
    [None, "Grand Total", None, None, None, 241499.01, None],
    [None, None, None, None, None, None, None],
    [None, "Terms & Conditions", None, None, None, None, None],
    ["1", "Water and power by the client", None, None, None, None, None],
]


def write_xlsx(path: Path) -> None:
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["A1"] = "Invented Towers — tender documents"
    ws = wb.create_sheet("BOQ")
    ws["A1"] = "Name of work: Waterproofing for Invented Towers"
    ws.merge_cells("A1:G1")
    ws["A2"] = "Client: Example Developers"
    ws.append([])
    ws.append(HEADER)
    for row in ROWS:
        ws.append(row)
    heading = 4 + ROWS.index(["5", "External areas", None, None, None, None, None]) + 1
    ws.merge_cells(f"B{heading}:G{heading}")
    wb.save(path)


def write_csv(path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Item No", "Specification", "UOM", "Quantity", "Unit Rate", "Total Amount"])
        w.writerow(["A", "Toilet waterproofing", "", "", "", ""])
        w.writerow(["A.1", "Acrylic polymer coating two coats in sunken slab", "sqm", "85", "",
                    ""])  # fmt: skip
        w.writerow(["A.2", "Pipe sleeve packing up to 110 mm", "nos", "16", "", ""])
        w.writerow(["", "", "", "", "", ""])
        w.writerow(["B", "Water tank waterproofing", "sqm", "RO", "", ""])
        w.writerow(["", "Total", "", "", "", "0"])


if __name__ == "__main__":
    write_xlsx(HERE / "client_boq.xlsx")
    write_csv(HERE / "client_boq.csv")
    print("written")
