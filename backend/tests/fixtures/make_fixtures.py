"""Writes rate_library_sample.xlsx: a tiny invented workbook in the same layout as the real
rate library export (no company data). Run from backend/:  python tests/fixtures/make_fixtures.py

What it exercises:
- an item whose description appears with different spacing/case in a line (still links)
- a sub-item that only links as "<parent item> — <description>"
- a unit spelled differently in the item ("Sq.Mt") and the line ("Sqmt")
- a non-numeric quantity ("NQ") that must land in qty_note
- a line with no matching item (stays unlinked)
- a duplicate source position and a blank description (both skipped, with reasons)
- a "Margin — ..." working row and a below-₹1 item (both excluded, not deleted)
- an item that only has other bidders' rates (competitor), and an extra competitor line on a
  normal item (left out of its stats)
- the same pipe outlet item quoted per nos and per sqm (for merge / unmerge)
"""

from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent

APP = "Providing and applying APP modified bitumen membrane 4 mm thick over terrace slab"
PIPE_PARENT = (
    "Sealing / packing the joint around the PVC pipe opening with polymer modified mortar"
)
PIPE = f"{PIPE_PARENT} — 110mm dia pipe"
CRYSTAL = "Integral crystalline admixture for RCC raft and retaining walls"
TOILET = "Toilet sunken slab waterproofing with two coat acrylic polymer cementitious coating"
MOBILISE = "Mobilisation and demobilisation of men and machinery"
MARGIN = "Margin — Pipe sleeve 110 mm"
HACKING = "Hacking and cleaning of old plaster surface"
COMPETITOR = "Crystalline coating to water tank walls with two coats"
OUTLET = "Treatment around pipe outlet with polymer modified mortar"
COMPARATIVE = "comparative sheet (may be other bidders)"

ITEMS_HEADER = [
    "Description", "Unit", "No. of BOQs", "Latest rate (₹)", "Min", "Median", "Max",
    "Latest client folder", "Latest source file", "Product / make", "Remarks",
    "From EESPL-priced file", "Check",
]  # fmt: skip
ITEMS = [
    [APP, "Sqm", 6, 450, 380, 420, 480, "CLIENT A", "a.xlsx", "Texsa", None, "Yes", None],
    [PIPE, "Nos", 9, 425, 300, 350, 425, "CLIENT B", "b.xlsx", None, None, "Yes", None],
    [CRYSTAL, "Kg", 4, 210.5, 180, 200, 430, "CLIENT A", "a.xlsx", None, None, "Yes",
     "rates vary more than 2x"],
    [TOILET, "Sq.Mt", 3, 260, 240, 255, 260, "CLIENT C", "c.xlsx", None, None, "No", None],
    [MOBILISE, "(blank)", 1, 15000, 15000, 15000, 15000, "CLIENT D", "d.xlsx", None, None,
     "Yes", None],
    [MARGIN, "(blank)", 1, 0.3, 0.3, 0.3, 0.3, "CLIENT D", "d.xlsx", None, None, "Yes",
     "below ₹1"],
    [HACKING, "Sqm", 2, 0.5, 0.5, 0.5, 0.5, "CLIENT D", "d.xlsx", None, None, "Yes",
     "below ₹1"],
    [COMPETITOR, "Sqm", 1, 150, 150, 150, 150, "CLIENT E", "e.xlsx", None, None, "No",
     COMPARATIVE],
    [OUTLET, "Nos", 2, 300, 250, 275, 300, "CLIENT F", "f.xlsx", None, None, "Yes", None],
    [OUTLET, "Sqm", 1, 1087, 1087, 1087, 1087, "CLIENT H", "h.xlsx", None, None, "Yes", None],
]  # fmt: skip

LINES_HEADER = [
    "Client folder", "File", "Sheet", "Row", "Item No", "Parent item", "Description",
    "Unit (as written)", "Unit", "Qty", "Qty note", "Rate (₹)", "Product / make", "Remarks",
    "EESPL file", "Check",
]  # fmt: skip
LINES = [
    ["CLIENT A", "CLIENT A/a.xlsx", "BOQ", 10, "1", None, APP, "SQM", "Sqm", 1200, None, 450,
     "Texsa", None, "Yes", None],
    ["CLIENT C", "CLIENT C/c.xlsx", "Sheet1", 5, "2.1", None, "  providing and applying APP "
     "modified bitumen membrane  4 mm thick over TERRACE slab ", "Sq.m.", "Sqm", 300.5, None,
     380, None, None, "No", None],
    ["CLIENT B", "CLIENT B/b.xlsx", "WP", 22, "3a", PIPE_PARENT, "110mm dia pipe", "Nos",
     "Nos", None, "QRO", 425, None, None, "Yes", None],
    ["CLIENT A", "CLIENT A/a.xlsx", "BOQ", 11, "2", None, CRYSTAL, "KG", "Kg", "NQ", None,
     210.5, None, None, "Yes", None],
    ["CLIENT D", "CLIENT D/d.xlsx", "BOQ", 3, None, None, "An item that is not in the rate "
     "library", "RMT", "Rmt", 50, None, 99, None, None, "Yes", None],
    ["CLIENT A", "CLIENT A/a.xlsx", "BOQ", 10, "1", None, APP, "SQM", "Sqm", 1200, None, 450,
     None, None, "Yes", None],  # same file/sheet/row as the first line: skipped
    ["CLIENT C", "CLIENT C/c.xlsx", "Sheet1", 9, "4", None, TOILET, "Sqmt", "Sqm", 85, None,
     260, None, None, "No", None],
    ["CLIENT C", "CLIENT C/c.xlsx", "Sheet1", 12, None, None, None, "Nos", "Nos", 1, None, 5,
     None, None, "No", None],  # blank description: skipped
    ["CLIENT D", "CLIENT D/d.xlsx", "Working", 20, None, "Margin", "Pipe sleeve 110 mm", None,
     None, None, None, 0.3, None, None, "Yes", "below ₹1"],
    ["CLIENT E", "CLIENT E/e.xlsx", "Comparative", 4, "1", None, COMPETITOR, "SQM", "Sqm", 500,
     None, 150, None, None, "No", COMPARATIVE],
    ["CLIENT E", "CLIENT E/e.xlsx", "Comparative", 5, "2", None, APP, "SQM", "Sqm", 900, None,
     999, None, None, "No", COMPARATIVE],  # other bidder's rate on a normal item
    ["CLIENT F", "CLIENT F/f.xlsx", "BOQ", 1, "5", None, OUTLET, "Nos", "Nos", 40, None, 300,
     None, None, "Yes", None],
    ["CLIENT G", "CLIENT G/g.xlsx", "BOQ", 2, "6", None, OUTLET, "No.", "Nos", 25, None, 250,
     None, None, "Yes", None],
    ["CLIENT H", "CLIENT H/h.xlsx", "BOQ", 3, "7", None, OUTLET, "Sqm", "Sqm", 12, None, 1087,
     None, None, "Yes", None],
]  # fmt: skip


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rate library"
    ws.append(ITEMS_HEADER)
    for row in ITEMS:
        ws.append(row)
    ws = wb.create_sheet("All lines")
    ws.append(LINES_HEADER)
    for row in LINES:
        ws.append(row)
    ws = wb.create_sheet("Files read")
    ws.append(["File", "Client folder", "Priced lines found", "Sheets used", "Note"])
    wb.save(HERE / "rate_library_sample.xlsx")


if __name__ == "__main__":
    main()
