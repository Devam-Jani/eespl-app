# ruff: noqa: E501, F811
"""Go-live readiness on invented data: the production setup's files, the secrets script, backup
retention and the missed-backup alert, the health page, the "not set" checks, the user import, the
carry-over of masters (dry run, idempotent, no DEMO rows) and the go-live database, the roles
report, help texts, and security (password rules, lockout and unlock, two-step sign-in, sessions
ending on a password change or after inactivity, the audit log's kinds, upload checks)."""

import io
import os
import re
import runpy
import stat
import time
from datetime import UTC, datetime, timedelta
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select, text

from app.auth import totp
from app.config import settings
from app.main import app
from app.models import AuditLog, RefreshToken, SecuritySettings, User
from app.ops import backup, golive, service
from app.portal.models import Notification
from tests.conftest import PASSWORD, login
from tests.test_finance import boss  # noqa: F401  (fixture)

REPO = Path("/repo")
PROD = REPO / "docker-compose.prod.yml"


# --- the production setup ----------------------------------------------------------------------------


@pytest.mark.skipif(not PROD.exists(), reason="the production files are mounted at /repo in the dev stack")
def test_the_production_compose_has_only_80_and_443_no_mailpit_no_dev_server_and_health_checks():
    text_ = PROD.read_text()
    ports = re.findall(r'^\s+- "(\d+):(\d+)"', text_, re.M)
    assert sorted(p[0] for p in ports) == ["443", "80"]
    for bad in ("mailpit", "--reload", "npm run dev", "5173", "5174", "8001", "5434", "lan-https", "./backend:/app"):
        assert bad not in text_.replace('MAILPIT_HOST: ""', ""), bad
    section = text_.split("\nservices:\n", 1)[1].split("\nvolumes:")[0]
    services = re.findall(r"^  ([a-z-]+):\n", section, re.M)
    assert {"db", "api", "worker", "web", "osrm", "caddy", "backup", "api-training", "worker-training"} <= set(services)
    blocks = re.split(r"^  (?=[a-z-]+:\n)", section, flags=re.M)[1:]
    for b in blocks:
        name = b.split(":")[0]
        assert "restart: always" in b, name
        assert "healthcheck:" in b, name
        assert "logging: *logging" in b, name
    assert "max-size" in text_ and "media:/media" in text_
    caddy = (REPO / "deploy/prod/Caddyfile").read_text()
    for header in ("Strict-Transport-Security", "Content-Security-Policy", "X-Frame-Options", "frame-ancestors 'none'"):
        assert header in caddy
    example = (REPO / ".env.prod.example").read_text()
    assert all(line.rstrip().endswith("=") for line in example.splitlines() if "=" in line and not line.startswith("#"))  # keys, no values


@pytest.mark.skipif(not (REPO / "scripts/make-prod-secrets.py").exists(), reason="mounted at /repo in the dev stack")
def test_the_secrets_script_makes_fresh_values_never_the_dev_ones_and_never_prints_them(tmp_path, capsys):
    dev = tmp_path / "dev.env"
    dev.write_text("JWT_SECRET=dev-secret-shared-in-chat-0123456789abcdef\nPOSTGRES_PASSWORD=devpass\nBACKUP_KEY=devkey\n")
    out = tmp_path / "prod.env"
    mod = runpy.run_path(str(REPO / "scripts/make-prod-secrets.py"))
    assert mod["main"](["--example", str(REPO / ".env.prod.example"), "--out", str(out), "--dev-env", str(dev)]) == 0
    printed = capsys.readouterr().out
    vals = dict(line.split("=", 1) for line in out.read_text().splitlines() if "=" in line and not line.startswith("#"))
    for key in ("JWT_SECRET", "POSTGRES_PASSWORD", "BACKUP_KEY", "DATA_KEY", "DEMO_USER_PASSWORD"):
        assert len(vals[key]) >= 16 and vals[key] not in ("dev-secret-shared-in-chat-0123456789abcdef", "devpass", "devkey")
        assert vals[key] not in printed  # never shown
    assert vals["APP_DOMAIN"] == "" and vals["EMAIL_ENABLED"] == "false"
    if os.name != "nt":
        assert stat.S_IMODE(out.stat().st_mode) == 0o600
    with pytest.raises(SystemExit):  # an existing .env is never overwritten
        mod["main"](["--example", str(REPO / ".env.prod.example"), "--out", str(out)])
    before = vals["JWT_SECRET"]
    assert mod["main"](["--out", str(out), "--rotate", "JWT_SECRET"]) == 0
    after = dict(line.split("=", 1) for line in out.read_text().splitlines() if "=" in line and not line.startswith("#"))
    assert after["JWT_SECRET"] != before and after["POSTGRES_PASSWORD"] == vals["POSTGRES_PASSWORD"] and after["JWT_SECRET"] not in capsys.readouterr().out


# --- backups ----------------------------------------------------------------------------------------------


def _names(days: int, start: datetime) -> list[str]:
    return [f"eespl-{(start - timedelta(days=i)):%Y%m%d}-0200.tar.gpg" for i in range(days)]


def test_retention_keeps_14_daily_8_weekly_12_monthly_and_removes_the_rest(tmp_path, monkeypatch):
    newest = datetime(2026, 10, 21, 2, 0)
    names = _names(500, newest)  # every night for 500 nights
    keep, remove = backup.plan(names + ["notes.txt"], 14, 8, 12)
    dates = sorted(backup.archive_date(n).date() for n in keep)
    assert len(keep) + len(remove) == 500 and "notes.txt" not in keep + remove
    assert set(_names(14, newest)) <= set(keep)  # the last 14 nights
    weeks = {d.isocalendar()[:2] for d in dates}
    assert len({w for w in weeks if w >= (newest.date() - timedelta(weeks=7)).isocalendar()[:2]}) == 8
    months = {(d.year, d.month) for d in dates}
    assert len(months) == 12 and min(dates) >= datetime(2025, 11, 1).date()  # nothing older than 12 months
    assert len(keep) < 14 + 8 + 12  # overlapping days count once
    # applied to a folder: only the archives outside the plan go
    folder = tmp_path / "backups"
    folder.mkdir()
    monkeypatch.setenv("BACKUP_DIR", str(folder))
    for n in _names(40, newest):
        (folder / n).write_bytes(b"x")
    kept, removed = backup.prune()
    assert sorted(p.name for p in folder.iterdir()) == sorted(kept) and len(removed) == 40 - len(kept)


def test_a_missed_backup_tells_the_director_and_the_office_admin(db, make_user, monkeypatch):
    monkeypatch.setattr(settings, "backup_enabled", True)
    director = make_user("director-bk@example.com", "director")
    office = make_user("office-bk@example.com", "office_admin")
    now = datetime(2026, 10, 21, 9, 0, tzinfo=UTC)
    service.put(db, "backup", {"ok": True, "archive": "eespl-20261020-0200.tar.gpg"}, at=now - timedelta(hours=7))
    db.commit()
    assert service.missed_backup_sweep(db, now) is False  # last night's is there
    service.put(db, "backup", {"ok": True, "archive": "eespl-20261019-0200.tar.gpg"}, at=now - timedelta(hours=31))
    db.commit()
    assert service.missed_backup_sweep(db, now) is True
    told = {n.user_id for n in db.scalars(select(Notification).where(Notification.rule_code == "backup_missed"))}
    assert {director.id, office.id} <= told
    assert service.missed_backup_sweep(db, now) is False  # once a day
    status = service.backup_status(db)
    assert status["missed"] and status["retention"] == {"daily": 14, "weekly": 8, "monthly": 12}


def test_a_failed_backup_is_missed_too(db, monkeypatch):
    monkeypatch.setattr(settings, "backup_enabled", True)
    service.put(db, "backup", {"ok": False, "error": "CalledProcessError: pg_dump"})
    db.commit()
    assert service.backup_status(db)["missed"]


# --- health and the go-live settings --------------------------------------------------------------------------


def test_the_health_page_reads_the_database_the_worker_and_osrm(db):
    c = TestClient(app)
    r = c.get("/api/health/full")
    assert r.json()["db"] == "ok" and r.json()["worker"] == "never ran" and r.status_code == 503
    service.heartbeat(db)
    out = c.get("/api/health/full").json()
    assert out["worker"] == "ok" and "osrm" in out and "secret" not in str(out).lower()


def test_the_not_set_banner_lists_what_go_live_needs_until_filled(db, login_as):
    c, h = login_as("office_admin", email="oa-golive@example.com")
    missing = {m["key"] for m in c.get("/api/ops/go-live-checks", headers=h).json()["missing"]}
    assert {"company", "gstin", "bank"} <= missing
    from app.masters.models import CompanyBankAccount, CompanyGstin, CompanyProfile

    p = db.get(CompanyProfile, 1) or CompanyProfile(id=1)
    p.legal_name, p.pan = "Invented Applicators Pvt Ltd", "AAACI1234A"
    db.merge(p)
    db.add_all([CompanyGstin(gstin="24AAACI1234A1Z5", state="Gujarat", address="Invented Road, Ahmedabad", is_default=True),
                CompanyBankAccount(account_name="Invented Applicators", account_number="000111222333", ifsc="INVN0000001", is_default=True)])  # fmt: skip
    db.commit()
    missing = {m["key"] for m in c.get("/api/ops/go-live-checks", headers=h).json()["missing"]}
    assert not {"company", "gstin", "bank"} & missing
    cs, hs = login_as("sales", email="sales-golive@example.com")
    assert cs.get("/api/ops/go-live-checks", headers=hs).status_code == 403


# --- users from the Excel template -----------------------------------------------------------------------------


def _sheet(rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Users"
    ws.append(["Name", "Email", "Phone", "Roles", "Manager (email)", "Sites (codes)", "Tracked (yes/no)", "Travel eligible (yes/no)", "Rate per km"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_the_user_import_reports_bad_rows_saves_nothing_until_all_pass_and_adds_nobody_twice(db, login_as):
    c, h = login_as("office_admin", email="oa-import@example.com")
    t = c.get("/api/ops/user-import/template.xlsx", headers=h)
    wb = openpyxl.load_workbook(io.BytesIO(t.content))
    assert wb.sheetnames == ["Users", "Roles", "Sites"] and "sales" in [r[0].value for r in wb["Roles"].iter_rows(min_row=2)]
    bad = _sheet([["Invented Planner One", "planner.one@example.com", "", "planning", "", "", "no", "no", ""],
                  ["Invented Sales One", "sales.one@example.com", "", "salesman", "nobody@example.com", "S-0000-9999", "maybe", "yes", ""],
                  ["Invented Twice", "planner.one@example.com", "", "billing", "", "", "", "", ""]])  # fmt: skip
    r = c.post("/api/ops/user-import/check", files={"file": ("users.xlsx", bad)}, headers=h).json()
    errors = {x["row"]: x["errors"] for x in r["rows"] if x["errors"]}
    assert r["errors"] == 2 and 2 not in errors
    assert any("no role 'salesman'" in e for e in errors[3]) and any("manager" in e for e in errors[3]) and any("no site" in e for e in errors[3])
    assert any("yes or no" in e for e in errors[3]) and any("rate per km" in e for e in errors[3]) and errors[4] == ["this email is in the file twice"]
    r = c.post("/api/ops/user-import/save", files={"file": ("users.xlsx", bad)}, headers=h)
    assert r.status_code == 422 and db.scalar(select(User).where(User.email == "planner.one@example.com")) is None  # nothing saved
    good = _sheet([["Invented Planner One", "planner.one@example.com", "+91 90000 00001", "planning", "", "", "no", "no", ""],
                   ["Invented Sales One", "sales.one@example.com", "", "Sales", "planner.one@example.com", "", "yes", "yes", "3.50"]])  # fmt: skip
    r = c.post("/api/ops/user-import/save", files={"file": ("users.xlsx", good)}, headers=h).json()
    links = {u["email"]: u["link"] for u in r["users"]}
    assert set(links) == {"planner.one@example.com", "sales.one@example.com"} and all(x.startswith("/set-password/") for x in links.values())
    sales = db.scalar(select(User).where(User.email == "sales.one@example.com"))
    assert [x.code for x in sales.roles] == ["sales"] and sales.password_hash is None
    from app.field.models import TrackingProfile, TravelRate

    prof = db.scalar(select(TrackingProfile).where(TrackingProfile.user_id == sales.id))
    assert prof.tracked and prof.reimbursable and prof.manager_id == db.scalar(select(User.id).where(User.email == "planner.one@example.com"))
    assert str(db.scalar(select(TravelRate.rate_per_km).where(TravelRate.user_id == sales.id))) == "3.50"
    again = c.post("/api/ops/user-import/save", files={"file": ("users.xlsx", good)}, headers=h).json()
    assert {u["status"] for u in again["users"]} == {"already a user"} and db.scalar(select(text("count(*)")).select_from(User).where(User.email.like("%.one@example.com"))) == 2
    # the link sets the password once
    token = links["sales.one@example.com"].rsplit("/", 1)[1]
    p = TestClient(app)
    assert p.get(f"/api/auth/password-link/{token}").json()["email"] == "sales.one@example.com"
    assert p.post(f"/api/auth/password-link/{token}", json={"password": "password123"}).status_code == 422  # too common
    assert p.post(f"/api/auth/password-link/{token}", json={"password": "Tide-Lamp-Orchid-42"}).status_code == 204
    assert p.post(f"/api/auth/password-link/{token}", json={"password": "Another-Fresh-Pass-7"}).status_code == 404  # once
    assert login(TestClient(app), "sales.one@example.com", "Tide-Lamp-Orchid-42")
    cs, hs = login_as("billing", email="billing-import@example.com")
    assert cs.post("/api/ops/user-import/check", files={"file": ("users.xlsx", good)}, headers=hs).status_code == 403


# --- the go-live database and the carry-over -----------------------------------------------------------------------


@pytest.fixture(scope="module")
def golive_db():
    from tests.conftest import _test_url

    name = f"{_test_url.database}_golive"
    old_url = settings.database_url
    settings.database_url = _test_url.render_as_string(hide_password=False)
    admin = create_engine(_test_url.set(database="postgres"), isolation_level="AUTOCOMMIT")
    with admin.connect() as conn:
        conn.execute(text(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)'))
    out = golive.run(name)
    yield name, out
    with admin.connect() as conn:
        conn.execute(text(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)'))
    admin.dispose()
    settings.database_url = old_url


def test_the_go_live_command_makes_a_database_with_the_seed_and_no_demo_rows(golive_db):
    name, out = golive_db
    assert out["ok"] and out["made_now"] and out["demo_rows"] == {}
    assert out["seed"]["roles"] >= 12 and out["seed"]["permissions"] >= 120 and out["seed"]["security_settings"] == 1
    assert golive.run(name)["made_now"] is False  # a second run changes nothing


def test_the_carry_over_counts_in_a_dry_run_leaves_demo_out_and_adds_nothing_twice(golive_db, db):
    from app.masters.models import Client, Product, Vendor
    from app.ops import carryover
    from tests.conftest import _test_url

    name, _ = golive_db
    from app.knowledge.models import ProductReview

    db.add_all([Product(code="INV-REAL-1", name="Invented real coating", unit="kg"), Product(code="DEMO-X", name="DEMO coating", unit="kg", is_demo=True),
                Vendor(name="Invented Vendor Pvt Ltd"), Client(name="Invented Client Builders"), Client(name="DEMO Client (demo client)", is_demo=True),
                ProductReview(name="Invented doubtful coating", status="pending"), ProductReview(name="Invented merged coating", status="rejected")])  # fmt: skip
    db.commit()
    source = _test_url.render_as_string(hide_password=False)
    target = create_engine(golive.url_for(name))
    dry = carryover.run(source, target, ["products", "vendors", "clients", "knowledge"], dry_run=True)
    by = {t["table"]: t for t in dry["tables"]}
    assert by["products"]["added"] == 1 and by["products"]["skipped"] == {"DEMO": 1}
    assert by["knowledge_product_reviews"]["added"] == 1 and by["knowledge_product_reviews"]["skipped"] == {"still to review: planning accepts, merges or rejects it first": 1}
    assert by["clients"]["added"] == 1 and by["clients"]["skipped"] == {"DEMO": 1} and by["vendors"]["added"] == 1
    with target.connect() as conn:
        assert conn.execute(text("SELECT count(*) FROM products WHERE code = 'INV-REAL-1'")).scalar() == 0  # the dry run changed nothing
    done = carryover.run(source, target, ["products", "vendors", "clients"])
    assert {t["table"]: t["added"] for t in done["tables"]}["products"] == 1
    again = carryover.run(source, target, ["products", "vendors", "clients"])
    assert again["added"] == 0  # idempotent
    with target.connect() as conn:
        assert conn.execute(text("SELECT count(*) FROM products WHERE name LIKE 'DEMO%' OR is_demo")).scalar() == 0
        assert conn.execute(text("SELECT count(*) FROM clients WHERE name = 'Invented Client Builders'")).scalar() == 1
    assert golive.demo_rows(name) == {}
    target.dispose()


# --- the roles report and help ---------------------------------------------------------------------------------------


def test_the_roles_report_lists_each_roles_pages_actions_and_what_it_never_sees(db, login_as):
    c, h = login_as("director", email="director-roles@example.com")
    r = c.get("/api/ops/roles-report", headers=h).json()
    by = {x["code"]: x for x in r["roles"]}
    assert "Salaries, wages and payroll" in by["site_supervisor"]["never"] and "Costs, purchase rates and margins" in by["billing"]["never"]
    assert "Billing" in by["billing"]["pages"] and "Stores" not in by["billing"]["pages"]
    assert by["director"]["never"] == [] or "Users, roles and permissions" in by["director"]["never"]
    pdf = c.get("/api/ops/roles-report.pdf", headers=h)
    assert pdf.content[:4] == b"%PDF"
    assert db.scalar(select(AuditLog).where(AuditLog.action == "export").order_by(AuditLog.id.desc())).after["path"] == "/api/ops/roles-report.pdf"  # the download is logged
    cs, hs = login_as("sales", email="sales-roles@example.com")
    assert cs.get("/api/ops/roles-report", headers=hs).status_code == 403


def test_help_is_for_the_page_and_the_role_and_the_office_admin_edits_it(db, login_as):
    ce, he = login_as("site_engineer", email="se-help@example.com")
    assert ce.get("/api/help?page=/sites", headers=he).json()["help"]["title"] == "Your sites"  # the engineer's own
    cb, hb = login_as("billing", email="billing-help@example.com")
    assert cb.get("/api/help?page=/sites", headers=hb).json()["help"]["title"] == "Sites for billing"
    co, ho = login_as("office_admin", email="oa-help@example.com")
    assert co.put("/api/help", json={"page": "/bills", "title": "Bills", "steps": ["one", "two"]}, headers=ho).status_code == 422  # 3 to 6
    r = co.put("/api/help", json={"page": "/bills", "title": "Bills (our way)", "steps": ["Open the bill.", "Check the quantities.", "Mark the step."]}, headers=ho)
    assert r.status_code == 200
    assert cb.get("/api/help?page=/bills", headers=hb).json()["help"]["title"] == "Bills (our way)"
    assert cb.put("/api/help", json={"page": "/bills", "title": "x", "steps": ["a", "b", "c"]}, headers=hb).status_code == 403
    co.delete("/api/help?page=/bills", headers=ho)
    assert cb.get("/api/help?page=/bills", headers=hb).json()["help"]["title"] == "Bills"


# --- security ------------------------------------------------------------------------------------------------------


def test_passwords_need_10_characters_and_not_a_common_one_or_the_persons_name():
    from app.auth.passwords import problem

    assert "10 characters" in problem("short1")
    assert "too common" in problem("password123") and "too common" in problem("Welcome@123")
    assert "email" in problem("ravi.patel2026!", email="ravi.patel@example.com")
    assert "plain sequence" in problem("qrstuvwxyz") and "too common" in problem("1234567890")
    assert problem("Tide-Lamp-Orchid-42") is None


def test_repeated_failures_lock_and_the_third_lock_waits_for_the_office_admin(db, make_user, login_as, monkeypatch):
    u = make_user("lock-me@example.com", "sales")
    c = TestClient(app)
    for lock in range(3):
        for _ in range(5):
            assert c.post("/api/auth/login", json={"email": u.email, "password": "wrong-password-x"}).status_code == 401
        db.refresh(u)
        assert u.locked_until is not None and u.lockouts == lock + 1
        u.locked_until = datetime.now(UTC) - timedelta(seconds=1)  # the 15 minutes passed
        db.commit()
    db.refresh(u)
    assert u.locked_by_failures
    assert c.post("/api/auth/login", json={"email": u.email, "password": PASSWORD}).status_code == 401  # even the right password
    co, ho = login_as("office_admin", email="oa-unlock@example.com")
    assert co.post(f"/api/admin/users/{u.id}/unlock", headers=ho).status_code == 200
    assert login(TestClient(app), u.email)
    assert db.scalar(select(AuditLog.id).where(AuditLog.action == "user.unlock")) is not None
    actions = {a.action for a in db.scalars(select(AuditLog).where(AuditLog.entity_id == str(u.id)))}
    assert {"auth.lockout", "auth.login_failed", "user.unlock"} <= actions
    assert co.get("/api/audit?category=logins", headers=ho).json()["total"] >= 15


def test_two_step_sign_in_is_required_for_the_roles_set_and_set_up_at_the_next_sign_in(db, make_user):
    s = db.get(SecuritySettings, 1)
    s.totp_required, s.totp_roles = True, ["director", "accounts"]
    db.commit()
    acc = make_user("accounts-2fa@example.com", "accounts")
    c = TestClient(app)
    r = c.post("/api/auth/login", json={"email": acc.email, "password": PASSWORD}).json()
    assert r["two_step"] == "enroll" and r["secret"] and "<svg" in r["qr"] and "access_token" not in r
    assert c.post("/api/auth/login/two-step", json={"ticket": r["ticket"], "code": "000000"}).status_code == 401
    ok = c.post("/api/auth/login/two-step", json={"ticket": r["ticket"], "code": totp.code_at(r["secret"], time.time())})
    assert ok.status_code == 200 and ok.json()["access_token"]
    db.refresh(acc)
    assert acc.totp_enabled_at is not None and acc.totp_secret and r["secret"] not in acc.totp_secret  # kept encrypted
    r2 = c.post("/api/auth/login", json={"email": acc.email, "password": PASSWORD}).json()
    assert r2["two_step"] == "code" and "secret" not in r2
    assert c.post("/api/auth/login/two-step", json={"ticket": r2["ticket"], "code": totp.code_at(r["secret"], time.time())}).status_code == 200
    h = {"Authorization": f"Bearer {ok.json()['access_token']}"}
    assert c.post("/api/auth/two-step/off", json={"code": totp.code_at(r["secret"], time.time())}, headers=h).status_code == 422  # required for the role
    sales = make_user("sales-2fa@example.com", "sales")
    assert "access_token" in c.post("/api/auth/login", json={"email": sales.email, "password": PASSWORD}).json()  # not for this role
    assert totp.verify(r["secret"], totp.code_at(r["secret"], time.time() - 30)) and not totp.verify(r["secret"], totp.code_at(r["secret"], time.time() - 120))


def test_a_password_change_ends_every_session(db, make_user):
    u = make_user("change-pw@example.com", "sales")
    c = TestClient(app)
    h = login(c, u.email)
    other = TestClient(app)
    h2 = login(other, u.email)
    time.sleep(1.1)  # tokens carry seconds
    assert c.post("/api/auth/password", json={"current": PASSWORD, "new": "password1"}, headers=h).status_code == 422
    assert c.post("/api/auth/password", json={"current": PASSWORD, "new": "Quiet-River-Stone-9"}, headers=h).status_code == 204
    assert other.get("/api/auth/me", headers=h2).status_code == 401  # the access token issued before
    assert other.post("/api/auth/refresh").status_code == 401  # and the refresh cookie
    assert db.scalar(select(AuditLog.id).where(AuditLog.action == "auth.password_change")) is not None
    assert login(TestClient(app), u.email, "Quiet-River-Stone-9")


def test_a_session_ends_after_the_inactivity_limit(db, make_user):
    u = make_user("idle@example.com", "sales")
    c = TestClient(app)
    login(c, u.email)
    assert c.post("/api/auth/refresh", headers={"X-Idle-Seconds": "60", "X-Device": "desktop"}).status_code == 200
    assert c.post("/api/auth/refresh", headers={"X-Idle-Seconds": str(9 * 3600), "X-Device": "desktop"}).status_code == 401  # 8 h on a desktop
    p = TestClient(app)
    assert p.post("/api/auth/login", json={"email": u.email, "password": PASSWORD}, headers={"X-Device": "phone"}).status_code == 200
    assert p.post("/api/auth/refresh", headers={"X-Idle-Seconds": str(9 * 3600), "X-Device": "phone"}).status_code == 200  # 12 h on a phone
    tok = db.scalars(select(RefreshToken).where(RefreshToken.user_id == u.id).order_by(RefreshToken.created_at.desc())).first()
    assert tok.device == "phone"


def test_uploads_are_checked_by_content_and_size(boss):
    c, h = boss
    site = c.post("/api/sites", json={"name": "Invented Upload Site", "state": "Gujarat"}, headers=h).json()["id"]
    r = c.post(f"/api/sites/{site}/drawings", data={"title": "Fake", "discipline": "architectural"}, files={"file": ("plan.pdf", b"MZ\x90\x00 not a pdf", "application/pdf")}, headers=h)
    assert r.status_code == 422 and "not a .pdf file" in r.json()["detail"]
    r = c.post(f"/api/sites/{site}/drawings", data={"title": "Real", "discipline": "architectural"}, files={"file": ("plan.pdf", b"%PDF-1.4 invented", "application/pdf")}, headers=h)
    assert r.status_code == 201


def test_every_download_is_in_the_audit_log_under_exports(db, boss):
    c, h = boss
    r = c.get("/api/audit/export", headers=h)
    assert r.status_code == 200
    row = db.scalars(select(AuditLog).where(AuditLog.action == "export").order_by(AuditLog.id.desc())).first()
    assert row.after["path"] == "/api/audit/export" and row.user_id is not None
    assert c.get("/api/audit?category=exports", headers=h).json()["total"] >= 1


def test_mode_says_whether_this_is_the_training_company(monkeypatch, login_as):
    c = TestClient(app)
    m = c.get("/api/ops/mode").json()
    assert {k: m[k] for k in ("training", "main_url", "training_url")} == {"training": False, "main_url": None, "training_url": None}
    monkeypatch.setattr(settings, "training_mode", True)
    monkeypatch.setattr(settings, "main_url", "https://app.example.in")
    assert c.get("/api/ops/mode").json()["training"] is True
    c, h = login_as("office_admin")
    assert c.get("/api/ops/go-live-checks", headers=h).json() == {"missing": []}  # never goes live


def test_the_training_reset_only_runs_on_the_training_company():
    from app.ops import training

    with pytest.raises(SystemExit):
        training.reset()



def test_each_roles_training_guide_has_help_for_its_pages_and_stays_under_12_pages(db, tmp_path):
    from app.ops import training_guide
    from app.ops.help import DEFAULTS

    for role, (_, _, pages) in training_guide.GUIDES.items():
        assert all(p in DEFAULTS for p in pages), role
    screens = tmp_path / "screens" / "site_engineer"
    screens.mkdir(parents=True)
    from PIL import Image

    for p in training_guide.GUIDES["site_engineer"][2]:  # the longest guide, with a picture on every page
        Image.new("RGB", (1600, 1000), "white").save(screens / f"{training_guide.slug(p)}.png")
    data, pages = training_guide.pdf(db, "site_engineer", tmp_path / "screens")
    assert data.startswith(b"%PDF") and pages < 12
    text_ = training_guide.html(db, "site_engineer", None)
    assert "Daily reports" in text_ and "training company" in text_
