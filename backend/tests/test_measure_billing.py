# ruff: noqa: E501
"""Measuring from the drawing, the measurement sheet, the RA bill from it, and old sheets read back.
Invented data only (the DEMO pool's shapes, drawn here)."""

import io
import math
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import select

from app.closeout.models import ClientWorkOrder, ClientWorkOrderLine
from app.config import settings
from app.finance.models import ClientContract, ContractLine, RaBill
from app.masters.models import Client, CompanyGstin
from app.measure import bill as mb
from app.measure import demo, excel, geom
from app.measure.models import MeasureArea, MeasureBilled, MeasureEdge, MeasureSheet
from app.models import User
from app.sites.models import Drawing, DrawingRevision, Site
from app.takeoff import pipeline
from app.takeoff.models import CadFile
from tests.test_finance import boss  # noqa: F401  (fixture)

D = Decimal


# --- the geometry ---------------------------------------------------------------------------------


def test_arc_area_and_periphery_within_a_square_mm_and_a_hundredth_of_a_mm():
    circle = [[0, 0, 1], [4, 0, 1]]  # radius 2, two half circles
    assert abs(geom.area(circle) - 4 * math.pi) < 1e-6
    assert abs(geom.periphery(circle) - 4 * math.pi) < 1e-5
    sh = demo.shapes()
    r, L = sh["r"], sh["L"]
    assert abs(geom.area(sh["pool"]) - (2 * r * L + math.pi * r * r)) < 1e-6  # a stadium, by its formula
    assert abs(geom.area(sh["pool"]) - 55.0027) < 1e-6 and abs(geom.periphery(sh["pool"]) - 30.119) < 1e-5
    assert abs(geom.area(sh["tank"]) - 16.9877) < 1e-6 and abs(geom.periphery(sh["tank"]) - 15.786) < 1e-5
    # a bulge turned the other way takes the arc's segment out instead of adding it
    sq_out = [[0, 0, 1], [2, 0, 0], [2, 2, 0], [0, 2, 0]]
    sq_in = [[0, 0, -1], [2, 0, 0], [2, 2, 0], [0, 2, 0]]
    assert abs(geom.area(sq_out) - (4 + math.pi / 2)) < 1e-6 and abs(geom.area(sq_in) - (4 - math.pi / 2)) < 1e-6
    # the flattened copy (the trim's) stays within a square millimetre of the arcs
    assert abs(geom.polygon(sh["pool"]).area - geom.area(sh["pool"])) < 1e-6


def test_overlap_trim_and_the_shared_edge():
    sh = demo.shapes()
    a, per, over = geom.trim(sh["slope"], [sh["pool"]])
    assert round(a, 4) == 6.1324 and round(per, 3) == 13.772 and over > 1
    assert geom.trim(sh["pool"], []) == (geom.area(sh["pool"]), geom.periphery(sh["pool"]), 0.0)
    edges = geom.shared_edges(sh["pool"], sh["tank"])
    assert len(edges) == 1 and abs(edges[0][0] - 5.19) < 1e-9
    assert geom.shared_edges(sh["pool"], sh["slope"]) == []  # overlapping, not side by side


def test_slope_rounding_and_sqft():
    f = geom.slope_factor(D("0.45"), D("1.051"))
    assert f == (D("0.45") ** 2 + D("1.051") ** 2).sqrt() / D("1.051")
    assert geom.round3(D("1.2345")) == D("1.235")  # half up, where truncation gives 1.234
    assert geom.round3(D("0.0005")) == D("0.001") and geom.round3(D("2.9999")) == D("3.000")
    assert geom.SQFT_PER_SQM == D("10.7639") and geom.to_sqft(D(1)) == D("10.7639")
    h, src = geom.height_from("Section Y-Y", "+1800", "+600")
    assert h == 1200 and src == "Section Y-Y: +1800 to +600 = 1200 mm"
    assert geom.level_mm("+1.800") == 1800 and geom.level_mm("FFL -600") == -600


# --- a site measured from its drawing ------------------------------------------------------------


@pytest.fixture
def pool(boss, db):  # noqa: F811
    c, h = boss
    me = db.scalar(select(User).where(User.email == "boss@example.com"))
    db.add(CompanyGstin(gstin="24AAACE1234A1Z1", state="Gujarat", address="Our office", is_default=True))
    cl = Client(name="Invented Pool Club", gstin="24AAACI1111A1Z1", state="Gujarat", city="Ahmedabad", address="Invented Road")
    db.add(cl)
    db.flush()
    site = Site(code="S-POOL-1", name="Invented pool", status="active", client_id=cl.id, state="Gujarat")
    db.add(site)
    db.flush()
    wo = ClientWorkOrder(code="CWO-T-1", client_ref="INV/WO/1", site_id=site.id, client_id=cl.id, status="accepted")
    for i, (no, desc, unit, qty, rate, per_layer, layers) in enumerate(
        [("1", "Coating", "sqft", "3000", "85", False, 1), ("2", "Pipe sleeve", "nos", "10", "1250", False, 1), ("3", "Brick bat per layer", "sqft", "1700", "42", True, 2)]
    ):
        wo.lines.append(ClientWorkOrderLine(sort_order=i, item_no=no, description=desc, unit=unit, qty=D(qty), rate=D(rate), per_layer=per_layer, layers=layers, sac="995454"))
    db.add(wo)
    db.flush()
    k = ClientContract(site_id=site.id, client_id=cl.id, contract_value=D(0), retention_percent=D(0), sac="9954", client_gstin=cl.gstin, place_of_supply="Gujarat")
    for w in wo.lines:
        k.lines.append(ContractLine(sort_order=w.sort_order, item_no=w.item_no, description=w.description, unit=w.unit, qty=w.qty, rate=w.rate, wo_line_id=w.id))
    db.add(k)
    rel = f"sites/{site.id}/drawings/pool.dxf"
    demo.dxf(Path(settings.media_dir) / rel)
    d = Drawing(site_id=site.id, title="Invented pool", discipline="architectural")
    db.add(d)
    db.flush()
    r = DrawingRevision(drawing_id=d.id, rev_no=1, stored_path=rel, filename="pool.dxf", size_bytes=1)
    db.add(r)
    db.flush()
    f = CadFile(revision_id=r.id, drawing_id=d.id, site_id=site.id, status="queued")
    db.add(f)
    db.commit()
    f = pipeline.process(db, f)
    assert f.status == "ready", f.error
    sh = demo.shapes()
    areas = {}
    for code in "AQR":
        x, y = sh["clicks"][code]
        res = c.post(f"/api/measure/sites/{site.id}/areas", json={"cad_file_id": f.id, "x": x, "y": y, "code": code, "name": f"Area {code}"}, headers=h)
        assert res.status_code == 201, res.text
        areas[code] = res.json()["id"]
    for code, top, bottom in (("A", "+1800", "+600"), ("R", "+600", "-900")):
        assert c.post(f"/api/measure/areas/{areas[code]}/height", json={"section": "Section Y-Y", "top": top, "bottom": bottom}, headers=h).status_code == 200
    assert c.post(f"/api/measure/areas/{areas['Q']}/slope", json={"section": "Section Y-Y", "top": "+2250", "bottom": "+1800", "run_mm": "1051"}, headers=h).status_code == 200
    wl = {w.item_no: w.id for w in wo.lines}
    for no, body in (("1", {"parts": ["floor", "walls", "tops", "slope"]}), ("2", {"parts": ["count"], "count_of": "layer:SLEEVE"}), ("3", {"parts": ["floor"], "layers": 2, "areas": [areas["A"], areas["R"]]})):
        assert c.put(f"/api/measure/sites/{site.id}/items/{wl[no]}", json=body, headers=h).status_code == 200
    edges = list(db.scalars(select(MeasureEdge).where(MeasureEdge.site_id == site.id)))
    edge = next(e for e in edges if {e.area_a_id, e.area_b_id} == {areas["A"], areas["R"]})
    # the slope trimmed to the pool shares the cut edge with it: the pool's wall stays there
    aq = next(e for e in edges if {e.area_a_id, e.area_b_id} == {areas["A"], areas["Q"]})
    assert c.post(f"/api/measure/edges/{aq.id}/answer", json={"answer": "both"}, headers=h).status_code == 200
    return {"c": c, "h": h, "me": me, "site": site, "wo": wo, "wl": wl, "areas": areas, "edge": edge, "contract": k, "client": cl, "file": f}


def _answer(p, answer, weir=None):
    r = p["c"].post(f"/api/measure/edges/{p['edge'].id}/answer", json={"answer": answer, "weir_mm": weir}, headers=p["h"])
    assert r.status_code == 200, r.text


def _sheet(p, on="2026-09-01"):
    r = p["c"].post(f"/api/measure/sites/{p['site'].id}/sheets", json={"measured_on": on}, headers=p["h"])
    assert r.status_code == 201, r.text
    return r.json()


def test_areas_picked_from_the_drawing_and_the_edge_asked_once(pool, db):
    st = pool["c"].get(f"/api/measure/sites/{pool['site'].id}", headers=pool["h"]).json()
    a = {x["code"]: x for x in st["areas"]}
    assert D(a["A"]["area_sqm"]).quantize(D("0.0001")) == D("55.0027") and D(a["A"]["periphery_m"]).quantize(D("0.001")) == D("30.119")
    assert D(a["Q"]["net_area_sqm"]).quantize(D("0.0001")) == D("6.1324") and D(a["Q"]["net_periphery_m"]).quantize(D("0.001")) == D("13.772")
    assert D(a["R"]["area_sqm"]).quantize(D("0.0001")) == D("16.9877") and D(a["R"]["periphery_m"]).quantize(D("0.001")) == D("15.786")
    assert a["Q"]["trim_with"] == [a["A"]["id"]] and a["A"]["counts"] == {"layer:SLEEVE": 4}
    assert a["A"]["height_source"] == "Section Y-Y: +1800 to +600 = 1200 mm"
    edges = {f"{e['a']}|{e['b']}": e for e in st["edges"]}
    assert set(edges) == {"A|R", "A|Q"}
    assert D(edges["A|R"]["length_m"]) == D("5.19") and edges["A|R"]["answer"] is None
    assert abs(D(edges["A|Q"]["length_m"]) - D("5.835038670")) < D("0.000001")  # the slope's cut edge
    s = _sheet(pool)
    assert any(q["kind"] == "edge" and "Wall on this edge?" in q["text"] for q in s["questions"])
    assert pool["c"].post(f"/api/measure/sheets/{s['id']}/send", headers=pool["h"]).status_code == 422  # not before the answer


def _walls_and_tops(lines, area_id):
    return sum(D(ln["total"]) for ln in lines if ln["area_id"] == area_id and ln["part"] in ("walls", "tops") and ln["wo_line_id"] and ln["total"])


def test_each_edge_answer_changes_the_deduction(pool, db):
    s = _sheet(pool)
    got = {}
    for answer, weir in (("both", None), ("open", None), ("weir", "450")):
        _answer(pool, answer, weir)
        lines = pool["c"].get(f"/api/measure/sheets/{s['id']}", headers=pool["h"]).json()["lines"]
        got[answer] = (_walls_and_tops(lines, pool["areas"]["A"]), _walls_and_tops(lines, pool["areas"]["R"]), [ln for ln in lines if ln["kind"] == "deduction"])
    a_both, a_open, a_weir = got["both"][0], got["open"][0], got["weir"][0]
    assert a_both > a_weir > a_open  # a wall on both faces, a weir up to 450 mm, no wall
    sq = geom.SQFT_PER_SQM
    # A: open takes the edge × height and the edge × wall top; a weir only the wall above it
    assert a_both - a_open == geom.round3(D("5.19") * D("1.2") * sq) + geom.round3(D("5.19") * D("0.25") * sq)
    assert a_both - a_weir == geom.round3(D("5.19") * D("0.75") * sq)
    # the wall top between A and R is counted once: R gives its share back on "both" and "weir"
    assert any("shared with A" in ln["particulars"] for ln in got["both"][2])
    assert all(ln["reason"] for ln in got["weir"][2])


def test_the_slope_line_shows_its_formula(pool, db):
    _answer(pool, "both")
    s = _sheet(pool)
    line = next(ln for ln in s["lines"] if ln["part"] == "slope")
    q = D("6.1324")
    f = geom.slope_factor(D("0.45"), D("1.051"))
    assert abs(D(line["qty"]) - q * f) < D("0.0001")
    assert "√(0.45² + 1.051²) / 1.051" in line["formula"]
    assert D(line["total"]) == geom.round3(D(line["qty"]) * geom.SQFT_PER_SQM)


def test_layers_double_only_when_the_rate_is_per_layer(pool, db):
    _answer(pool, "both")
    s = _sheet(pool)
    bb = [ln for ln in s["lines"] if ln["wo_line_id"] == pool["wl"]["3"] and ln["kind"] == "line"]
    assert sorted({ln["layer"] for ln in bb}) == [1, 2]
    w = db.get(ClientWorkOrderLine, pool["wl"]["3"])
    w.per_layer = False
    db.commit()
    s = pool["c"].get(f"/api/measure/sheets/{s['id']}", headers=pool["h"]).json()
    assert any(q["kind"] == "layers" for q in s["questions"])
    assert pool["c"].put(f"/api/measure/sites/{pool['site'].id}/items/{w.id}", json={"parts": ["floor"], "layers": 2, "layers_answer": "once"}, headers=pool["h"]).status_code == 200
    s = pool["c"].get(f"/api/measure/sheets/{s['id']}", headers=pool["h"]).json()
    assert not s["questions"] and {ln["layer"] for ln in s["lines"] if ln["wo_line_id"] == w.id and ln["kind"] == "line"} == {1}


def test_the_excel_formulas_give_the_apps_totals(pool, db):
    _answer(pool, "weir", "450")
    s = _sheet(pool)
    r = pool["c"].get(f"/api/measure/sheets/{s['id']}/xlsx", headers=pool["h"])
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert ws["H3"].value == 10.7639
    app_lines = [ln for ln in s["lines"] if ln["kind"] != "area"]
    rows = [rr for rr in range(7, ws.max_row + 1) if isinstance(ws.cell(rr, 8).value, str) and ws.cell(rr, 8).value.startswith("=ROUND(G")]
    assert len(rows) == len(app_lines)
    for rr, ln in zip(rows, app_lines, strict=True):
        assert D(repr(excel.evaluate(ws, rr, 8))) == D(ln["total"]), (rr, ln["particulars"])
    this = {x["wo_line_id"]: D(x["this"]) for x in s["cumulative"]}
    ab = next(rr for rr in range(1, ws.max_row + 1) if ws.cell(rr, 1).value == "Abstract") + 2
    for i, x in enumerate(s["cumulative"]):
        assert abs(D(repr(excel.evaluate(ws, ab + i, 7))) - this[x["wo_line_id"]]) < D("0.0000001")


def _send_and_bill(p, db, sheet_id, **body):
    r = p["c"].post(f"/api/measure/sheets/{sheet_id}/send", headers=p["h"])
    assert r.status_code == 200, r.text
    r = p["c"].post(f"/api/measure/sheets/{sheet_id}/bill", json=body, headers=p["h"])
    assert r.status_code == 201, r.text
    return r.json()


def test_ra_bills_from_sheets_previous_this_cumulative_and_no_double_billing(pool, db):
    _answer(pool, "weir", "450")
    s1 = _sheet(pool)
    for wid, layer in ((pool["wl"]["1"], 1), (pool["wl"]["2"], 1), (pool["wl"]["3"], 1), (pool["wl"]["3"], 2)):
        pool["c"].post(f"/api/measure/sheets/{s1['id']}/not-done", json={"area_id": pool["areas"]["R"], "wo_line_id": wid, "layer": layer}, headers=pool["h"])
    s1 = pool["c"].get(f"/api/measure/sheets/{s1['id']}", headers=pool["h"]).json()
    nd = [ln for ln in s1["lines"] if ln["kind"] == "not_done"]
    assert len(nd) == 4 and all(ln["total"] is None for ln in nd)  # "–" on the sheet
    b1 = _send_and_bill(pool, db, s1["id"], proforma_on="2026-09-03")
    assert b1["checklist"] and all(x["ok"] for x in b1["checklist"]), b1["checklist"]
    assert pool["c"].post(f"/api/finance/ra-bills/{b1['id']}/submit", headers=pool["h"]).status_code == 200
    s2 = _sheet(pool, "2026-10-01")
    # A and Q are not on sheet 2: billed on sheet 1 for those items and layers
    assert not any(ln["area_id"] in (pool["areas"]["A"], pool["areas"]["Q"]) and ln["kind"] == "line" and ln["wo_line_id"] == pool["wl"]["1"] for ln in s2["lines"])
    b2 = _send_and_bill(pool, db, s2["id"], proforma_on="2026-09-25")
    for ln in b2["lines"]:
        assert ln["previous"] + ln["this"] == ln["cumulative"] if isinstance(ln["previous"], Decimal) else D(ln["previous"]) + D(ln["this"]) == D(ln["cumulative"])
    one = {ln["item_no"]: D(ln["this"]) for ln in b1["lines"]}
    two = {ln["item_no"]: D(ln["previous"]) for ln in b2["lines"]}
    assert all(two[k] == one[k] for k in one)
    assert b2["warnings"] and "before the measurement" in b2["warnings"][0]  # dated 25-09, measured 01-10
    # the same area, item and layer cannot go on another sheet: the key is unique
    db.add(MeasureBilled(area_id=pool["areas"]["A"], wo_line_id=pool["wl"]["1"], layer=1, sheet_id=s2["id"]))
    with pytest.raises(Exception):  # noqa: B017
        db.flush()
    db.rollback()


def test_gst_split_round_off_and_rates_from_the_work_order(pool, db):
    _answer(pool, "both")
    s = _sheet(pool)
    # a contract rate typed differently is ignored: the bill takes the work order's
    cl = db.scalar(select(ContractLine).where(ContractLine.wo_line_id == pool["wl"]["1"]))
    cl.rate = D("99")
    db.commit()
    b = _send_and_bill(pool, db, s["id"])
    assert {ln["item_no"]: D(ln["rate"]) for ln in b["lines"]}["1"] == D("85")
    t = b["tax"]
    assert not t["interstate"] and D(t["cgst_percent"]) == 9 and D(t["sgst_percent"]) == 9 and D(t["igst"]) == 0
    basic = sum(D(ln["amount"]) for ln in b["lines"])
    assert D(t["taxable"]) == basic
    before = D(t["taxable"]) + D(t["cgst"]) + D(t["sgst"])
    assert D(t["total"]) == before.quantize(D(1), rounding="ROUND_HALF_UP") and D(t["round_off"]) == D(t["total"]) - before
    # a client outside Gujarat (state code 27): IGST 18 %
    k = db.get(ClientContract, pool["contract"].id)
    k.client_gstin = "27AAACI1111A1Z1"
    db.commit()
    t2 = pool["c"].get(f"/api/measure/bills/{b['id']}", headers=pool["h"]).json()["tax"]
    assert t2["interstate"] and D(t2["igst_percent"]) == 18 and D(t2["cgst"]) == 0 and D(t2["igst"]) == (basic * D("0.18")).quantize(D("0.01"))
    # every line carries a SAC; the header is spelled right; the buyer from the client master
    assert all(ln["sac"] == "995454" for ln in b["lines"]) and b["title"] == "PROFORMA INVOICE" and b["buyer"]["name"] == "Invented Pool Club"
    assert b["proforma_no"].startswith("PI/")
    # a rate changed on the bill is caught by the checklist, and the bill cannot be sent
    ra = db.get(RaBill, b["id"])
    ra.lines[0].rate = D("90")
    db.commit()
    chk = {x["key"]: x for x in pool["c"].get(f"/api/measure/bills/{b['id']}", headers=pool["h"]).json()["checklist"]}
    assert not chk["rates"]["ok"]
    assert pool["c"].post(f"/api/finance/ra-bills/{b['id']}/submit", headers=pool["h"]).status_code == 422


def test_proforma_to_tax_invoice_in_one_click_keeps_its_lines(pool, db):
    _answer(pool, "both")
    b = _send_and_bill(pool, db, _sheet(pool)["id"])
    assert pool["c"].post(f"/api/finance/ra-bills/{b['id']}/submit", headers=pool["h"]).status_code == 200
    r = pool["c"].post(f"/api/measure/bills/{b['id']}/tax-invoice", json={}, headers=pool["h"])
    assert r.status_code == 201, r.text
    inv = r.json()
    assert inv["lines"] == len([ln for ln in b["lines"] if D(ln["amount"])]) and D(inv["total"]) == D(b["tax"]["total"])
    assert pool["c"].post(f"/api/measure/bills/{b['id']}/tax-invoice", json={}, headers=pool["h"]).status_code == 409


def test_a_verbal_order_needs_who_and_when(pool, db):
    _answer(pool, "both")
    s = _sheet(pool)
    assert pool["c"].post(f"/api/measure/sheets/{s['id']}/send", headers=pool["h"]).status_code == 200
    r = pool["c"].post(f"/api/measure/sheets/{s['id']}/bill", json={"basis": "verbal"}, headers=pool["h"])
    assert r.status_code == 422
    r = pool["c"].post(f"/api/measure/sheets/{s['id']}/bill", json={"basis": "verbal", "verbal_by": "Client's site engineer", "verbal_on": "2026-08-30"}, headers=pool["h"])
    assert r.status_code == 201 and r.json()["order"]["text"] == "Verbal order confirmed by Client's site engineer on 30-08-2026"


def test_client_corrections_are_logged_and_billed(pool, db):
    from app.closeout import measure

    site = db.get(Site, pool["site"].id)
    site.measurements_need_approval = True
    db.commit()
    _answer(pool, "both")
    s = _sheet(pool)
    r = pool["c"].post(f"/api/measure/sheets/{s['id']}/send", headers=pool["h"])
    assert r.status_code == 200 and r.json()["status"] == "sent"
    sheet = db.get(MeasureSheet, s["id"])
    book = mb.book_lines(db, sheet)
    me = pool["me"]
    for m in book:
        if "Pipe sleeve" in m.description:
            ch = measure.client_correct(db, m, "qty", "5", "One sleeve is not fixed", me)
            measure.answer(db, m, ch, True, None, me)
        else:
            measure.client_approve(db, m, me)
    db.commit()
    out = pool["c"].get(f"/api/measure/sheets/{s['id']}", headers=pool["h"]).json()
    assert out["status"] == "approved" and len(out["corrections"]) == 1 and out["corrections"][0]["remark"] == "One sleeve is not fixed"
    b = pool["c"].post(f"/api/measure/sheets/{s['id']}/bill", json={}, headers=pool["h"]).json()
    assert {ln["item_no"]: D(ln["this"]) for ln in b["lines"]}["2"] == 5  # the corrected figure is billed
    # the sheet's Abstract, the bill, the Excel and the PDF say the same, the drawing figure beside it
    out = pool["c"].get(f"/api/measure/sheets/{s['id']}", headers=pool["h"]).json()
    row = next(r for r in out["cumulative"] if r["item_no"] == "2")
    assert D(row["this"]) == 5 and D(row["drawing"]) == 6 and D(row["approved"]) == 5 and D(row["upto"]) == 5
    assert row["note"] == "6 on the drawing, 5 approved by the client"
    bills = {ln["item_no"]: ln for ln in b["lines"]}
    for r in out["cumulative"]:
        assert D(r["this"]) == D(bills[r["item_no"]]["this"]) and D(r["upto"]) == D(bills[r["item_no"]]["cumulative"])
    ws = openpyxl.load_workbook(io.BytesIO(pool["c"].get(f"/api/measure/sheets/{s['id']}/xlsx", headers=pool["h"]).content)).active
    ab = next(rr for rr in range(1, ws.max_row + 1) if ws.cell(rr, 1).value == "Abstract") + 2
    for i, r in enumerate(out["cumulative"]):
        assert abs(D(repr(excel.evaluate(ws, ab + i, 7))) - D(r["this"])) < D("0.0000001"), r["item_no"]
    assert excel.evaluate(ws, ab + 1, 5) == 6.0 and excel.evaluate(ws, ab + 1, 7) == 5.0
    from pypdf import PdfReader

    pdf = pool["c"].get(f"/api/measure/sheets/{s['id']}/pdf", headers=pool["h"]).content
    text = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "6 on the drawing, 5 approved by the client" in " ".join(text.split())


def test_import_gives_the_same_totals_as_the_sheet(pool, db):
    _answer(pool, "weir", "450")
    s = _sheet(pool)
    data = pool["c"].get(f"/api/measure/sheets/{s['id']}/xlsx", headers=pool["h"]).content
    want = {x["item_no"]: D(x["this"]) for x in s["cumulative"]}
    read = excel.read_sheet(data)
    assert {k: v["qty"] for k, v in read["items"].items()} == want
    assert read["skipped"] == [], read["skipped"]
    db.query(MeasureSheet).filter(MeasureSheet.id == s["id"]).delete()
    db.commit()
    # an old proforma (values as typed) and a row the reader cannot place
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["PROFORMA INVOICE"])
    ws.append(["Item", "Description", "Unit", "Qty", "Rate", "Amount"])
    for k, q in want.items():
        rate = {"1": 85, "2": 1250, "3": 42}[k]
        ws.append([k, f"Item {k}", "x", float(q), rate, float((q * rate).quantize(D("0.01")))])
    ws.append([None, "Transport", None, 1, 500, 500])
    ws.append([None, "Basic", None, None, None, 1])
    buf = io.BytesIO()
    wb.save(buf)
    files = {"sheet": ("old-sheet.xlsx", data), "proforma": ("old-proforma.xlsx", buf.getvalue())}
    r = pool["c"].post(f"/api/measure/sites/{pool['site'].id}/import", data={"measured_on": "2026-08-01"}, files=files, headers=pool["h"])
    assert r.status_code == 201, r.text
    out = r.json()
    assert {x["item_no"]: D(x["qty"]) for x in out["items"]} == want
    assert any("Transport" in x for x in out["skipped"])
    # they are the next sheet's and the next bill's previous
    s2 = _sheet(pool, "2026-09-01")
    assert {x["item_no"]: D(x["previous"]) for x in s2["cumulative"]} == want
    assert db.get(RaBill, out["ra_bill_id"]).imported


def test_a_billed_area_is_locked(pool, db):
    _answer(pool, "both")
    s = _sheet(pool)
    assert pool["c"].post(f"/api/measure/sheets/{s['id']}/send", headers=pool["h"]).status_code == 200
    a = pool["areas"]["A"]
    assert pool["c"].post(f"/api/measure/areas/{a}/height", json={"section": "Y", "top": "+1900", "bottom": "+600"}, headers=pool["h"]).status_code == 409
    assert pool["c"].delete(f"/api/measure/areas/{a}", headers=pool["h"]).status_code == 409
    assert db.get(MeasureArea, a) is not None and date.today()


def test_an_old_row_without_an_item_is_listed_never_silently_added():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Sr / WO item", "Particulars", "No", "Length", "Width", "Height / Depth", "Qty", "Total qty", "Unit"])
    ws.append([None, "A - Pool"])
    ws.append(["1", "Floor", 1, None, None, None, 10, 107.639, "sqft"])
    ws.append([None, "Walls (the same item, no number)", 1, 10, None, 1, 10, 107.639, "sqft"])
    ws.append([None, "Lift pit (extra)"])  # an area's header inside the item: the item goes on
    ws.append([None, "Floor", 1, 2, 2, None, 4, 43.056, "sqft"])
    buf = io.BytesIO()
    wb.save(buf)
    read = excel.read_sheet(buf.getvalue())
    assert read["items"]["1"]["qty"] == D("258.334")
    assert len(read["assumed"]) == 2 and all("counted under item 1" in a for a in read["assumed"])  # listed to check
    # the item's own total row confirms the lines under it: nothing left to check
    ws.append([None, "Total Quantity =", None, None, None, None, 24, 258.334])
    ws.append([None, None, None, None, None, None, "Sq Mt", "Sq Ft"])
    buf = io.BytesIO()
    wb.save(buf)
    read = excel.read_sheet(buf.getvalue())
    assert read["items"]["1"]["qty"] == D("258.334") and read["items"]["1"]["unit"] == "sqft" and read["items"]["1"]["sqm"] == 24
    assert read["assumed"] == [] and read["skipped"] == []
    # a row with a number before any item is never placed
    wb2 = openpyxl.Workbook()
    w2 = wb2.active
    w2.append(["Sr / WO item", "Particulars", "No", "Length", "Width", "Height / Depth", "Qty", "Total qty", "Unit"])
    w2.append([None, "Lift pit (extra)", 1, 2, 2, None, 4, 43.056, "sqft"])
    buf = io.BytesIO()
    wb2.save(buf)
    read = excel.read_sheet(buf.getvalue())
    assert read["items"] == {} and "43.056 with no work-order item" in read["skipped"][0]


def test_a_shared_edge_along_an_arc_is_found_with_its_true_length():
    # A: a half disc of radius 2 on the x axis; B: a frame around it whose inner edge is that arc
    a = [[2, 0, 1], [-2, 0, 0]]
    b = [[-2, 0, -1], [2, 0, 0], [3, 0, 0], [3, 3, 0], [-3, 3, 0], [-3, 0, 0]]
    found = geom.shared_paths(a, b)
    assert len(found) == 1
    length, path = found[0]
    assert abs(length - 2 * math.pi) < 1e-5  # within a hundredth of a millimetre
    assert len(path) >= 8 and abs(math.hypot(path[len(path) // 2 // 2 * 2], path[len(path) // 2 // 2 * 2 + 1]) - 2) < 1e-6
    assert geom.shared_edges(a, b) == []  # the straight-edge finder alone missed it
    # a slope drawn into the pool and trimmed to it: the edge is the pool's arc inside the slope
    pool = [[0, 0, 0], [6, 0, 1], [6, 4, 0], [0, 4, 1]]  # straight sides, round ends of radius 2
    slope = [[6.5, 1, 0], [9, 1, 0], [9, 3, 0], [6.5, 3, 0]]  # over the right-hand end
    (length, _), = geom.shared_paths(pool, slope, a_trims_b=True)
    half = math.asin(1 / 2)  # the arc of radius 2 between y = 1 and y = 3 around its centre (6, 2)
    assert abs(length - 2 * 2 * half) < 1e-5
