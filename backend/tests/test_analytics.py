# ruff: noqa: E501  (fixtures and expectations read better on one line)
"""Dashboards, analytics, alerts, reports and the demo company. Invented fixtures with known
numbers; dates are relative to today so the tests hold on any day."""

from datetime import UTC, datetime, time, timedelta
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select, text

from app.analytics import alerts as alert_svc
from app.analytics import demo, reports, summary, worker
from app.analytics.common import add_months, month_start, today
from app.analytics.models import (
    Alert,
    AlertRecipient,
    InvoiceSummary,
    MonthlySummary,
    ProgressSnapshot,
    SiteSummary,
    WeeklyReport,
)
from app.config import settings
from app.crm.models import Lead
from app.execution.models import Dpr, SiteBudget, SiteCost
from app.execution.service import IST
from app.finance import service as fsvc
from app.finance.models import (
    ClientContract,
    Payment,
    PaymentAllocation,
    PettyCashAccount,
    PettyCashEntry,
    Receipt,
    ReceiptAllocation,
    TaxInvoice,
    VendorBill,
)
from app.main import app
from app.masters.models import Client, Product, Vendor
from app.material.models import PurchaseOrder, StockLedger, Store
from app.models import Base, User
from app.portal.models import Notification
from app.sites.models import Site, SiteMember
from app.tenders.models import Tender, TenderRevision
from tests.conftest import login

D = Decimal


@pytest.fixture(autouse=True)
def media(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "media_dir", str(tmp_path))


def at(d, hour=12):
    return datetime.combine(d, time(hour), tzinfo=IST).astimezone(UTC)


def uid(db, email):
    return db.scalar(select(User.id).where(User.email == email))


@pytest.fixture
def boss(login_as):
    return login_as("super_admin", email="boss@example.com")


@pytest.fixture
def company(boss, db, make_user):
    """Two clients, four sites with contracts, invoices, a credit note, receipts, costs, tenders,
    supplier bills. The expected numbers are worked out in each test from these."""
    day = today()
    m0 = month_start(day)
    prev0 = add_months(m0, -1)
    a = Client(name="Invented Builders", type="builder", state="Gujarat")
    b = Client(name="Invented Developers", type="developer", state="Maharashtra")
    v = Vendor(name="Invented Supplier", type="material_supplier")
    sup = make_user("sup-1@example.com", "site_supervisor", name="Sup One")
    db.add_all([a, b, v])
    db.flush()
    s1 = Site(
        code="S-T-1",
        name="Behind time",
        client_id=a.id,
        status="active",
        state="Gujarat",
        start_date=day - timedelta(days=100),
        target_date=day + timedelta(days=100),
        progress_percent=30,
        site_incharge_id=sup.id,
    )
    s2 = Site(
        code="S-T-2",
        name="On time",
        client_id=b.id,
        status="active",
        state="Maharashtra",
        start_date=day - timedelta(days=50),
        target_date=day + timedelta(days=150),
        progress_percent=30,
    )
    s3 = Site(
        code="S-T-3",
        name="Past its end",
        client_id=a.id,
        status="active",
        state="Gujarat",
        start_date=day - timedelta(days=200),
        target_date=day - timedelta(days=5),
        progress_percent=80,
    )
    s4 = Site(
        code="S-T-4",
        name="Done",
        client_id=b.id,
        status="completed",
        state="Maharashtra",
        start_date=day - timedelta(days=400),
        target_date=day - timedelta(days=100),
        progress_percent=100,
    )
    db.add_all([s1, s2, s3, s4])
    db.flush()
    db.add(SiteMember(site_id=s1.id, user_id=sup.id, role_on_site="supervisor"))
    contracts = {}
    for s, value in ((s1, 1_000_000), (s2, 2_000_000), (s3, 500_000), (s4, 300_000)):
        c = ClientContract(site_id=s.id, client_id=s.client_id, contract_value=value)
        db.add(c)
        db.flush()
        contracts[s.id] = c.id

    def invoice(number, s, on, taxable, retention=0, kind="invoice", against=None):
        gst = D(taxable) * D("0.09")
        inv = TaxInvoice(
            number=number,
            kind=kind,
            invoice_date=on,
            due_date=on + timedelta(days=30),
            site_id=s.id,
            client_id=s.client_id,
            contract_id=contracts[s.id],
            taxable=taxable,
            cgst=gst,
            sgst=gst,
            total=D(taxable) + 2 * gst,
            retention=retention,
            against_id=against,
        )
        db.add(inv)
        db.flush()
        return inv

    def receipt(number, s, on, amount, inv):
        r = Receipt(
            number=number,
            client_id=s.client_id,
            site_id=s.id,
            on_date=on,
            mode="neft",
            amount=amount,
        )
        db.add(r)
        db.flush()
        db.add(ReceiptAllocation(receipt_id=r.id, invoice_id=inv.id, amount=amount))

    i1 = invoice("T/0001", s1, day - timedelta(days=90), 200_000, 10_000)  # total 236,000
    receipt("R-1", s1, day - timedelta(days=70), 150_000, i1)
    i2 = invoice("T/0002", s2, day, 100_000, 5_000)  # total 118,000
    receipt("R-2", s2, day, 50_000, i2)
    invoice("T/CN1", s2, day, 10_000, kind="credit_note", against=i2.id)  # total 11,800
    i3 = invoice("T/0003", s3, prev0, 50_000)  # total 59,000, unpaid
    for s, amount in ((s1, 150_000), (s2, 120_000), (s3, 20_000)):
        db.add(
            SiteCost(
                site_id=s.id, head="material", on_date=day, amount=amount, description="invented"
            )
        )
    owner = uid(db, "boss@example.com")
    t1 = Tender(
        code="T-T-1",
        name="Won now",
        client_id=a.id,
        status="won",
        quoted_total=1_000_000,
        owner_id=owner,
        decided_at=at(day),
    )
    t2 = Tender(
        code="T-T-2",
        name="Won last month",
        client_id=b.id,
        status="won",
        quoted_total=400_000,
        owner_id=owner,
        decided_at=at(prev0),
    )
    t3 = Tender(
        code="T-T-3",
        name="Lost now",
        client_id=a.id,
        status="lost",
        lost_reason="price",
        quoted_total=300_000,
        owner_id=owner,
        decided_at=at(day),
    )
    t4 = Tender(
        code="T-T-4",
        name="Submitted now",
        client_id=b.id,
        status="submitted",
        quoted_total=200_000,
        owner_id=owner,
    )
    db.add_all([t1, t2, t3, t4])
    db.flush()
    for t, when in (
        (t1, day - timedelta(days=70)),
        (t2, day - timedelta(days=70)),
        (t3, day - timedelta(days=70)),
        (t4, day),
    ):
        db.add(TenderRevision(tender_id=t.id, rev_no=1, snapshot={}, submitted_at=at(when)))
    soon = VendorBill(
        number="VB-T-1",
        kind="material",
        vendor_id=v.id,
        bill_no="X1",
        bill_date=day,
        due_date=day + timedelta(days=3),
        taxable=10_000,
        total=10_000,
        payable=10_000,
        status="approved",
    )
    late = VendorBill(
        number="VB-T-2",
        kind="material",
        vendor_id=v.id,
        bill_no="X2",
        bill_date=day - timedelta(days=40),
        due_date=day - timedelta(days=5),
        taxable=20_000,
        total=20_000,
        payable=20_000,
        status="partly_paid",
    )
    db.add_all([soon, late])
    db.flush()
    pay = Payment(
        number="PAY-T-1", vendor_id=v.id, on_date=day, mode="neft", amount=5_000, status="paid"
    )
    db.add(pay)
    db.flush()
    db.add(PaymentAllocation(payment_id=pay.id, vendor_bill_id=late.id, amount=5_000))
    db.add(Dpr(site_id=s2.id, on_date=day, status="submitted", work_done="Primer"))
    db.commit()
    summary.refresh(db)
    return {
        "sites": [s1.id, s2.id, s3.id, s4.id],
        "clients": [a.id, b.id],
        "invoices": [i1.id, i2.id, i3.id],
        "vendor": v.id,
        "supervisor": sup,
        "m0": m0,
        "prev0": prev0,
    }


def tiles(data):
    return {t["key"]: t for sec in data["sections"] for t in sec["tiles"]}


# --- headline numbers ----------------------------------------------------------------------------


def test_management_kpis_from_fixed_fixtures(boss, company):
    client, h = boss
    day = today()
    d = client.get("/api/dashboard/company", headers=h).json()
    t = tiles(d)
    # order book: S1-S3 (S4 is completed): 3.5 M contracts, billed 200,000 + 90,000 (less the credit note) + 50,000
    assert D(str(t["order_value"]["value"])) == 3_500_000
    assert D(str(t["order_billed"]["value"])) == 340_000
    assert D(str(t["order_balance"]["value"])) == 3_160_000
    # this month: invoice 100,000 less credit note 10,000; received 50,000; one won (1,000,000), one submitted
    assert (
        D(str(t["billed_month"]["value"])) == 90_000 and D(str(t["billed_month"]["prev"])) == 50_000
    )
    assert (
        D(str(t["received_month"]["value"])) == 50_000 and D(str(t["received_month"]["prev"])) == 0
    )
    assert (
        D(str(t["won_month"]["value"])) == 1_000_000 and D(str(t["won_month"]["prev"])) == 400_000
    )
    assert t["submitted_month"]["value"] == 1
    # receivables: i1 236,000 - 150,000 = 86,000 (10,000 retention); i2 118,000 - 50,000 - 11,800 = 56,200 (5,000);
    # i3 59,000
    assert D(str(t["outstanding"]["value"])) == 86_000 + 56_200 + 59_000
    over = 76_000 + (59_000 if company["prev0"] < day - timedelta(days=60) else 0)
    assert D(str(t["overdue_60"]["value"])) == over
    assert D(str(t["retention_held"]["value"])) == 15_000
    # payables: 10,000 due in 3 days; 20,000 - 5,000 paid overdue
    assert (
        D(str(t["payables_week"]["value"])) == 10_000
        and D(str(t["payables_overdue"]["value"])) == 15_000
    )
    # sites: three active; S1 is 20 points behind time (> 15), S3 is past its planned end; S2 has today's DPR
    assert (t["active_sites"]["value"], t["delayed_sites"]["value"], t["dpr_missing"]["value"]) == (
        3,
        2,
        2,
    )
    # margin: (340,000 - 290,000) / 340,000 = 14.7 %; S2 costs more than it billed
    m = {x["key"]: x for x in d["margin"]["tiles"]}
    assert D(str(m["margin"]["value"])) == D("14.7") and m["over_cost"]["value"] == 1
    assert [r["site_id"] for r in d["margin"]["lowest"]][:1] == [company["sites"][1]]
    assert all(x["as_of"] for x in (t["order_value"], t["outstanding"], m["margin"]))
    # every tile drills down to its records
    rows = client.get(
        "/api/analytics/drill",
        params={"kind": "invoices", "filter": "overdue", "days": 60},
        headers=h,
    ).json()["rows"]
    assert {r["number"] for r in rows} >= {"T/0001"} and all(r["age"] > 60 for r in rows)
    delayed = client.get(
        "/api/analytics/drill", params={"kind": "sites", "filter": "delayed"}, headers=h
    ).json()["rows"]
    assert {r["code"] for r in delayed} == {"S-T-1", "S-T-3"}
    month = client.get(
        "/api/analytics/drill",
        params={"kind": "invoices", "filter": "month", "month": str(company["m0"])},
        headers=h,
    ).json()
    assert sum(D(str(r["taxable"])) for r in month["rows"]) == 90_000


def test_margin_is_hidden_without_tender_margin(company, db, make_user):
    from app.models import Role, RolePermission

    role = Role(
        code="mgmt_no_cost",
        name="Management without cost",
        permissions=[
            RolePermission(permission_code="dashboard.company", scope="all"),
            RolePermission(permission_code="reports.export", scope="all"),
        ],
    )
    db.add(role)
    db.commit()
    make_user("viewer@example.com", "mgmt_no_cost")
    c = TestClient(app)
    h = login(c, "viewer@example.com")
    d = c.get("/api/dashboard/company", headers=h).json()
    assert d["margin"] is None and "cost" not in str(tiles(d))
    assert (
        c.get(
            "/api/analytics/drill", params={"kind": "sites", "filter": "margin"}, headers=h
        ).status_code
        == 403
    )
    assert c.get("/api/analytics/pages/profitability", headers=h).status_code == 403
    rows = c.get(
        "/api/analytics/drill", params={"kind": "sites", "filter": "active"}, headers=h
    ).json()["rows"]
    assert rows and all("margin" not in r and "cost" not in r for r in rows)
    pdf = c.get("/api/analytics/pdf/management", headers=h)
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"


def test_summary_tables_match_the_live_numbers(company, db):
    for site_id in company["sites"]:
        live = fsvc.site_profit(db, db.get(Site, site_id), True)
        s = db.get(SiteSummary, site_id)
        assert (D(s.billed), D(s.cost) + D(s.salary_cost)) == (
            D(live["billed"]),
            D(live["cost_total"]),
        ), site_id
        assert D(s.received) == D(live["received"]) and D(s.retention_held) == D(
            live["retention_held"]
        )
    for inv_id in company["invoices"]:
        inv = db.get(TaxInvoice, inv_id)
        assert D(db.get(InvoiceSummary, inv_id).outstanding) == fsvc.outstanding(db, inv)
    billed = db.scalar(
        select(func.sum(MonthlySummary.billed)).where(MonthlySummary.is_demo.is_(False))
    )
    assert D(billed) == sum(D(db.get(SiteSummary, s).billed) for s in company["sites"])
    received = db.scalar(select(func.sum(MonthlySummary.received)))
    assert D(received) == db.scalar(select(func.sum(Receipt.amount)))


# --- role dashboards -----------------------------------------------------------------------------


def test_each_role_gets_only_its_dashboard_and_data(company, db, make_user, login_as):
    sales_a = make_user("sales-a@example.com", "sales", name="Sales A")
    sales_b = make_user("sales-b@example.com", "sales", name="Sales B")
    day = today()
    for i, (owner, status, follow) in enumerate(
        [
            (sales_a, "new", day),
            (sales_a, "quoted", day - timedelta(days=3)),
            (sales_b, "contacted", day),
            (sales_b, "site_visit", None),
        ]
    ):
        db.add(
            Lead(
                code=f"L-T-{i}",
                contact_name=f"Contact {i}",
                phone=f"+9190000000{i}",
                lead_source="call",
                status=status,
                owner_id=owner.id,
                est_value=100_000 * (i + 1),
                next_follow_up=follow,
            )
        )
    db.commit()
    ca = TestClient(app)
    ha = login(ca, "sales-a@example.com")
    home = ca.get("/api/dashboard", headers=ha).json()
    assert [d["key"] for d in home["dashboards"]] == ["sales"] and home["dashboards"][0][
        "scope"
    ] == "own"
    s = ca.get("/api/dashboard/sales", headers=ha).json()
    assert (
        sum(r["count"] for r in s["pipeline"] if r["kind"] == "lead") == 2
    )  # only their own leads
    t = {x["key"]: x["value"] for x in s["tiles"]}
    assert (t["followups_today"], t["followups_overdue"]) == (1, 1)
    leads = ca.get(
        "/api/analytics/drill", params={"kind": "leads", "filter": "follow_today"}, headers=ha
    ).json()["rows"]
    assert [r["contact"] for r in leads] == ["Contact 0"]
    assert ca.get("/api/dashboard/company", headers=ha).status_code == 403
    assert ca.get("/api/analytics/pages/finance", headers=ha).status_code == 403
    # office admin: all leads
    co, ho = login_as("office_admin", email="office@example.com")
    s = co.get("/api/dashboard/sales", headers=ho).json()
    assert sum(r["count"] for r in s["pipeline"] if r["kind"] == "lead") == 4
    # supervisor: only their assigned site
    cs = TestClient(app)
    hs = login(cs, "sup-1@example.com")
    mine = cs.get("/api/dashboard/site", headers=hs).json()
    assert [x["code"] for x in mine["sites"]] == ["S-T-1"]
    todo = {x["key"]: x for x in mine["sites"][0]["todos"]}
    assert todo["dpr"]["done"] is False and todo["dpr"]["link"].endswith("tab=dpr")
    pages = cs.get("/api/analytics/pages/sites", headers=hs).json()
    assert {r["code"] for r in pages["tables"]["forecast"]["rows"]} == {"S-T-1"}
    assert cs.get("/api/dashboard/company", headers=hs).status_code == 403
    assert cs.get("/api/dashboard/finance", headers=hs).status_code == 403
    # accounts and store / purchase
    cf, hf = login_as("accounts", email="accounts@example.com")
    fin = cf.get("/api/dashboard/finance", headers=hf).json()
    ageing = {r["bucket"]: D(str(r["amount"])) for r in fin["ageing"]}
    assert sum(ageing.values()) == D(76_000 + 51_200 + 59_000)
    assert cf.get("/api/dashboard/purchase", headers=hf).status_code == 403
    cp, hp = login_as("store_purchase", email="store@example.com")
    assert cp.get("/api/dashboard/purchase", headers=hp).status_code == 200
    assert cp.get("/api/dashboard/sales", headers=hp).status_code == 403


def test_client_logins_get_403_everywhere(company, login_as):
    c, h = login_as("client", email="client@example.com")
    for path in (
        "/api/dashboard",
        "/api/dashboard/company",
        "/api/analytics/pages/sites",
        "/api/analytics/strategy/funnel",
        "/api/analytics/alerts",
        "/api/analytics/drill?kind=sites",
    ):
        assert c.get(path, headers=h).status_code == 403, path


def test_purchase_dashboard_low_stock_and_late_pos(company, db, login_as):
    day = today()
    store = db.scalar(select(Store).where(Store.kind == "godown"))
    low = Product(code="P-LOW", name="Invented primer", unit="kg", reorder_level=100)
    ok = Product(code="P-OK", name="Invented coat", unit="kg", reorder_level=10)
    db.add_all([low, ok])
    db.flush()
    for p, qty in ((low, 40), (ok, 50)):
        db.add(
            StockLedger(
                store_id=store.id,
                product_id=p.id,
                qty=qty,
                unit="kg",
                rate=1,
                value=qty,
                ref_type="opening",
            )
        )
    db.add(
        PurchaseOrder(
            code="PO-T-1",
            vendor_id=company["vendor"],
            store_id=store.id,
            po_date=day - timedelta(days=20),
            expected_delivery=day - timedelta(days=2),
            status="sent",
        )
    )
    db.add(
        PurchaseOrder(
            code="PO-T-2",
            vendor_id=company["vendor"],
            store_id=store.id,
            po_date=day,
            expected_delivery=day + timedelta(days=5),
            status="approved",
        )
    )
    db.commit()
    c, h = login_as("store_purchase", email="store@example.com")
    t = {x["key"]: x["value"] for x in c.get("/api/dashboard/purchase", headers=h).json()["tiles"]}
    assert (t["pos_awaiting"], t["pos_late"], t["low_stock"]) == (2, 1, 1)
    rows = c.get(
        "/api/analytics/drill", params={"kind": "products", "filter": "low_stock"}, headers=h
    ).json()["rows"]
    assert [r["code"] for r in rows] == ["P-LOW"]


# --- forecasts, funnel, win / loss ---------------------------------------------------------------


def test_forecast_end_from_progress_history(company, db):
    day = today()
    rate, end = summary.forecast(
        [(day - timedelta(days=30), D(20))], D(50), day - timedelta(days=200), day
    )
    assert rate == D("1.0000") and end == day + timedelta(days=50)  # 30 points in 30 days, 50 to go
    rate, end = summary.forecast(
        [], D(40), day - timedelta(days=100), day
    )  # no history: since the start
    assert rate == D("0.4000") and end == day + timedelta(days=150)
    rate, end = summary.forecast(
        [(day - timedelta(days=3), D(10))], D(12), day - timedelta(days=100), day
    )
    assert rate == D("0.1200")  # under 7 days of history: the average since the start
    assert summary.forecast([], D(100), None, day)[1] == day
    # too slow, too little history: no date
    assert summary.forecast([], D(4), day - timedelta(days=100), day) == (
        D("0.0400"),
        None,
    )  # 0.04 % a day
    assert summary.forecast([], D(5), day - timedelta(days=10), day) == (
        None,
        None,
    )  # started 10 days ago
    assert summary.forecast_text(None, day, 40) == ("not enough progress to forecast", None)
    assert summary.forecast_text(day + timedelta(days=800), day, 40) == (
        "more than 2 years late",
        None,
    )
    assert summary.forecast_text(day + timedelta(days=20), day, 40) == (
        f"{day + timedelta(days=20):%d %b %Y}",
        20,
    )
    # end to end through the nightly snapshots
    s1 = company["sites"][0]
    db.add(ProgressSnapshot(site_id=s1, on_date=day - timedelta(days=30), percent=15))
    db.commit()
    summary.refresh(db)
    s = db.get(SiteSummary, s1)
    assert s.rate_per_day == D("0.5000") and s.forecast_end == day + timedelta(
        days=140
    )  # (30-15)/30 a day, 70 to go
    assert db.get(ProgressSnapshot, (s1, day)).percent == 30


def make_tenders(db, rows):
    """rows: (code, client_type, status, quoted, lead source or None, decided days ago)"""
    day = today()
    out = []
    clients = {}
    for code, ctype, status, quoted, source, ago in rows:
        if ctype not in clients:
            c = Client(name=f"Client {ctype}", type=ctype, state="Gujarat")
            db.add(c)
            db.flush()
            clients[ctype] = c
        t = Tender(
            code=code,
            name=code,
            client_id=clients[ctype].id,
            status=status,
            quoted_total=quoted,
            received_on=day - timedelta(days=60),
            lost_reason="price" if status == "lost" else None,
            site_state="Gujarat",
            decided_at=at(day - timedelta(days=ago)) if status in ("won", "lost") else None,
        )
        db.add(t)
        db.flush()
        if status != "draft":
            db.add(
                TenderRevision(
                    tender_id=t.id, rev_no=1, snapshot={}, submitted_at=at(day - timedelta(days=50))
                )
            )
        if source:
            db.add(
                Lead(
                    code=f"L-{code}",
                    contact_name=code,
                    phone="+919000000001",
                    lead_source=source,
                    status="quoted",
                    tender_id=t.id,
                    client_id=clients[ctype].id,
                )
            )
        out.append(t)
    db.commit()
    return out


def test_funnel_conversion_maths(boss, db):
    client, h = boss
    make_tenders(
        db,
        [
            ("F-1", "builder", "won", 100, "call", 5),
            ("F-2", "builder", "lost", 200, "call", 5),
            ("F-3", "developer", "draft", 300, None, 0),
        ],
    )
    db.add(
        Lead(
            code="L-OPEN",
            contact_name="open",
            phone="+919000000002",
            lead_source="website",
            status="new",
        )
    )
    db.add(
        Lead(
            code="L-JUNK",
            contact_name="junk",
            phone="+919000000003",
            lead_source="website",
            status="junk",
        )
    )
    db.commit()
    f = client.get("/api/analytics/strategy/funnel", headers=h).json()["total"]
    steps = {s["step"]: s for s in f["steps"]}
    # leads 3 (junk out) -> tenders 3 (100 %) -> submitted 2 (66.7 %) -> won 1 (50 %); win rate 1 / (1 + 1)
    assert [steps[k]["count"] for k in ("leads", "tenders", "submitted", "won")] == [3, 3, 2, 1]
    assert [steps[k]["conversion"] for k in ("tenders", "submitted", "won")] == [100.0, 66.7, 50.0]
    assert (
        D(str(steps["won"]["value"])) == 100 and f["lost"]["count"] == 1 and f["win_rate"] == 50.0
    )
    by = client.get(
        "/api/analytics/strategy/funnel", params={"group_by": "source"}, headers=h
    ).json()["groups"]
    g = {x["key"]: {s["step"]: s["count"] for s in x["steps"]} for x in by}
    assert g["call"] == {"leads": 2, "tenders": 2, "submitted": 2, "won": 1}
    assert g["direct (no lead)"]["tenders"] == 1 and g["website"]["leads"] == 1


def test_win_rate_by_segment_and_lost_reasons(boss, db):
    client, h = boss
    make_tenders(
        db,
        [
            ("W-1", "builder", "won", 500_000, None, 5),
            ("W-2", "builder", "won", 3_000_000, None, 5),
            ("W-3", "builder", "lost", 900_000, None, 5),
            ("W-4", "developer", "lost", 800_000, None, 5),
            ("W-5", "developer", "lost", 800_000, None, 900),
        ],
    )  # W-5: decided before the range
    w = client.get("/api/analytics/strategy/winloss", headers=h).json()
    seg = {x["key"]: x for x in w["segments"]["client_type"]}
    assert (seg["builder"]["won"], seg["builder"]["lost"], seg["builder"]["win_rate"]) == (
        2,
        1,
        66.7,
    )
    assert (seg["developer"]["won"], seg["developer"]["lost"], seg["developer"]["win_rate"]) == (
        0,
        1,
        0.0,
    )
    bands = {x["key"]: x["win_rate"] for x in w["segments"]["band"]}
    assert (
        bands["< ₹10 L"] == 33.3 and bands["₹10–50 L"] == 100.0
    )  # W-1 won; W-3, W-4 lost; W-2 won
    assert w["decided"] == 4 and w["win_rate"] == 50.0
    reasons = {r["reason"]: r for r in w["lost_reasons"]}
    assert reasons["price"]["tenders"] == 2
    drill = client.get(
        "/api/analytics/drill",
        params={"kind": "tenders", "filter": "segment", "client_type": "builder", "status": "won"},
        headers=h,
    ).json()
    assert {r["code"] for r in drill["rows"]} == {"W-1", "W-2"}


def test_lost_needs_a_reason_on_leads_and_tenders(boss, db):
    client, h = boss
    lead = Lead(
        code="L-LOST",
        contact_name="x",
        phone="+919000000009",
        lead_source="call",
        status="contacted",
    )
    c = Client(name="X", type="builder")
    db.add_all([lead, c])
    db.flush()
    t = Tender(code="T-LOST", name="x", client_id=c.id, status="submitted")
    db.add(t)
    db.commit()
    assert (
        client.patch(f"/api/leads/{lead.id}", json={"status": "lost"}, headers=h).status_code == 422
    )
    assert (
        client.patch(
            f"/api/leads/{lead.id}", json={"status": "lost", "lost_reason": "cheap"}, headers=h
        ).status_code
        == 422
    )
    r = client.patch(
        f"/api/leads/{lead.id}",
        json={"status": "lost", "lost_reason": "competitor", "lost_to": "Rival"},
        headers=h,
    )
    assert r.status_code == 200 and (r.json()["lost_reason"], r.json()["lost_to"]) == (
        "competitor",
        "Rival",
    )
    assert (
        client.patch(f"/api/tenders/{t.id}", json={"status": "lost"}, headers=h).status_code == 422
    )
    r = client.patch(
        f"/api/tenders/{t.id}", json={"status": "lost", "lost_reason": "spec"}, headers=h
    )
    assert r.status_code == 200 and r.json()["decided_at"]
    r = client.patch(f"/api/leads/{lead.id}", json={"status": "contacted"}, headers=h)
    assert r.json()["lost_reason"] is None  # cleared when reopened


# --- alerts --------------------------------------------------------------------------------------


@pytest.fixture
def alert_world(company, db, make_user):
    day = today()
    s1 = company["sites"][0]
    make_user("office@example.com", "office_admin")
    make_user("accounts@example.com", "accounts")
    make_user("store@example.com", "store_purchase")
    db.add(
        SiteBudget(site_id=s1, head="material", amount=160_000, source="manual")
    )  # cost 150,000 = 94 %
    i = db.get(TaxInvoice, company["invoices"][2])
    i.due_date = day - timedelta(days=35)  # unpaid, 35 days late
    store = db.scalar(select(Store).where(Store.kind == "godown"))
    p = Product(code="P-AL", name="Invented sealant", unit="kg", reorder_level=50)
    db.add(p)
    db.flush()
    db.add(
        StockLedger(
            store_id=store.id,
            product_id=p.id,
            qty=5,
            unit="kg",
            rate=1,
            value=5,
            ref_type="opening",
        )
    )
    db.add(
        PurchaseOrder(
            code="PO-AL-1",
            vendor_id=company["vendor"],
            store_id=store.id,
            po_date=day - timedelta(days=10),
            expected_delivery=day - timedelta(days=1),
            status="sent",
        )
    )
    c = Client(name="Tender client", type="builder")
    db.add(c)
    db.flush()
    db.add(
        Tender(
            code="T-DUE",
            name="Due tomorrow",
            client_id=c.id,
            status="draft",
            due_on=day + timedelta(days=1),
            owner_id=uid(db, "boss@example.com"),
        )
    )
    acc = PettyCashAccount(user_id=company["supervisor"].id)
    db.add(acc)
    db.flush()
    db.add(
        PettyCashEntry(
            number="PC-T-1",
            account_id=acc.id,
            kind="expense",
            on_date=day,
            amount=500,
            status="approved",
        )
    )
    # demo data never raises alerts
    db.add(
        Site(
            code="S-DEMO",
            name="DEMO",
            status="active",
            is_demo=True,
            start_date=day - timedelta(days=100),
            target_date=day - timedelta(days=1),
            progress_percent=10,
        )
    )
    db.commit()
    return company


def test_each_alert_rule_fires_once_per_item_per_day_to_the_right_people(alert_world, db):
    evening = at(today(), 21)
    raised = alert_svc.run(db, evening)
    assert {k for k, v in raised.items() if v} == {
        "dpr_missing",
        "behind_schedule",
        "budget_head",
        "invoice_overdue",
        "po_late",
        "low_stock",
        "petty_negative",
        "tender_due",
    }
    assert raised["behind_schedule"] == 2 and raised["kylas_failing"] == 0  # Kylas is off in tests
    assert alert_svc.run(db, evening + timedelta(minutes=50)) == dict.fromkeys(
        raised, 0
    )  # once per item per day
    assert not db.scalar(
        select(func.count()).select_from(Alert).where(Alert.title.contains("S-DEMO"))
    )

    def who(rule):
        ids = db.scalars(
            select(AlertRecipient.user_id)
            .join(Alert, Alert.id == AlertRecipient.alert_id)
            .where(Alert.rule == rule)
        )
        return {db.get(User, i).email for i in ids}

    assert who("po_late") == {"store@example.com"} and who("low_stock") == {"store@example.com"}
    assert who("invoice_overdue") == {"accounts@example.com", "office@example.com"}
    assert {"boss@example.com", "office@example.com", "sup-1@example.com"} <= who("behind_schedule")
    assert "store@example.com" not in who("behind_schedule")
    assert who("petty_negative") == {"accounts@example.com", "sup-1@example.com"}  # and the holder
    assert who("tender_due") == {"office@example.com", "boss@example.com"}
    title = db.scalar(
        select(Alert.title).where(Alert.rule == "invoice_overdue", Alert.title.contains("T/0003"))
    )
    assert "T/0003" in title and "35 days" in title
    store_user = uid(db, "store@example.com")
    assert (
        db.scalar(
            select(func.count()).select_from(Notification).where(Notification.user_id == store_user)
        )
        == 2
    )
    # the morning run does not raise the 8 pm DPR rule
    db.execute(text("DELETE FROM alert_recipients"))
    db.execute(text("DELETE FROM alerts"))
    db.commit()
    assert alert_svc.run(db, at(today(), 9))["dpr_missing"] == 0


def test_acknowledge_snoozes_the_item_and_others_come_back_next_day(alert_world, db, login_as):
    alert_svc.run(db, at(today(), 21))
    c, h = login_as("store_purchase", email="store2@example.com")
    assert c.get("/api/analytics/alerts", headers=h).json()["alerts"] == []  # not theirs
    c = TestClient(app)
    h = login(c, "store@example.com")
    mine = c.get("/api/analytics/alerts", headers=h).json()["alerts"]
    assert {a["rule"] for a in mine} == {"po_late", "low_stock"}
    po = next(a for a in mine if a["rule"] == "po_late")
    r = c.post(f"/api/analytics/alerts/{po['id']}/ack", headers=h)
    assert r.status_code == 200 and r.json()["acknowledged_by"]
    assert {a["rule"] for a in c.get("/api/analytics/alerts", headers=h).json()["alerts"]} == {
        "low_stock"
    }
    other = db.scalar(select(Alert.id).where(Alert.rule == "budget_head"))
    assert (
        c.post(f"/api/analytics/alerts/{other}/ack", headers=h).status_code == 404
    )  # not sent to them
    tomorrow = alert_svc.run(db, at(today() + timedelta(days=1), 21))
    assert (
        tomorrow["po_late"] == 0 and tomorrow["low_stock"] == 1
    )  # acknowledged: snoozed; the rest again
    old = db.scalars(select(Alert).where(Alert.rule == "low_stock").order_by(Alert.id)).all()
    assert len(old) == 2 and old[0].closed_at is not None and old[1].closed_at is None


# --- reports, exports, worker --------------------------------------------------------------------


def test_exports_need_reports_export_and_pdfs_render(company, boss, login_as, db):
    client, h = boss
    r = client.get(
        "/api/analytics/drill",
        params={"kind": "invoices", "filter": "outstanding", "format": "xlsx"},
        headers=h,
    )
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert client.get("/api/dashboard/company/xlsx", headers=h).content[:2] == b"PK"
    assert (
        client.get("/api/analytics/pages/finance", params={"format": "xlsx"}, headers=h).content[:2]
        == b"PK"
    )
    for path in (
        "/api/analytics/pdf/management",
        f"/api/analytics/pdf/site/{company['sites'][0]}",
        "/api/analytics/pdf/receivables",
        "/api/analytics/pdf/pipeline",
    ):
        r = client.get(path, headers=h)
        assert r.status_code == 200 and r.content[:4] == b"%PDF", path
    cs = TestClient(app)
    hs = login(cs, "sup-1@example.com")
    assert cs.get("/api/dashboard/site", headers=hs).status_code == 200
    assert cs.get("/api/dashboard/site/xlsx", headers=hs).status_code == 403  # no reports.export
    assert (
        db.scalar(
            select(func.count())
            .select_from(Base.metadata.tables["audit_log"])
            .where(text("action = 'report.export'"))
        )
        >= 7
    )


def test_weekly_summary_once_a_week_into_the_bell(company, db, make_user):
    make_user("office@example.com", "office_admin")
    r = reports.weekly(db)
    assert r is not None and r.week_start.weekday() == 0
    assert reports.weekly(db) is None
    office = uid(db, "office@example.com")
    assert (
        db.scalar(
            select(Notification.link).where(
                Notification.user_id == office, Notification.kind == "weekly_report"
            )
        )
        == "/reports"
    )
    from app.portal.models import NotificationOutbox

    assert (
        db.scalar(
            select(NotificationOutbox.channel)
            .join(Notification, Notification.id == NotificationOutbox.notification_id)
            .where(Notification.kind == "weekly_report")
            .limit(1)
        )
        == "in_app"
    )
    assert db.scalar(select(func.count()).select_from(WeeklyReport)) == 1


def test_worker_schedule():
    monday = datetime(2026, 10, 12, 8, 30, tzinfo=IST)
    jobs = worker.due(monday)
    assert {"refresh", "alerts", "weekly"} <= set(
        jobs
    )  # never refreshed or run yet; Monday 8:30 IST
    sunday = datetime(2026, 10, 11, 8, 30, tzinfo=IST)
    assert "weekly" not in worker.due(sunday)


# --- the demo company ----------------------------------------------------------------------------


def test_demo_seed_and_purge_leave_real_data_untouched(company, db, boss):
    client, h = boss
    real = client.get("/api/dashboard/company", headers=h).json()
    skip = {
        "site_summaries",
        "monthly_summaries",
        "invoice_summaries",
        "progress_snapshots",
        "kpi_snapshots",
        "analytics_settings",
    }
    tables = [t for t in sorted(Base.metadata.tables) if t not in skip]

    def counts():
        return {t: db.execute(text(f'SELECT count(*) FROM "{t}"')).scalar() for t in tables}

    before = counts()
    made = demo.seed(db, sites=4, leads=30)
    assert made["sites"] >= 1 and made["leads"] == 30
    with pytest.raises(ValueError):
        demo.seed(db, sites=4, leads=30)
    # the real dashboard does not move; the demo one has its own numbers
    again = client.get("/api/dashboard/company", headers=h).json()
    assert tiles(again)["order_value"]["value"] == tiles(real)["order_value"]["value"]
    assert tiles(again)["active_sites"]["value"] == tiles(real)["active_sites"]["value"]
    d = client.get("/api/dashboard/company", params={"demo": "true"}, headers=h).json()
    assert tiles(d)["order_value"]["value"] != tiles(real)["order_value"]["value"]
    assert client.get("/api/dashboard", headers=h).json()["demo_available"] is True
    assert not db.scalar(
        select(func.count()).select_from(Site).where(Site.is_demo, Site.code.not_like("DEMO-%"))
    )
    demo.purge(db)
    db.expire_all()
    after = counts()
    assert {t: (before[t], after[t]) for t in tables if before[t] != after[t]} == {}
    assert client.get("/api/dashboard", headers=h).json()["demo_available"] is False
    assert db.scalar(select(func.count()).select_from(SiteSummary).where(SiteSummary.is_demo)) == 0


def test_funnel_and_winloss_state_their_basis_also_in_excel(boss, db):
    from io import BytesIO

    from openpyxl import load_workbook

    client, h = boss
    make_tenders(db, [("B-1", "builder", "won", 100, "call", 5)])
    f = client.get("/api/analytics/strategy/funnel", headers=h).json()
    w = client.get("/api/analytics/strategy/winloss", headers=h).json()
    assert f["basis"].startswith("Basis: leads created") and w["basis"].startswith(
        "Basis: tenders decided"
    )
    for part, basis in (("funnel", f["basis"]), ("winloss", w["basis"])):
        r = client.get(f"/api/analytics/strategy/{part}", params={"format": "xlsx"}, headers=h)
        wb = load_workbook(BytesIO(r.content))
        assert all(ws["A1"].value == basis for ws in wb.worksheets), part


def test_compact_money_for_tiles():
    from app.analytics.common import inr_compact

    assert [inr_compact(v) for v in (294909829.72, 789420, 72228.4, -4500000)] == [
        "₹29.49 Cr",
        "₹7.89 L",
        "₹72,228",
        "-₹45.00 L",
    ]
